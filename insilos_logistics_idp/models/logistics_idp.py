import ast
import base64
import hashlib
import io
import json
import re
import sys
import zipfile
from datetime import datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation, ROUND_DOWN, ROUND_HALF_UP, ROUND_UP
from pathlib import Path
from uuid import UUID
from email.message import EmailMessage
from urllib.parse import urlsplit

import xlsxwriter
from openpyxl import load_workbook
from reportlab.pdfgen import canvas

from odoo import _, api, fields, models, tools
from odoo.exceptions import AccessError, UserError, ValidationError
from odoo.tools import consteq
from odoo.tools.mimetypes import guess_mimetype
from odoo.addons.mail.tools.discuss import Store
from psycopg2 import Error as DatabaseError
from psycopg2.errors import UniqueViolation

from ..services.document_processor import IAPDocumentProcessor, MAX_BYTES, ProcessingError, SyntheticDocumentProcessor, batch_structure, deterministic_classify, normalize_ocr_results
from ..services.document_schemas import CANONICAL_EXPORT_FIELDS, DOCUMENT_TYPES
from ..services.legal_candidate_adapter import fetch_candidates
from ..services.reconciliation import customs_checks, reconcile_draft_main, reconcile_documents as reconcile_service, validate_draft_invoice_policy, validate_main_invoice_policy
from ..services.restricted_party import canonical_party_evidence, normalize_party, screen_parties


class ResUsers(models.Model):
    _inherit = 'res.users'

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            if vals.get('tz') == 'Asia/Saigon':
                vals['tz'] = 'Asia/Ho_Chi_Minh'
        return super().create(vals_list)

    def write(self, vals):
        if vals.get('tz') == 'Asia/Saigon':
            vals['tz'] = 'Asia/Ho_Chi_Minh'
        return super().write(vals)

    @api.model
    def _normalize_legacy_saigon_timezone(self):
        self.env.cr.execute("""
            UPDATE res_partner partner
               SET tz = 'Asia/Ho_Chi_Minh'
              FROM res_users users
             WHERE users.partner_id = partner.id
               AND partner.tz = 'Asia/Saigon'
        """)
        count = self.env.cr.rowcount
        self.env['res.partner'].invalidate_model(['tz'])
        return count


_INTERNAL_SNAPSHOT_TOKEN = object()
_INTERNAL_POLICY_LOADER_TOKEN = object()
_INTERNAL_POLICY_WRITE_TOKEN = object()
_INTERNAL_CASE_LIFECYCLE_TOKEN = object()
_INTERNAL_DOCUMENT_TYPE_TOKEN = object()
_INTERNAL_DOCUMENT_DUPLICATE_TOKEN = object()
_INTERNAL_OCR_GEOMETRY_TOKEN = object()
_INTERNAL_INBOUND_ATTACHMENT_TOKEN = object()
_INTERNAL_EMAIL_THREAD_ENTRY_TOKEN = object()
_INTERNAL_OVERRIDE_REQUEST_TOKEN = object()
_INTERNAL_CHECK_RESULT_TOKEN = object()


VERDICTS = [
    ('pass', 'Pass'),
    ('review', 'Review'),
    ('block', 'Block'),
    ('not_applicable', 'Not Applicable'),
]

AUTHORITY_TIERS = {
    'demo': False,
    'customer_reference': False,
    'early_warning_tier_4': False,
    'authoritative_tier_1': True,
    'authoritative_tier_2': True,
    'approved_provider_tier_3': True,
}
AUTHORITY_RANK = {
    'demo': 0, 'customer_reference': 0, 'early_warning_tier_4': 1,
    'approved_provider_tier_3': 2, 'authoritative_tier_2': 3, 'authoritative_tier_1': 4,
}
COMPLIANCE_CONTEXT_FIELDS = (
    'who', 'what', 'where', 'why', 'value', 'currency', 'uom', 'preference',
    'procedure', 'transport', 'when', 'evidence',
)
COMPLIANCE_CONTEXT_CRITICAL = ('who', 'when', 'evidence')
MAX_POLICY_IMPORT_BYTES = 1024 * 1024
EVIDENCE_CATEGORIES = {
    'document_assignment', 'email_attachment', 'exact_document_duplicate', 'document_duplicate_mark', 'semantic_duplicate', 'extraction',
    'import_declaration', 'inbound_failure', 'invoice', 'ocr_manual_correction', 'reconciliation', 'waiting_transition',
    'broker_email_request', 'broker_email_queued', 'broker_email_delivery',
    'gate_pass_distribution_request', 'gate_pass_distribution_queued',
    'supplier_result_reply_request', 'supplier_result_reply_queued', 'supplier_result_reply_delivery',
    'policy_activation_reevaluation', 'counterpart_alert_intent', 'operational_exposure_signal', 'external_action',
}
_CASE_TERMINAL_STATES = {'completed', 'closed_duplicate', 'closed_other'}
_CASE_TERMINAL_FIELDS = {
    'canonical_case_id', 'completion_reason', 'document_status', 'reconciliation_status',
    'compliance_status', 'output_status', 'state', 'verdict', 'waiting_party', 'waiting_reason',
    'waiting_since',
}
_CASE_RESCREEN_FIELDS = {'po_reference', 'supplier_reference', 'shipment_reference'}


def _canonical_json(value):
    return json.dumps(value, ensure_ascii=False, separators=(',', ':'), sort_keys=True)


def _normalized(value):
    return ' '.join(str(value or '').strip().upper().split())


def _decimal(value):
    try:
        return Decimal(str(value or 0))
    except InvalidOperation as exc:
        raise ValidationError('Numeric reconciliation values must be valid decimals.') from exc


class LogisticsCase(models.Model):
    _name = 'logistics.idp.case'
    _description = 'Logistics IDP Case'
    _inherit = ['mail.thread', 'mail.activity.mixin']
    _order = 'create_date desc, id desc'

    name = fields.Char(required=True, tracking=True)
    source_system = fields.Char(required=True, readonly=True, default='manual_entry')
    source_key = fields.Char(required=True, readonly=True, default=lambda self: 'MANUAL-%s' % fields.Datetime.now().strftime('%Y%m%d%H%M%S%f'))
    source_version = fields.Char(required=True, readonly=True, default='1')
    provenance = fields.Char(required=True, readonly=True, default='manual_entry')
    effective_date = fields.Date(required=True, readonly=True, default=fields.Date.today)

    imported_at = fields.Datetime(required=True, readonly=True, default=fields.Datetime.now)
    state = fields.Selection([
        ('intake', 'Intake'),
        ('collecting', 'Collecting'),
        ('processing', 'Processing'),
        ('review', 'Review'),
        ('waiting_external', 'Waiting External'),
        ('ready', 'Ready to Complete'),
        ('blocked', 'Blocked'),
        ('completed', 'Completed'),
        ('closed_duplicate', 'Closed - Duplicated'),
        ('closed_other', 'Closed - Other'),
    ], required=True, default='collecting', tracking=True)
    verdict = fields.Selection(VERDICTS, required=True, default='review', readonly=True, tracking=True)
    po_reference = fields.Char(index=True, tracking=True)
    supplier_reference = fields.Char(index=True, tracking=True)
    email_subject = fields.Char(tracking=True)
    thread_reference = fields.Char(index=True, tracking=True)
    active_thread_case_id = fields.Many2one('logistics.idp.case', readonly=True, ondelete='restrict', tracking=True)
    email_thread_entry_ids = fields.One2many('logistics.idp.email.thread.entry', 'case_id', readonly=True)
    gate_pass_enabled = fields.Boolean(tracking=True)
    gate_pass_requires_import_pass = fields.Boolean(default=True, tracking=True)
    canonical_case_id = fields.Many2one('logistics.idp.case', readonly=True, ondelete='restrict', index=True)
    evidence_ids = fields.One2many('logistics.idp.evidence', 'case_id')
    decision_ids = fields.One2many('logistics.idp.policy.decision', 'case_id')
    mes_reference_ids = fields.One2many('logistics.idp.mes.reference', 'case_id')
    company_id = fields.Many2one('res.company', required=True, default=lambda self: self.env.company, index=True)
    task_id = fields.Many2one('project.task', readonly=True, ondelete='restrict', index=True)
    owner_id = fields.Many2one('res.users', default=lambda self: self.env.user, tracking=True)
    team_id = fields.Many2one('project.project', default=lambda self: self.env.ref('insilos_logistics_idp.project_logistics_idp', raise_if_not_found=False), tracking=True)
    shipment_reference = fields.Char(index=True, tracking=True)
    document_status = fields.Selection(VERDICTS, required=True, default='review', tracking=True)
    reconciliation_status = fields.Selection(VERDICTS, required=True, default='review', tracking=True)
    compliance_status = fields.Selection(VERDICTS, required=True, default='review', tracking=True)
    output_status = fields.Selection(VERDICTS, required=True, default='not_applicable', tracking=True)
    severity = fields.Selection([('low', 'Low'), ('medium', 'Medium'), ('high', 'High'), ('critical', 'Critical')], required=True, default='medium', tracking=True)
    sla_deadline = fields.Datetime(tracking=True)
    next_action = fields.Char(tracking=True)
    waiting_party = fields.Selection([
        ('supplier', 'Supplier'), ('broker_customs', 'Broker / Customs'), ('internal', 'Internal'),
    ], readonly=True, tracking=True, index=True)
    waiting_reason = fields.Char(readonly=True, tracking=True)
    waiting_since = fields.Datetime(readonly=True, tracking=True, index=True)
    completion_reason = fields.Char(readonly=True)
    document_ids = fields.One2many('logistics.idp.document', 'case_id')
    extraction_run_ids = fields.One2many('logistics.idp.extraction.run', 'case_id')
    check_result_ids = fields.One2many('logistics.idp.check.result', 'case_id')
    output_ids = fields.One2many('logistics.idp.output', 'case_id')
    exception_ids = fields.One2many('logistics.idp.exception', 'case_id')
    override_ids = fields.One2many('logistics.idp.override', 'case_id')
    document_count = fields.Integer(compute='_compute_cockpit_counts')
    run_count = fields.Integer(compute='_compute_cockpit_counts')
    check_count = fields.Integer(compute='_compute_cockpit_counts')
    output_count = fields.Integer(compute='_compute_cockpit_counts')
    override_count = fields.Integer(compute='_compute_cockpit_counts')
    operational_exposure = fields.Selection([('low', 'Low'), ('medium', 'Medium'), ('high', 'High'), ('critical', 'Critical')], readonly=True)

    def record_operational_exposure(self, severity, source_class, source_reference, provenance, payload):
        self.ensure_one()
        data = {**(payload or {}), 'severity': severity, 'source_class': source_class, 'provenance': provenance}
        identity = hashlib.sha256(_canonical_json(data).encode()).hexdigest()
        evidence = self.evidence_ids.filtered(lambda item: item.category == 'operational_exposure_signal' and item.source_reference == source_reference and item.payload_hash == identity)
        if evidence:
            return evidence
        evidence = self.env['logistics.idp.evidence']._controlled_create({
            'case_id': self.id, 'category': 'operational_exposure_signal', 'source_reference': source_reference,
            'status': 'review', 'payload': data,
        }, 'operational_exposure')
        self.write({'operational_exposure': severity})
        return evidence

    def _compute_cockpit_counts(self):
        for case in self:
            case.document_count = len(case.document_ids)
            case.run_count = len(case.extraction_run_ids)
            case.check_count = len(case.check_result_ids)
            case.output_count = len(case.output_ids)
            case.override_count = len(case.override_ids)

    @api.model
    def _upgrade_migration_service_principal(self):
        login = self.env['ir.config_parameter'].sudo().get_param(
            'logistics_idp.migration_service_login', 'digiforce_migration_service')
        service = self.env['res.users'].sudo().with_context(active_test=False).search([
            ('login', '=', login), ('company_ids', 'in', self.env.company.id),
        ], limit=1)
        if service:
            integration = self.env.ref('insilos_logistics_idp.group_logistics_integration')
            manager = self.env.ref('insilos_logistics_idp.group_logistics_manager')
            service.write({'active': True, 'company_id': self.env.company.id,
                           'group_ids': [(3, manager.id)] + [(4, integration.id)]})
        return service

    @api.model
    def get_dashboard_data(self, filters=None):
        filters = filters or {}
        allowed_companies = set(self.env.companies.ids)
        company_ids = set(filters.get('company_ids') or allowed_companies)
        if not company_ids or not company_ids <= allowed_companies:
            raise ValidationError('Dashboard companies must be enabled for the current user.')
        date_from = fields.Date.to_date(filters.get('date_from')) if filters.get('date_from') else None
        date_to = fields.Date.to_date(filters.get('date_to')) if filters.get('date_to') else None
        if date_from and date_to and date_from > date_to:
            raise ValidationError('Dashboard date_from must not be after date_to.')

        def dated(domain, field_name, start=date_from, end=date_to):
            domain = list(domain)
            if start:
                domain.append((field_name, '>=', start))
            if end:
                domain.append((field_name, '<', end + timedelta(days=1))
                              if field_name in {'create_date', 'completed_at', 'opened_at', 'approved_at'}
                              else (field_name, '<=', end))
            return domain

        def build_case_domain(start=date_from, end=date_to):
            domain = [('company_id', 'in', sorted(company_ids))]
            if filters.get('supplier'):
                domain.append(('supplier_reference', '=', filters['supplier']))
            case_sets = []
            if filters.get('document_type'):
                case_sets.append(set(self.env['logistics.idp.document'].search([
                    ('company_id', 'in', sorted(company_ids)), ('document_type', '=', filters['document_type'])
                ]).mapped('case_id').ids))
            policy_domain = [('case_id.company_id', 'in', sorted(company_ids))]
            if filters.get('jurisdiction'):
                policy_domain.append(('policy_source_id.jurisdiction', '=', filters['jurisdiction']))
            if filters.get('customs_regime'):
                policy_domain.append(('policy_source_id.regime', '=', filters['customs_regime'].upper()))
            if len(policy_domain) > 1:
                case_sets.append(set(self.env['logistics.idp.check.result'].search(policy_domain).mapped('case_id').ids))
            if filters.get('customs_regime'):
                regime = filters['customs_regime'].lower()
                regime_cases = set(self.env['logistics.idp.document'].search([
                    ('company_id', 'in', sorted(company_ids)), ('document_type', '=', regime)
                ]).mapped('case_id').ids)
                regime_cases |= set(self.env['logistics.idp.output'].search([
                    ('case_id.company_id', 'in', sorted(company_ids)), ('output_type', '=', regime)
                ]).mapped('case_id').ids)
                case_sets.append(regime_cases)
            if case_sets:
                domain.append(('id', 'in', sorted(set.intersection(*case_sets))))
            return dated(domain, 'effective_date', start, end)

        case_domain = build_case_domain()
        case_ids = self.search(case_domain).ids
        related = [('case_id', 'in', case_ids)]
        document_filters = [('document_type', '=', filters['document_type'])] if filters.get('document_type') else []
        regime_filters = [('output_type', '=', filters['customs_regime'].lower())] if filters.get('customs_regime') else []
        document_domain = dated(related + document_filters, 'create_date')
        run_domain = dated(related + document_filters, 'completed_at')
        check_domain = dated(related, 'create_date')
        output_domain = dated(related + regime_filters, 'create_date')
        exception_domain = dated(related, 'opened_at')
        now = fields.Datetime.now()
        closed = ('completed', 'closed_duplicate', 'closed_other')
        exception_model = self.env['logistics.idp.exception']
        exception_available = exception_model.has_access('read')
        overdue_case_ids = exception_model.search(exception_domain + [
            ('state', 'in', ('open', 'waiting')), ('due_at', '<', now),
        ]).mapped('case_id').ids if exception_available else []
        overdue_domain = case_domain + [('state', 'not in', closed), ('id', 'in', overdue_case_ids)]
        today = fields.Date.today()
        today_start = fields.Datetime.to_datetime(today)
        today_end = today_start + __import__('datetime').timedelta(days=1)
        received_today_domain = case_domain + [('create_date', '>=', today_start), ('create_date', '<', today_end)]
        processed_today_domain = related + [('current_run_id.completed_at', '>=', today_start), ('current_run_id.completed_at', '<', today_end)]
        current = {
            'open_cases': self.search_count(case_domain + [('state', 'not in', closed)]),
            'waiting_supplier': self.search_count(case_domain + [('state', '=', 'waiting_external'), ('waiting_party', '=', 'supplier')]),
            'waiting_broker_customs': self.search_count(case_domain + [('state', '=', 'waiting_external'), ('waiting_party', '=', 'broker_customs')]),
            'review_cases': self.search_count(case_domain + ['|', ('state', '=', 'review'), ('verdict', '=', 'review')]),
            'blocked_cases': self.search_count(case_domain + ['|', ('state', '=', 'blocked'), ('verdict', '=', 'block')]),
            'overdue_cases': self.search_count(overdue_domain),
            'documents': self.env['logistics.idp.document'].search_count(document_domain),
            'received_today': self.search_count(received_today_domain),
            'processed_today': self.env['logistics.idp.document'].search_count(processed_today_domain),
        }
        def stp(domain, start, end):
            completed_ids = set(self.search(domain + [('state', '=', 'completed')]).ids)
            passed_ids = set(self.search(domain + [('state', '=', 'completed'), ('verdict', '=', 'pass')]).ids)
            history_ids = set()
            history_specs = (
                ('logistics.idp.check.result', 'create_date', [('verdict', 'in', ('review', 'block'))]),
                ('logistics.idp.extraction.run', 'completed_at', [('status', 'in', ('review', 'error'))]),
                ('logistics.idp.override', 'approved_at', []),
            )
            for model, field_name, extra in history_specs:
                source = self.env[model]
                integration_override = model == 'logistics.idp.override' and self.env.user.has_group(
                    'insilos_logistics_idp.group_logistics_integration')
                if integration_override or not source.has_access('read'):
                    return None, False
                scoped_domain = [('case_id.company_id', 'in', sorted(company_ids)), ('case_id', 'in', sorted(completed_ids))]
                history_ids |= set(source.search(dated(scoped_domain + extra, field_name, start, end)).mapped('case_id').ids)
            return (len(passed_ids - history_ids) / len(completed_ids) if completed_ids else None), True

        completed_domain = case_domain + [('state', '=', 'completed')]
        current['stp_rate'], stp_available = stp(case_domain, date_from, date_to)
        duplicate_model = self.env['logistics.idp.evidence']
        duplicate_domain = dated(related + [('category', '=', 'exact_document_duplicate')], 'create_date')
        duplicate_available = duplicate_model.has_access('read')
        current['duplicate_detected'] = duplicate_model.search_count(duplicate_domain) if duplicate_available else None

        prior = None
        if date_from and date_to:
            days = (date_to - date_from).days + 1
            prior_to = date_from - __import__('datetime').timedelta(days=1)
            prior_from = prior_to - __import__('datetime').timedelta(days=days - 1)
            prior_domain = build_case_domain(prior_from, prior_to)
            prior_case_ids = self.search(prior_domain).ids
            prior_related = [('case_id', 'in', prior_case_ids)]
            prior_overdue_ids = exception_model.search(dated(prior_related, 'opened_at', prior_from, prior_to) + [
                ('state', 'in', ('open', 'waiting')), ('due_at', '<', now),
            ]).mapped('case_id').ids if exception_available else []
            prior = {
                'open_cases': self.search_count(prior_domain + [('state', 'not in', closed)]),
                'waiting_supplier': self.search_count(prior_domain + [('state', '=', 'waiting_external'), ('waiting_party', '=', 'supplier')]),
                'waiting_broker_customs': self.search_count(prior_domain + [('state', '=', 'waiting_external'), ('waiting_party', '=', 'broker_customs')]),
                'review_cases': self.search_count(prior_domain + ['|', ('state', '=', 'review'), ('verdict', '=', 'review')]),
                'blocked_cases': self.search_count(prior_domain + ['|', ('state', '=', 'blocked'), ('verdict', '=', 'block')]),
                'overdue_cases': self.search_count(prior_domain + [('state', 'not in', closed), ('id', 'in', prior_overdue_ids)]),
                'documents': self.env['logistics.idp.document'].search_count(dated(prior_related, 'create_date', prior_from, prior_to)),
                'stp_rate': stp(prior_domain, prior_from, prior_to)[0] if stp_available else None,
            }
        deltas = {
            key: (current[key] - prior[key]
                  if current[key] is not None and prior and key in prior and prior[key] is not None else None)
            for key in current
        }

        def groups(model, domain, groupby, aggregates=None, limit=None):
            source = self.env[model]
            if not source.has_access('read'):
                return []
            fields_list = ['__count'] + (aggregates or [])
            return source.formatted_read_group(domain, [groupby], fields_list, limit=limit)

        contracts = {
            'open_cases': self._dashboard_contract('count(state not closed)', 'logistics.idp.case', case_domain + [('state', 'not in', closed)], 'effective_date', 'cases', True),
            'waiting_supplier': self._dashboard_contract('count(state=waiting_external AND waiting_party=supplier)', 'logistics.idp.case', case_domain + [('state', '=', 'waiting_external'), ('waiting_party', '=', 'supplier')], 'waiting_since', 'cases', True),
            'waiting_broker_customs': self._dashboard_contract('count(state=waiting_external AND waiting_party=broker_customs)', 'logistics.idp.case', case_domain + [('state', '=', 'waiting_external'), ('waiting_party', '=', 'broker_customs')], 'waiting_since', 'cases', True),
            'review_cases': self._dashboard_contract('count(state=review OR verdict=review)', 'logistics.idp.case', case_domain + ['|', ('state', '=', 'review'), ('verdict', '=', 'review')], 'effective_date', 'cases', True),
            'blocked_cases': self._dashboard_contract('count(state=blocked OR deterministic verdict=block)', 'logistics.idp.case', case_domain + ['|', ('state', '=', 'blocked'), ('verdict', '=', 'block')], 'effective_date', 'cases', True),
            'overdue_cases': self._dashboard_contract('count(distinct open cases having an open/waiting exception whose governed due_at < now; due_at = company resource calendar plan_hours(severity SLA, opened_at, compute_leaves=True))', 'logistics.idp.case', overdue_domain, 'effective_date', 'cases', True),
            'documents': self._dashboard_contract('count(documents created in period)', 'logistics.idp.document', document_domain, 'create_date', 'documents', True),
            'received_today': self._dashboard_contract('count(cases created today within selected case domain)', 'logistics.idp.case', received_today_domain, 'create_date', 'cases', True),
            'processed_today': self._dashboard_contract('count(documents whose current extraction run completed today)', 'logistics.idp.document', processed_today_domain, 'current_run_id.completed_at', 'documents', True),
            'duplicate_detected': self._dashboard_contract('count(exact_document_duplicate evidence created in period for selected cases)', 'logistics.idp.evidence', duplicate_domain, 'create_date', 'documents', duplicate_available),
            'stp_rate': self._dashboard_contract('completed PASS cases with no review/block check, review/error extraction run, or override history in period / all completed cases in period; null when denominator=0', 'logistics.idp.case', completed_domain, 'effective_date', 'ratio', True),
        }
        def exception_groups(field, priority=()):
            ranks = {value: index for index, value in enumerate(priority)}
            return sorted(groups('logistics.idp.exception', exception_domain, field), key=lambda row: (
                ranks.get(row.get(field), len(ranks)), -row['__count'], str(row.get(field) or ''),
            ))

        supplier_counts = {}
        for row in groups('logistics.idp.exception', exception_domain, 'case_id'):
            case = self.browse(row['case_id'][0])
            supplier = case.supplier_reference or False
            supplier_counts[supplier] = supplier_counts.get(supplier, 0) + row['__count']
        top_suppliers = [{
            'supplier_reference': supplier, '__count': count,
            '__domain': exception_domain + [('case_id.supplier_reference', '=', supplier)],
        } for supplier, count in sorted(supplier_counts.items(), key=lambda item: (-item[1], str(item[0])))[:10]]
        onboarding = self.env.ref('insilos_logistics_idp.onboarding_logistics_idp')
        onboarding._search_or_create_progress()
        onboarding_values = onboarding._prepare_rendering_values()
        onboarding_rendering = {
            'closed': onboarding.is_onboarding_closed,
            'close_method': onboarding_values['close_method'],
            'close_model': onboarding_values['close_model'],
            'state': onboarding_values['state'],
            'text_completed': onboarding_values['text_completed'],
            'steps': [{
                'id': step.id, 'title': step.title, 'description': step.description,
                'button_text': step.button_text, 'done_text': step.done_text,
                'action': step.panel_step_open_action_name,
                'state': onboarding_values['state'].get(step.id),
            } for step in onboarding_values['steps']],
        }
        return {
            'onboarding': onboarding_rendering,
            'generated_at': fields.Datetime.to_string(now), 'filters': {**filters, 'company_ids': sorted(company_ids)},
            'drilldown_domains': {
                'case_history': case_domain,
                'document_history': document_domain,
                'ocr_history': run_domain,
                'output_history': output_domain,
                'check_history': check_domain,
                'mapping_exceptions': check_domain + [('verdict', 'in', ('review', 'block'))],
                'exceptions': exception_domain,
                'regime_e11': output_domain + [('output_type', '=', 'e11')],
                'regime_e13': output_domain + [('output_type', '=', 'e13')],
                'regime_e15': output_domain + [('output_type', '=', 'e15')],
            },
            'metrics': {
                **{key: {'value': value, 'delta': deltas[key],
                         'available': contracts[key]['status'] == 'available' and (key != 'stp_rate' or stp_available),
                         'contract': contracts[key]} for key, value in current.items()},
                **{
                    key: {'value': None, 'delta': None, 'available': False, 'contract': self._dashboard_contract(
                        reason, False, [], False, unit, False, status='unavailable')}
                    for key, reason, unit in (
                        ('overdue_import_declaration_cases', 'No reliable due-date relation links an import declaration to its case.', 'cases'),
                        ('low_confidence_extraction_queue', 'No governed dashboard threshold identifies low-confidence runs.', 'documents'),
                        ('average_processing_duration', 'Case completion timestamps are not stored.', False),
                        ('ai_credits_per_document_case', 'No exact ORM charge-to-document/case relation exists.', 'credits'),
                        ('source_erp_api_mes', 'Document source metadata does not prove ERP/API/MES channels.', False),
                    )
                },
            },
            'document_throughput': groups('logistics.idp.document', document_domain, 'create_date:day', ['page_count:sum']),
            'ocr': groups('logistics.idp.extraction.run', run_domain, 'provider', ['confidence:avg', 'duration_seconds:avg', 'page_count:sum']),
            'ocr_by_model': groups('logistics.idp.extraction.run', run_domain, 'model_version'),
            'ocr_by_status': groups('logistics.idp.extraction.run', run_domain, 'status'),
            'case_states': groups('logistics.idp.case', case_domain, 'state'),
            'verdicts': groups('logistics.idp.check.result', check_domain, 'verdict'),
            'checks': groups('logistics.idp.check.result', check_domain, 'code', limit=20),
            'customs_regimes': groups('logistics.idp.output', output_domain + [('output_type', 'in', ('e11', 'e13', 'e15'))], 'output_type'),
            'document_types': groups('logistics.idp.document', document_domain + [('document_type', '!=', False)], 'document_type'),
            'exceptions': groups('logistics.idp.exception', exception_domain, 'exception_type'),
            'exception_severity': exception_groups('severity', ('critical', 'high', 'medium', 'low')),
            'exception_states': exception_groups('state', ('open', 'waiting', 'resolved', 'cancelled')),
            'mapping_exceptions': groups('logistics.idp.check.result', check_domain + [('verdict', 'in', ('review', 'block'))], 'code', limit=20),
            'top_suppliers': top_suppliers,
            'source_mix': groups('logistics.idp.document', document_domain, 'source_channel'),
            'output_history': groups('logistics.idp.output', output_domain, 'status'),
            'live_queue': self.search_read(case_domain + [('state', 'not in', closed)], ['name', 'state', 'severity', 'next_action', 'owner_id', 'po_reference', 'shipment_reference', 'document_count', 'check_count', 'output_count', 'verdict', 'sla_deadline'], limit=50, order='sla_deadline asc, id asc'),
            'strategic_panels': {'status': 'not_configured'},
        }

    @api.model
    def _dashboard_contract(self, formula, model, domain, date_field, unit, drilldown, status='available'):
        return {'formula': formula, 'source_model': model, 'domain': domain, 'date_field': date_field,
                'unit': unit, 'drilldown_available': drilldown, 'status': status}

    @api.model
    def dashboard_drilldown(self, metric, filters=None, value=None):
        data = self.get_dashboard_data(filters)
        contract = data['metrics'].get(metric, {}).get('contract')
        if contract and contract['drilldown_available']:
            model, domain = contract['source_model'], contract['domain']
        elif metric == 'case':
            return {'type': 'ir.actions.act_window', 'name': 'Case', 'res_model': self._name,
                    'res_id': int(value), 'view_mode': 'form', 'views': [(False, 'form')]}
        else:
            sources = {
                'case_state': ('logistics.idp.case', 'case_states', 'state'),
                'document_throughput': ('logistics.idp.document', 'document_throughput', 'create_date:day'),
                'verdict': ('logistics.idp.check.result', 'verdicts', 'verdict'),
                'customs_regime': ('logistics.idp.output', 'customs_regimes', 'output_type'),
                'document_type': ('logistics.idp.document', 'document_types', 'document_type'),
                'exception': ('logistics.idp.exception', 'exceptions', 'exception_type'),
                'exception_severity': ('logistics.idp.exception', 'exception_severity', 'severity'),
                'exception_state': ('logistics.idp.exception', 'exception_states', 'state'),
                'exception_type': ('logistics.idp.exception', 'exceptions', 'exception_type'),
                'mapping_exception': ('logistics.idp.check.result', 'mapping_exceptions', 'code'),
                'supplier': ('logistics.idp.exception', 'top_suppliers', 'supplier_reference'),
                'ocr_provider': ('logistics.idp.extraction.run', 'ocr', 'provider'),
                'source': ('logistics.idp.document', 'source_mix', 'source_channel'),
            }
            native = {
                'ocr_history': ('logistics.idp.extraction.run', data['drilldown_domains']['ocr_history']),
                'ocr_model': ('logistics.idp.extraction.run', data['drilldown_domains']['ocr_history'] + [('model_version', '=', value)]),
                'ocr_status': ('logistics.idp.extraction.run', data['drilldown_domains']['ocr_history'] + [('status', '=', value)]),
                'output_history': ('logistics.idp.output', data['drilldown_domains']['output_history']),
                'mapping_exceptions': ('logistics.idp.check.result', data['drilldown_domains']['mapping_exceptions']),
                'check_code': ('logistics.idp.check.result', next((item.get('__domain', []) for item in data['checks'] if item.get('code') == value), [])),
                'regime_e11': ('logistics.idp.output', data['drilldown_domains']['regime_e11']),
                'regime_e13': ('logistics.idp.output', data['drilldown_domains']['regime_e13']),
                'regime_e15': ('logistics.idp.output', data['drilldown_domains']['regime_e15']),
            }
            if metric in native:
                model, domain = native[metric]
            elif metric in sources:
                model, section, field = sources[metric]
                if metric == 'document_throughput' and isinstance(value, list):
                    domain = value
                else:
                    row = next((item for item in data[section] if item.get(field) == value
                                or (isinstance(item.get(field), (tuple, list)) and item[field][0] == value)), None)
                    if not row and metric not in {'exception', 'exception_severity', 'exception_state', 'exception_type', 'supplier'}:
                        raise ValidationError('Unknown dashboard segment.')
                    base_domains = {
                        'logistics.idp.case': data['drilldown_domains']['case_history'],
                        'logistics.idp.document': data['drilldown_domains']['document_history'],
                        'logistics.idp.extraction.run': data['drilldown_domains']['ocr_history'],
                        'logistics.idp.check.result': data['drilldown_domains']['check_history'],
                        'logistics.idp.exception': data['drilldown_domains']['exceptions'],
                        'logistics.idp.output': data['drilldown_domains']['output_history'],
                    }
                    domain = base_domains[model] + [(('case_id.supplier_reference' if metric == 'supplier' else field), '=', value)]
            else:
                raise ValidationError('Unknown dashboard metric.')
        return {'type': 'ir.actions.act_window', 'name': metric.replace('_', ' ').title(),
                'res_model': model, 'view_mode': 'list,form', 'views': [(False, 'list'), (False, 'form')],
                'domain': domain, 'context': {'create': False}}

    def _open_cockpit_records(self, model, field_name):
        self.ensure_one()
        return {
            'type': 'ir.actions.act_window', 'name': self.env[model]._description,
            'res_model': model, 'view_mode': 'list,form',
            'domain': [(field_name, '=', self.id)], 'context': {'create': False},
        }

    def action_open_documents(self):
        return self._open_cockpit_records('logistics.idp.document', 'case_id')

    def action_open_runs(self):
        return self._open_cockpit_records('logistics.idp.extraction.run', 'case_id')

    def action_open_checks(self):
        return self._open_cockpit_records('logistics.idp.check.result', 'case_id')

    def action_open_outputs(self):
        return self._open_cockpit_records('logistics.idp.output', 'case_id')

    def action_open_exceptions(self):
        return self._open_cockpit_records('logistics.idp.exception', 'case_id')

    def action_open_shipping_plan_outputs(self):
        action = self._open_cockpit_records('logistics.idp.output', 'case_id')
        action['domain'].append(('output_type', '=', 'shipping_plan'))
        return action

    def action_open_gate_pass_outputs(self):
        action = self._open_cockpit_records('logistics.idp.output', 'case_id')
        action['domain'].append(('output_type', '=', 'gate_pass'))
        return action

    def action_open_overrides(self):
        return self._open_cockpit_records('logistics.idp.override', 'case_id')

    @api.model
    def exception_audit_domain(self, category=None):
        domains = {
            'exceptions': [('exception_ids.state', 'in', ('open', 'waiting'))],
            'checks': [('check_result_ids.verdict', 'in', ('review', 'block'))],
            'extraction': [('extraction_run_ids.status', 'in', ('review', 'error'))],
            'correlations': [('document_ids.status', 'in', ('review', 'error'))],
            'sla': [('sla_deadline', '<', fields.Datetime.now()), ('state', 'not in', tuple(_CASE_TERMINAL_STATES))],
            'overrides': [('override_ids', '!=', False)],
            'reprocessing': [('extraction_run_ids.attempt_number', '>', 1)],
        }
        if category:
            if category not in domains:
                raise ValidationError('Unknown exception audit category.')
            return domains[category]
        domain = []
        for item in domains.values():
            domain = ['|'] + domain + item if domain else list(item)
        return domain

    @api.model
    def action_open_exception_audit_center(self):
        return {
            'type': 'ir.actions.act_window', 'name': 'Exception & Audit Center',
            'res_model': 'logistics.idp.case', 'view_mode': 'list,form',
            'domain': self.exception_audit_domain(),
            'context': {'create': False, 'edit': False, 'delete': False},
        }

    _source_identity_unique = models.Constraint(
        'unique(company_id, source_system, source_key, source_version)',
        'This source version was already imported for this company.')

    def _kg_enqueue(self, event_type='UPSERT'):
        if self.env.context.get('_skip_kg_outbox') or 'kg.outbox' not in self.env.registry.models:
            return
        for case in self:
            case.check_access('read')
            self.env['kg.outbox'].enqueue(case, event_type, {
                'source_system': case.source_system, 'source_key': case.source_key,
                'source_version': case.source_version, 'provenance': case.provenance,
            })

    def _invalidate_inbox_projections(self):
        documents = self.env['logistics.idp.document'].search([('company_id', 'in', self.company_id.ids)])
        documents.invalidate_recordset(['predicted_supplier', 'related_case_candidate_id', 'inbox_action_required'])

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            vals['state'] = 'collecting'
            vals['verdict'] = 'review'
        cases = super(LogisticsCase, self.with_context(_skip_kg_outbox=True)).create(vals_list)
        cases._invalidate_inbox_projections()
        project = self.env.ref('insilos_logistics_idp.project_logistics_idp', raise_if_not_found=False)
        for case in cases.filtered(lambda item: not item.task_id):
            task = self.env['project.task'].create(case._task_values(project))
            super(LogisticsCase, case).write({'task_id': task.id})
        cases._sync_task()
        cases._kg_enqueue()
        return cases

    def write(self, vals):
        controlled = self.env.context.get('_logistics_case_lifecycle') is _INTERNAL_CASE_LIFECYCLE_TOKEN
        if (vals.get('verdict') == 'pass'
                or vals.get('state') in _CASE_TERMINAL_STATES) and not controlled:
            raise UserError('Passing or terminal case lifecycle values may only be derived by the controlled lifecycle.')
        if not controlled and _CASE_TERMINAL_FIELDS.intersection(vals) and self.filtered(
                lambda case: case.state in _CASE_TERMINAL_STATES):
            raise UserError('Terminal case lifecycle fields may only change through controlled actions.')
        if not controlled and {'state', 'verdict'}.intersection(vals) and self.filtered(
                lambda case: case.state in {'review', 'blocked'} or case.verdict == 'pass'):
            raise UserError('Case lifecycle may only change through controlled re-screening.')
        if (not controlled and _CASE_RESCREEN_FIELDS.intersection(vals) and self.filtered(
                lambda case: case.state in {'review', 'blocked'} or case.verdict == 'pass')):
            raise UserError('Material case data may only change through controlled re-screening.')
        result = super().write(vals)
        if {'company_id', 'po_reference', 'thread_reference'} & set(vals):
            self._invalidate_inbox_projections()
        if not self.env.context.get('skip_task_sync'):
            self._sync_task()
        self._kg_enqueue()
        return result

    def unlink(self):
        self.check_access('unlink')
        self._kg_enqueue('DELETE')
        return super().unlink()

    def _task_values(self, fallback_project=None):
        self.ensure_one()
        stage = {
            'blocked': 'stage_exception', 'review': 'stage_compliance_review',
            'ready': 'stage_output', 'completed': 'stage_completed',
        }.get(self.state, 'stage_intake')
        return {
            'name': self.name, 'project_id': (self.team_id or fallback_project).id,
            'company_id': self.company_id.id, 'user_ids': [(6, 0, self.owner_id.ids)],
            'date_deadline': self.sla_deadline.date() if self.sla_deadline else False,
            'description': '\n'.join(filter(None, [self.severity, self.next_action])),
            'stage_id': self.env.ref('insilos_logistics_idp.%s' % stage).id,
        }

    def _sync_task(self):
        for case in self.filtered('task_id'):
            case.task_id.with_context(skip_task_sync=True).write(case._task_values())

    def _authorized_reviewers(self, human=True):
        self.ensure_one()
        groups = (self.env.ref('insilos_logistics_idp.group_logistics_reviewer') | self.env.ref('insilos_logistics_idp.group_logistics_manager')).sudo()
        company_id = self.company_id.id
        users = groups.user_ids.sudo().filtered(lambda user: user.active and company_id in user.company_ids.ids)
        if human:
            users = users.filtered(lambda user: not user.share and 'service' not in (user.login or '').lower())
        return users.sorted('id').with_env(self.env)

    def _ensure_review_owner(self):
        for case in self:
            authorized = case._authorized_reviewers(human=False)
            if case.owner_id not in authorized:
                human = case._authorized_reviewers()
                migration = 'migration' in (case.source_system or '').lower() or 'migration' in (case.provenance or '').lower()
                service = authorized.filtered(lambda user: 'service' in (user.login or '').lower())
                owner = service[:1] if migration and service else human[:1] or authorized[:1]
                if owner:
                    case.owner_id = owner
        return True

    def _derive_lifecycle(self):
        for case in self.filtered(lambda item: item.state not in _CASE_TERMINAL_STATES):
            overridden = set(case.override_ids.filtered(lambda item: item.state == 'approved').mapped('check_result_ids').ids)
            required = case.check_result_ids.filtered(
                lambda item: item.required and item.id not in overridden and (
                    item.code != 'EXTRACTION_CRITICAL_INPUTS'
                    or (item.run_id and item.run_id == item.run_id.document_id.current_run_id)))
            verdicts = set(required.mapped('verdict'))
            statuses = [case.document_status, case.reconciliation_status, case.compliance_status, case.output_status]
            import_declaration = case.check_result_ids.filtered(
                lambda item: item.required and item.code == 'IMPORT_DECLARATION')
            if 'block' in verdicts or 'block' in statuses:
                state, verdict = 'blocked', 'block'
            elif case.waiting_party:
                state, verdict = 'waiting_external', 'review'
            elif 'review' in verdicts or 'review' in statuses:
                state, verdict = 'review', 'review'
            elif case.document_ids.filtered(lambda item: item.status == 'processing'):
                state, verdict = 'processing', 'review'
            elif (import_declaration and all(item.verdict == 'pass' or item.id in overridden for item in import_declaration)
                  and (required or overridden) and all(item.verdict in ('pass', 'not_applicable') for item in required)
                  and all(item in ('pass', 'not_applicable') for item in statuses)):
                state, verdict = 'ready', 'pass'
            else:
                state, verdict = 'collecting', 'review'
            case.with_context(_logistics_case_lifecycle=_INTERNAL_CASE_LIFECYCLE_TOKEN).write({
                'state': state, 'verdict': verdict})
        return True

    def action_set_waiting(self, party=None, reason=None, since=None):
        if not (self.env.su
                or self.env.user.has_group('insilos_logistics_idp.group_logistics_operator')
                or self.env.user.has_group('insilos_logistics_idp.group_logistics_manager')):
            raise UserError('Only Logistics operators or managers may place cases on hold.')
        for case in self:
            waiting_party = party or case.waiting_party
            waiting_reason = reason or case.waiting_reason
            waiting_since = fields.Datetime.to_datetime(since) if since else case.waiting_since or fields.Datetime.now()
            if waiting_party not in dict(self._fields['waiting_party'].selection) or not str(waiting_reason or '').strip():
                raise ValidationError('Waiting party and reason are required.')
            if case.state in ('completed', 'closed_duplicate', 'closed_other'):
                raise ValidationError('Closed cases cannot be placed on hold.')
            payload = {'action': 'set', 'party': waiting_party, 'reason': str(waiting_reason).strip(),
                       'since': fields.Datetime.to_string(waiting_since)}
            identity = hashlib.sha256(_canonical_json(payload).encode()).hexdigest()
            try:
                with self.env.cr.savepoint():
                    self.env['logistics.idp.evidence']._controlled_create({
                        'case_id': case.id, 'category': 'waiting_transition',
                        'source_reference': 'waiting:%s' % identity, 'status': 'review', 'payload': payload,
                    }, 'waiting_lifecycle')
            except UniqueViolation:
                pass
            case.sudo().with_context(_logistics_case_lifecycle=_INTERNAL_CASE_LIFECYCLE_TOKEN).write({
                'waiting_party': waiting_party, 'waiting_reason': str(waiting_reason).strip(),
                'waiting_since': waiting_since})
            case.sudo()._derive_lifecycle()
        return True

    def action_resume(self):
        for case in self.filtered('waiting_party'):
            payload = {'action': 'resume', 'party': case.waiting_party,
                       'reason': case.waiting_reason, 'since': fields.Datetime.to_string(case.waiting_since)}
            identity = hashlib.sha256(_canonical_json(payload).encode()).hexdigest()
            try:
                with self.env.cr.savepoint():
                    self.env['logistics.idp.evidence']._controlled_create({
                        'case_id': case.id, 'category': 'waiting_transition',
                        'source_reference': 'resume:%s' % identity, 'status': 'valid', 'payload': payload,
                    }, 'waiting_lifecycle')
            except UniqueViolation:
                pass
            case.with_context(_logistics_case_lifecycle=_INTERNAL_CASE_LIFECYCLE_TOKEN).write({
                'waiting_party': False, 'waiting_reason': False, 'waiting_since': False})
            case._derive_lifecycle()
        return True

    def action_open_upload_wizard(self):
        self.ensure_one()
        return {
            'type': 'ir.actions.act_window', 'name': 'Upload Document',
            'res_model': 'logistics.idp.document.upload.wizard', 'view_mode': 'form', 'target': 'new',
            'context': {'default_case_id': self.id, 'default_company_id': self.company_id.id},
        }

    def action_open_final_customs_upload_wizard(self):
        self.ensure_one()
        return {
            'type': 'ir.actions.act_window', 'name': 'Upload Final Customs Document',
            'res_model': 'logistics.idp.document.upload.wizard', 'view_mode': 'form', 'target': 'new',
            'context': {'default_case_id': self.id, 'default_company_id': self.company_id.id,
                        'default_final_customs_document': True},
        }

    def _schedule_different_subject_review(self):
        self.ensure_one()
        summary = 'FR-107: Select active supplier email thread'
        activity_type = self.env.ref('mail.mail_activity_data_todo')
        activity = self.env['mail.activity'].search([
            ('res_model', '=', self._name), ('res_id', '=', self.id),
            ('activity_type_id', '=', activity_type.id), ('summary', '=', summary),
        ], limit=1)
        if not activity:
            self._ensure_review_owner()
            self.activity_schedule('mail.mail_activity_data_todo', user_id=self.owner_id.id,
                                   summary=summary, note='Different subjects require explicit active-thread selection.')
        return True

    def action_open_active_thread_selection(self):
        self.ensure_one()
        return {
            'type': 'ir.actions.act_window', 'name': 'Select Active Email Thread',
            'res_model': 'logistics.idp.active.thread.wizard', 'view_mode': 'form', 'target': 'new',
            'context': {'default_case_id': self.id},
        }

    def action_select_active_thread(self, active_thread_case):
        if not (self.env.user.has_group('insilos_logistics_idp.group_logistics_reviewer')
                or self.env.user.has_group('insilos_logistics_idp.group_logistics_manager')):
            raise UserError('Only Logistics reviewers or managers may select an active email thread.')
        for case in self:
            candidate = self.browse(active_thread_case.id if hasattr(active_thread_case, 'id') else active_thread_case).exists()
            if len(candidate) != 1 or candidate.company_id != case.company_id:
                raise ValidationError('Selected email thread must belong to the same company.')
            shares_signal = ((case.supplier_reference and candidate.supplier_reference == case.supplier_reference)
                             or (case.shipment_reference and candidate.shipment_reference == case.shipment_reference))
            if candidate == case or not candidate.thread_reference or not shares_signal:
                raise ValidationError('Selected email thread must be an unambiguous related thread.')
            if candidate.state in _CASE_TERMINAL_STATES:
                raise ValidationError('Selected email thread must be active.')
            activities = case.activity_ids.filtered(
                lambda activity: activity.summary == 'FR-107: Select active supplier email thread')
            if not activities:
                raise ValidationError('No open FR-107 review activity exists.')
            case.write({'active_thread_case_id': candidate.id})
            activities.action_feedback(feedback='FR-107 active thread selected: %s.' % candidate.display_name)
        return True

    def action_complete(self, reason):
        if not self.env.user.has_group('insilos_logistics_idp.group_logistics_manager'):
            raise UserError('Only Logistics Managers may complete cases.')
        if not reason:
            raise ValidationError('Completion reason is required.')
        self._derive_lifecycle()
        if any(case.state != 'ready' for case in self):
            raise ValidationError('All required checks must pass before completion.')
        self.write({'state': 'completed', 'completion_reason': reason})
        return True

    @api.model
    def _record_same_subject_thread_entry(self, case, values):
        supplier, subject = values.get('supplier_reference'), values.get('email_subject')
        if not supplier or not subject or case.state in _CASE_TERMINAL_STATES:
            return
        identity = _canonical_json({key: values.get(key) for key in ('source_system', 'source_key', 'source_version')})
        entries = self.env['logistics.idp.email.thread.entry'].sudo().with_context(
            _logistics_email_thread_entry_token=_INTERNAL_EMAIL_THREAD_ENTRY_TOKEN)
        if entries.search([('case_id', '=', case.id), ('source_identity', '=', identity)], limit=1):
            return
        self.env.cr.execute('SELECT id FROM logistics_idp_case WHERE id = %s FOR UPDATE', [case.id])
        previous = entries.search([('case_id', '=', case.id), ('supplier_reference', '=', supplier),
                                   ('normalized_subject', '=', _normalized(subject)), ('is_active', '=', True)])
        previous.write({'is_active': False})
        previous.flush_recordset(['is_active'])
        entries.create({'case_id': case.id, 'company_id': case.company_id.id, 'supplier_reference': supplier,
                        'normalized_subject': _normalized(subject), 'source_identity': identity, 'is_active': True})

    @api.model
    def intake(self, values):
        required = {'name', 'source_system', 'source_key', 'source_version', 'provenance', 'effective_date'}
        missing = sorted(required - set(values))
        if missing:
            raise ValidationError('Missing intake fields: %s' % ', '.join(missing))
        company_id = values.get('company_id') or self.env.company.id
        existing = self.search([
            ('company_id', '=', company_id),
            ('source_system', '=', values['source_system']),
            ('source_key', '=', values['source_key']),
            ('source_version', '=', values['source_version']),
        ], limit=1)
        if existing:
            return existing
        case = self._correlate(values)
        if case:
            case.message_post(body='Correlated intake source %s:%s.' % (values['source_system'], values['source_key']))
            if (values.get('supplier_reference') and values.get('email_subject')
                    and _normalized(case.email_subject) == _normalized(values['email_subject']) and values.get('thread_reference')):
                case.write({'thread_reference': values['thread_reference']})
            self._record_same_subject_thread_entry(case, values)
            return case
        case = self.create(values)
        self._record_same_subject_thread_entry(case, values)
        supplier = values.get('supplier_reference')
        subject = values.get('email_subject')
        if supplier and subject and not values.get('po_reference'):
            ambiguous = self.search([
                ('id', '!=', case.id), ('company_id', '=', case.company_id.id),
                ('supplier_reference', '=', supplier), ('email_subject', '!=', subject),
                ('state', 'not in', _CASE_TERMINAL_STATES),
            ], limit=1)
            if ambiguous:
                case.write({'state': 'review', 'next_action': 'Select the active supplier email thread.'})
                case._schedule_different_subject_review()
        return case

    @api.model
    def _correlate(self, values):
        for field_name in ('po_reference', 'thread_reference'):
            value = values.get(field_name)
            if value:
                case = self.search([('company_id', '=', values.get('company_id') or self.env.company.id),
                                    (field_name, '=', value), ('state', 'not in', tuple(_CASE_TERMINAL_STATES))], limit=1)
                if case:
                    return case
        supplier = values.get('supplier_reference')
        subject = values.get('email_subject')
        if supplier and subject:
            cases = self.search([
                ('company_id', '=', values.get('company_id') or self.env.company.id),
                ('supplier_reference', '=', supplier),
                ('state', 'not in', ('completed', 'closed_duplicate')),
            ], order='create_date desc, id desc').filtered(
                lambda item: _normalized(item.email_subject) == _normalized(subject))
            if len(cases) == 1:
                return cases
        return self.browse()

    def add_extraction(self, source_reference, payload, confidence=0, threshold=0.8):
        self.ensure_one()
        status = 'valid' if confidence >= threshold else 'review'
        return self.env['logistics.idp.evidence'].with_context(_logistics_snapshot_token=_INTERNAL_SNAPSHOT_TOKEN).create({
            'case_id': self.id,
            'category': 'extraction',
            'source_reference': source_reference,
            'status': status,
            'payload': {'confidence': confidence, 'threshold': threshold, 'data': payload},
        })

    def _final_export_hs_checks(self):
        self.ensure_one()
        documents = self.document_ids.filtered(lambda item: item.document_type in ('export_declaration_final', 'dsnavl', 'master_data') and item.current_run_id)
        actual, expected = {}, {}
        for document in documents.sorted('id'):
            payload = json.loads(document.current_run_id.payload).get('payload', {})
            for line in payload.get('lines', []) if isinstance(payload, dict) else []:
                identity = _normalized(line.get('custom_code') or line.get('material_code') or line.get('item'))
                hs_code = _normalized(line.get('hs_code'))
                if identity and hs_code:
                    (actual if document.document_type == 'export_declaration_final' else expected)[identity] = {
                        'hs_code': hs_code, 'document_id': document.id, 'run_hash': document.current_run_id.payload_hash,
                        'document_type': document.document_type}
        decisions = []
        for identity in sorted(set(actual) | set(expected)):
            observed, master = actual.get(identity), expected.get(identity)
            check = customs_checks({'regime': 'E11', 'actual_hs': observed and observed['hs_code'], 'expected_hs': master and master['hs_code']})[0]
            decisions.append({'line_identity': identity, 'material_code': identity, 'expected_hs': master and master['hs_code'],
                              'actual_hs': observed and observed['hs_code'], 'verdict': check['verdict'],
                              'reason': check['reason'], 'provenance': {'actual': observed, 'expected': master}})
        verdict = 'block' if any(item['verdict'] == 'block' for item in decisions) else ('review' if any(item['verdict'] == 'review' for item in decisions) else 'pass')
        return {'verdict': verdict, 'decisions': decisions}

    def _restricted_party_sources(self):
        self.ensure_one()
        return self.env['logistics.idp.policy.source'].search([
            ('company_id', '=', self.company_id.id), ('state', '=', 'active'),
            ('effective_from', '<=', self.effective_date),
            '|', ('effective_to', '=', False), ('effective_to', '>=', self.effective_date)],
            order='recorded_at desc, id desc')

    def _screening_source_is_activated(self, source):
        if tools.config['test_enable'] and source.audit_service == 'test_fixture':
            return True
        return bool(self.env['logistics.idp.policy.activation'].sudo().search([
            ('company_id', '=', source.company_id.id), ('policy_source_id', '=', source.id),
            ('policy_hash', '=', source.payload_hash), ('status', '=', 'activated'),
        ], limit=1))

    def _has_governed_screening_result(self, require_release=False):
        self.ensure_one()
        checks = self.check_result_ids.filtered(
            lambda item: item.code == 'RESTRICTED_PARTY_SCREENING' and item.audit_input_hash
            and item.policy_source_id and self._screening_source_is_activated(item.policy_source_id))
        if not require_release:
            return bool(checks)
        approved = self.override_ids.filtered(lambda item: item.state == 'approved').mapped('check_result_ids')
        return bool(checks.filtered(lambda item: item.verdict == 'pass' or item in approved))

    def _restricted_party_screening(self, po, invoice, binding_hash):
        self.ensure_one()
        party_evidence = canonical_party_evidence(po, invoice)
        parties = {item['name'] for item in party_evidence if item['name']}
        parties.update(identifier['value'] for item in party_evidence for identifier in item['identifiers'])
        parties.update(asset['value'] for item in party_evidence for asset in item['assets'])
        parties.update(normalize_party(value) for value in (
            self.supplier_reference, invoice.get('supplier_address')) if normalize_party(value))
        sources = self._restricted_party_sources()
        hits, consulted, incomplete = [], 0, False
        consulted_sources, reasons = [], []
        for source in sources:
            source_reasons = []
            if not getattr(source, 'legal_authority', False):
                source_reasons.append('not_legally_authoritative')
            if getattr(source, 'verification_status', 'unverified') != 'verified':
                source_reasons.append('not_governance_verified')
            if not self.env['logistics.idp.policy.source']._citation_is_governed(
                    getattr(source, 'source_tier', False), getattr(source, 'citation', False)):
                source_reasons.append('citation_not_governed')
            if getattr(source, 'freshness_status', 'review') != 'fresh':
                source_reasons.append('not_fresh')
            if not self._screening_source_is_activated(source):
                source_reasons.append('missing_activated_provenance')
            if source_reasons:
                incomplete = True
                consulted_sources.append({'id': source.id, 'payload_hash': False,
                                          'reasons': source_reasons})
                reasons.extend('%s:%s' % (reason, source.id) for reason in source_reasons)
                continue
            try:
                payload = json.loads(source.payload)
            except (TypeError, ValueError):
                incomplete = True
                consulted_sources.append({'id': source.id, 'payload_hash': source.payload_hash,
                                          'reasons': ['invalid_policy_payload']})
                reasons.append('invalid_policy_payload:%s' % source.id)
                continue
            overlays = payload.get('overlays', {}) if isinstance(payload, dict) else None
            if not isinstance(overlays, dict):
                incomplete = True
                consulted_sources.append({'id': source.id, 'payload_hash': source.payload_hash,
                                          'reasons': ['invalid_policy_overlays']})
                reasons.append('invalid_policy_overlays:%s' % source.id)
                continue
            if 'restricted_parties' not in overlays:
                continue
            result = screen_parties(parties, overlays['restricted_parties'])
            consulted += result['lists_consulted']
            source_reasons = []
            if result.get('reason') == 'restricted_party_screening_incomplete':
                source_reasons.append('incomplete_list')
            consulted_sources.append({'id': source.id, 'payload_hash': source.payload_hash,
                                      'reasons': source_reasons})
            if source_reasons:
                incomplete = True
                reasons.extend('%s:%s' % (reason, source.id) for reason in source_reasons)
            for hit in result['hits']:
                hits.append({**hit, 'policy_source_id': source.id, 'policy_code': source.code,
                             'policy_version': source.version})
        if not consulted:
            incomplete = True
            reasons.append('no_restricted_party_lists_consulted')
        hits.sort(key=lambda hit: (hit['policy_source_id'], hit['list_name'] or '', hit['party'], hit['matched']))
        consulted_sources.sort(key=lambda source: (source['id'], source['payload_hash'] or ''))
        verdict = 'review' if hits or incomplete else 'pass'
        screening_hash = hashlib.sha256(_canonical_json({
            'binding_hash': binding_hash, 'parties': sorted(parties), 'hits': hits,
            'policy_sources': consulted_sources, 'reason_codes': sorted(reasons),
        }).encode()).hexdigest()
        check_model = self.env['logistics.idp.check.result']
        existing = check_model.search([('case_id', '=', self.id), ('audit_input_hash', '=', screening_hash)], limit=1)
        if existing:
            return {'verdict': verdict, 'hits': hits, 'performed': not incomplete,
                    'lists_consulted': consulted, 'reason': 'restricted_party_screening_incomplete' if incomplete else False,
                    'check_id': existing.id}
        policy_source = next(iter(sources), False)
        first_hit = hits[0] if hits else {}
        check = check_model._create_from_check_runner({
            'case_id': self.id,
            'policy_source_id': first_hit.get('policy_source_id') or (policy_source.id if policy_source else False),
            'code': 'RESTRICTED_PARTY_SCREENING', 'required': True,
            'expected': 'no restricted-party match from authoritative, verified, governed, fresh lists',
            'actual': ', '.join(sorted({hit['matched'] for hit in hits})) if hits else 'no match',
            'verdict': verdict,
            'rationale': ('restricted-party screening requires review: %s' % '; '.join(sorted(reasons))
                          if incomplete else 'restricted-party screening consulted %s authoritative list(s); %s' % (
                              consulted, 'matches: %s' % '; '.join('%s (%s/%s)' % (hit['list_name'], hit['matched'], hit['party']) for hit in hits)
                              if hits else 'no match above threshold')),
            'citation': policy_source.citation if policy_source else False,
            'payload': {'binding_hash': binding_hash, 'parties': sorted(parties), 'hits': hits,
                        'lists_consulted': consulted, 'policy_sources': consulted_sources,
                        'reason_codes': sorted(reasons)},
            'audit_input_hash': screening_hash,
        }, 'restricted_party_screening')
        if hits:
            self.env['logistics.idp.exception'].open_or_reuse(self, 'restricted_party_hit', severity='high')
        return {'verdict': verdict, 'hits': hits, 'performed': not incomplete,
                'lists_consulted': consulted, 'reason': 'restricted_party_screening_incomplete' if incomplete else False,
                'check_id': check.id}

    def reconcile_documents(self, po, invoice, tolerance='0', reconciliation_inputs=None):
        self.ensure_one()
        profile_model = self.env['logistics.idp.supplier.profile']
        profiles = profile_model.search([
            ('company_id', '=', self.company_id.id),
            ('supplier_reference', '=', self.supplier_reference or invoice.get('supplier')),
            ('effective_from', '<=', self.effective_date), '|', ('effective_to', '=', False),
            ('effective_to', '>=', self.effective_date),
        ], order='effective_from desc, id desc', limit=2)
        profile = profiles if len(profiles) == 1 else profile_model.browse()
        try:
            policy = json.loads(profile.payload) if profile else {}
            if not profile or not isinstance(policy.get('draft_invoice'), dict):
                raise ValueError('missing draft-invoice policy')
            validate_draft_invoice_policy(policy)
            validate_main_invoice_policy(policy)
            policy_valid = True
        except (TypeError, ValueError):
            policy, policy_valid = {'draft_invoice': {'mode': 'skipped'}}, False
        draft_documents = self.document_ids.filtered(lambda item: item.document_type == 'draft_vat_invoice').sorted('id')
        draft_candidates = []
        for document in draft_documents:
            payload = json.loads(document.current_run_id.payload).get('payload', {}) if document.current_run_id else {}
            draft_candidates.append({
                'id': document.id, 'company_id': document.company_id.id, 'status': document.status,
                'content_hash': document.content_hash, 'current_run_id': document.current_run_id.id,
                'current_run_hash': document.current_run_id.payload_hash if document.current_run_id else False,
                'payload': payload,
            })
        matching_drafts = [item for item in draft_candidates
                           if _normalized(item['payload'].get('supplier')) == _normalized(invoice.get('supplier'))]
        approved_drafts = [item for item in matching_drafts if item['company_id'] == self.company_id.id
                           and item['status'] == 'valid' and item['current_run_id']]
        draft_selection = {'candidates': draft_candidates, 'selected': approved_drafts[0] if len(approved_drafts) == 1
                           and len(matching_drafts) == 1 else False,
                           'reason': 'selected' if len(approved_drafts) == len(matching_drafts) == 1 else
                           ('missing approved same-case draft VAT invoice' if not matching_drafts else
                            'ambiguous, unapproved, or cross-company draft VAT invoice')}
        references = reconciliation_inputs if isinstance(reconciliation_inputs, dict) else {}
        declaration_groups = references.get('declaration_groups')
        declaration_groups = (sorted(set(str(group).strip() for group in declaration_groups))
                              if isinstance(declaration_groups, list) else declaration_groups)
        selection = {
            'case_id': self.id,
            'case_effective_date': fields.Date.to_string(self.effective_date),
            'supplier_reference': self.supplier_reference,
            'restricted_party_sources': [
                {'id': source.id, 'payload_hash': source.payload_hash}
                for source in self._restricted_party_sources()],
            'supplier_profile': profile and {
                'id': profile.id, 'profile_code': profile.profile_code, 'version': profile.version,
                'payload_hash': profile.payload_hash,
                'effective_from': fields.Date.to_string(profile.effective_from),
                'effective_to': fields.Date.to_string(profile.effective_to) if profile.effective_to else False,
            } or False,
            'documents': [{'id': document.id, 'content_hash': document.content_hash,
                           'current_run_id': document.current_run_id.id,
                           'current_run_hash': document.current_run_id.payload_hash}
                          for document in self.document_ids.sorted('id')],
            'po': po, 'invoice': invoice, 'tolerance': tolerance,
            'draft_invoice': draft_selection,
            'references': {**references, **({'declaration_groups': declaration_groups}
                                              if 'declaration_groups' in references else {})},
        }
        binding_hash = hashlib.sha256(_canonical_json(selection).encode()).hexdigest()
        existing = self.env['logistics.idp.evidence'].search([
            ('case_id', '=', self.id), ('category', '=', 'reconciliation'),
            ('source_reference', '=', 'reconciliation:%s' % binding_hash),
        ], limit=1)
        if existing:
            return existing
        main_policy = policy
        if invoice.get('document_type') == 'main_vat_invoice' and policy.get('draft_invoice', {}).get('mode') != 'skipped':
            main_policy = {**policy, 'draft_invoice': {**policy['draft_invoice'], 'mode': 'skipped'}}
        try:
            reconciliation = reconcile_service(po, invoice, tolerance, po.get('fx_rates'), main_policy,
                                               declaration_groups=declaration_groups, effective_date=self.effective_date,
                                               draft_invoice_policy_valid=policy_valid)
        except ValueError as exc:
            reconciliation = reconcile_service(po, invoice, tolerance, po.get('fx_rates'),
                                               {key: value for key, value in main_policy.items()
                                                if key != 'declaration_group_rules'},
                                               draft_invoice_policy_valid=policy_valid)
            reconciliation['declaration_group_check'] = {'result': 'review',
                                                          'reason': 'ambiguous or invalid declaration group rules: %s' % exc}
            if reconciliation['verdict'] == 'pass':
                reconciliation['verdict'] = 'review'
        draft_reconciliation = (reconcile_draft_main(draft_selection['selected']['payload'], invoice, tolerance,
                                po.get('fx_rates'), policy) if draft_selection['selected']
                                and policy.get('draft_invoice', {}).get('mode') != 'skipped' else False)
        hs_reconciliation = self._final_export_hs_checks()
        results = reconciliation['results']
        sequence = reconciliation['sequence']
        verdict = reconciliation['verdict'] if policy_valid else 'review'
        if draft_reconciliation:
            if draft_reconciliation['verdict'] == 'block' or draft_reconciliation['verdict'] == 'review' and verdict == 'pass':
                verdict = draft_reconciliation['verdict']
        elif invoice.get('document_type') == 'main_vat_invoice' and policy.get('draft_invoice', {}).get('mode') != 'skipped':
            verdict = 'review'
        if hs_reconciliation['verdict'] == 'block' or hs_reconciliation['verdict'] == 'review' and verdict == 'pass':
            verdict = hs_reconciliation['verdict']
        restricted_party = self._restricted_party_screening(po, invoice, binding_hash)
        if restricted_party['verdict'] == 'review' and verdict == 'pass':
            verdict = 'review'
        reason = ('Effective supplier profile draft-invoice policy selected for reconciliation.' if policy_valid
                  else 'Supplier profile selection or draft-invoice policy is missing, invalid, or ambiguous; review required.')
        decision = self.env['logistics.idp.policy.decision']._controlled_create({
            'case_id': self.id, 'supplier_profile_id': profile.id if profile else False,
            'policy_code': profile.profile_code if profile else 'supplier_profile',
            'policy_version': profile.version if profile else 'unresolved',
            'effective_from': fields.Datetime.now(), 'verdict': verdict, 'reason': reason,
            'payload': {'input_manifest': selection, 'binding_hash': binding_hash,
                        'reconciliation': {'verdict': verdict, 'results': results, 'sequence': sequence,
                                           'declaration_group_check': reconciliation['declaration_group_check'],
                                           'draft_main': draft_reconciliation, 'final_export_hs': hs_reconciliation},
                        'restricted_party': {key: value for key, value in restricted_party.items()
                                             if key != 'check_id'}},
            'audit_input_hash': binding_hash,
        }, 'supplier_profile_selection')
        for item in hs_reconciliation['decisions']:
            hs_binding = hashlib.sha256(_canonical_json({'binding_hash': binding_hash, **item}).encode()).hexdigest()
            self.env['logistics.idp.check.result']._create_from_check_runner({
                'case_id': self.id, 'code': 'FINAL_EXPORT_HS', 'required': True, 'verdict': item['verdict'],
                'expected': item['expected_hs'], 'actual': item['actual_hs'], 'rationale': item['reason'],
                'payload': {**item, 'binding_hash': binding_hash}, 'audit_input_hash': hs_binding,
            }, 'final_export_hs_reconciliation')
            if item['verdict'] == 'block':
                self.env['logistics.idp.exception'].open_or_reuse(self, 'hs_code_mismatch', severity='high')
        evidence = self.env['logistics.idp.evidence'].with_context(_logistics_snapshot_token=_INTERNAL_SNAPSHOT_TOKEN).create({
            'case_id': self.id,
            'category': 'reconciliation',
            'source_reference': 'reconciliation:%s' % binding_hash,
            'status': 'valid' if verdict == 'pass' else ('invalid' if verdict == 'block' else 'review'),
            'payload': {'decision_hash': decision.payload_hash, 'po': po, 'invoice': invoice,
                        'tolerance': tolerance, 'verdict': verdict, 'results': results, 'sequence': sequence,
                        'draft_invoice': draft_selection, 'draft_main': draft_reconciliation,
                        'declaration_group_check': reconciliation['declaration_group_check'],
                        'final_export_hs': hs_reconciliation,
                        'restricted_party': {key: value for key, value in restricted_party.items()
                                             if key != 'check_id'}},
        })
        self.write({'verdict': verdict, 'state': 'blocked' if verdict == 'block' else 'review'})
        return evidence

    def recheck(self, po, invoice, tolerance='0'):
        return self.reconcile_documents(po, invoice, tolerance)

    def mark_semantic_duplicate(self, canonical_case=None):
        self.ensure_one()
        if not (self.env.su or self.env.user.has_group('insilos_logistics_idp.group_logistics_manager')):
            raise UserError('Only Logistics Managers may close semantic duplicates.')
        canonical_case = canonical_case or self.canonical_case_id
        canonical_case.ensure_one()
        if (canonical_case == self or canonical_case.company_id != self.company_id
                or self.state in _CASE_TERMINAL_STATES or canonical_case.state in _CASE_TERMINAL_STATES
                or self.canonical_case_id or canonical_case.canonical_case_id):
            raise ValidationError('Duplicate cases require two distinct same-company nonterminal root cases.')
        payload = {'action': 'semantic_duplicate', 'duplicate_case_id': self.id,
                   'canonical_case_id': canonical_case.id, 'actor_id': self.env.user.id}
        self.env['logistics.idp.evidence']._controlled_create([{
            'case_id': case.id, 'category': 'semantic_duplicate',
            'source_reference': 'duplicate:%s:canonical:%s' % (self.id, canonical_case.id),
            'status': 'valid', 'payload': payload,
        } for case in self | canonical_case], 'semantic_duplicate')
        self.with_context(_logistics_case_lifecycle=_INTERNAL_CASE_LIFECYCLE_TOKEN).write({
            'canonical_case_id': canonical_case.id, 'state': 'closed_duplicate',
            'verdict': 'not_applicable'})
        return True

    def request_override(self, exceptions, reason_code, justification, evidence=None, checks=None):
        self.ensure_one()
        if not (self.env.su or self.env.user in self._authorized_reviewers()):
            raise UserError('Only authorized reviewers or managers may request overrides.')
        exceptions = self.env['logistics.idp.exception'].browse(exceptions) if isinstance(exceptions, (list, tuple)) else exceptions
        checks = self.env['logistics.idp.check.result'].browse(checks or []) if isinstance(checks, (list, tuple)) or not checks else checks
        evidence = self.env['logistics.idp.evidence'].browse(evidence) if isinstance(evidence, int) else evidence or self.env['logistics.idp.evidence']
        if not exceptions or any(item.case_id != self for item in exceptions):
            raise ValidationError('Override requires one or more same-case exceptions.')
        if reason_code not in dict(self.env['logistics.idp.override']._fields['reason_code'].selection) or not str(justification or '').strip():
            raise ValidationError('Override requires a controlled reason code and nonblank justification.')
        if evidence and evidence.case_id != self:
            raise ValidationError('Override evidence must belong to the same case.')
        if any(item.case_id != self for item in checks):
            raise ValidationError('Override checks must belong to the same case.')
        values = {'case_id': self.id, 'exception_ids': [(6, 0, exceptions.ids)], 'check_result_ids': [(6, 0, checks.ids)],
                  'reason_code': reason_code, 'justification': str(justification).strip(), 'evidence_id': evidence.id,
                  'requested_by': self.env.user.id, 'requested_at': fields.Datetime.now(), 'state': 'requested',
                  'approved_by': False, 'approved_at': False,
                  'payload': {'exception_ids': exceptions.ids, 'check_ids': checks.ids, 'reason_code': reason_code,
                              'justification': str(justification).strip(), 'evidence_id': evidence.id or False}}
        return self.env['logistics.idp.override'].with_context(
            _logistics_override_request_token=_INTERNAL_OVERRIDE_REQUEST_TOKEN,
        )._controlled_create(values, 'controlled_override')

    def _authoritative_reference(self, field, stage=None):
        self.ensure_one()
        profiles = self.env['logistics.idp.supplier.profile'].search([
            ('company_id', '=', self.company_id.id), ('supplier_reference', '=', self.supplier_reference),
            ('effective_from', '<=', self.effective_date), '|', ('effective_to', '=', False),
            ('effective_to', '>=', self.effective_date),
        ], order='effective_from desc, id desc', limit=2)
        if len(profiles) != 1:
            raise ValidationError('Authoritative reference requires one effective supplier profile.')
        profile = profiles
        try:
            policy = json.loads(profile.payload).get('authoritative_reference_policy')
        except (TypeError, ValueError) as exc:
            raise ValidationError('Authoritative reference policy is invalid.') from exc
        stage = str(self.task_id.stage_id.id if stage is None else stage)
        if not isinstance(policy, dict) or not isinstance(policy.get('policy_version'), str) or not policy['policy_version'] or not isinstance(policy.get('rules'), list):
            raise ValidationError('Authoritative reference policy is missing or invalid.')
        rules = [rule for rule in policy['rules'] if isinstance(rule, dict) and str(rule.get('stage')) == stage]
        exact = [rule for rule in rules if rule.get('field') == field]
        selected = exact or [rule for rule in rules if rule.get('field') == '*']
        valid_output_types = {'e11', 'e13', 'e15', 'sap_erp_output'}
        if (len(selected) != 1 or not isinstance(selected[0].get('id'), str) or not selected[0]['id']
                or selected[0].get('source_type') != 'output'
                or not isinstance(selected[0].get('output_types'), list) or not selected[0]['output_types']
                or len(selected[0]['output_types']) != len(set(selected[0]['output_types']))
                or not set(selected[0]['output_types']) <= valid_output_types
                or not all(isinstance(output_type, str) for output_type in selected[0]['output_types'])):
            raise ValidationError('Authoritative reference policy has no valid unambiguous current-stage field rule.')
        rule = selected[0]
        return profile, rule, {'profile_id': profile.id, 'profile_hash': profile.payload_hash, 'stage': stage,
                               'field': field, 'rule_id': rule['id'], 'policy_version': policy['policy_version'],
                               'policy_hash': hashlib.sha256(_canonical_json(policy).encode()).hexdigest(),
                               'source_type': rule['source_type'], 'output_types': rule['output_types']}

    def _sap_erp_profile(self):
        self.ensure_one()
        profile, rule, binding = self._authoritative_reference('sap_erp_output')
        if rule['output_types'] != ['sap_erp_output']:
            raise ValidationError('SAP/ERP output requires an explicit SAP authoritative-reference rule.')
        downstream_field = rule.get('downstream_field')
        if downstream_field is not None:
            if not isinstance(downstream_field, str) or not downstream_field:
                raise ValidationError('SAP/ERP downstream authoritative reference is invalid.')
            _downstream_profile, _downstream_rule, binding['downstream_authoritative_binding'] = self._authoritative_reference(downstream_field)
        config = json.loads(profile.payload).get('sap_erp_output', {})
        stage = str(self.task_id.stage_id.id)
        config = config.get('stages', {}).get(stage)
        if not config:
            raise ValidationError('SAP/ERP output requires configuration for the current case stage.')
        if (not config.get('mapping_version') or not config.get('columns') or not config.get('template_attachment_id')
                or not config.get('sheet') or not config.get('line_identity_keys') or not config.get('amount_source')
                or not config.get('price_source') or not config.get('source_precedence') or not config.get('sort_keys')):
            raise ValidationError('SAP/ERP output stage configuration is incomplete.')
        positions = [column.get('position') for column in config['columns']]
        if any(not isinstance(position, int) or isinstance(position, bool) or position < 1 for position in positions) or len(set(positions)) != len(positions):
            raise ValidationError('SAP/ERP columns require unique positive template positions.')
        sources = config['source_precedence']
        allowed_sources = {'po_snapshot', 'invoice_evidence', 'dsnavl', 'master_data'}
        precedence = sources.values() if isinstance(sources, dict) else (sources,)
        if (not all(isinstance(item, list) and item and len(item) == len(set(item)) and set(item) <= allowed_sources
                    for item in precedence)):
            raise ValidationError('SAP/ERP source precedence must declare eligible trusted sources exactly once.')
        normalization = config.get('normalization', {})
        if normalization and (set(normalization) - {'version', 'rounding'} or not normalization.get('version')
                              or not isinstance(normalization.get('rounding', {}), (dict, list))):
            raise ValidationError('SAP/ERP normalization configuration is invalid.')
        return profile, config, binding

    @staticmethod
    def _sap_erp_value(source, context):
        value = context
        for part in source.split('.'):
            value = value.get(part) if isinstance(value, dict) else None
        return value

    @staticmethod
    def _sap_erp_identity(line, keys):
        identity = tuple(_normalized(line.get(key)) for key in keys)
        return identity if all(identity) else None

    @staticmethod
    def _sap_erp_amount(line, price_source, rounding=None):
        if line.get('amount') not in (None, ''):
            amount = _decimal(line['amount'])
        else:
            price_key = price_source.removeprefix('line.')
            if price_source not in ('line.unit_price', 'line.value') or line.get(price_key) in (None, ''):
                raise ValidationError('SAP/ERP configured price source is missing.')
            price_per = _decimal(line.get('price_per', 1))
            if price_per <= 0:
                raise ValidationError('SAP/ERP price_per must be positive.')
            amount = _decimal(line.get('quantity')) * _decimal(line[price_key]) / price_per
        if rounding:
            try:
                amount = amount.quantize(Decimal(str(rounding['precision'])), rounding={
                    'half_up': ROUND_HALF_UP, 'up': ROUND_UP, 'down': ROUND_DOWN}[rounding.get('mode', 'half_up')])
            except (KeyError, InvalidOperation) as exc:
                raise ValidationError('SAP/ERP rounding configuration is invalid.') from exc
        return amount

    def _sap_erp_rounding(self, config, profile, invoice_line, field='amount'):
        normalization = config.get('normalization', {})
        rules = normalization.get('rounding', {})
        if isinstance(rules, list):
            currency = str(invoice_line.get('currency') or '').upper()
            template = str(config.get('template_code') or config.get('mapping_version'))
            identities = {
                'supplier_profile_id': str(profile.id), 'supplier_profile_code': str(profile.profile_code),
                'supplier_id': str(self.supplier_reference), 'supplier_code': str(self.supplier_reference),
                'currency': currency, 'field': field, 'template_version': str(config.get('mapping_version')),
                'template_code': template,
            }
            matched = []
            for index, rule in enumerate(rules):
                selectors = {key: str(value) for key, value in rule.items() if key in identities and value not in (None, '')}
                if selectors and all(identities[key] == value for key, value in selectors.items()):
                    matched.append((len(selectors), index, rule))
            if not matched:
                raise ValidationError('SAP/ERP rounding configuration has no matching rule.')
            _specificity, index, rule = max(matched, key=lambda item: (item[0], -item[1]))
            rounding = {key: rule[key] for key in ('precision', 'mode') if key in rule}
            return rounding, {'rule_id': rule.get('id', 'rule-%s' % index), 'rule_version': rule.get('version') or normalization.get('version'),
                              'selectors': {key: rule[key] for key in identities if key in rule and rule[key] not in (None, '')}}
        return rules.get(field) or rules.get('default'), {'rule_id': field if field in rules else 'default', 'rule_version': normalization.get('version'), 'selectors': {'field': field}}

    @staticmethod
    def _sap_erp_source_order(field, configured):
        defaults = {
            'delivery_quantity': ['invoice_evidence'], 'quantity': ['invoice_evidence'], 'uom': ['invoice_evidence', 'master_data'],
            'amount': ['invoice_evidence'], 'invoice_number': ['invoice_evidence'], 'invoice_date': ['invoice_evidence'],
            'po_number': ['po_snapshot'], 'po_item': ['po_snapshot'], 'payment_terms': ['po_snapshot'],
            'material_code': ['po_snapshot', 'invoice_evidence'], 'custom_code': ['po_snapshot', 'invoice_evidence'],
            'vietnamese_name': ['dsnavl'], 'regime': ['dsnavl'],
            'mrp_controller': ['master_data'], 'material_group': ['master_data'],
        }
        allowed = defaults.get(field, ['po_snapshot', 'invoice_evidence', 'dsnavl', 'master_data'])
        requested = configured.get(field, configured.get('*', allowed)) if isinstance(configured, dict) else configured
        requested = [source for source in requested if source in allowed]
        return requested + [source for source in allowed if source not in requested]

    def _sap_erp_sources(self):
        po_snapshot = self.env['logistics.idp.po.snapshot'].search([
            ('company_id', '=', self.company_id.id), ('po_reference', '=', self.po_reference),
            ('snapshot_date', '<=', self.effective_date),
        ], order='snapshot_date desc, id desc', limit=1)
        invoice_evidence = self.env['logistics.idp.evidence'].search([
            ('case_id', '=', self.id), ('category', '=', 'invoice'), ('status', '=', 'valid'),
        ], order='id desc', limit=1)
        references = {}
        records = []
        for reference_type in ('dsnavl', 'master_data'):
            snapshot = self.env['logistics.idp.reference.snapshot'].search([
                ('company_id', '=', self.company_id.id), ('reference_type', '=', reference_type),
                ('effective_date', '<=', self.effective_date),
            ], order='effective_date desc, id desc', limit=1)
            if snapshot:
                payload = json.loads(snapshot.payload)
                items = payload.get('lines', payload.get('materials', payload.get('material_codes', [])))
                references[reference_type] = {
                    _normalized(item.get('material_code')): item for item in items if isinstance(item, dict) and item.get('material_code')
                }
                records.append(snapshot)
        if not po_snapshot or not invoice_evidence:
            raise ValidationError('SAP/ERP output requires effective PO snapshot and valid case invoice evidence.')
        return po_snapshot, invoice_evidence, references, records

    def generate_sap_erp_output(self, *legacy_payload):
        self.ensure_one()
        if legacy_payload:
            raise ValidationError('SAP/ERP output accepts no caller-provided source payloads.')
        profile, config, policy_binding = self._sap_erp_profile()
        po_snapshot, invoice_evidence, references, reference_records = self._sap_erp_sources()
        po = json.loads(po_snapshot.payload)
        invoice = json.loads(invoice_evidence.payload)
        po_lines, invoice_lines = po.get('lines', []), invoice.get('lines', [])
        if not invoice_lines:
            raise ValidationError('SAP/ERP output requires invoice lines.')
        identity_keys = config['line_identity_keys']
        po_by_identity = {}
        for po_line in po_lines:
            identity = self._sap_erp_identity(po_line, identity_keys)
            if not identity or identity in po_by_identity:
                raise ValidationError('SAP/ERP PO lines require unique configured identity keys.')
            po_by_identity[identity] = po_line
        columns, rows, failures = config['columns'], [], []
        used_reference_records = self.env['logistics.idp.reference.snapshot']
        sort_keys = config['sort_keys']
        for invoice_line in sorted(invoice_lines, key=lambda line: tuple(str(line.get(key, '')) for key in sort_keys)):
            identity = self._sap_erp_identity(invoice_line, identity_keys)
            po_line = po_by_identity.get(identity) if identity else None
            if not po_line:
                failures.append('%s: no PO match' % (identity,))
                continue
            records_by_name = {
                'po_snapshot': po_snapshot, 'invoice_evidence': invoice_evidence,
                'dsnavl': next((record for record in reference_records if record.reference_type == 'dsnavl'), False),
                'master_data': next((record for record in reference_records if record.reference_type == 'master_data'), False),
            }
            source_lines = {'po_snapshot': po_line, 'invoice_evidence': invoice_line}
            line, line_records = {}, {}
            field_precedence = config['source_precedence']
            material = _normalized(invoice_line.get('material_code') or po_line.get('material_code'))
            source_lines.update({name: references.get(name, {}).get(material, {}) for name in ('dsnavl', 'master_data')})
            fields_to_resolve = set().union(*(source.keys() for source in source_lines.values()))
            for field in fields_to_resolve:
                for source_name in self._sap_erp_source_order(field, field_precedence):
                    value = source_lines[source_name].get(field)
                    if value not in (None, ''):
                        line[field], line_records[field] = value, records_by_name[source_name]
                        if source_name in ('dsnavl', 'master_data'):
                            used_reference_records |= records_by_name[source_name]
                        break
            if not line.get('custom_code') and config.get('custom_code_template'):
                try:
                    line['custom_code'] = config['custom_code_template'].format(
                        **line, po_number=self.po_reference, po_item=po_line.get('line_key'))
                except KeyError as exc:
                    failures.append('%s: custom_code missing %s' % (identity, exc.args[0]))
            elif not line.get('custom_code') and not line.get('material_code'):
                custom = config.get('custom_code', {})
                line['custom_code'] = '%s%s%s%s%s' % (custom.get('prefix', 'PO'), custom.get('separator', '-'), self.po_reference,
                                                       custom.get('separator', '-'), format(po_line.get('line_key'), custom.get('item_format', '')))
            rounding, rounding_evidence = self._sap_erp_rounding(config, profile, invoice_line)
            line['amount'] = str(self._sap_erp_amount(invoice_line, config['price_source'], rounding))
            line_records['amount'] = invoice_evidence
            material_code = _normalized(line.get('material_code'))
            context = {'case': {'id': self.id, 'po_reference': self.po_reference, 'supplier_reference': self.supplier_reference},
                       'po': po, 'invoice': invoice, 'line': line, 'po_line': po_line,
                       **{name: values.get(material_code, {}) for name, values in references.items()}}
            row, sources, missing = {}, {}, []
            record_by_source = {'po': po_snapshot, 'po_line': po_snapshot, 'invoice': invoice_evidence,
                                'line': False,
                                'dsnavl': records_by_name['dsnavl'] if context.get('dsnavl') else False,
                                'master_data': records_by_name['master_data'] if context.get('master_data') else False}
            for column in columns:
                name, source = column.get('name'), column.get('source')
                if not name or not source:
                    raise ValidationError('SAP/ERP columns require name and source.')
                value = self._sap_erp_value(source, context)
                root, field = source.split('.', 1) if '.' in source else (source, '')
                record = line_records.get(field) if root == 'line' else record_by_source.get(root)
                row[name] = value
                sources[name] = {'mapping_version': config['mapping_version'], 'source': source,
                                 'record_type': record._name if record else False, 'record_id': record.id if record else False,
                                 'record_hash': record.payload_hash if record else False}
                if column.get('required', True) and value in (None, ''):
                    diagnostic = 'no matching material' if root in ('dsnavl', 'master_data') else (record.display_name if record else 'no source record')
                    missing.append('%s (%s: %s)' % (name, source, diagnostic))
            if missing:
                failures.append('%s: %s' % (identity, ', '.join(sorted(missing))))
            rows.append({'values': row, 'sources': sources,
                         'normalization': {'rounding': rounding, 'rule': rounding_evidence, 'result': line['amount']}})
        if failures:
            raise ValidationError('SAP/ERP required source fields missing: %s.' % '; '.join(failures))
        amount_column = config.get('amount_column')
        if not amount_column or next((column for column in columns if column['name'] == amount_column and column['source'] == config['amount_source']), None) is None:
            raise ValidationError('SAP/ERP amount column must use configured amount source.')
        invoice_total = sum((self._sap_erp_amount(
            line, config['price_source'], self._sap_erp_rounding(config, profile, line)[0]) for line in invoice_lines), Decimal())
        output_total = sum((_decimal(row['values'][amount_column]) for row in rows), Decimal())
        variance = abs(invoice_total - output_total)
        tolerance = _decimal(config.get('total_tolerance', '0'))
        if variance > tolerance:
            raise ValidationError('SAP/ERP invoice/output total variance: invoice=%s output=%s variance=%s tolerance=%s.' % (invoice_total, output_total, variance, tolerance))
        template = self.env['ir.attachment'].browse(config['template_attachment_id']).exists()
        if not template or template.company_id not in (self.company_id, self.env['res.company']):
            raise ValidationError('SAP/ERP template attachment is unavailable for this company.')
        template_hash = hashlib.sha256(template.raw).hexdigest()
        return self.env['logistics.idp.output'].generate(self, 'sap_erp_output', {
            'mapping_version': config['mapping_version'], 'authoritative_policy_binding': policy_binding,
            'template_attachment_id': template.id, 'template_hash': template_hash,
            'sheet': config['sheet'], 'start_row': config.get('start_row', 2), 'supplier_profile_id': profile.id,
            'source_precedence': config['source_precedence'], 'source_records': [
                {'id': record.id, 'hash': record.payload_hash} for record in [po_snapshot, invoice_evidence, *used_reference_records]],
            'columns': columns, 'rows': rows,
            'lines': [{'material_code': row['values'].get('material'), 'hs_code': row['values'].get('hs_code'),
                       'quantity': row['values'].get('quantity'), 'uom': row['values'].get('uom'),
                       'line_total': row['values'].get(amount_column)} for row in rows],
            'invoice_total': str(invoice_total), 'output_total': str(output_total),
            'po_reference': self.po_reference, 'supplier_reference': po.get('supplier') or self.supplier_reference,
            'invoice_number': invoice.get('invoice_number'), 'bill_number': invoice.get('bill_number'),
            'variance': str(variance), 'tolerance': str(tolerance),
            'normalization': {'version': config.get('normalization', {}).get('version'),
                              'rounding': [row['normalization'] for row in rows]},
        })

    def generate_customs_output(self, regime, lines, mapping_version='v1'):
        self.ensure_one()
        if regime not in ('E13', 'E15'):
            raise ValidationError('Vietnam customs output regime must be E13 or E15.')
        _profile, rule, binding = self._authoritative_reference(regime.lower())
        if rule['output_types'] != [regime.lower()]:
            raise ValidationError('%s requires an explicit customs authoritative-reference rule.' % regime)
        return self.env['logistics.idp.output'].generate(
            self, regime.lower(), {'regime': regime, 'mapping_version': mapping_version,
                                   'authoritative_policy_binding': binding, 'lines': lines})

    def generate_broker_package(self, *args, **kwargs):
        self.ensure_one()
        if args or kwargs:
            raise ValidationError('Broker package policy does not accept caller-provided requirements or exclusions.')
        profiles = self.env['logistics.idp.supplier.profile'].search([
            ('company_id', '=', self.company_id.id), ('supplier_reference', '=', self.supplier_reference),
            ('effective_from', '<=', self.effective_date), '|', ('effective_to', '=', False),
            ('effective_to', '>=', self.effective_date),
        ], order='effective_from desc, id desc', limit=2)
        if len(profiles) != 1:
            raise ValidationError('Broker package requires one effective supplier profile.')
        profile = profiles
        policy = json.loads(profile.payload).get('broker_package')
        if not isinstance(policy, dict):
            raise ValidationError('Broker package requires configured supplier profile policy.')
        qdtq_types = policy.get('qdtq_document_types')
        excluded_types = policy.get('exclude_document_types_without_qdtq')
        required = policy.get('required_categories')
        required_checks = policy.get('required_checks', [])
        if (not isinstance(qdtq_types, list) or not qdtq_types or not all(isinstance(item, str) for item in qdtq_types)
                or not isinstance(excluded_types, bool) or not isinstance(required, dict)
                or not all(isinstance(name, str) and isinstance(rule, dict) and set(rule) <= {'document_types', 'output_types'}
                           and any(isinstance(rule.get(key, []), list) and rule.get(key, []) for key in ('document_types', 'output_types'))
                           and all(isinstance(value, str) for key in ('document_types', 'output_types') for value in rule.get(key, []))
                           for name, rule in required.items())
                or not isinstance(required_checks, list) or not all(isinstance(code, str) for code in required_checks)):
            raise ValidationError('Broker package supplier profile policy is invalid.')
        documents = self.document_ids.filtered(lambda item: item.attachment_id)
        selected_documents, exclusions = self.env['logistics.idp.document'], []
        for document in documents.sorted('id'):
            if excluded_types and document.document_type not in qdtq_types:
                exclusions.append({'document_id': document.id, 'rule': 'exclude_document_types_without_qdtq',
                                   'reason': 'document_type_not_in_qdtq_document_types'})
            else:
                selected_documents |= document
        generated = self.output_ids.filtered(lambda item: item.status == 'generated' and item.output_type != 'broker_package' and item.attachment_id)
        available_documents = set(selected_documents.mapped('document_type'))
        available_outputs = set(generated.mapped('output_type'))
        missing = [name for name, rule in sorted(required.items()) if not (
            set(rule.get('document_types', [])) <= available_documents and set(rule.get('output_types', [])) <= available_outputs)]
        effective_checks = {}
        for check in self.check_result_ids:
            effective_checks.setdefault(check.code, []).append(check)
        invalid_checks = sorted(code for code in required_checks
                                if len(effective_checks.get(code, [])) != 1
                                or effective_checks[code][0].verdict != 'pass')
        if missing or invalid_checks:
            raise ValidationError('Broker package blocked: missing required categories %s; invalid required checks %s.' % (
                ', '.join(missing) or 'none', ', '.join(invalid_checks) or 'none'))
        entries = []
        for document in selected_documents:
            entries.append({'name': 'documents/%s-%s' % (document.id, document.attachment_id.name),
                            'attachment_id': document.attachment_id.id, 'source_record_id': document.id,
                            'source_hash': document.content_hash, 'document_type': document.document_type})
        for output in generated.sorted(lambda item: (item.output_type, item.version, item.id)):
            entries.append({'name': 'outputs/%s-v%s-%s' % (output.output_type, output.version, output.attachment_id.name),
                            'attachment_id': output.attachment_id.id, 'source_record_id': output.id,
                            'source_hash': output.artifact_sha256, 'output_type': output.output_type})
        output = self.env['logistics.idp.output']
        return output.with_context(_logistics_broker_package_capability=output._broker_package_capability)._generate_broker_package(self, {
            'profile': {'id': profile.id, 'version': profile.version, 'hash': profile.payload_hash},
            'entries': entries, 'exclusions': exclusions,
        })

    def queue_broker_email(self, *args, **kwargs):
        self.ensure_one()
        if args or kwargs:
            raise ValidationError('Broker email policy does not accept caller-provided values.')
        if not (self.env.su or self.env.user.has_group('insilos_logistics_idp.group_logistics_manager')):
            raise UserError('Only Logistics Managers may queue broker email.')
        profiles = self.env['logistics.idp.supplier.profile'].search([
            ('company_id', '=', self.company_id.id), ('supplier_reference', '=', self.supplier_reference),
            ('effective_from', '<=', self.effective_date), '|', ('effective_to', '=', False),
            ('effective_to', '>=', self.effective_date),
        ], order='effective_from desc, id desc', limit=2)
        if len(profiles) != 1:
            raise ValidationError('Broker email requires one effective supplier profile.')
        profile = profiles
        config = json.loads(profile.payload).get('broker_package', {}).get('email')
        if not isinstance(config, dict) or set(config) - {'template_id', 'to', 'cc', 'reply_to'}:
            raise ValidationError('Broker email supplier profile configuration is invalid.')
        template = self.env['mail.template'].browse(config.get('template_id')).exists()
        if len(template) != 1 or template.model != self._name:
            raise ValidationError('Broker email requires a native mail template linked to Logistics Case.')
        def recipients(value, required=False):
            values = value if isinstance(value, list) else re.split(r'[,;\n]+', value or '') if isinstance(value, str) else []
            raw = [item.strip() for item in values if isinstance(item, str) and item.strip()]
            normalized = sorted({item.lower() for item in raw if re.fullmatch(r'[^\s@,;]+@[^\s@,;]+\.[^\s@,;]+', item)})
            if (required and not normalized) or len(normalized) != len({item.lower() for item in raw}):
                raise ValidationError('Broker email recipients must be a safe valid list.')
            return ','.join(normalized)
        email_to, email_cc = recipients(config.get('to'), True), recipients(config.get('cc'))
        reply_to = recipients(config.get('reply_to'))
        if reply_to and ',' in reply_to:
            raise ValidationError('Broker email reply_to must contain one safe address.')
        package = self.output_ids.filtered(lambda item: item.output_type == 'broker_package' and item.status == 'generated' and item.attachment_id).sorted('id', reverse=True)[:1]
        if not package:
            raise ValidationError('Broker email requires a current generated broker package.')
        rendered = template._generate_template([self.id], ('subject', 'body_html'))[self.id]
        template_version = fields.Datetime.to_string(template.write_date)
        key = hashlib.sha256(_canonical_json({'profile_hash': profile.payload_hash, 'template_version': template_version, 'content': rendered, 'package': package.artifact_sha256, 'to': email_to, 'cc': email_cc, 'reply_to': reply_to}).encode()).hexdigest()
        self.env.cr.execute("SELECT pg_advisory_xact_lock(hashtextextended(%s, 0))", ['logistics-idp-broker-email:%s' % key])
        existing = self.env['mail.mail'].sudo().search([('headers', 'ilike', key), ('state', 'in', ('outgoing', 'sent', 'exception'))], limit=1)
        if existing:
            return existing
        with self.env.cr.savepoint():
            mail_id = template.send_mail(self.id, force_send=False, email_values={'email_to': email_to, 'email_cc': email_cc or False, 'reply_to': reply_to or False, 'auto_delete': False, 'attachment_ids': [(4, package.attachment_id.id)], 'headers': repr({'X-Logistics-Broker-Email-Key': key})})
            mail = self.env['mail.mail'].sudo().browse(mail_id)
            request = {'key': key, 'profile_hash': profile.payload_hash, 'template_id': template.id, 'template_version': template_version, 'package_output_id': package.id, 'package_artifact_sha256': package.artifact_sha256, 'recipients': {'to': email_to, 'cc': email_cc, 'reply_to': reply_to}, 'rendered': rendered}
            try:
                prepared = mail._prepare_outgoing_list(mail.mail_server_id)
                if len(prepared) != 1:
                    raise ValidationError('Broker email must prepare exactly one outgoing message.')
                item = prepared[0]
                message = (mail.mail_server_id or self.env['ir.mail_server'])._build_email__(
                    email_from=item['email_from'], email_to=item['email_to'], subject=item['subject'], body=item['body'],
                    body_alternative=item['body_alternative'], email_cc=item['email_cc'], reply_to=item['reply_to'],
                    attachments=item['attachments'], message_id=item['message_id'], references=item['references'],
                    object_id=item['object_id'], subtype='html', subtype_alternative='plain', headers=item['headers'])
                raw = message.as_bytes()
                snapshot = self.env['ir.attachment'].sudo().create({
                    'name': 'broker-email-%s.eml' % key, 'raw': raw, 'mimetype': 'message/rfc822',
                    'res_model': self._name, 'res_id': self.id, 'company_id': self.company_id.id,
                })
                request['mime'] = {'available': True, 'attachment_id': snapshot.id,
                    'sha256': hashlib.sha256(raw).hexdigest(), 'size': len(raw),
                    'parts': ['%s/%s' % (part.get_content_maintype(), part.get_content_subtype()) for part in message.walk()],
                    'charset': 'utf-8', 'attachment_sha256': package.artifact_sha256}
            except AssertionError:
                request['mime'] = {'available': False, 'reason': 'native_sender_unavailable'}
            evidence = self.env['logistics.idp.evidence']
            evidence._controlled_create({'case_id': self.id, 'category': 'broker_email_request', 'source_reference': 'broker-email:%s' % key, 'status': 'review', 'payload': request}, 'broker_email_queue')
            evidence._controlled_create({'case_id': self.id, 'category': 'broker_email_queued', 'source_reference': 'broker-email:%s' % key, 'status': 'valid', 'payload': {**request, 'mail_id': mail.id, 'mail_state': mail.state}}, 'broker_email_queue')
            self.message_post(body='Broker package email queued: %s.' % key)
            return mail

    def queue_supplier_result_reply(self, *args, **kwargs):
        self.ensure_one()
        if args or kwargs:
            raise ValidationError('Supplier result reply does not accept caller-provided values.')
        if not (self.env.su or self.env.user.has_group('insilos_logistics_idp.group_logistics_manager')):
            raise UserError('Only Logistics Managers may queue supplier result replies.')
        profiles = self.env['logistics.idp.supplier.profile'].search([
            ('company_id', '=', self.company_id.id), ('supplier_reference', '=', self.supplier_reference),
            ('effective_from', '<=', self.effective_date), '|', ('effective_to', '=', False),
            ('effective_to', '>=', self.effective_date),
        ], order='effective_from desc, id desc', limit=2)
        if len(profiles) != 1:
            raise ValidationError('Supplier result reply requires one effective supplier profile.')
        profile = profiles
        config = json.loads(profile.payload).get('draft_invoice', {}).get('result_email')
        if not isinstance(config, dict) or set(config) - {'template_id', 'cc'}:
            raise ValidationError('Supplier result reply configuration is invalid.')
        template = self.env['mail.template'].browse(config.get('template_id')).exists()
        if len(template) != 1 or template.model != self._name:
            raise ValidationError('Supplier result reply requires a native mail template linked to Logistics Case.')
        def recipients(value):
            values = value if isinstance(value, list) else re.split(r'[,;\n]+', value or '') if isinstance(value, str) else []
            raw = [item.strip() for item in values if isinstance(item, str) and item.strip()]
            normalized = sorted({item.lower() for item in raw if re.fullmatch(r'[^\s@,;]+@[^\s@,;]+\.[^\s@,;]+', item)})
            if len(normalized) != len({item.lower() for item in raw}):
                raise ValidationError('Supplier result reply recipients must be a safe valid list.')
            return ','.join(normalized)
        email_cc = recipients(config.get('cc'))
        trusted_domains = json.loads(profile.payload).get('email_domains')
        if trusted_domains is not None and (not isinstance(trusted_domains, list) or not all(isinstance(domain, str) and domain for domain in trusted_domains)):
            raise ValidationError('Supplier result reply trusted sender configuration is invalid.')
        trusted_documents = self.document_ids.filtered(lambda item: item.document_type == 'draft_vat_invoice'
            and item.source_channel == 'queue' and item.source_sender and item.source_message_reference
            and self.evidence_ids.filtered(lambda evidence: evidence.category == 'email_attachment'
                and evidence.status == 'valid' and evidence.payload_hash
                and json.loads(evidence.payload).get('audit_service') == 'intake_supplier_email'
                and json.loads(evidence.payload).get('source') == 'intake_supplier_email'
                and json.loads(evidence.payload).get('document_id') == item.id
                and json.loads(evidence.payload).get('content_hash') == item.content_hash
                and self.env['logistics.idp.inbound.job'].sudo().browse(json.loads(evidence.payload).get('inbound_job_id')).exists().filtered(
                    lambda job: job.case_id == self and job.source_system == 'supplier_email' and job.state == 'done'))).sorted('id')
        if len(trusted_documents) != 1:
            raise ValidationError('Supplier result reply requires exactly one trusted supplier-email draft invoice source.')
        source = trusted_documents
        email_to = recipients(source.source_sender)
        reference = source.source_message_reference.strip()
        if not email_to or not re.fullmatch(r'<[^<>\r\n]+>', reference):
            raise ValidationError('Supplier result reply source metadata is invalid.')
        if trusted_domains and email_to.rsplit('@', 1)[1] not in {domain.lower() for domain in trusted_domains}:
            raise ValidationError('Supplier result reply source sender is not trusted for this supplier profile.')
        rendered = template._generate_template([self.id], ('subject', 'body_html'))[self.id]
        template_version = fields.Datetime.to_string(template.write_date)
        key = hashlib.sha256(_canonical_json({'profile_hash': profile.payload_hash, 'template_version': template_version,
            'content': rendered, 'source_document_id': source.id, 'source_hash': source.content_hash,
            'sender': email_to, 'reference': reference, 'cc': email_cc}).encode()).hexdigest()
        self.env.cr.execute("SELECT pg_advisory_xact_lock(hashtextextended(%s, 0))", ['logistics-idp-supplier-result-reply:%s' % key])
        existing = self.env['mail.mail'].sudo().search([('headers', 'ilike', key), ('state', 'in', ('outgoing', 'sent', 'exception'))], limit=1)
        if existing:
            return existing
        headers = {'References': reference, 'In-Reply-To': reference, 'X-Logistics-Supplier-Result-Reply-Key': key}
        mail_id = template.send_mail(self.id, force_send=False, email_values={
            'email_to': email_to, 'email_cc': email_cc or False, 'auto_delete': False, 'headers': repr(headers),
        })
        mail = self.env['mail.mail'].sudo().browse(mail_id)
        request = {'key': key, 'profile_hash': profile.payload_hash, 'template_id': template.id,
            'template_version': template_version, 'source_document_id': source.id, 'source_content_hash': source.content_hash,
            'source_sender': email_to, 'source_message_reference': reference, 'cc': email_cc, 'rendered': rendered}
        evidence = self.env['logistics.idp.evidence']
        evidence._controlled_create({'case_id': self.id, 'category': 'supplier_result_reply_request',
            'source_reference': 'supplier-result-reply:%s' % key, 'status': 'review', 'payload': request}, 'supplier_result_reply_queue')
        evidence._controlled_create({'case_id': self.id, 'category': 'supplier_result_reply_queued',
            'source_reference': 'supplier-result-reply:%s' % key, 'status': 'valid',
            'payload': {**request, 'mail_id': mail.id, 'mail_state': mail.state, 'headers': headers}}, 'supplier_result_reply_queue')
        self.message_post(body='Supplier result reply queued: %s.' % key)
        return mail

    def record_broker_email_delivery(self, mail):
        self.ensure_one()
        mail = self.env['mail.mail'].sudo().browse(mail.id if hasattr(mail, 'id') else mail).exists()
        if len(mail) != 1 or mail.model != self._name or mail.res_id != self.id or mail.state not in ('sent', 'exception'):
            return False
        try:
            headers = ast.literal_eval(mail.headers or '{}')
        except (SyntaxError, TypeError, ValueError):
            return False
        key = headers.get('X-Logistics-Broker-Email-Key') if isinstance(headers, dict) else False
        if not isinstance(key, str) or not re.fullmatch(r'[0-9a-f]{64}', key):
            return False
        reference = 'broker-email:%s:%s' % (key, mail.state)
        self.env.cr.execute("SELECT pg_advisory_xact_lock(hashtextextended(%s, 0))", [
            'logistics-idp-broker-email-delivery:%s:%s' % (key, mail.state)])
        if not self.env['logistics.idp.evidence'].search_count([
            ('case_id', '=', self.id), ('category', '=', 'broker_email_delivery'),
            ('source_reference', '=', reference),
        ]):
            self.env['logistics.idp.evidence']._controlled_create({'case_id': self.id, 'category': 'broker_email_delivery', 'source_reference': reference, 'status': 'valid' if mail.state == 'sent' else 'invalid', 'payload': {'key': key, 'mail_id': mail.id, 'message_id': mail.message_id, 'state': mail.state, 'failure_type': mail.failure_type, 'failure_code': 'broker_delivery_failed' if mail.state == 'exception' else False}}, 'broker_email_delivery')
        return True

    def record_supplier_result_reply_delivery(self, mail):
        self.ensure_one()
        mail = self.env['mail.mail'].sudo().browse(mail.id if hasattr(mail, 'id') else mail).exists()
        if len(mail) != 1 or mail.model != self._name or mail.res_id != self.id or mail.state not in ('sent', 'exception'):
            return False
        try:
            headers = ast.literal_eval(mail.headers or '{}')
        except (SyntaxError, TypeError, ValueError):
            return False
        key = headers.get('X-Logistics-Supplier-Result-Reply-Key') if isinstance(headers, dict) else False
        if not isinstance(key, str) or not re.fullmatch(r'[0-9a-f]{64}', key):
            return False
        reference = 'supplier-result-reply:%s:%s' % (key, mail.state)
        self.env.cr.execute("SELECT pg_advisory_xact_lock(hashtextextended(%s, 0))", [
            'logistics-idp-supplier-result-reply-delivery:%s:%s' % (key, mail.state)])
        evidence = self.env['logistics.idp.evidence']
        if not evidence.search_count([('case_id', '=', self.id), ('category', '=', 'supplier_result_reply_delivery'), ('source_reference', '=', reference)]):
            evidence._controlled_create({'case_id': self.id, 'category': 'supplier_result_reply_delivery',
                'source_reference': reference, 'status': 'valid' if mail.state == 'sent' else 'invalid',
                'payload': {'key': key, 'mail_id': mail.id, 'message_id': mail.message_id, 'state': mail.state,
                            'failure_type': mail.failure_type, 'failure_code': 'supplier_result_reply_delivery_failed' if mail.state == 'exception' else False}}, 'supplier_result_reply_delivery')
        return True

    def reconcile_import_declaration(self, declaration, tolerance='0'):
        self.ensure_one()
        declaration = declaration.exists() if hasattr(declaration, 'exists') else self.env['logistics.idp.document']
        if len(declaration) != 1 or declaration.case_id != self or declaration.document_type != 'import_declaration' or declaration.status not in ('valid', 'review') or not declaration.current_run_id:
            raise ValidationError('Import declaration reconciliation requires one valid canonical import-declaration document.')
        extracted = json.loads(declaration.current_run_id.payload).get('payload', {})
        if not isinstance(extracted, dict):
            raise ValidationError('Import declaration canonical output is invalid.')
        profile, reference_rule, policy_binding = self._authoritative_reference('import_declaration')
        if reference_rule['output_types'] == ['sap_erp_output']:
            regimes = ('sap_erp_output',)
        elif set(reference_rule['output_types']) == {'e13', 'e15'}:
            regimes = tuple(reference_rule['output_types'])
        else:
            raise ValidationError('Import declaration reconciliation requires an explicit E13/E15 authoritative-reference rule.')
        references = self.output_ids.filtered(lambda item: item.output_type in regimes and item.status == 'generated')
        missing_types = sorted(set(regimes) - set(references.mapped('output_type')))
        if missing_types:
            raise ValidationError('Import declaration reconciliation blocked: missing generated %s reference dataset(s); no reconciliation evidence created.' % '/'.join(item.upper() for item in missing_types))
        latest, selected_types, output_payloads = self.env['logistics.idp.output'], set(), []
        for output in references.sorted(lambda item: (item.version, item.id), reverse=True):
            payload = json.loads(output.payload)
            binding = payload.get('authoritative_policy_binding')
            _output_profile, _output_rule, expected_binding = self._authoritative_reference(output.output_type)
            if (not isinstance(binding, dict) or any(binding.get(key) != expected_binding[key] for key in (
                    'profile_id', 'profile_hash', 'stage', 'field', 'rule_id', 'policy_version', 'policy_hash',
                    'source_type', 'output_types'))):
                raise ValidationError('Import declaration reconciliation blocked: unbound or stale authoritative output; no reconciliation evidence created.')
            if output.output_type not in selected_types:
                selected_types.add(output.output_type)
                latest |= output
                output_payloads.append((output, payload))

        def identity(line):
            material, custom = _normalized(line.get('material_code')), _normalized(line.get('custom_code'))
            return ('material_code', material) if material else (('custom_code', custom) if custom else (None, None))

        def aggregate(lines):
            grouped, errors = {}, []
            for line in lines:
                kind, key = identity(line)
                if not key:
                    errors.append('missing material/custom code')
                    continue
                item = grouped.setdefault((kind, key, _normalized(line.get('regime'))), {'quantity': Decimal(), 'line_total': Decimal(), 'lines': []})
                try:
                    item['quantity'] += _decimal(line.get('quantity'))
                    item['line_total'] += _decimal(line.get('line_total', line.get('amount')))
                except (InvalidOperation, TypeError, ValueError):
                    errors.append('%s: invalid quantity or line total' % key)
                item['lines'].append(line)
            return grouped, errors

        expected_lines, output_values = [], {}
        for output, payload in output_payloads:
            lines = payload.get('lines', [])
            if output.output_type == 'sap_erp_output':
                lines = [{
                    'material_code': row.get('values', {}).get('material'),
                    'custom_code': row.get('values', {}).get('custom'),
                    'quantity': row.get('values', {}).get('quantity'),
                    'uom': row.get('values', {}).get('uom'),
                    'line_total': row.get('values', {}).get('amount'),
                } for row in payload.get('rows', [])]
            expected_lines.extend(lines)
            output_values.update({key: value for key, value in payload.items() if key != 'lines' and value not in (None, '')})
            for line in lines:
                line.setdefault('regime', payload.get('regime'))
                for field in ('supplier', 'importer', 'currency'):
                    if field not in output_values and line.get(field) not in (None, ''):
                        output_values[field] = line[field]
        expected, expected_errors = aggregate(expected_lines)
        actual, actual_errors = aggregate(extracted.get('lines', []))
        results, missing = [], expected_errors + actual_errors
        declaration_number = extracted.get('declaration_number')
        if not declaration_number:
            missing.append('declaration number')
        for field in ('supplier', 'importer', 'currency'):
            expected_value, actual_value = output_values.get(field), extracted.get(field)
            if expected_value not in (None, '') or actual_value not in (None, ''):
                if expected_value in (None, '') or actual_value in (None, ''):
                    missing.append(field)
                else:
                    results.append({'field': field, 'expected': expected_value, 'actual': actual_value,
                                    'result': 'pass' if _normalized(expected_value) == _normalized(actual_value) else 'block'})
        if output_values.get('total_invoice_value') is None:
            output_values['total_invoice_value'] = str(sum((item['line_total'] for item in expected.values()), Decimal()))
        if extracted.get('total_invoice_value') not in (None, ''):
            results.append({'field': 'total_invoice_value', 'expected': output_values['total_invoice_value'], 'actual': extracted['total_invoice_value'],
                            'result': 'pass' if _decimal(output_values['total_invoice_value']) == _decimal(extracted['total_invoice_value']) else 'block'})
        for key in sorted(set(expected) | set(actual)):
            expected_item, actual_item = expected.get(key), actual.get(key)
            if not expected_item or not actual_item:
                results.append({'field': key[0], 'identity': key[1], 'result': 'block', 'reason': 'unmatched'})
                continue
            expected_line, actual_line = expected_item['lines'][0], actual_item['lines'][0]
            comparisons = {'uom': (_normalized(expected_line.get('uom')), _normalized(actual_line.get('uom'))),
                           'quantity': (expected_item['quantity'], actual_item['quantity']),
                           'line_total': (expected_item['line_total'], actual_item['line_total'])}
            if regimes != ('sap_erp_output',):
                comparisons.update({'regime': (_normalized(expected_line.get('regime')), _normalized(actual_line.get('regime'))),
                                    'hs_code': (_normalized(expected_line.get('hs_code')), _normalized(actual_line.get('hs_code')))})
            for field, (expected_value, actual_value) in comparisons.items():
                if expected_value in (None, '') or actual_value in (None, ''):
                    missing.append('%s:%s' % (key[1], field))
                else:
                    results.append({'field': field, 'identity': key[1], 'expected': str(expected_value), 'actual': str(actual_value),
                                    'result': 'pass' if expected_value == actual_value else 'block'})
        verdict = 'review' if missing else ('block' if any(item['result'] == 'block' for item in results) else 'pass')
        evidence = self.env['logistics.idp.evidence']._controlled_create({
            'case_id': self.id, 'category': 'import_declaration', 'source_reference': str(declaration_number or 'declaration:unknown'),
            'status': {'pass': 'valid', 'block': 'invalid', 'review': 'review'}[verdict],
            'payload': {'reference': '/'.join(item.upper() for item in regimes), 'verdict': verdict, 'results': results,
                        'missing': sorted(set(missing)), 'authoritative_policy_binding': policy_binding,
                        'declaration_document': {'id': declaration.id, 'hash': declaration.content_hash, 'run_hash': declaration.current_run_id.payload_hash},
                        'output_records': [{'id': output.id, 'hash': output.payload_hash} for output in latest]},
        }, 'import_declaration_reconciliation')
        self.env['logistics.idp.check.result']._create_from_check_runner({
            'case_id': self.id, 'code': 'IMPORT_DECLARATION', 'required': True, 'verdict': verdict,
            'rationale': 'E13/E15 import declaration reconciliation', 'citation': evidence.source_reference,
            'payload': {'evidence_id': evidence.id, 'evidence_hash': evidence.payload_hash},
        }, 'import_declaration_reconciliation')
        self._derive_lifecycle()
        return evidence

    def _local_executor_profile(self, output_type):
        self.ensure_one()
        profiles = self.env['logistics.idp.supplier.profile'].search([
            ('company_id', '=', self.company_id.id), ('supplier_reference', '=', self.supplier_reference),
            ('effective_from', '<=', self.effective_date), '|', ('effective_to', '=', False),
            ('effective_to', '>=', self.effective_date),
        ], order='effective_from desc, id desc', limit=2)
        if len(profiles) != 1:
            raise ValidationError('%s requires exactly one effective same-company supplier profile.' % output_type.replace('_', ' ').title())
        config = json.loads(profiles.payload).get(output_type)
        if not isinstance(config, dict):
            raise ValidationError('%s supplier profile contract is missing.' % output_type.replace('_', ' ').title())
        fields_config = config.get('fields')
        if (not isinstance(config.get('mapping_version'), str) or not config['mapping_version']
                or not isinstance(config.get('template_attachment_id'), int) or not isinstance(config.get('sheet'), str)
                or not isinstance(config.get('logical_key'), str) or not isinstance(config.get('declaration_optional_pre_customs'), bool)
                or not isinstance(fields_config, list) or not fields_config):
            raise ValidationError('%s supplier profile contract is invalid.' % output_type.replace('_', ' ').title())
        positions = [item.get('position') for item in fields_config if isinstance(item, dict)]
        if (len(positions) != len(fields_config) or any(not isinstance(item.get('name'), str) or not item['name'] or not isinstance(item.get('source'), str) or not item['source'] or not isinstance(item.get('required', True), bool) or not isinstance(item.get('position'), int) or isinstance(item['position'], bool) or item['position'] < 1 for item in fields_config) or len(set(positions)) != len(positions)):
            raise ValidationError('%s mapped fields require names, sources, required flags and unique positive positions.' % output_type.replace('_', ' ').title())
        mapped = {item['name']: item for item in fields_config}
        required_names = ({'sequence', 'delivery_date', 'declaration_number', 'po_number', 'invoice_number', 'bill_number', 'supplier', 'goods_description', 'package_count', 'gross_weight', 'net_weight', 'carrier', 'note'} if output_type == 'shipping_plan' else {'number', 'validity', 'vehicle_or_carrier', 'approval_or_signature'} if output_type == 'gate_pass' else set())
        optional_declaration = output_type == 'shipping_plan' and config.get('declaration_optional_pre_customs')
        if any(name not in mapped or (not mapped[name]['required'] and not (name == 'declaration_number' and optional_declaration)) for name in required_names):
            raise ValidationError('%s semantic required fields are invalid.' % output_type.replace('_', ' ').title())
        if output_type == 'gate_pass' and (not isinstance(config.get('required_checks', []), list) or not all(isinstance(code, str) and code for code in config['required_checks'])):
            raise ValidationError('Gate Pass required checks are invalid.')
        template = self.env['ir.attachment'].browse(config['template_attachment_id']).exists()
        if not template or template.company_id != self.company_id:
            raise ValidationError('%s template attachment is unavailable for this company.' % output_type.replace('_', ' ').title())
        return profiles, config, template

    @staticmethod
    def _local_executor_value(source, context):
        allowed = {
            'case': {'source_key', 'po_reference', 'supplier_reference', 'shipment_reference', 'effective_date'},
            'sap': {'invoice_total', 'output_total', 'logical_key', 'mapped_fields', 'po_reference', 'supplier_reference', 'invoice_number', 'bill_number'},
            'customs': {'regime', 'lines', 'logical_key', 'mapped_fields'},
            'declaration': {'declaration_number', 'logical_key', 'mapped_fields'},
            'invoice': {'goods_description', 'package_count', 'gross_weight', 'net_weight', 'carrier', 'note'},
            'shipping_plan': {'logical_key', 'mapped_fields'},
        }
        parts = source.split('.')
        if len(parts) < 2 or parts[0] not in allowed or parts[1] not in allowed[parts[0]]:
            raise ValidationError('Local direct output source is not allowed: %s.' % source)
        value = context[parts[0]]
        for part in parts[1:]:
            if isinstance(value, dict):
                value = value.get(part)
            elif isinstance(value, (list, tuple)) and part.isdigit() and int(part) < len(value):
                value = value[int(part)]
            else:
                return None
        return value

    def _current_semantic_checks(self, code, output_id=None):
        self.ensure_one()
        checks = []
        for check in self.check_result_ids.filtered(lambda item: item.code == code):
            payload = json.loads(check.payload)
            if output_id is None or payload.get('output_id') == output_id:
                checks.append(check)
        return checks

    def _local_executor_sources(self, declaration_optional):
        sap = self.output_ids.filtered(lambda item: item.output_type == 'sap_erp_output' and item.status == 'generated').sorted(lambda item: (item.version, item.id), reverse=True)[:1]
        if not sap:
            raise ValidationError('Local direct output requires current generated SAP/ERP output.')
        customs = self.output_ids.filtered(lambda item: item.output_type in ('e13', 'e15') and item.status == 'generated').sorted(lambda item: (item.version, item.id), reverse=True)[:1]
        declaration = self.env['logistics.idp.evidence']
        reconciliations = {item.payload_hash for item in self.evidence_ids.filtered(
            lambda item: item.category == 'reconciliation' and item.status == 'valid')}
        for item in self.evidence_ids.filtered(lambda evidence: evidence.category == 'import_declaration' and evidence.status == 'valid').sorted('id', reverse=True):
            payload = json.loads(item.payload)
            if payload.get('reconciliation_hash') in reconciliations:
                declaration = item
                break
        invoice = self.evidence_ids.filtered(lambda item: item.category == 'invoice' and item.status == 'valid').sorted('id', reverse=True)[:1]
        if not declaration_optional and (not customs or not declaration):
            raise ValidationError('Local direct output requires canonical customs output and valid declaration reconciliation evidence.')
        records = [sap] + list(customs) + list(declaration) + list(invoice)
        return sap, customs, declaration, invoice, records

    def _generate_local_direct_output(self, output_type):
        self.ensure_one()
        profile, config, template = self._local_executor_profile(output_type)
        sap, customs, declaration, invoice, records = self._local_executor_sources(config['declaration_optional_pre_customs'])
        if output_type == 'gate_pass':
            if not config.get('enabled'):
                raise ValidationError('Gate Pass is disabled by its supplier profile.')
            plan = self.output_ids.filtered(lambda item: item.output_type == 'shipping_plan' and item.status == 'generated').sorted(lambda item: (item.version, item.id), reverse=True)[:1]
            approvals = self._current_semantic_checks('shipping_plan_approval', plan.id if plan else None)
            invalid_checks = sorted(code for code in config['required_checks']
                                    if len(self._current_semantic_checks(code)) != 1
                                    or self._current_semantic_checks(code)[0].verdict not in ('pass', 'not_applicable'))
            if not plan or len(approvals) != 1 or approvals[0].verdict != 'pass' or invalid_checks:
                raise ValidationError('Gate Pass requires an enabled current approved Shipping Plan and valid required checks %s.' % (', '.join(invalid_checks) or 'none'))
            if config.get('requires_validated_declaration') and not declaration:
                raise ValidationError('Gate Pass requires valid declaration evidence.')
            records.append(plan)
        context = {'case': {'source_key': self.source_key, 'po_reference': self.po_reference,
                            'supplier_reference': self.supplier_reference, 'shipment_reference': self.shipment_reference,
                            'effective_date': fields.Date.to_string(self.effective_date)},
                   'sap': json.loads(sap.payload), 'customs': json.loads(customs.payload) if customs else {},
                   'declaration': json.loads(declaration.payload).get('declaration', json.loads(declaration.payload)) if declaration else {},
                   'invoice': json.loads(invoice.payload) if invoice else {},
                   'shipping_plan': json.loads(plan.payload) if output_type == 'gate_pass' else {}}
        mapped, provenance, missing = {}, {}, []
        for field in config['fields']:
            value = self._local_executor_value(field['source'], context)
            mapped[field['name']] = value
            provenance[field['name']] = {'mapping_version': config['mapping_version'], 'source': field['source'],
                                         'record_type': 'logistics.idp.output' if field['source'].startswith(('sap.', 'customs.', 'shipping_plan.')) else 'logistics.idp.evidence' if field['source'].startswith(('declaration.', 'invoice.')) else 'logistics.idp.case',
                                         'record_id': sap.id if field['source'].startswith('sap.') else customs.id if field['source'].startswith('customs.') and customs else plan.id if field['source'].startswith('shipping_plan.') else declaration.id if field['source'].startswith('declaration.') and declaration else invoice.id if field['source'].startswith('invoice.') and invoice else self.id,
                                         'record_hash': sap.payload_hash if field['source'].startswith('sap.') else customs.payload_hash if field['source'].startswith('customs.') and customs else plan.payload_hash if field['source'].startswith('shipping_plan.') else declaration.payload_hash if field['source'].startswith('declaration.') and declaration else invoice.payload_hash if field['source'].startswith('invoice.') and invoice else False}
            if field.get('required', True) and value in (None, ''):
                missing.append(field['name'])
        if missing:
            raise ValidationError('%s required mapped fields are missing: %s.' % (output_type.replace('_', ' ').title(), ', '.join(missing)))
        logical_key = self._local_executor_value(config['logical_key'], context)
        if logical_key in (None, ''):
            raise ValidationError('%s logical key is missing.' % output_type.replace('_', ' ').title())
        values = {'mapping_version': config['mapping_version'], 'logical_key': str(logical_key), 'template_attachment_id': template.id,
                  'template_hash': hashlib.sha256(template.raw).hexdigest(), 'sheet': config['sheet'], 'fields': config['fields'],
                  'mapped_fields': mapped, 'field_provenance': provenance, 'supplier_profile_id': profile.id,
                  'supplier_profile_hash': profile.payload_hash, 'source_records': [{'model': item._name, 'id': item.id, 'hash': item.payload_hash} for item in records]}
        return self.env['logistics.idp.output']._generate(self, output_type, values)

    def generate_shipping_plan(self, *args, **kwargs):
        if args or kwargs:
            raise ValidationError('Shipping Plan accepts no caller-provided source payloads.')
        return self._generate_local_direct_output('shipping_plan')

    def generate_gate_pass(self, *args, **kwargs):
        if args or kwargs:
            raise ValidationError('Gate Pass accepts no caller-provided source payloads.')
        return self._generate_local_direct_output('gate_pass')

    def queue_gate_pass_distribution(self, *args, **kwargs):
        self.ensure_one()
        if args or kwargs:
            raise ValidationError('Gate Pass distribution accepts no caller-provided values.')
        if not (self.env.su or self.env.user.has_group('insilos_logistics_idp.group_logistics_manager')):
            raise UserError('Only Logistics Managers may queue Gate Pass distribution.')
        profiles = self.env['logistics.idp.supplier.profile'].search([
            ('company_id', '=', self.company_id.id), ('supplier_reference', '=', self.supplier_reference),
            ('effective_from', '<=', self.effective_date), '|', ('effective_to', '=', False),
            ('effective_to', '>=', self.effective_date),
        ], order='effective_from desc, id desc', limit=2)
        if len(profiles) != 1:
            raise ValidationError('Gate Pass distribution requires one effective same-company supplier profile.')
        profile = profiles
        config = json.loads(profile.payload).get('gate_pass', {}).get('distribution')
        if not isinstance(config, dict) or set(config) != {'channel', 'template_id', 'to'} or config['channel'] != 'native_mail':
            raise ValidationError('Gate Pass distribution configuration is invalid.')
        raw = config['to'] if isinstance(config['to'], list) else re.split(r'[,;\n]+', config['to'] or '') if isinstance(config['to'], str) else []
        recipients = sorted({item.strip().lower() for item in raw if isinstance(item, str) and re.fullmatch(r'[^\s@,;]+@[^\s@,;]+\.[^\s@,;]+', item.strip())})
        if not recipients or len(recipients) != len({item.strip().lower() for item in raw if isinstance(item, str) and item.strip()}):
            raise ValidationError('Gate Pass distribution recipients must be a safe valid list.')
        template = self.env['mail.template'].browse(config['template_id']).exists()
        if len(template) != 1 or template.model != self._name:
            raise ValidationError('Gate Pass distribution requires a native mail template linked to Logistics Case.')
        gate_pass = self.output_ids.filtered(lambda item: item.output_type == 'gate_pass' and item.status == 'generated' and item.attachment_id).sorted(lambda item: (item.version, item.id), reverse=True)[:1]
        if not gate_pass:
            raise ValidationError('Gate Pass distribution requires an active generated Gate Pass.')
        source_records = json.loads(gate_pass.payload).get('source_records', [])
        plans = [item for item in source_records if item.get('model') == 'logistics.idp.output' and item.get('id') and item.get('hash')]
        plan = self.env['logistics.idp.output'].browse(plans[-1]['id']).exists() if plans else self.env['logistics.idp.output']
        if len(plan) != 1 or plan.case_id != self or plan.output_type != 'shipping_plan' or plan.status != 'generated' or plan.payload_hash != plans[-1]['hash']:
            raise ValidationError('Gate Pass distribution requires its active current Shipping Plan.')
        approvals = self._current_semantic_checks('shipping_plan_approval', plan.id)
        if len(approvals) != 1 or approvals[0].verdict != 'pass':
            raise ValidationError('Gate Pass distribution requires its approved Shipping Plan.')
        rendered = template._generate_template([self.id], ('subject', 'body_html'))[self.id]
        template_version = fields.Datetime.to_string(template.write_date)
        recipient_text = ','.join(recipients)
        key = hashlib.sha256(_canonical_json({'output_id': gate_pass.id, 'output_hash': gate_pass.artifact_sha256,
            'output_version': gate_pass.version, 'shipping_plan_id': plan.id, 'shipping_plan_hash': plan.payload_hash,
            'profile_hash': profile.payload_hash, 'template_id': template.id, 'template_version': template_version,
            'channel': config['channel'], 'recipients': recipient_text, 'rendered': rendered}).encode()).hexdigest()
        self.env.cr.execute("SELECT pg_advisory_xact_lock(hashtextextended(%s, 0))", ['logistics-idp-gate-pass-distribution:%s' % key])
        existing = self.env['mail.mail'].sudo().search([('headers', 'ilike', key), ('state', 'in', ('outgoing', 'sent', 'exception'))], limit=1)
        if existing:
            return existing
        request = {'key': key, 'channel': config['channel'], 'profile_hash': profile.payload_hash,
            'template_id': template.id, 'template_version': template_version, 'gate_pass': {'id': gate_pass.id,
            'hash': gate_pass.artifact_sha256, 'version': gate_pass.version}, 'shipping_plan': {'id': plan.id,
            'hash': plan.payload_hash}, 'recipients': recipient_text}
        try:
            with self.env.cr.savepoint():
                mail_id = template.send_mail(self.id, force_send=False, email_values={'email_to': recipient_text,
                    'auto_delete': False, 'attachment_ids': [(4, gate_pass.attachment_id.id)],
                    'headers': repr({'X-Logistics-Gate-Pass-Distribution-Key': key})})
                mail = self.env['mail.mail'].sudo().browse(mail_id)
                self.env['logistics.idp.evidence']._controlled_create({'case_id': self.id,
                    'category': 'gate_pass_distribution_request', 'source_reference': 'gate-pass-distribution:%s' % key,
                    'status': 'review', 'payload': request}, 'gate_pass_distribution_queue')
                self.env['logistics.idp.evidence']._controlled_create({'case_id': self.id,
                    'category': 'gate_pass_distribution_queued', 'source_reference': 'gate-pass-distribution:%s' % key,
                    'status': 'valid', 'payload': {**request, 'mail_id': mail.id, 'mail_state': mail.state}}, 'gate_pass_distribution_queue')
                self.message_post(body='Gate Pass distribution queued: %s.' % key)
                return mail
        except Exception as exc:
            if isinstance(exc, (UserError, ValidationError)):
                raise
            raise ValidationError('Gate Pass distribution queue failed.') from None

    @api.model
    def evaluate_taxable_value(self, invoice_value, invoice_currency, taxable_currency, effective_fx=None):
        invoice_currency = _normalized(invoice_currency)
        taxable_currency = _normalized(taxable_currency)
        missing = [name for name, value in (
            ('invoice_value', invoice_value), ('invoice_currency', invoice_currency),
            ('taxable_currency', taxable_currency), ('effective_fx', effective_fx),
        ) if value in (None, '')]
        if missing:
            return {'verdict': 'review', 'reason': 'INSUFFICIENT_CONTEXT', 'missing_critical': missing}
        if not isinstance(effective_fx, dict) or any(effective_fx.get(key) in (None, '') for key in ('rate', 'effective_date', 'base_currency', 'quote_currency')):
            return {'verdict': 'review', 'reason': 'INSUFFICIENT_FX_CONTEXT', 'missing_critical': ['effective_fx']}
        base = _normalized(effective_fx['base_currency'])
        quote = _normalized(effective_fx['quote_currency'])
        rate = _decimal(effective_fx['rate'])
        if rate <= 0 or (base, quote) != (invoice_currency, taxable_currency):
            return {'verdict': 'review', 'reason': 'INVALID_FX_CONTEXT', 'missing_critical': ['effective_fx']}
        value = _decimal(invoice_value) * rate
        return {'verdict': 'pass', 'invoice_value': str(_decimal(invoice_value)),
                'invoice_currency': invoice_currency, 'taxable_value': str(value),
                'taxable_currency': taxable_currency, 'effective_fx': {
                    'rate': str(rate), 'effective_date': fields.Date.to_string(
                        fields.Date.to_date(effective_fx['effective_date'])),
                    'base_currency': base, 'quote_currency': quote,
                }}

    @api.model
    def evaluate_counterpart_deadline(self, start_date, today, rule, counterpart_present=False):
        start_date = fields.Date.to_date(start_date)
        today = fields.Date.to_date(today)
        deadline = fields.Date.add(start_date, days=rule['deadline_days'])
        if counterpart_present:
            status = 'satisfied'
        elif today >= deadline:
            status = 'overdue'
        elif (deadline - today).days in rule['warning_offsets']:
            status = 'warning'
        else:
            status = 'pending'
        return {'status': status, 'deadline': deadline}

    def _record_counterpart_alert_intent(self, rule_code, status, cutoff, communication):
        self.ensure_one()
        if self.state in communication['no_send_states']:
            return self.env['logistics.idp.evidence']
        existing = self.evidence_ids.filtered(lambda item: item.category == 'counterpart_alert_intent')
        payloads = [json.loads(item.payload) for item in existing]
        if any(item.get('rule_code') == rule_code and item.get('cutoff') == fields.Date.to_string(cutoff)
               for item in payloads):
            return self.env['logistics.idp.evidence']
        sent = sum(item.get('rule_code') == rule_code for item in payloads)
        if sent >= communication['max_reminders']:
            return self.env['logistics.idp.evidence']
        payload = {
            'rule_code': rule_code, 'alert_status': status, 'cutoff': fields.Date.to_string(cutoff),
            'cadence_slot': sent, 'send_enabled': bool(communication.get('send_enabled')),
        }
        return self.env['logistics.idp.evidence']._controlled_create({
            'case_id': self.id, 'category': 'counterpart_alert_intent',
            'source_reference': '%s:%s:%s' % (rule_code, fields.Date.to_string(cutoff), sent),
            'status': 'valid', 'payload': payload,
        }, 'counterpart_deadline_alert')

    def evaluate_compliance_deadlines(self, today=None):
        today = fields.Date.to_date(today or fields.Date.context_today(self))
        policy_model = self.env['logistics.idp.policy.source'].sudo()
        evaluated = self.env['logistics.idp.check.result']
        for case in self.filtered(lambda item: item.company_id in self.env.companies):
            policy = policy_model.select_effective_pack(
                case.company_id, 'trade_compliance', None, 'VN', 'ALL', case.effective_date)
            if not policy:
                continue
            payload = policy_model.validate_policy_payload(
                policy.payload, 'trade_compliance', None, 'VN', 'ALL')
            pack = payload['vertical']['packs'][payload['vertical']['pack']]
            aliases = pack['batch']['lifecycle_aliases']
            documents = [(document, aliases.get(document.document_type, document.document_type))
                         for document in case.document_ids]
            roles = {role for _document, role in documents}
            evidence_snapshot = [{
                'id': evidence.id, 'category': evidence.category, 'payload_hash': evidence.payload_hash,
                'status': evidence.status,
            } for evidence in case.evidence_ids.filtered(
                lambda item: item.category != 'counterpart_alert_intent').sorted('id')]
            for rule in pack['counterpart_rules']:
                start_roles = set(rule['start_roles'])
                triggers = [document for document, role in documents if role in start_roles]
                if not triggers:
                    continue
                expected = set(rule['expected_roles'])
                dated_triggers = []
                for document in triggers:
                    try:
                        raw_payload = (document.current_run_id.payload if document.current_run_id else None)
                        if not raw_payload and document.attachment_id:
                            raw_payload = document.attachment_id.raw.decode('utf-8')
                        extracted = json.loads(raw_payload or '{}')
                        for _ in range(5):
                            if isinstance(extracted, str):
                                extracted = json.loads(extracted)
                            elif isinstance(extracted, dict) and isinstance(extracted.get('payload'), (dict, str)):
                                extracted = extracted['payload']
                            else:
                                break
                        field_name = rule['start_date']['document_field']
                        def find_field(value, name):
                            if isinstance(value, dict):
                                if value.get(name) is not None:
                                    return value[name]
                                for child in value.values():
                                    found = find_field(child, name)
                                    if found is not None:
                                        return found
                            elif isinstance(value, list):
                                for child in value:
                                    found = find_field(child, name)
                                    if found is not None:
                                        return found
                            return None
                        value = find_field(extracted, field_name)
                        if value is None and document.attachment_id:
                            source = json.loads(document.attachment_id.raw.decode('utf-8'))
                            value = find_field(source, field_name)
                        date = fields.Date.to_date(value) if value is not None else None
                    except (TypeError, ValueError):
                        date = None
                    dated_triggers.append((document, date))
                valid_dates = [date for _document, date in dated_triggers if date]
                counterpart_present = bool(roles.intersection(expected))
                missing_context = [] if valid_dates or counterpart_present else ['when']
                start_date = min(valid_dates) if valid_dates else case.effective_date
                result = self.evaluate_counterpart_deadline(
                    start_date, today, rule, counterpart_present)
                if missing_context:
                    result['status'] = 'insufficient_context'
                rule_code = rule['code']
                check_code = '%s_%s' % (rule_code, result['status'].upper())
                context = {
                    'case': {'company_id': case.company_id.id, 'source_system': case.source_system,
                             'source_key': case.source_key, 'source_version': case.source_version},
                    'roles': sorted(roles),
                    'triggers': [{'document_id': document.id, 'content_hash': document.content_hash,
                                  'date': fields.Date.to_string(date) if date else None}
                                 for document, date in sorted(dated_triggers, key=lambda item: item[0].id)],
                    'evidence': evidence_snapshot,
                }
                audit_input = {
                    'policy': {'id': policy.id, 'code': policy.code, 'version': policy.version,
                               'payload_hash': policy.payload_hash, 'source_tier': policy.source_tier,
                               'citation': policy.citation},
                    'rule': {'code': rule_code, 'version': rule['version']},
                    'context': context, 'evaluation_cutoff': fields.Date.to_string(today),
                }
                audit_input_hash = hashlib.sha256(_canonical_json(audit_input).encode()).hexdigest()
                open_exceptions = self.env['logistics.idp.exception'].search([
                    ('case_id', '=', case.id), ('exception_type', '=', 'policy'),
                    ('state', 'in', ('open', 'waiting')),
                    ('message_ids.body', 'ilike', '[%s]' % rule_code),
                ])
                activities = case.sudo().activity_ids.filtered(
                    lambda activity: activity.summary and activity.summary.startswith(rule_code + '_'))
                if result['status'] == 'satisfied':
                    open_exceptions.action_resolve('Configured counterpart evidence present for %s.' % rule_code)
                    activities.action_feedback(feedback='Configured counterpart evidence present.')
                    continue
                if result['status'] not in ('warning', 'overdue', 'insufficient_context'):
                    continue
                if result['status'] in ('warning', 'overdue'):
                    case._record_counterpart_alert_intent(
                        rule_code, result['status'], today, payload['supplier_communication'])
                check_model = self.env['logistics.idp.check.result'].sudo()
                existing = check_model.search([('audit_input_hash', '=', audit_input_hash)], limit=1)
                if existing:
                    evaluated |= existing
                    continue
                legal_authority = AUTHORITY_TIERS[policy.source_tier]
                verdict = ('block' if result['status'] == 'overdue' and legal_authority else 'review')
                severity = 'high' if result['status'] == 'overdue' else 'medium'
                rationale = ('Counterpart declaration %s; deadline %s.' % (result['status'], result['deadline'])
                             if not missing_context else 'INSUFFICIENT_CONTEXT: missing valid trigger date.')
                values = {
                    'case_id': case.id, 'policy_source_id': policy.id, 'code': check_code,
                    'expected': json.dumps(rule['expected_roles'], sort_keys=True),
                    'actual': json.dumps(sorted(roles)), 'verdict': verdict,
                    'rationale': rationale, 'citation': policy.citation,
                    'audit_input_hash': audit_input_hash,
                    'payload': {**result, 'rule_code': rule_code, 'rule_version': rule['version'],
                                'start_date': fields.Date.to_string(start_date),
                                'deadline': fields.Date.to_string(result['deadline']),
                                'missing_critical': missing_context, 'legal_authority': legal_authority,
                                'disposition': 'legal' if legal_authority else 'customer_policy',
                                'audit_input': audit_input},
                }
                try:
                    with self.env.cr.savepoint():
                        check = check_model._create_from_check_runner(values, 'compliance_deadline_evaluation')
                except UniqueViolation:
                    check = check_model.search([('audit_input_hash', '=', audit_input_hash)], limit=1)
                evaluated |= check
                if not open_exceptions:
                    exception = self.env['logistics.idp.exception'].create({
                        'case_id': case.id, 'exception_type': 'policy', 'severity': severity,
                        'due_at': fields.Datetime.to_datetime(result['deadline']),
                    })
                    exception.message_post(body='[%s] Counterpart declaration policy exception.' % rule_code)
                if activities:
                    activities.write({'summary': check_code, 'date_deadline': result['deadline'],
                                      'note': rule['review']})
                else:
                    case.sudo().activity_schedule(
                        'mail.mail_activity_data_todo', date_deadline=result['deadline'],
                        summary=check_code, note=rule['review'],
                        user_id=case.owner_id.id or self.env.user.id)
        return evaluated

    @api.model
    def _cron_evaluate_compliance_deadlines(self):
        domain = [('company_id', 'in', self.env.companies.ids),
                  ('state', 'not in', ('completed', 'closed_duplicate', 'closed_other'))]
        self.search(domain).evaluate_compliance_deadlines()
        return True

    def reconcile(self):
        for case in self:
            evidence = case.evidence_ids
            verdict = 'review'
            if evidence and all(item.status == 'valid' for item in evidence):
                verdict = 'pass'
            if any(item.status == 'invalid' for item in evidence):
                verdict = 'block'
            case.with_context(_logistics_case_lifecycle=_INTERNAL_CASE_LIFECYCLE_TOKEN).write({
                'verdict': verdict, 'state': 'blocked' if verdict == 'block' else 'review'})
        return True


class ImmutableSnapshot(models.AbstractModel):
    _name = 'logistics.idp.immutable.snapshot'
    _description = 'Immutable Logistics IDP Snapshot'

    payload = fields.Text(required=True, readonly=True)
    payload_hash = fields.Char(required=True, readonly=True, index=True)
    audit_actor_id = fields.Many2one('res.users', required=True, readonly=True, default=lambda self: self.env.user)
    audit_service = fields.Char(required=True, readonly=True, default='domain_engine')
    audit_input_hash = fields.Char(readonly=True, index=True)

    @api.model
    def _controlled_create(self, values, service):
        values = [dict(item, audit_actor_id=self.env.user.id, audit_service=service) for item in values] if isinstance(values, list) else dict(values, audit_actor_id=self.env.user.id, audit_service=service)
        snapshots = self.with_context(_logistics_snapshot_token=_INTERNAL_SNAPSHOT_TOKEN).create(values)
        return snapshots.with_context(_logistics_snapshot_token=None)

    @api.model
    def load(self, fields, data):
        # ponytail: trust only the core CSV caller; new loader paths need explicit review.
        if (self.env.su and not self.env.registry.ready
                and self.env.context.get('install_mode')
                and not self.env.context.get('import_file')
                and sys._getframe(1).f_code is tools.convert.convert_csv_import.__code__):
            self = self.with_context(_logistics_snapshot_token=_INTERNAL_SNAPSHOT_TOKEN)
        return super().load(fields, data)

    def _load_records_create(self, vals_list):
        if (self.env.su and not self.env.registry.ready
                and self.env.context.get('install_mode')
                and not self.env.context.get('import_file')
                and '_import_current_module' not in self.env.context):
            loader = self.with_context(
                _logistics_snapshot_token=_INTERNAL_SNAPSHOT_TOKEN,
                _logistics_policy_loader_token=_INTERNAL_POLICY_LOADER_TOKEN,
            )
            return super(ImmutableSnapshot, loader)._load_records_create(vals_list).with_context(
                _logistics_snapshot_token=None, _logistics_policy_loader_token=None)
        return super()._load_records_create(vals_list)

    @api.model_create_multi
    def create(self, vals_list):
        if (self.env.context.get('_logistics_snapshot_token') is not _INTERNAL_SNAPSHOT_TOKEN
                and not tools.config['test_enable']):
            raise UserError('Snapshots may only be created by a controlled Logistics service.')
        for vals in vals_list:
            vals.setdefault('audit_actor_id', self.env.user.id)
            vals.setdefault('audit_service', self.env.context.get('_logistics_snapshot_service', 'domain_engine'))
            try:
                payload = json.loads(vals['payload']) if isinstance(vals['payload'], str) else vals['payload']
            except (TypeError, ValueError) as exc:
                raise ValidationError('Snapshot payload must be valid JSON.') from exc
            vals['payload'] = _canonical_json(payload)
            vals['payload_hash'] = hashlib.sha256(vals['payload'].encode()).hexdigest()
        return super().create(vals_list)

    def write(self, vals):
        if self.env.context.get('_logistics_snapshot_token') is _INTERNAL_SNAPSHOT_TOKEN:
            vals = dict(vals)
            if 'payload' in vals:
                payload = json.loads(vals['payload']) if isinstance(vals['payload'], str) else vals['payload']
                vals['payload'] = _canonical_json(payload)
                vals['payload_hash'] = hashlib.sha256(vals['payload'].encode()).hexdigest()
            return super().write(vals)
        raise UserError('Snapshots are immutable; append a new version instead.')

    def unlink(self):
        raise UserError('Snapshots are immutable and cannot be deleted.')


class LogisticsEmailThreadEntry(models.Model):
    _name = 'logistics.idp.email.thread.entry'
    _description = 'Immutable Supplier Email Thread Entry'
    _order = 'id desc'

    case_id = fields.Many2one('logistics.idp.case', required=True, readonly=True, ondelete='restrict', index=True)
    company_id = fields.Many2one('res.company', required=True, readonly=True, index=True)
    supplier_reference = fields.Char(required=True, readonly=True, index=True)
    normalized_subject = fields.Char(required=True, readonly=True, index=True)
    source_identity = fields.Char(required=True, readonly=True, index=True)
    is_active = fields.Boolean(required=True, readonly=True, default=False, index=True)

    _source_identity_unique = models.Constraint(
        'unique(case_id, source_identity)', 'This email source is already recorded for the case.')

    @api.model_create_multi
    def create(self, vals_list):
        if self.env.context.get('_logistics_email_thread_entry_token') is not _INTERNAL_EMAIL_THREAD_ENTRY_TOKEN:
            raise AccessError('Email thread history is controlled by intake.')
        return super().create(vals_list)

    def write(self, vals):
        if set(vals) - {'is_active'} or self.env.context.get('_logistics_email_thread_entry_token') is not _INTERNAL_EMAIL_THREAD_ENTRY_TOKEN:
            raise AccessError('Email thread history is immutable.')
        return super().write(vals)

    def unlink(self):
        raise AccessError('Email thread history is immutable.')

    def init(self):
        self.env.cr.execute('DROP INDEX IF EXISTS logistics_idp_email_thread_entry_one_active_uniq')
        self.env.cr.execute('''
            CREATE UNIQUE INDEX logistics_idp_email_thread_entry_one_active_uniq
            ON logistics_idp_email_thread_entry (case_id, supplier_reference, normalized_subject)
            WHERE is_active
        ''')


class LogisticsInboundJob(models.Model):
    _name = 'logistics.idp.inbound.job'
    _description = 'Logistics IDP Inbound Job'
    _inherit = ['mail.thread']
    _order = 'next_attempt_at, id'

    profile_code = fields.Char(required=True, readonly=True)
    adapter_type = fields.Selection([('generic_inbound', 'Generic Inbound')], required=True, default='generic_inbound', readonly=True)
    source_system = fields.Char(required=True, readonly=True)
    source_key = fields.Char(required=True, readonly=True)
    source_version = fields.Char(required=True, readonly=True)
    payload = fields.Text(required=True, readonly=True)
    company_id = fields.Many2one('res.company', required=True, readonly=True, default=lambda self: self.env.company, index=True)
    attachment_ids = fields.Many2many('ir.attachment', 'logistics_idp_inbound_job_attachment_rel', 'job_id', 'attachment_id', readonly=True)
    state = fields.Selection([
        ('pending', 'Pending'), ('processing', 'Processing'), ('retry', 'Retry'),
        ('done', 'Done'), ('dead', 'Dead Letter'), ('cancelled', 'Cancelled'),
    ], required=True, default='pending', readonly=True, index=True, tracking=True)
    attempt_count = fields.Integer(default=0, readonly=True)
    max_attempts = fields.Integer(default=3, required=True, readonly=True)
    next_attempt_at = fields.Datetime(default=fields.Datetime.now, required=True, readonly=True, index=True)
    last_attempt_at = fields.Datetime(readonly=True)
    completed_at = fields.Datetime(readonly=True)
    last_error = fields.Text(readonly=True)
    case_id = fields.Many2one('logistics.idp.case', readonly=True, ondelete='restrict')

    _source_identity_unique = models.Constraint(
        'unique(company_id, source_system, source_key, source_version)',
        'This inbound source version is already queued for this company.')
    _attempts_positive = models.Constraint('check(max_attempts > 0)', 'Maximum attempts must be positive.')

    def _validate_attachment_link(self, vals):
        if 'attachment_ids' in vals and self.env.context.get('_logistics_inbound_attachment_token') is not _INTERNAL_INBOUND_ATTACHMENT_TOKEN:
            raise ValidationError('Inbound job attachments require controlled supplier-email intake.')

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            self._validate_attachment_link(vals)
            if vals.get('source_system') != 'supplier_email' and vals.get('payload'):
                try:
                    attachments = json.loads(vals['payload']).get('attachments')
                except (TypeError, ValueError):
                    attachments = None
                if attachments:
                    raise ValidationError('Generic inbound jobs cannot carry attachment metadata.')
        return super().create(vals_list)

    def write(self, vals):
        self._validate_attachment_link(vals)
        return super().write(vals)

    @api.model
    def intake_supplier_email(self, message, attachments):
        """Queue one sanitized supplier email. Caller supplies bytes; URLs are never accepted."""
        message_id = str(message.get('message_id') or '').strip()
        company = self.env['res.company'].browse(message.get('company_id')) if message.get('company_id') else self.env.company
        if not message_id or not company.exists() or company not in self.env.companies:
            raise ValidationError('Email Message-ID and an allowed company are required.')
        sender_domain = str(message.get('sender_domain') or '').strip().lower()
        profiles = self.env['logistics.idp.supplier.profile'].sudo().search([
            ('company_id', '=', company.id), ('effective_from', '<=', fields.Date.today()),
            '|', ('effective_to', '=', False), ('effective_to', '>=', fields.Date.today()),
        ])
        matches = profiles.filtered(lambda profile: sender_domain and sender_domain in {
            str(domain).lower() for domain in json.loads(profile.payload).get('email_domains', [])})
        supplier = matches.supplier_reference if len(matches) == 1 else False
        thread = str(message.get('thread_reference') or message_id).strip()[:255]
        po_reference = str(message.get('po_reference') or '').strip()[:255]
        subject = str(message.get('subject') or '').strip()[:255]
        payload = {
            'name': subject or 'Supplier inbound email', 'provenance': 'supplier_email',
            'effective_date': str(message.get('effective_date') or fields.Date.today()),
            'company_id': company.id, 'thread_reference': thread, 'email_subject': subject,
            'po_reference': po_reference or None, 'supplier_reference': supplier or None,
            'source_sender': str(message.get('source_sender') or '').strip()[:255] or None,
            'source_message_reference': message_id[:255],
            'supplier_ambiguous': len(matches) != 1, 'attachments': [],
        }
        source_key = hashlib.sha256(('%s\0%s' % (company.id, message_id)).encode()).hexdigest()
        domain = [('source_system', '=', 'supplier_email'), ('source_key', '=', source_key), ('source_version', '=', '1')]
        existing = self.search(domain, limit=1)
        if existing:
            return existing
        try:
            with self.env.cr.savepoint():
                job = self.create({
                    'profile_code': 'supplier_email', 'source_system': 'supplier_email',
                    'source_key': source_key, 'source_version': '1', 'company_id': company.id,
                    'payload': _canonical_json(payload),
                })
        except UniqueViolation:
            return self.search(domain, limit=1)
        attachment_ids = []
        for index, item in enumerate(attachments or []):
            content = item.get('content')
            filename = str(item.get('filename') or '')[:255]
            if not isinstance(content, bytes) or not content or len(content) > MAX_BYTES:
                status = 'missing' if not isinstance(content, bytes) else 'empty' if not content else 'oversized'
                entry = {'index': index, 'status': status,
                         'manifest_identity': hashlib.sha256(content).hexdigest() if isinstance(content, bytes)
                         else hashlib.sha256(_canonical_json({'index': index, 'status': status}).encode()).hexdigest()}
                payload['attachments'].append(entry)
                continue
            detected = guess_mimetype(content, default='application/octet-stream')
            declared = str(item.get('mimetype') or '').split(';', 1)[0].lower()
            allowed = detected in {'application/pdf', 'application/vnd.ms-excel',
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'text/csv',
                'application/msword', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'}
            entry = {'filename': filename, 'index': index, 'mimetype': detected,
                     'declared_mimetype': declared, 'mime_valid': not declared or declared == detected,
                     'status': 'queued' if allowed else 'unsupported',
                     'manifest_identity': hashlib.sha256(content).hexdigest()}
            if allowed:
                attachment = self.env['ir.attachment'].sudo().create({
                    'name': filename or 'supplier-email-%s' % index, 'raw': content, 'mimetype': detected,
                    'res_model': self._name, 'res_id': job.id, 'company_id': company.id,
                })
                entry.update({'attachment_id': attachment.id, 'content_hash': hashlib.sha256(content).hexdigest()})
                attachment_ids.append(attachment.id)
            payload['attachments'].append(entry)
        job.with_context(_logistics_inbound_attachment_token=_INTERNAL_INBOUND_ATTACHMENT_TOKEN).write({
            'payload': _canonical_json(payload), 'attachment_ids': [(6, 0, attachment_ids)],
        })
        return job

    def _check_recovery_access(self):
        if not (self.env.user.has_group('insilos_logistics_idp.group_logistics_manager')
                or self.env.user.has_group('insilos_logistics_idp.group_logistics_admin')):
            raise AccessError(_('Only Logistics managers may recover inbound jobs.'))

    def _recover(self, action):
        self.ensure_one()
        self._check_recovery_access()
        allowed = {'retry': {'dead'}, 'requeue': {'dead', 'cancelled'},
                   'cancel': {'pending', 'processing', 'retry', 'dead'}}
        invalid = self.filtered(lambda job: job.state not in allowed[action])
        if invalid:
            raise UserError(_('Selected job state does not allow this recovery action.'))
        now = fields.Datetime.now()
        values = {'retry': {'state': 'retry', 'attempt_count': self.max_attempts - 1},
                  'requeue': {'state': 'pending', 'attempt_count': 0},
                  'cancel': {'state': 'cancelled'}}[action]
        values.update({'next_attempt_at': now, 'completed_at': now if action == 'cancel' else False,
                       'last_error': False if action != 'cancel' else self.last_error})
        self.write(values)
        for job in self:
            job.message_post(body=_('Inbound job recovery: %(action)s by %(user)s.',
                                    action=action, user=self.env.user.display_name))
        return True

    def action_retry(self):
        return self._recover('retry')

    def action_requeue(self):
        return self._recover('requeue')

    def action_cancel(self):
        return self._recover('cancel')

    @api.model
    def operational_health(self, stale_minutes=30):
        if not (self.env.user.has_group('insilos_logistics_idp.group_logistics_manager')
                or self.env.user.has_group('insilos_logistics_idp.group_logistics_admin')
                or self.env.user.has_group('insilos_logistics_idp.group_logistics_integration')):
            raise AccessError(_('Only Logistics managers, administrators, and integration services may read operational health.'))
        if not isinstance(stale_minutes, int) or isinstance(stale_minutes, bool) or stale_minutes <= 0:
            raise ValidationError('Stale processing threshold must be a positive integer.')
        now = fields.Datetime.now()
        company = self.env.company
        jobs = self.sudo().with_company(company)
        company_domain = [('company_id', '=', company.id)]
        queued_domain = company_domain + [('state', 'in', ('pending', 'retry'))]
        oldest = jobs.search(queued_domain, order='next_attempt_at asc, id asc', limit=1)
        queue_age_seconds = max(0, int((now - oldest.next_attempt_at).total_seconds())) if oldest else None
        stale_domain = company_domain + [
            ('state', '=', 'processing'),
            ('last_attempt_at', '<=', fields.Datetime.subtract(now, minutes=stale_minutes)),
        ]
        policies = self.env['logistics.idp.policy.source'].sudo().with_company(company).search([
            ('state', '=', 'active'), ('company_id', 'in', (False, company.id)),
            ('effective_from', '<=', fields.Date.today()),
            '|', ('effective_to', '=', False), ('effective_to', '>=', fields.Date.today()),
        ])
        domain = self.env['openrouter.domain'].sudo().search([
            ('code', '=', 'logistics_ocr'), ('active', '=', True),
        ], limit=1)
        models = (domain.primary_model_id | domain.fallback_model_ids).filtered(
            lambda model: model.active and model.supports_vision and model.supports_json)
        return {
            'schema_version': '1.0',
            'observed_at': fields.Datetime.to_string(now),
            'company_id': company.id,
            'queue': {
                'queued_count': jobs.search_count(queued_domain),
                'oldest_age_seconds': queue_age_seconds,
                'dead_count': jobs.search_count(company_domain + [('state', '=', 'dead')]),
                'stale_processing_count': jobs.search_count(stale_domain),
                'stale_after_seconds': stale_minutes * 60,
            },
            'policy': {
                'active': bool(policies),
                'active_count': len(policies),
                'fresh_count': len(policies.filtered(lambda policy: policy.freshness_status == 'fresh')),
            },
            'provider': {
                'active': bool(domain and models),
                'provider': domain.provider if domain else None,
                'capable_model_count': len(models),
            },
        }

    @api.model
    def _recover_stale_processing(self, stale_minutes=30):
        cutoff = fields.Datetime.subtract(fields.Datetime.now(), minutes=stale_minutes)
        stale = self.search([('state', '=', 'processing'), ('last_attempt_at', '<=', cutoff)])
        for job in stale:
            job.write({'state': 'retry', 'next_attempt_at': fields.Datetime.now(),
                       'last_error': 'INBOUND_FAILURE:StaleProcessing'})
            job.message_post(body=_('Inbound job automatically recovered from stale processing.'))
        return len(stale)

    @api.model
    def _claim_due(self, limit=20):
        self.env.cr.execute("""SELECT id FROM logistics_idp_inbound_job WHERE state IN ('pending', 'retry') AND next_attempt_at <= NOW() ORDER BY next_attempt_at, id FOR UPDATE SKIP LOCKED LIMIT %s""", [limit])
        jobs = self.browse(row[0] for row in self.env.cr.fetchall())
        jobs.write({'state': 'processing', 'last_attempt_at': fields.Datetime.now()})
        return jobs

    @api.model
    def _cron_process(self, limit=20):
        self._recover_stale_processing()
        for job in self._claim_due(limit):
            job._process_one()
        return True

    def _record_terminal_failure(self, error):
        error_code = 'INBOUND_FAILURE:%s' % type(error).__name__
        case = self.case_id or self.env['logistics.idp.case'].with_company(self.company_id).intake({
            'name': 'Inbound failure', 'provenance': 'inbound_failure', 'effective_date': fields.Date.today(),
            'company_id': self.company_id.id, 'source_system': self.source_system,
            'source_key': self.source_key, 'source_version': self.source_version,
        })
        try:
            attachments = json.loads(self.payload).get('attachments', [])
        except (TypeError, ValueError):
            attachments = []
        if not attachments:
            attachments = [{'index': index, 'status': 'stored',
                            'manifest_identity': hashlib.sha256(attachment.raw).hexdigest()}
                           for index, attachment in enumerate(self.attachment_ids)]
        for attachment in attachments:
            index = attachment.get('index', 0)
            reference = '%s:%s' % (self.source_key, index)
            if not case.evidence_ids.filtered(lambda item: item.category == 'inbound_failure' and item.source_reference == reference):
                self.env['logistics.idp.evidence']._controlled_create({
                    'case_id': case.id, 'category': 'inbound_failure', 'source_reference': reference,
                    'status': 'invalid', 'payload': {'error_code': error_code, 'index': index,
                                                      'status': attachment.get('status'),
                                                      'manifest_identity': attachment.get('manifest_identity')},
                }, 'inbound_terminal_failure')
        exception = self.env['logistics.idp.exception'].search([
            ('case_id', '=', case.id), ('exception_type', '=', 'inbound_failure'),
            ('state', 'in', ('open', 'waiting')),
        ], limit=1)
        if not exception:
            self.env['logistics.idp.exception'].create({
                'case_id': case.id, 'exception_type': 'inbound_failure', 'severity': 'medium',
            })
        activity_type = self.env.ref('mail.mail_activity_data_todo')
        if not self.env['mail.activity'].search([
            ('res_model', '=', case._name), ('res_id', '=', case.id),
            ('activity_type_id', '=', activity_type.id), ('summary', '=', error_code),
        ], limit=1):
            case.activity_schedule('mail.mail_activity_data_todo', user_id=case.owner_id.id, summary=error_code)
        case.write({'state': 'review', 'next_action': error_code})
        return case

    def _process_one(self):
        self.ensure_one()
        if self.state in ('done', 'dead', 'cancelled'):
            return self.state == 'done'
        now = fields.Datetime.now()
        attempts = self.attempt_count + 1
        try:
            values = json.loads(self.payload)
            required = ('name', 'provenance', 'effective_date')
            if any(not values.get(field) for field in required):
                raise ValidationError('Inbound payload is missing required context.')
            company = self.env['res.company'].browse(values['company_id']) if values.get('company_id') else self.env.company
            case = self.env['logistics.idp.case'].with_company(company).intake({
                **{field: values[field] for field in required},
                'company_id': values.get('company_id') or self.env.company.id,
                'source_system': self.source_system, 'source_key': self.source_key, 'source_version': self.source_version,
                **{field: values[field] for field in ('po_reference', 'supplier_reference', 'thread_reference', 'email_subject') if values.get(field)},
            })
            if self.source_system == 'supplier_email':
                policy = self.env['logistics.idp.policy.source'].sudo().select_effective_pack(
                    case.company_id, 'trade_compliance', None, 'VN', 'ALL', case.effective_date)
                config = self.env['logistics.idp.policy.source'].validate_policy_payload(policy.payload) if policy else None
                for attachment in values.get('attachments', []):
                    evidence_payload = {key: attachment.get(key) for key in (
                        'filename', 'index', 'mimetype', 'declared_mimetype', 'mime_valid', 'status')}
                    evidence_payload.update({'audit_service': 'intake_supplier_email', 'source': 'intake_supplier_email'})
                    if attachment.get('status') == 'queued':
                        stored = self.attachment_ids.filtered(lambda item: item.id == attachment.get('attachment_id'))
                        if len(stored) != 1 or stored.company_id != case.company_id or stored.res_model != self._name or stored.res_id != self.id:
                            raise ValidationError('Inbound attachment is not owned by this job and company.')
                        content = stored.raw
                        if hashlib.sha256(content).hexdigest() != attachment.get('content_hash'):
                            raise ValidationError('Inbound attachment hash mismatch.')
                        evidence_payload.update({'document_type': 'unknown', 'classification_confidence': 0})
                        document = self.env['logistics.idp.document'].intake_content(
                            case, content, attachment['mimetype'], {
                                'filename': attachment['filename'], 'source_sender': values.get('source_sender'),
                                'source_message_reference': values.get('source_message_reference'),
                                'action_required': 'Review classification', 'status': 'review',
                            }, source_channel='queue', process=False)
                        evidence_payload.update({'content_hash': document.content_hash, 'document_id': document.id,
                                                 'inbound_job_id': self.id})
                    reference = '%s:%s:%s' % (self.source_key, attachment.get('index', 0),
                                               attachment.get('content_hash') or attachment.get('filename') or 'unnamed')
                    if not case.evidence_ids.filtered(lambda item: item.category == 'email_attachment' and item.source_reference == reference):
                        self.env['logistics.idp.evidence']._controlled_create({
                            'case_id': case.id, 'category': 'email_attachment', 'source_reference': reference,
                            'status': 'review' if attachment.get('status') != 'queued' or values.get('supplier_ambiguous') else 'valid',
                            'payload': evidence_payload,
                        }, 'supplier_email_intake')
                completeness = batch_structure([
                    {'document_type': document.document_type} for document in case.document_ids], config)
                if values.get('supplier_ambiguous') or completeness['missing']:
                    reason = 'Ambiguous supplier identity.' if values.get('supplier_ambiguous') else 'Missing roles: %s.' % ', '.join(completeness['missing'])
                    exception_type = 'supplier' if values.get('supplier_ambiguous') else 'policy'
                    exception = self.env['logistics.idp.exception'].search([
                        ('case_id', '=', case.id), ('exception_type', '=', exception_type),
                        ('state', 'not in', ('resolved', 'cancelled')),
                    ], limit=1)
                    if not exception:
                        self.env['logistics.idp.exception'].create({
                            'case_id': case.id, 'exception_type': exception_type, 'severity': 'medium',
                        }).message_post(body=reason)
                    activity_type = self.env.ref('mail.mail_activity_data_todo')
                    activity = self.env['mail.activity'].search([
                        ('res_model', '=', case._name), ('res_id', '=', case.id),
                        ('activity_type_id', '=', activity_type.id), ('summary', '=', reason),
                    ], limit=1)
                    if not activity:
                        case.activity_schedule('mail.mail_activity_data_todo', user_id=case.owner_id.id, summary=reason)
                    case.write({'state': 'review', 'next_action': reason})
        except Exception as exc:
            terminal = attempts >= self.max_attempts
            case = self._record_terminal_failure(exc) if terminal else self.case_id
            self.write({
                'attempt_count': attempts, 'last_attempt_at': now,
                'last_error': 'INBOUND_FAILURE:%s' % type(exc).__name__,
                'state': 'dead' if terminal else 'retry',
                'next_attempt_at': fields.Datetime.add(now, minutes=2 ** (attempts - 1)),
                'case_id': case.id if case else False,
            })
            return False
        self.write({'attempt_count': attempts, 'last_attempt_at': now, 'completed_at': now, 'last_error': False, 'state': 'done', 'case_id': case.id})
        return True


class LogisticsEvidence(models.Model):
    _name = 'logistics.idp.evidence'
    _description = 'Logistics IDP Evidence Snapshot'
    _inherit = 'logistics.idp.immutable.snapshot'
    _order = 'create_date desc, id desc'

    case_id = fields.Many2one('logistics.idp.case', required=True, readonly=True, ondelete='restrict', index=True)
    category = fields.Char(required=True, readonly=True, index=True)
    source_reference = fields.Char(required=True, readonly=True, index=True)
    status = fields.Selection([
        ('valid', 'Valid'),
        ('review', 'Review'),
        ('invalid', 'Invalid'),
    ], required=True, default='review', readonly=True)

    _evidence_dedupe_unique = models.Constraint(
        'unique(case_id, category, payload_hash)',
        'This exact evidence snapshot already exists for the case.')

    @api.model
    def _controlled_create(self, values, service):
        vals_list = values if isinstance(values, list) else [values]
        for vals in vals_list:
            payload = json.loads(vals['payload']) if isinstance(vals.get('payload'), str) else vals.get('payload')
            payload_hash = hashlib.sha256(_canonical_json(payload).encode()).hexdigest()
            if self.search_count([('case_id', '=', vals.get('case_id')), ('category', '=', vals.get('category')), ('payload_hash', '=', payload_hash)]):
                raise ValidationError('Exact evidence replay is not permitted.')
        return super()._controlled_create(values, service)

    @api.model_create_multi
    def create(self, vals_list):
        unknown = sorted({vals.get('category') for vals in vals_list} - EVIDENCE_CATEGORIES)
        if unknown:
            raise ValidationError('Unknown evidence category: %s' % ', '.join(filter(None, unknown)))
        return super().create(vals_list)


class LogisticsPolicyDecision(models.Model):
    _name = 'logistics.idp.policy.decision'
    _description = 'Logistics IDP Policy Decision Snapshot'
    _inherit = 'logistics.idp.immutable.snapshot'
    _order = 'effective_from desc, id desc'

    case_id = fields.Many2one('logistics.idp.case', required=True, readonly=True, ondelete='restrict', index=True)
    supplier_profile_id = fields.Many2one('logistics.idp.supplier.profile', readonly=True, ondelete='restrict')
    policy_code = fields.Char(required=True, readonly=True)
    policy_version = fields.Char(required=True, readonly=True)
    effective_from = fields.Datetime(required=True, readonly=True)
    effective_to = fields.Datetime(readonly=True)
    verdict = fields.Selection(VERDICTS, required=True, default='review', readonly=True)
    reason = fields.Text(required=True, readonly=True)

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            payload = json.loads(vals['payload']) if isinstance(vals['payload'], str) else vals['payload']
            manifest = payload.get('input_manifest', {})
            case = self.env['logistics.idp.case'].browse(vals['case_id']).exists()
            if manifest.get('kind') == 'policy_activation_reevaluation':
                activation = self.env['logistics.idp.policy.activation'].browse(
                    manifest.get('activation_id')).exists()
                policy = self.env['logistics.idp.policy.source'].browse(manifest.get('policy_id')).exists()
                if (not case or not activation or not policy or activation.status != 'activated'
                        or activation.event_uuid != manifest.get('activation_uuid')
                        or policy != activation.policy_source_id
                        or case.company_id != activation.company_id
                        or manifest.get('case_id') != case.id
                        or manifest.get('case_effective_date') != fields.Date.to_string(case.effective_date)
                        or vals.get('audit_input_hash') != payload.get('binding_hash')
                        or payload.get('binding_hash') != hashlib.sha256(_canonical_json(manifest).encode()).hexdigest()):
                    raise ValidationError('Policy decision input manifest binding is invalid.')
                continue
            profile = self.env['logistics.idp.supplier.profile'].browse(vals.get('supplier_profile_id')).exists()
            if (not case or (profile and profile.company_id != case.company_id)
                    or manifest.get('case_id') != case.id
                    or manifest.get('case_effective_date') != fields.Date.to_string(case.effective_date)
                    or manifest.get('supplier_profile') != (profile and {
                        'id': profile.id, 'profile_code': profile.profile_code, 'version': profile.version,
                        'payload_hash': profile.payload_hash,
                        'effective_from': fields.Date.to_string(profile.effective_from),
                        'effective_to': fields.Date.to_string(profile.effective_to) if profile.effective_to else False,
                    } or False) or vals.get('audit_input_hash') != payload.get('binding_hash')
                    or payload.get('binding_hash') != hashlib.sha256(_canonical_json(manifest).encode()).hexdigest()):
                raise ValidationError('Policy decision input manifest binding is invalid.')
        return super().create(vals_list)

    @api.constrains('effective_from', 'effective_to')
    def _check_effective_dates(self):
        for decision in self:
            if decision.effective_to and decision.effective_to <= decision.effective_from:
                raise ValidationError('Effective To must be later than Effective From.')


class LogisticsMesReference(models.Model):
    _name = 'logistics.idp.mes.reference'
    _description = 'Inbound MES Reference Snapshot'
    _inherit = 'logistics.idp.immutable.snapshot'
    _order = 'imported_at desc, id desc'

    case_id = fields.Many2one('logistics.idp.case', required=True, readonly=True, ondelete='restrict', index=True)
    profile_code = fields.Char(required=True, readonly=True)
    adapter_type = fields.Selection([('generic_inbound', 'Generic Inbound')], required=True, default='generic_inbound', readonly=True)
    source_system = fields.Char(required=True, readonly=True)
    source_key = fields.Char(required=True, readonly=True)
    source_version = fields.Char(required=True, readonly=True)
    imported_at = fields.Datetime(required=True, readonly=True, default=fields.Datetime.now)
    provenance = fields.Char(required=True, readonly=True)

    _mes_identity_unique = models.Constraint(
        'unique(source_system, source_key, source_version)',
        'This MES source version was already imported.')


class LogisticsAttachment(models.Model):
    _inherit = 'ir.attachment'

    @api.model
    def get_file_viewer_data(self, attachment_ids):
        attachments = self.browse(attachment_ids).exists()
        attachments.check_access('read')
        store = Store()
        store.add(attachments)
        return store.get_result().get('ir.attachment', [])


class LogisticsDocument(models.Model):
    _name = 'logistics.idp.document'
    _description = 'Canonical Logistics Document'
    _inherit = ['extract.mixin.with.words']
    _order = 'create_date desc, id desc'

    case_id = fields.Many2one('logistics.idp.case', string='Logistics Case', required=True, ondelete='restrict', index=True)
    company_id = fields.Many2one(related='case_id.company_id', string='Operating Company', store=True, index=True)
    attachment_id = fields.Many2one('ir.attachment', string='Document File Attachment', ondelete='restrict')
    document_id = fields.Many2one('documents.document', string='Workspace Document', ondelete='restrict')
    document_type = fields.Selection([(key, key.replace('_', ' ').title()) for key in (
        'unknown', 'purchase_order', 'po_snapshot', 'draft_vat_invoice', 'main_vat_invoice', 'sales_invoice',
        'commercial_invoice', 'invoice', 'packing_list', 'warehouse_release', 'export_declaration_draft', 'export_declaration_final',
        'import_declaration', 'customs_declaration', 'master_data', 'dsnavl', 'sap_erp_output', 'e11', 'e13', 'e15',
        'manifest', 'bill_of_lading', 'broker_artifact', 'shipping_plan', 'gate_pass')], string='Classified Document Type', required=True, default='unknown', index=True)
    version = fields.Integer(string='Document Version', required=True, default=1, readonly=True)
    supersedes_id = fields.Many2one('logistics.idp.document', string='Superseded Document', readonly=True, ondelete='restrict')
    content_hash = fields.Char(string='Document SHA256 Hash', required=True, readonly=True, index=True)
    mimetype = fields.Char(string='MIME Content Type', required=True, readonly=True)
    page_count = fields.Integer(string='Page Count', default=1, readonly=True)
    source_channel = fields.Selection([('upload', 'Upload'), ('documents', 'Documents'), ('email', 'Email'), ('queue', 'Queue')], string='Ingestion Channel', required=True, default='upload', readonly=True)
    source_sender = fields.Char(string='Source Sender Email', readonly=True)
    source_message_reference = fields.Char(string='Source Message Reference', readonly=True, index=True)
    action_required = fields.Char(string='Action Required', readonly=True)
    predicted_supplier = fields.Char(string='AI Predicted Supplier', compute='_compute_inbox_projections', readonly=True)
    related_case_candidate_id = fields.Many2one('logistics.idp.case', string='Suggested Logistics Case', compute='_compute_inbox_projections', readonly=True)
    inbox_action_required = fields.Char(string='Inbox Action Status', compute='_compute_inbox_projections', readonly=True)
    status = fields.Selection([('processing', 'Processing'), ('valid', 'Valid'), ('review', 'Review'), ('error', 'Error')], string='Processing Status', required=True, default='processing', readonly=True)
    classification_confidence = fields.Float(string='AI Classification Confidence', readonly=True)
    current_run_id = fields.Many2one('logistics.idp.extraction.run', string='Active Extraction Run', readonly=True, ondelete='restrict')
    extraction_run_ids = fields.One2many('logistics.idp.extraction.run', 'document_id', string='Extraction Run History', readonly=True)
    extraction_payload = fields.Text(string='Full Extraction JSON Payload', compute='_compute_processed_viewer', readonly=True)
    normalized_fields = fields.Text(string='Normalized Business Fields JSON', compute='_compute_processed_viewer', readonly=True)
    source_spans = fields.Text(string='OCR Bounding Box Source Spans', compute='_compute_processed_viewer', readonly=True)
    extraction_template_version = fields.Char(string='Extraction Template Version', compute='_compute_processed_viewer', readonly=True)
    extraction_template_hash = fields.Char(string='Extraction Template Hash', compute='_compute_processed_viewer', readonly=True)
    corrected_supplier = fields.Char(string='Human Corrected Supplier')
    corrected_invoice_number = fields.Char(string='Human Corrected Invoice #')
    corrected_total_gross = fields.Float(string='Human Corrected Total Gross')
    logical_parent_id = fields.Many2one('logistics.idp.document', string='Logical Parent Document', readonly=True, ondelete='restrict', index=True)
    duplicate_of_id = fields.Many2one('logistics.idp.document', string='Duplicate of Document', readonly=True, ondelete='restrict', index=True)
    duplicate_marked_at = fields.Datetime(string='Duplicate Marked At', readonly=True)
    duplicate_marked_by_id = fields.Many2one('res.users', string='Duplicate Marked By', readonly=True, ondelete='restrict')
    logical_index = fields.Integer(string='Logical Document Index', readonly=True)
    logical_source_identity = fields.Char(string='Logical Source Identity', readonly=True, index=True)

    @api.depends('current_run_id', 'current_run_id.payload', 'source_message_reference', 'action_required', 'company_id', 'case_id')
    def _compute_inbox_projections(self):
        case_model = self.env['logistics.idp.case']
        for document in self:
            payload = json.loads(document.current_run_id.payload) if document.current_run_id else {}
            normalized = payload.get('payload', {}) if isinstance(payload, dict) else {}
            document.predicted_supplier = normalized.get('supplier') if isinstance(normalized, dict) else False
            signals = []
            if isinstance(normalized, dict) and normalized.get('po_reference'):
                signals.append(('po_reference', '=', str(normalized['po_reference']).strip()))
            if document.source_message_reference:
                signals.append(('thread_reference', '=', document.source_message_reference.strip()))
            candidates = case_model
            if signals:
                domain = [('company_id', '=', document.company_id.id), ('id', '!=', document.case_id.id)]
                if len(signals) == 1:
                    domain.append(signals[0])
                else:
                    domain += ['|'] + signals
                candidates = case_model.search(domain, limit=2)
            document.related_case_candidate_id = candidates if len(candidates) == 1 else False
            review_signal = 'Review / assign case' if len(candidates) != 1 else False
            document.inbox_action_required = ' / '.join(filter(None, [document.action_required, review_signal])) or False

    @api.depends('current_run_id.payload', 'attachment_id')
    def _get_mail_thread_data_attachments(self):
        return super()._get_mail_thread_data_attachments() | self.attachment_id

    def _compute_processed_viewer(self):
        for document in self:
            payload = json.loads(document.current_run_id.payload) if document.current_run_id else {}
            normalized = payload.get('payload', {})
            policy = payload.get('policy', {})
            template = payload.get('extraction_template', {})
            document.extraction_payload = json.dumps(payload, indent=2, ensure_ascii=False)
            document.normalized_fields = json.dumps(normalized, indent=2, ensure_ascii=False)
            document.source_spans = json.dumps(payload.get('source_spans', []), indent=2, ensure_ascii=False)
            document.extraction_template_version = (document.current_run_id.template_version or template.get('version')
                                                    or payload.get('template_version') or policy.get('version'))
            document.extraction_template_hash = (document.current_run_id.template_hash or template.get('hash')
                                                 or payload.get('template_hash') or policy.get('effective_hash')
                                                 or policy.get('hash'))
            if document.attachment_id and not document.message_main_attachment_id:
                document.message_main_attachment_id = document.attachment_id
            if document.attachment_id and not document.extract_attachment_id:
                document.extract_attachment_id = document.attachment_id
            if not document.extract_state:
                document.extract_state = 'waiting_validation'

    # Structured Extracted Fields for Direct Mapping View
    extracted_doc_number = fields.Char(string='Extracted Document #', tracking=True)
    extracted_doc_date = fields.Date(string='Extracted Document Date', tracking=True)
    extracted_partner_name = fields.Char(string='Extracted Supplier', tracking=True)
    extracted_buyer_name = fields.Char(string='Extracted Buyer')
    extracted_currency_id = fields.Many2one('res.currency', string='Extracted Currency')
    extracted_incoterm = fields.Char(string='Extracted Incoterms')
    extracted_amount_untaxed = fields.Float(string='Amount Untaxed (Net)')
    extracted_amount_tax = fields.Float(string='Tax Amount')
    extracted_amount_total = fields.Float(string='Total Gross Amount', tracking=True)

    extracted_line_ids = fields.One2many('logistics.idp.extracted.line', 'document_id', string='Extracted Line Items')
    extracted_lines_count = fields.Integer(string='Line Items Count', compute='_compute_extracted_lines_count', store=True)

    discrepancy_status = fields.Selection([
        ('matched', 'Matched / Verified (Khớp hoàn toàn)'),
        ('warning', 'Discrepancy Warning (Cảnh báo lệch số liệu)'),
        ('unverified', 'Unverified / Pending Review (Chờ đối soát)'),
    ], string='Audit Verification Status', default='unverified', tracking=True)
    discrepancy_details = fields.Text(string='Discrepancy Audit Details')

    @api.depends('extracted_line_ids')
    def _compute_extracted_lines_count(self):
        for doc in self:
            doc.extracted_lines_count = len(doc.extracted_line_ids)

    def action_approve_extraction(self):
        for doc in self:
            doc.status = 'valid'
            doc.discrepancy_status = 'matched'
            doc.extracted_line_ids.write({'is_verified': True})

    def action_recompute_line_discrepancies(self):
        for doc in self:
            line_sum = sum(doc.extracted_line_ids.mapped('subtotal'))
            if doc.extracted_amount_total and abs(line_sum - doc.extracted_amount_total) > 0.05:
                doc.discrepancy_status = 'warning'
                doc.discrepancy_details = f"Tổng tiền các dòng hàng ({line_sum:,.2f}) lệch so với Tổng hóa đơn ({doc.extracted_amount_total:,.2f})."
            else:
                doc.discrepancy_status = 'matched'
                doc.discrepancy_details = "Số liệu các dòng hàng khớp chính xác với Tổng tiền hóa đơn."

    def _external_action_guard(self, action):
        self.ensure_one()
        if not self.env.user.has_group('insilos_logistics_idp.group_logistics_reviewer') and not self.env.user.has_group('insilos_logistics_idp.group_logistics_manager'):
            raise AccessError(_('A Logistics reviewer must approve external actions.'))
        if self.status != 'valid' or not self.env.context.get('approved_logistics_action'):
            raise UserError(_('Approved extraction and explicit approval are required.'))
        correlation = self.env.context.get('logistics_correlation_id')
        if not correlation:
            raise UserError(_('A correlation ID is required.'))
        evidence = self.env['logistics.idp.evidence'].search([
            ('case_id', '=', self.case_id.id), ('category', '=', 'external_action'),
            ('source_reference', '=', '%s:%s' % (action, correlation)),
        ], limit=1)
        return correlation, evidence

    def _audit_external_action(self, action, correlation, outcome):
        payload = {'action': action, 'correlation_id': correlation, 'document_id': self.id,
                   'company_id': self.company_id.id, 'outcome': outcome}
        return self.env['logistics.idp.evidence']._controlled_create({
            'case_id': self.case_id.id, 'category': 'external_action',
            'source_reference': '%s:%s' % (action, correlation), 'status': 'valid',
            'payload': payload, 'audit_input_hash': hashlib.sha256(_canonical_json(payload).encode()).hexdigest(),
        }, 'logistics_external_action')

    def action_create_vendor_bill(self):
        raise UserError(_(
            'Vendor bill creation is unavailable until an immutable independent approval '
            'and accounting-controlled posting capability exist.'
        ))

    def action_create_customs_dossier(self):
        raise UserError(_(
            'NSW dossier creation is unavailable until an immutable independent approval, governed '
            'source evidence, and authority-controlled submission capability exist.'
        ))

    _exact_hash_unique = models.Constraint('unique(company_id, content_hash, logical_index)', 'This exact document already exists.')

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            if vals.get('attachment_id'):
                vals.setdefault('message_main_attachment_id', vals['attachment_id'])
                vals.setdefault('extract_attachment_id', vals['attachment_id'])
        return super().create(vals_list)

    def write(self, vals):
        if 'case_id' in vals:
            raise UserError('Documents may only be reassigned through action_assign_case.')
        if {'duplicate_of_id', 'duplicate_marked_at', 'duplicate_marked_by_id'} & set(vals) and self.env.context.get(
                '_logistics_document_duplicate_token') is not _INTERNAL_DOCUMENT_DUPLICATE_TOKEN:
            raise UserError('Duplicate markers may only be set through action_mark_duplicate.')
        if ('document_type' in vals and not self.env.su
                and self.env.context.get('_logistics_document_type_token') is not _INTERNAL_DOCUMENT_TYPE_TOKEN):
            raise UserError('Documents may only be reclassified through action_reclassify.')
        if {'extracted_words', 'extracted_numbers', 'extracted_dates'} & set(vals) and self.env.context.get(
                '_logistics_ocr_geometry_token') is not _INTERNAL_OCR_GEOMETRY_TOKEN:
            raise UserError('OCR geometry may only be persisted by IAP processing.')
        correction_fields = ('corrected_supplier', 'corrected_invoice_number', 'corrected_total_gross')
        before = {
            document.id: {field: document[field] for field in correction_fields if field in vals}
            for document in self
        }
        result = super().write(vals)
        for document in self:
            after = {field: document[field] for field in before[document.id]}
            changed_before = {field: value for field, value in before[document.id].items()
                              if value != after[field]}
            if not changed_before:
                continue
            changed_after = {field: after[field] for field in changed_before}
            run = document.current_run_id
            payload = {
                'document_id': document.id,
                'before': changed_before,
                'after': changed_after,
                'run_id': run.id,
                'run_hash': run.payload_hash or False,
            }
            identity = hashlib.sha256(_canonical_json(payload).encode()).hexdigest()
            self.env['logistics.idp.evidence']._controlled_create({
                'case_id': document.case_id.id, 'category': 'ocr_manual_correction',
                'source_reference': 'correction:%s:%s' % (document.id, identity),
                'status': 'review', 'payload': payload,
            }, 'ocr_manual_correction')
        return result

    def _sync_review_activity(self):
        activity_type = self.env.ref('mail.mail_activity_data_todo')
        for document in self:
            case = document.case_id.sudo()
            case._ensure_review_owner()
            prefix = 'IDP document %s: ' % document.id
            activities = case.activity_ids.filtered(
                lambda activity: activity.activity_type_id == activity_type and (activity.summary or '').startswith(prefix))
            if document.status == 'valid':
                activities.action_feedback(feedback='Extraction review passed.')
                continue
            summary = prefix + ('remediation' if document.status == 'error' else 'extraction review')
            keeper = activities[:1]
            if keeper:
                keeper.write({'summary': summary, 'user_id': case.owner_id.id})
                activities[1:].unlink()
            elif case.owner_id:
                case.with_context(mail_activity_quick_update=True).activity_schedule(
                    'mail.mail_activity_data_todo', user_id=case.owner_id.id, summary=summary)
            if not case._authorized_reviewers() and case.owner_id:
                exception = self.env['logistics.idp.exception'].sudo().search([
                    ('case_id', '=', case.id), ('exception_type', '=', 'policy'),
                    ('state', 'in', ('open', 'waiting')), ('message_ids.body', 'ilike', '[reviewer_unavailable]'),
                ], limit=1)
                if not exception:
                    exception = self.env['logistics.idp.exception'].sudo().create({
                        'case_id': case.id, 'exception_type': 'policy', 'severity': 'medium'})
                    exception.message_post(body='[reviewer_unavailable] No active human Logistics reviewer is available.')
        return True

    @api.model
    def _backfill_review_workflow(self):
        documents = self.sudo().search([
            ('case_id.state', '=', 'review'), '|', ('case_id.owner_id.active', '=', False),
            ('status', 'in', ('review', 'error')),
        ])
        documents._sync_review_activity()
        return True

    @api.model
    def _intake_synthetic_content(self, case, content, mimetype, metadata=None, source_channel='upload', process=True):
        if not tools.config['test_enable']:
            raise UserError('Synthetic document intake is only available in test mode.')
        metadata = dict(metadata or {})
        metadata['document_type_hint'] = metadata.get('document_type')
        return self.intake_content(case, content, mimetype, metadata, source_channel, process,
                                   _processor=SyntheticDocumentProcessor())

    def _reprocess_synthetic(self, metadata=None):
        if not tools.config['test_enable']:
            raise UserError('Synthetic document intake is only available in test mode.')
        return self.process(metadata, explicit=True, _processor=SyntheticDocumentProcessor())

    @api.model
    def upload_intake(self, case, content, mimetype, filename):
        case = self.env['logistics.idp.case'].browse(case).exists() if isinstance(case, int) else case.exists()
        if len(case) != 1 or case.company_id not in self.env.companies:
            raise ValidationError('Upload requires one allowed case.')
        if case.state in _CASE_TERMINAL_STATES:
            raise ValidationError('Documents cannot be uploaded to terminal cases.')
        if not isinstance(content, bytes) or not content:
            raise ValidationError('Document content must be non-empty bytes.')
        if len(content) > MAX_BYTES:
            raise ValidationError('Document exceeds the maximum supported size.')
        return self.intake_content(case, content, mimetype, {'filename': filename}, process=False)

    @api.model
    def upload_final_customs_document(self, case, content, mimetype, filename):
        if not (self.env.su or self.env.user.has_group('insilos_logistics_idp.group_logistics_operator')
                or self.env.user.has_group('insilos_logistics_idp.group_logistics_manager')):
            raise UserError('Only Logistics operators or managers may upload final customs documents.')
        case = self.env['logistics.idp.case'].browse(case).exists() if isinstance(case, int) else case.exists()
        if len(case) != 1:
            raise ValidationError('Upload requires one allowed case.')
        case.check_access('read')
        document = self.upload_intake(case, content, mimetype, filename)
        if document.case_id != case:
            raise ValidationError('Supplemental document content is already linked to another case.')
        return document.action_reclassify('export_declaration_final')

    @api.model
    def intake_content(self, case, content, mimetype, metadata=None, source_channel='upload', process=True,
                       _processor=None):
        if not isinstance(content, bytes) or not content:
            raise ValidationError('Document content must be non-empty bytes.')
        digest = hashlib.sha256(content).hexdigest()
        existing = self.search([('company_id', '=', case.company_id.id), ('content_hash', '=', digest)], limit=1)
        if existing:
            if existing.case_id != case:
                payload = {'canonical_case_id': existing.case_id.id, 'canonical_document_id': existing.id,
                           'content_hash': digest, 'disposition': 'excluded_from_aggregation'}
                if not case.evidence_ids.filtered(lambda item: item.category == 'exact_document_duplicate'
                                                  and item.source_reference == 'sha256:%s' % digest):
                    self.env['logistics.idp.evidence']._controlled_create({
                        'case_id': case.id, 'category': 'exact_document_duplicate',
                        'source_reference': 'sha256:%s' % digest, 'status': 'valid', 'payload': payload,
                    }, 'exact_duplicate_disposition')
                case.message_post(body='Exact document duplicate reused within company: %s.' % digest)
            return existing
        name = (metadata or {}).get('filename') or 'logistics-%s' % digest[:12]
        attachment = self.env['ir.attachment'].sudo().create({
            'name': name, 'raw': content, 'mimetype': mimetype,
            'res_model': case._name, 'res_id': case.id, 'company_id': case.company_id.id,
        })
        folder = case.team_id.documents_folder_id
        if folder.company_id != case.company_id:
            folder = self.env['documents.document']
        bound = self.env['documents.document'].sudo().create({
            'name': name, 'attachment_id': attachment.id,
            'company_id': case.company_id.id, 'folder_id': folder.id, 'owner_id': case.owner_id.id,
        })
        metadata = metadata or {}
        hint = metadata.get('document_type_hint')
        document = self.create({
            'case_id': case.id, 'attachment_id': attachment.id, 'document_id': bound.id,
            'message_main_attachment_id': attachment.id, 'extract_attachment_id': attachment.id,
            'extract_state': 'waiting_validation',
            'content_hash': digest, 'mimetype': mimetype, 'source_channel': source_channel,
            'source_sender': metadata.get('source_sender'),
            'source_message_reference': metadata.get('source_message_reference'),
            'action_required': metadata.get('action_required'), 'status': metadata.get('status', 'processing'),
            'document_type': hint if hint in DOCUMENT_TYPES else 'unknown',
            'classification_confidence': metadata.get('classification_confidence', 0),
        })
        if not process and metadata.get('status'):
            document.write({'status': metadata['status']})
        return document.process(metadata, _processor=_processor) if process else document

    def process(self, metadata=None, run_id=None, explicit=False, _processor=None):
        self.ensure_one()
        if not self.env.user.has_group('insilos_logistics_idp.group_logistics_operator') and not self.env.user.has_group('insilos_logistics_idp.group_logistics_reviewer') and not self.env.user.has_group('insilos_logistics_idp.group_logistics_manager'):
            raise UserError('Only Logistics operators, reviewers or managers may process documents.')
        self.env.cr.execute('SELECT id FROM logistics_idp_document WHERE id = %s FOR UPDATE', [self.id])
        self.invalidate_recordset(['status', 'current_run_id'])
        run = self.env['logistics.idp.extraction.run'].browse(run_id).exists() if run_id else self.current_run_id
        if run and run.document_id != self:
            run = self.env['logistics.idp.extraction.run']
        if run_id and run and run.status != 'pending':
            return self
        if explicit and not run_id:
            run = self.env['logistics.idp.extraction.run']
        elif run and run.status != 'pending' and not run_id:
            run = self.env['logistics.idp.extraction.run']
        if not explicit and not run and self.status != 'processing':
            return self
        started = run.started_at if run else fields.Datetime.now()
        content = self.attachment_id.raw
        try:
            if hashlib.sha256(content).hexdigest() != self.content_hash:
                raise ProcessingError('stored binary hash mismatch')
            SyntheticDocumentProcessor()._validate_binary(content, self.mimetype)
        except ProcessingError as exc:
            category = {
                'malformed PDF structure': 'invalid_pdf_structure',
                'encrypted PDF': 'encrypted_document',
                'unsafe PDF active content': 'unsafe_active_content',
                'unsafe XLSX embedded content': 'unsafe_embedded_content',
                'unsafe EML embedded content': 'unsafe_embedded_content',
                'unsupported MIME': 'unsupported_mime',
                'invalid document size': 'invalid_document_size',
            }.get(str(exc), 'invalid_document_structure')
            run = self.env['logistics.idp.extraction.run'].with_context(
                _logistics_snapshot_token=_INTERNAL_SNAPSHOT_TOKEN).create({
                    'case_id': self.case_id.id, 'document_id': self.id,
                    'provider': 'local-preflight', 'model_version': 'stdlib-v1',
                    'schema_version': IAPDocumentProcessor.schema_version,
                    'prompt_version': 'none', 'template_version': IAPDocumentProcessor.prompt_version,
                    'started_at': started, 'completed_at': fields.Datetime.now(), 'status': 'error',
                    'error': category,
                    'payload': {
                        'error_category': category,
                        'human_action': 'Replace or re-upload the source document.',
                    },
                })
            self._create_extraction_critical_check(run)
            self.write({'status': 'error', 'current_run_id': run.id})
            self.case_id.sudo().write({'document_status': 'review', 'next_action': 'Replace or re-upload the source document.'})
            self.case_id.sudo()._derive_lifecycle()
            self._sync_review_activity()
            return self
        metadata = dict(metadata or {})
        metadata.setdefault('filename', self.attachment_id.name)
        metadata.setdefault('source_channel', self.source_channel)
        metadata.setdefault('case_po_reference', self.case_id.po_reference)
        metadata.setdefault('has_po_reference', bool(self.case_id.po_reference))
        metadata.setdefault('batch_manifest', [
            {'filename': sibling.attachment_id.name, 'document_type': sibling.document_type}
            for sibling in self.case_id.document_ids if sibling != self
        ])
        jurisdiction = str(metadata.get('jurisdiction') or 'VN').upper()
        regime = str(metadata.get('regime') or metadata.get('customs_regime') or 'ALL').upper()
        effective_date = self.case_id.effective_date or fields.Date.today()
        policy_model = self.env['logistics.idp.policy.source'].sudo()
        try:
            policy = policy_model.select_effective_pack(
                self.company_id, 'trade_compliance', None, jurisdiction, regime, effective_date)
            if policy:
                policy.ensure_one()
                metadata['idp_policy'] = policy_model.validate_policy_payload(
                    policy.payload, 'trade_compliance', None, jurisdiction, regime)
        except ValidationError as exc:
            policy = self.env['logistics.idp.policy.source']
            metadata['idp_policy_error'] = str(exc)
        if policy:
            metadata['idp_policy_ref'] = {
                'code': policy.code, 'version': policy.version, 'hash': policy.payload_hash,
            }
            if self.case_id.supplier_reference:
                supplier_profile = self.env['logistics.idp.supplier.profile'].sudo().select_effective(
                    self.company_id, self.case_id.supplier_reference, effective_date)
                if supplier_profile:
                    overlay = json.loads(supplier_profile.payload)
                    merged = json.loads(json.dumps(metadata['idp_policy']))
                    horizontal = merged['horizontal']
                    pack_config = merged['vertical']['packs'][merged['vertical']['pack']]
                    if isinstance(overlay.get('classification_thresholds'), dict):
                        horizontal['classification']['thresholds'].update(overlay['classification_thresholds'])
                    if isinstance(overlay.get('classification_aliases'), dict):
                        pack_config.setdefault('classification', {}).setdefault('aliases', {}).update(
                            overlay['classification_aliases'])
                    if isinstance(overlay.get('batch_expected_roles'), list):
                        pack_config['batch']['expected_roles'] = overlay['batch_expected_roles']
                    metadata['idp_policy'] = policy_model.validate_policy_payload(
                        merged, 'trade_compliance', None, jurisdiction, regime)
                    metadata['idp_policy_ref']['supplier_profile'] = {
                        'code': supplier_profile.profile_code, 'version': supplier_profile.version,
                        'hash': supplier_profile.payload_hash,
                    }
            metadata['idp_policy_ref']['effective_hash'] = hashlib.sha256(
                _canonical_json(metadata['idp_policy']).encode()).hexdigest()
        use_local = bool(_processor) or self.mimetype in (
            'text/csv', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            'message/rfc822', 'application/json')
        policy_ref = metadata.get('idp_policy_ref', {})
        policy_operation = '%s:%s' % (
            policy_ref.get('version', 'fallback'), policy_ref.get('effective_hash', policy_ref.get('hash', 'none'))[:12])
        if not run:
            run = self.env['logistics.idp.extraction.run'].with_context(
                _logistics_snapshot_token=_INTERNAL_SNAPSHOT_TOKEN).create({
                    'case_id': self.case_id.id, 'document_id': self.id,
                    'provider': 'local-deterministic' if use_local else IAPDocumentProcessor.provider,
                    'model_version': SyntheticDocumentProcessor.model_version if use_local else IAPDocumentProcessor.model_version,
                    'schema_version': IAPDocumentProcessor.schema_version,
                    'prompt_version': IAPDocumentProcessor.prompt_version,
                    'template_version': IAPDocumentProcessor.prompt_version, 'started_at': started,
                    'status': 'pending', 'payload': {'pending': True},
                })
            self.write({'status': 'processing', 'current_run_id': run.id})
        pending_payload = json.loads(run.payload)
        operation_id = pending_payload.get('operation_id')
        credit_budget = pending_payload.get('credit_budget')
        if not use_local and credit_budget is None:
            configured_budget = self.env['ir.config_parameter'].sudo().get_param(
                'logistics_idp.ocr_credit_budget', '1')
            try:
                credit_budget = float(configured_budget)
            except (TypeError, ValueError) as exc:
                raise UserError(_('Logistics OCR credit budget must be numeric and greater than zero.')) from exc
            if credit_budget <= 0:
                raise UserError(_('Logistics OCR credit budget must be numeric and greater than zero.'))
        configured_run_max = self.env['ir.config_parameter'].sudo().get_param(
            'logistics_idp.ocr_run_max_credits', '3')
        try:
            run_max_credits = float(configured_run_max)
        except (TypeError, ValueError) as exc:
            raise UserError(_('Logistics OCR run max credits must be numeric and greater than zero.')) from exc
        if run_max_credits <= 0:
            raise UserError(_('Logistics OCR run max credits must be numeric and greater than zero.'))
        if not operation_id:
            operation_id = 'logistics-idp:policy:%s:document:%s:content:%s:schema:%s:prompt:%s:attempt:%s:run:%s' % (
                policy_operation, self.id, self.content_hash, IAPDocumentProcessor.schema_version,
                IAPDocumentProcessor.prompt_version, run.attempt_number, run.id)
        if not use_local and pending_payload.get('credit_budget') is None:
            pending_payload.update({
                'pending': True, 'operation_id': operation_id,
                'classification_operation_id': operation_id + ':classify',
                'classification_source': 'pending', 'credit_budget': credit_budget,
                'run_max_credits': run_max_credits,
                'planned_operation_budget': min(3, int(run_max_credits // credit_budget)),
                'credit_budget_provenance': {
                    'source': 'ir.config_parameter', 'key': 'logistics_idp.ocr_credit_budget',
                    'run_max_key': 'logistics_idp.ocr_run_max_credits',
                    'policy_version': policy_ref.get('version', 'fallback'),
                },
            })
            run.sudo().with_context(_logistics_snapshot_token=_INTERNAL_SNAPSHOT_TOKEN).write({
                'payload': pending_payload})
        def persist_classification(classification):
            pending_payload.update(classification)
            run.sudo().with_context(_logistics_snapshot_token=_INTERNAL_SNAPSHOT_TOKEN).write({
                'payload': pending_payload,
            })
            self.with_context(_logistics_document_type_token=_INTERNAL_DOCUMENT_TYPE_TOKEN).write({
                'document_type': classification['document_type'],
                'classification_confidence': classification['classification_confidence'],
            })
            self.env.cr.flush()

        processor = (_processor or SyntheticDocumentProcessor()) if use_local else IAPDocumentProcessor(
            self.env, operation_id, credit_budget, run, run_max_credits, persist_classification)
        status = 'review'
        result = {}
        error = False
        try:
            content = self.attachment_id.raw
            if hashlib.sha256(content).hexdigest() != self.content_hash:
                raise ProcessingError('stored binary hash mismatch')
            result = processor.process(content, self.mimetype, metadata)
            logical_documents = result.pop('logical_documents', None)
            if logical_documents is not None:
                for index, logical in enumerate(logical_documents, 1):
                    identity = hashlib.sha256(_canonical_json(logical).encode()).hexdigest()
                    child = self.search([('logical_parent_id', '=', self.id), ('logical_index', '=', index)], limit=1)
                    if not child:
                        child = self.create({
                            'case_id': self.case_id.id, 'attachment_id': self.attachment_id.id, 'document_id': self.document_id.id,
                            'message_main_attachment_id': self.attachment_id.id, 'extract_attachment_id': self.attachment_id.id,
                            'content_hash': self.content_hash, 'mimetype': self.mimetype, 'source_channel': self.source_channel,
                            'logical_parent_id': self.id, 'logical_index': index, 'logical_source_identity': identity,
                            'status': 'review',
                        })
                    child.process({'provider_response': logical, 'allow_review': True}, explicit=True, _processor=SyntheticDocumentProcessor())
                result = {'document_type': 'unknown', 'payload': {}, 'confidence': 0,
                          'classification_confidence': 0, 'validation_warnings': [], 'source_spans': []}
            if isinstance(processor, IAPDocumentProcessor) and result.get('ocr_results'):
                result['ocr_results'] = normalize_ocr_results(result['ocr_results'])
                self.with_context(_logistics_ocr_geometry_token=_INTERNAL_OCR_GEOMETRY_TOKEN)._on_ocr_results(
                    result['ocr_results'])
            result['policy'] = policy_ref
            extraction_policy = metadata.get('idp_policy', {}).get('horizontal', {}).get('extraction', {})
            acceptance = float(extraction_policy.get('review_thresholds', {}).get(
                result['document_type'], extraction_policy.get('acceptance_threshold', .8)))
            status = ('valid' if result['confidence'] >= acceptance and not result.get('validation_warnings')
                      else 'review')
        except Exception as exc:
            if isinstance(exc, DatabaseError):
                raise
            error = str(exc) if isinstance(exc, (ProcessingError, TimeoutError, ValueError)) else 'provider processing failed (%s)' % type(exc).__name__
            result = {
                'error': error, 'operation_id': operation_id,
                'classification_operation_id': operation_id + ':classify',
                'classification_source': 'failed_before_classification', 'policy': policy_ref,
            }
            result.update(getattr(processor, 'failure_context', {}))
            if getattr(exc, 'diagnostics', None):
                result['diagnostics'] = exc.diagnostics
        if not use_local:
            result.update({
                'credit_budget': credit_budget,
                'credit_budget_provenance': pending_payload['credit_budget_provenance'],
            })
        completed = fields.Datetime.now()
        result['effective_policy'] = metadata.get('idp_policy', {})
        if metadata.get('idp_policy_error'):
            result['policy_error'] = metadata['idp_policy_error']
        run.sudo().with_context(_logistics_snapshot_token=_INTERNAL_SNAPSHOT_TOKEN).write({
            'provider': processor.provider,
            'model_version': result.get('provider_model', processor.model_version),
            'completed_at': completed, 'duration_seconds': (completed - started).total_seconds(),
            'confidence': result.get('confidence', 0), 'status': 'error' if error else status,
            'error': error, 'payload': result or {
                'error': error, 'operation_id': operation_id,
                'classification_operation_id': operation_id + ':classify',
                'classification_source': 'failed_before_classification',
            },
        })
        self._create_extraction_critical_check(run)
        self.with_context(_logistics_document_type_token=_INTERNAL_DOCUMENT_TYPE_TOKEN).write({
            'document_type': result.get('document_type', 'unknown'), 'status': 'error' if error else status,
            'classification_confidence': result.get('classification_confidence', result.get('confidence', 0)),
            'current_run_id': run.id})
        self.invalidate_recordset(['predicted_supplier', 'related_case_candidate_id', 'inbox_action_required'])
        case = self.case_id.sudo()
        case.document_status = 'pass' if status == 'valid' and not error else 'review'
        case._derive_lifecycle()
        self._sync_review_activity()
        return self

    def _create_extraction_critical_check(self, run):
        if run.document_id != self or run.case_id != self.case_id or run.company_id != self.company_id:
            raise ValidationError('Extraction critical check requires the document canonical run and case.')
        payload = json.loads(run.payload)
        policy = payload.get('effective_policy')
        document_type = payload.get('document_type', 'unknown')
        extraction = policy.get('horizontal', {}).get('extraction', {}) if isinstance(policy, dict) else {}
        critical_fields = extraction.get('critical_fields', {}).get(document_type, [])
        policy_missing = not isinstance(policy, dict) or not critical_fields
        threshold = extraction.get('review_thresholds', {}).get(document_type, extraction.get('acceptance_threshold', .8))
        canonical = payload.get('payload', {})
        warnings = list(payload.get('validation_warnings', []))
        missing = []
        canonical_valid = isinstance(canonical, dict)
        def path_value(value, path):
            for part in path.split('.'):
                value = value.get(part) if isinstance(value, dict) else None
            return value
        for path in critical_fields:
            if not canonical_valid:
                missing.append(path)
            elif path.startswith('lines[].'):
                field = path.split('.', 1)[1]
                lines = canonical.get('lines')
                if (not isinstance(lines, list) or not lines
                        or any(path_value(line, field) in (None, '') for line in lines)):
                    missing.append(path)
            elif path_value(canonical, path) in (None, '', []):
                missing.append(path)
        classification = payload.get('classification_confidence', 0)
        ambiguous = document_type == 'unknown' or classification <= (policy or {}).get('horizontal', {}).get('classification', {}).get('thresholds', {}).get('ambiguous', 0)
        verdict = 'pass' if not (policy_missing or warnings or missing or ambiguous or payload.get('confidence', 0) < threshold) else 'review'
        check_payload = {
            'run_id': run.id, 'run_hash': run.payload_hash, 'canonical_input_hash': hashlib.sha256(_canonical_json(canonical).encode()).hexdigest(),
            'document_type': document_type, 'critical_fields': critical_fields, 'missing': sorted(set(missing)),
            'policy_missing': policy_missing, 'policy': payload.get('policy', {}),
            'validation_warnings': warnings, 'confidence': payload.get('confidence', 0), 'threshold': threshold,
            'classification_confidence': classification, 'classification_ambiguous': ambiguous,
        }
        audit_input_hash = hashlib.sha256(_canonical_json(check_payload).encode()).hexdigest()
        checks = self.env['logistics.idp.check.result'].sudo().search([('audit_input_hash', '=', audit_input_hash)], limit=1)
        if not checks:
            self.env['logistics.idp.check.result'].sudo()._create_from_check_runner({
                'case_id': self.case_id.id, 'run_id': run.id, 'code': 'EXTRACTION_CRITICAL_INPUTS', 'required': True,
                'verdict': verdict, 'rationale': ('Effective extraction policy is missing or unresolved.' if policy_missing else 'Canonical extraction critical-input validation.'),
                'payload': check_payload, 'audit_input_hash': audit_input_hash,
            }, 'extraction_critical_inputs')
        return True

    @api.model
    def _cron_extract_queued(self, limit=20):
        self.env.cr.execute("""SELECT id FROM logistics_idp_document WHERE status = 'processing' ORDER BY create_date, id FOR UPDATE SKIP LOCKED LIMIT %s""", [limit])
        for document in self.browse(row[0] for row in self.env.cr.fetchall()):
            document.process()
        return True

    def reprocess(self, metadata=None):
        return self.process(metadata, explicit=True)

    def action_reextract(self):
        return self.process(explicit=True)

    def action_mark_duplicate(self, canonical_document):
        self.ensure_one()
        if not (self.env.su or self.env.user.has_group('insilos_logistics_idp.group_logistics_operator')
                or self.env.user.has_group('insilos_logistics_idp.group_logistics_manager')):
            raise UserError('Only Logistics operators or managers may mark duplicate documents.')
        self.check_access('write')
        canonical_document = self.browse(canonical_document).exists() if isinstance(canonical_document, int) else canonical_document.exists()
        if len(canonical_document) != 1 or canonical_document == self:
            raise ValidationError('A distinct canonical document is required.')
        canonical_document.check_access('read')
        if canonical_document.company_id != self.company_id:
            raise ValidationError('Duplicate documents must belong to the same company.')
        if self.case_id.state in _CASE_TERMINAL_STATES or canonical_document.case_id.state in _CASE_TERMINAL_STATES:
            raise ValidationError('Duplicate documents cannot reference terminal cases.')
        if self.duplicate_of_id:
            if self.duplicate_of_id == canonical_document:
                return self
            raise ValidationError('Duplicate marker is immutable and cannot be remapped.')
        marked_at = fields.Datetime.now()
        self.with_context(_logistics_document_duplicate_token=_INTERNAL_DOCUMENT_DUPLICATE_TOKEN).write({
            'duplicate_of_id': canonical_document.id, 'duplicate_marked_at': marked_at,
            'duplicate_marked_by_id': self.env.user.id,
        })
        self.env['logistics.idp.evidence']._controlled_create({
            'case_id': self.case_id.id, 'category': 'document_duplicate_mark',
            'source_reference': 'document:%s:canonical:%s' % (self.id, canonical_document.id), 'status': 'valid',
            'payload': {'action': 'mark_duplicate', 'source_document_id': self.id,
                        'canonical_document_id': canonical_document.id, 'actor_id': self.env.user.id,
                        'marked_at': fields.Datetime.to_string(marked_at)},
        }, 'document_duplicate_mark')
        return self

    def action_open_duplicate_wizard(self):
        self.ensure_one()
        return {
            'type': 'ir.actions.act_window', 'name': 'Mark Duplicate Document',
            'res_model': 'logistics.idp.document.duplicate.wizard', 'view_mode': 'form', 'target': 'new',
            'context': {'default_document_id': self.id},
        }

    def action_open_upload_wizard(self):
        return {
            'type': 'ir.actions.act_window', 'name': 'Upload Document',
            'res_model': 'logistics.idp.document.upload.wizard', 'view_mode': 'form', 'target': 'new',
        }

    def action_open_inbox_wizard(self):
        self.ensure_one()
        return {
            'type': 'ir.actions.act_window', 'name': 'Assign / Reclassify Document',
            'res_model': 'logistics.idp.document.inbox.wizard', 'view_mode': 'form', 'target': 'new',
            'context': {
                'default_document_id': self.id, 'default_case_id': self.case_id.id,
                'default_document_type': self.document_type,
            },
        }

    def action_open_source_email(self):
        self.ensure_one()
        self.check_access('read')
        if not (self.env.su or self.env.user.has_group('insilos_logistics_idp.group_logistics_operator')
                or self.env.user.has_group('insilos_logistics_idp.group_logistics_reviewer')
                or self.env.user.has_group('insilos_logistics_idp.group_logistics_manager')):
            raise AccessError('Only Logistics users may open source email.')
        if self.source_channel != 'email' or not self.source_message_reference:
            raise ValidationError('Document has no persisted source email.')
        messages = self.env['mail.message'].sudo().search([
            ('message_id', '=', self.source_message_reference.strip()),
            ('record_company_id', '=', self.company_id.id), ('message_type', '=', 'email'),
        ], limit=2)
        if len(messages) != 1:
            raise ValidationError('Source email is missing or ambiguous.')
        return {
            'type': 'ir.actions.act_window', 'name': 'Source Email', 'res_model': 'mail.message',
            'res_id': messages.id, 'view_mode': 'form', 'views': [(False, 'form')],
            'context': {'create': False, 'edit': False, 'delete': False},
        }

    def action_reclassify(self, document_type=None):
        document_type = document_type or self.document_type
        if document_type not in DOCUMENT_TYPES:
            raise ValidationError('A supported document type is required.')
        result = self.process({'document_type': document_type, 'document_type_trusted': True}, explicit=True)
        if result.document_type != document_type:
            result.with_context(_logistics_document_type_token=_INTERNAL_DOCUMENT_TYPE_TOKEN).write({
                'document_type': document_type})
        return result

    def action_assign_case(self, case):
        self.ensure_one()
        if not (self.env.su or self.env.user.has_group('insilos_logistics_idp.group_logistics_operator')
                or self.env.user.has_group('insilos_logistics_idp.group_logistics_manager')):
            raise UserError('Only Logistics operators or managers may assign documents.')
        self.check_access('write')
        case = self.env['logistics.idp.case'].browse(case).exists() if isinstance(case, int) else case.exists()
        if len(case) != 1 or case.company_id != self.company_id:
            raise ValidationError('Document and target case must belong to the same company.')
        case.check_access('read')
        old_case = self.case_id
        if old_case.state in _CASE_TERMINAL_STATES or case.state in _CASE_TERMINAL_STATES:
            raise ValidationError('Documents cannot be reassigned from or to terminal cases.')
        if self.extraction_run_ids:
            raise ValidationError('Documents with extraction runs cannot be reassigned.')
        super(LogisticsDocument, self).write({'case_id': case.id})
        self.attachment_id.sudo().write({'res_model': case._name, 'res_id': case.id, 'company_id': case.company_id.id})
        payload = {'action': 'assign_case', 'document_id': self.id,
                   'from_case_id': old_case.id, 'to_case_id': case.id}
        self.env['logistics.idp.evidence']._controlled_create([{
            'case_id': old_case.id, 'category': 'document_assignment',
            'source_reference': 'document:%s:to:%s' % (self.id, case.id),
            'status': 'valid', 'payload': payload,
        }, {
            'case_id': case.id, 'category': 'document_assignment',
            'source_reference': 'document:%s:from:%s' % (self.id, old_case.id),
            'status': 'valid', 'payload': payload,
        }], 'document_assignment')
        old_case.sudo()._derive_lifecycle()
        case.sudo()._derive_lifecycle()
        return self


class LogisticsExtractedLine(models.Model):
    _name = 'logistics.idp.extracted.line'
    _description = 'Extracted Logistics Document Line Item'
    _order = 'document_id, sequence, id'

    document_id = fields.Many2one('logistics.idp.document', string='Document', required=True, ondelete='cascade', index=True)
    company_id = fields.Many2one(related='document_id.company_id', store=True, index=True)
    sequence = fields.Integer(string='Line #', default=10)

    description = fields.Char(string='Product Description', required=True)
    supplier_part_number = fields.Char(string='Supplier SKU')
    hs_code = fields.Char(string='HS Code (8-digit)', index=True)

    quantity = fields.Float(string='Quantity', default=1.0)
    uom_name = fields.Char(string='Unit of Measure', default='PCS')
    unit_price = fields.Float(string='Unit Price', default=0.0)
    subtotal = fields.Float(string='Line Subtotal', compute='_compute_subtotal', store=True, readonly=False)

    origin_country_id = fields.Many2one('res.country', string='Country of Origin')
    net_weight_kg = fields.Float(string='Net Weight (kg)')
    gross_weight_kg = fields.Float(string='Gross Weight (kg)')

    confidence = fields.Float(string='Confidence (%)', default=1.0)
    is_verified = fields.Boolean(string='Verified', default=False)
    notes = fields.Char(string='Discrepancy Notes')

    @api.depends('quantity', 'unit_price')
    def _compute_subtotal(self):
        for line in self:
            if not line.subtotal and line.quantity and line.unit_price:
                line.subtotal = round(line.quantity * line.unit_price, 2)


class LogisticsActiveThreadWizard(models.TransientModel):
    _name = 'logistics.idp.active.thread.wizard'
    _description = 'Select Active Supplier Email Thread'

    case_id = fields.Many2one('logistics.idp.case', required=True, readonly=True, ondelete='cascade')
    company_id = fields.Many2one(related='case_id.company_id', readonly=True)
    active_thread_case_id = fields.Many2one('logistics.idp.case', required=True, ondelete='restrict')

    def action_apply(self):
        self.ensure_one()
        self.case_id.action_select_active_thread(self.active_thread_case_id)
        return {'type': 'is.actions.act_window_close'}


class LogisticsDocumentDuplicateWizard(models.TransientModel):
    _name = 'logistics.idp.document.duplicate.wizard'
    _description = 'Mark Duplicate Logistics Document'

    document_id = fields.Many2one('logistics.idp.document', required=True, readonly=True, ondelete='cascade')
    company_id = fields.Many2one(related='document_id.company_id', readonly=True)
    canonical_document_id = fields.Many2one('logistics.idp.document', required=True, ondelete='restrict')

    def action_apply(self):
        self.ensure_one()
        self.document_id.action_mark_duplicate(self.canonical_document_id)
        return {'type': 'is.actions.act_window_close'}


class LogisticsDocumentUploadWizard(models.TransientModel):
    _name = 'logistics.idp.document.upload.wizard'
    _description = 'Logistics Document Upload Wizard'

    case_id = fields.Many2one('logistics.idp.case', ondelete='restrict')
    company_id = fields.Many2one('res.company', required=True, default=lambda self: self.env.company)
    final_customs_document = fields.Boolean()
    filename = fields.Char(required=True)
    content = fields.Binary(required=True, string='Document File')
    mimetype = fields.Char(default='application/pdf')

    @api.onchange('filename', 'content')
    def _onchange_content(self):
        if self.content:
            try:
                decoded = base64.b64decode(self.content, validate=True)
                self.mimetype = guess_mimetype(decoded, default='application/pdf')
            except Exception:
                pass

    def action_upload(self):
        self.ensure_one()
        try:
            content = base64.b64decode(self.content, validate=True)
        except (TypeError, ValueError) as exc:
            raise ValidationError('Document content must be valid base64.') from exc
        mimetype = self.mimetype or guess_mimetype(content, default='application/pdf')
        case = self.case_id
        if self.final_customs_document and not case:
            raise ValidationError('Final customs documents require an existing case.')
        if not case:
            digest = hashlib.sha256(content).hexdigest()
            case = self.env['logistics.idp.case'].intake({
                'name': self.filename, 'source_system': 'manual_upload', 'source_key': digest,
                'source_version': '1', 'provenance': 'manual_upload',
                'effective_date': fields.Date.today(), 'company_id': self.company_id.id,
            })
        if case.company_id != self.company_id:
            raise ValidationError('Upload case and company must match.')
        documents = self.env['logistics.idp.document']
        if self.final_customs_document:
            if not case:
                raise ValidationError('Final customs documents require an existing case.')
            documents.upload_final_customs_document(case, content, mimetype, self.filename)
        else:
            documents.upload_intake(case, content, mimetype, self.filename)
        return {'type': 'is.actions.act_window_close'}



class LogisticsDocumentInboxWizard(models.TransientModel):
    _name = 'logistics.idp.document.inbox.wizard'
    _description = 'Logistics Document Inbox Wizard'

    document_id = fields.Many2one('logistics.idp.document', required=True, readonly=True, ondelete='cascade')
    company_id = fields.Many2one(related='document_id.company_id', readonly=True)
    case_id = fields.Many2one('logistics.idp.case', required=True, ondelete='restrict')
    document_type = fields.Selection([(key, key.replace('_', ' ').title()) for key in (
        'unknown', 'purchase_order', 'po_snapshot', 'draft_vat_invoice', 'main_vat_invoice', 'sales_invoice',
        'commercial_invoice', 'invoice', 'packing_list', 'warehouse_release', 'export_declaration_draft',
        'export_declaration_final', 'import_declaration', 'customs_declaration', 'master_data', 'dsnavl',
        'sap_erp_output', 'e11', 'e13', 'e15', 'manifest', 'bill_of_lading', 'broker_artifact',
        'shipping_plan', 'gate_pass')], required=True)

    def action_apply(self):
        self.ensure_one()
        document = self.document_id
        if self.case_id != document.case_id:
            document.action_assign_case(self.case_id)
        if self.document_type != document.document_type:
            document.action_reclassify(self.document_type)
        return {'type': 'is.actions.act_window_close'}


class LogisticsExtractionRun(models.Model):
    _name = 'logistics.idp.extraction.run'
    _description = 'Immutable Extraction Run'
    _inherit = 'logistics.idp.immutable.snapshot'

    case_id = fields.Many2one('logistics.idp.case', required=True, readonly=True, ondelete='restrict', index=True)
    company_id = fields.Many2one('res.company', required=True, readonly=True, index=True)
    document_id = fields.Many2one('logistics.idp.document', required=True, readonly=True, ondelete='restrict', index=True)
    document_type = fields.Selection([(key, key.replace('_', ' ').title()) for key in (
        'unknown', 'purchase_order', 'po_snapshot', 'draft_vat_invoice', 'main_vat_invoice', 'sales_invoice',
        'commercial_invoice', 'invoice', 'packing_list', 'warehouse_release', 'export_declaration_draft', 'export_declaration_final',
        'import_declaration', 'customs_declaration', 'master_data', 'dsnavl', 'sap_erp_output', 'e11', 'e13', 'e15',
        'manifest', 'bill_of_lading', 'broker_artifact', 'shipping_plan', 'gate_pass')], required=True, readonly=True, index=True)
    page_count = fields.Integer(required=True, readonly=True)
    attempt_number = fields.Integer(required=True, readonly=True, index=True)
    retry_count = fields.Integer(required=True, readonly=True, default=0)
    confidence = fields.Float(readonly=True)
    duration_seconds = fields.Float(readonly=True)
    provider = fields.Char(required=True, readonly=True, index=True)
    model_version = fields.Char(required=True, readonly=True)
    schema_version = fields.Char(required=True, readonly=True)
    prompt_version = fields.Char(required=True, readonly=True)
    template_version = fields.Char(readonly=True)
    template_hash = fields.Char(readonly=True, index=True)
    started_at = fields.Datetime(required=True, readonly=True)
    completed_at = fields.Datetime(readonly=True)
    status = fields.Selection([
        ('pending', 'Pending'), ('valid', 'Valid'), ('review', 'Review'), ('error', 'Error'),
    ], required=True, readonly=True)
    error = fields.Text(readonly=True)

    _document_attempt_uniq = models.Constraint(
        'UNIQUE(document_id, attempt_number)', 'Extraction attempt already exists for this document.')

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            document = self.env['logistics.idp.document'].browse(vals['document_id']).exists()
            if not document:
                raise ValidationError('Extraction run requires an existing document.')
            self.env.cr.execute('SELECT id FROM logistics_idp_document WHERE id = %s FOR UPDATE', [document.id])
            self.env.cr.execute(
                'SELECT COALESCE(MAX(attempt_number), 0) + 1 '
                'FROM logistics_idp_extraction_run WHERE document_id = %s', [document.id])
            attempt_number = self.env.cr.fetchone()[0]
            if vals.get('attempt_number') not in (None, False, attempt_number):
                raise ValidationError('Extraction attempt number is allocated by the document.')
            if vals.get('retry_count') not in (None, False, attempt_number - 1):
                raise ValidationError('Extraction retry count is allocated by the document.')
            vals.update(attempt_number=attempt_number, retry_count=attempt_number - 1,
                        company_id=document.company_id.id, document_type=document.document_type,
                        page_count=document.page_count)
        return super().create(vals_list)


class LogisticsCheckResult(models.Model):
    _name = 'logistics.idp.check.result'
    _description = 'Immutable Logistics Check Result'
    _inherit = 'logistics.idp.immutable.snapshot'

    case_id = fields.Many2one('logistics.idp.case', required=True, readonly=True, ondelete='restrict', index=True)
    run_id = fields.Many2one('logistics.idp.extraction.run', readonly=True, ondelete='restrict')
    policy_source_id = fields.Many2one('logistics.idp.policy.source', readonly=True, ondelete='restrict')
    policy_version = fields.Char(related='policy_source_id.version', string='Policy Framework Version', readonly=True)
    policy_hash = fields.Char(related='policy_source_id.payload_hash', string='Policy Source Hash', readonly=True)
    policy_source_tier = fields.Selection(related='policy_source_id.source_tier', string='Policy Governance Tier', readonly=True)
    policy_citation = fields.Char(related='policy_source_id.citation', string='Policy Authority Citation', readonly=True)
    code = fields.Char(string='Check Rule Code', required=True, readonly=True, index=True)
    required = fields.Boolean(string='Mandatory Check', default=True, readonly=True)
    expected = fields.Text(string='Expected Specification', readonly=True)
    actual = fields.Text(string='Actual Observed Value', readonly=True)
    tolerance = fields.Char(string='Allowable Discrepancy Tolerance', readonly=True)
    verdict = fields.Selection(VERDICTS, string='Compliance Verdict', required=True, default='review', readonly=True)
    rationale = fields.Text(string='Compliance Rationale & Reasoning', required=True, readonly=True)
    citation = fields.Char(string='Rule Legal Citation', readonly=True)

    _audit_input_unique = models.Constraint(
        'unique(audit_input_hash)', 'This compliance input was already evaluated.')

    @api.model
    def _create_from_check_runner(self, values, service):
        runners = {
            LogisticsCase._restricted_party_screening.__code__, LogisticsCase.reconcile_documents.__code__,
            LogisticsCase.reconcile_import_declaration.__code__, LogisticsCase.evaluate_compliance_deadlines.__code__,
            LogisticsDocument._create_extraction_critical_check.__code__, LogisticsOutput.action_approve_shipping_plan.__code__,
            LogisticsPolicySource.evaluate_freshness.__code__,
        }
        if sys._getframe(1).f_code not in runners:
            raise UserError('Check results may only be created by an internal check runner.')
        return self.with_context(_logistics_check_result_token=_INTERNAL_CHECK_RESULT_TOKEN)._controlled_create(values, service)

    @api.model
    def _controlled_create(self, values, service):
        if self.env.context.get('_logistics_check_result_token') is not _INTERNAL_CHECK_RESULT_TOKEN:
            raise UserError('Check results may only be created by an internal check runner.')
        return super()._controlled_create(values, service)

    @api.model_create_multi
    def create(self, vals_list):
        if (self.env.context.get('_logistics_check_result_token') is not _INTERNAL_CHECK_RESULT_TOKEN
                and not tools.config['test_enable']):
            raise UserError('Check results may only be created by an internal check runner.')
        return super().create(vals_list)

    def write(self, vals):
        if self.env.context.get('_logistics_check_result_token') is not _INTERNAL_CHECK_RESULT_TOKEN:
            raise UserError('Check results are immutable; run a new check instead.')
        return super().write(vals)


class LogisticsOutput(models.Model):
    _broker_package_capability = object()

    _name = 'logistics.idp.output'
    _description = 'Immutable Logistics Output'
    _inherit = 'logistics.idp.immutable.snapshot'

    case_id = fields.Many2one('logistics.idp.case', required=True, readonly=True, ondelete='restrict', index=True)
    run_id = fields.Many2one('logistics.idp.extraction.run', readonly=True, ondelete='restrict')
    output_type = fields.Selection([(key, key.replace('_', ' ').title()) for key in ('e11', 'e13', 'e15', 'broker_package', 'import_declaration', 'shipping_plan', 'gate_pass', 'report_email', 'sap_erp_output')], required=True, readonly=True)
    version = fields.Integer(required=True, default=1, readonly=True)
    status = fields.Selection([('generated', 'Generated'), ('review', 'Review'), ('superseded', 'Superseded')], required=True, readonly=True)
    idempotency_key = fields.Char(required=True, readonly=True, index=True)
    supersedes_id = fields.Many2one('logistics.idp.output', readonly=True, ondelete='restrict')
    attachment_id = fields.Many2one('ir.attachment', readonly=True, ondelete='restrict')
    artifact_sha256 = fields.Char(readonly=True, index=True)
    artifact_mimetype = fields.Char(readonly=True)

    _output_key_unique = models.Constraint('unique(case_id, output_type, idempotency_key)', 'This output was already generated.')

    def action_approve_shipping_plan(self):
        self.ensure_one()
        if not (self.env.su or self.env.user.has_group('insilos_logistics_idp.group_logistics_manager')):
            raise UserError('Only Logistics Managers may approve Shipping Plans.')
        if self.output_type != 'shipping_plan' or self.status != 'generated':
            raise ValidationError('Only generated Shipping Plans may be approved.')
        current = self.case_id._current_semantic_checks('shipping_plan_approval', self.id)
        if len(current) == 1 and current.verdict == 'pass':
            return current
        return self.env['logistics.idp.check.result']._create_from_check_runner({
            'case_id': self.case_id.id, 'code': 'shipping_plan_approval', 'required': True,
            'verdict': 'pass', 'rationale': 'Shipping Plan approved by Logistics Manager.',
            'payload': {'output_id': self.id, 'artifact_sha256': self.artifact_sha256, 'actor_id': self.env.user.id},
        }, 'shipping_plan_approval')

    @api.model
    def _artifact(self, case, output_type, values):
        payload = _canonical_json(values)
        if output_type in ('sap_erp_output', 'shipping_plan', 'gate_pass'):
            template = self.env['ir.attachment'].browse(values.get('template_attachment_id')).exists()
            if not template:
                raise ValidationError('%s template attachment is unavailable.' % output_type.replace('_', ' ').title())
            workbook = load_workbook(io.BytesIO(template.raw))
            if values.get('sheet') not in workbook.sheetnames:
                raise ValidationError('%s template sheet does not exist.' % output_type.replace('_', ' ').title())
            stream = io.BytesIO()
            sheet = workbook[values['sheet']]
            if output_type == 'sap_erp_output':
                for row_index, row in enumerate(values['rows'], values.get('start_row', sheet.max_row + 1)):
                    for column in values['columns']:
                        sheet.cell(row_index, column['position'], row['values'].get(column['name']))
            else:
                for field in values['fields']:
                    sheet.cell(field['position'], 1, values['mapped_fields'].get(field['name']))
            workbook.save(stream)
            return '%s-v%s.xlsx' % (output_type.replace('_', '-'), values['mapping_version']), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', stream.getvalue()
        if output_type in ('e11', 'e13', 'e15'):
            stream = io.BytesIO()
            workbook = xlsxwriter.Workbook(stream, {'in_memory': True})
            workbook.set_properties({'title': output_type, 'created': __import__('datetime').datetime(2000, 1, 1)})
            sheet = workbook.add_worksheet('Data')
            sheet.write_row(0, 0, ('field', 'value'))
            for row, (key, value) in enumerate(sorted(values.items()), 1):
                sheet.write_row(row, 0, (key, _canonical_json(value) if isinstance(value, (dict, list)) else value))
            workbook.close()
            return '%s-v%s.xlsx' % (output_type, values.get('mapping_version', '1')), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', stream.getvalue()
        if output_type == 'import_declaration':
            stream = io.BytesIO()
            pdf = canvas.Canvas(stream, invariant=1)
            pdf.setTitle('Import Declaration')
            for row, line in enumerate(payload.splitlines() or [payload]):
                pdf.drawString(40, 800 - row * 14, line[:110])
            pdf.save()
            return 'import-declaration.pdf', 'application/pdf', stream.getvalue()
        if output_type == 'broker_package':
            stream = io.BytesIO()
            with zipfile.ZipFile(stream, 'w', zipfile.ZIP_DEFLATED) as archive:
                def write(name, content):
                    info = zipfile.ZipInfo(name, (2000, 1, 1, 0, 0, 0))
                    info.compress_type = zipfile.ZIP_DEFLATED
                    archive.writestr(info, content)
                write('manifest.json', payload.encode())
                for entry in values['entries']:
                    attachment = self.env['ir.attachment'].browse(entry['attachment_id']).exists()
                    if not attachment:
                        raise ValidationError('Broker package source attachment is unavailable.')
                    write(entry['name'], attachment.raw)
            return 'broker-package.zip', 'application/zip', stream.getvalue()
        message = EmailMessage()
        message['From'] = values.get('from', 'logistics@example.invalid')
        message['To'] = values.get('to', 'operations@example.invalid')
        message['Subject'] = values.get('subject', 'Logistics IDP report')
        message['Date'] = 'Sat, 01 Jan 2000 00:00:00 +0000'
        message['Message-ID'] = '<%s@logistics.invalid>' % hashlib.sha256(payload.encode()).hexdigest()[:24]
        message.set_content(payload)
        return 'logistics-report.eml', 'message/rfc822', message.as_bytes()

    @api.model
    def generate(self, case, output_type, values, run=None):
        if output_type in ('broker_package', 'shipping_plan', 'gate_pass'):
            raise ValidationError('%s must be generated through the case policy gate.' % output_type.replace('_', ' ').title())
        return self._generate(case, output_type, values, run)

    @api.model
    def _generate_broker_package(self, case, values):
        if self.env.context.get('_logistics_broker_package_capability') is not self._broker_package_capability:
            raise ValidationError('Broker packages must be generated through the case policy gate.')
        return self._generate(case, 'broker_package', values)

    @api.model
    def _generate(self, case, output_type, values, run=None):
        key = hashlib.sha256(_canonical_json(values).encode()).hexdigest()
        existing = self.search([('case_id', '=', case.id), ('output_type', '=', output_type), ('idempotency_key', '=', key)], limit=1)
        if existing:
            return existing
        previous = self.search([('case_id', '=', case.id), ('output_type', '=', output_type),
                                ('status', '=', 'generated')], order='version desc, id desc', limit=1)
        status = 'review' if values.get('missing') else 'generated'
        filename, mimetype, artifact = self._artifact(case, output_type, values)
        artifact_hash = hashlib.sha256(artifact).hexdigest()
        attachment = self.env['ir.attachment'].sudo().create({
            'name': filename, 'raw': artifact, 'mimetype': mimetype,
            'res_model': case._name, 'res_id': case.id, 'company_id': case.company_id.id,
        })
        output = self._controlled_create({
            'case_id': case.id, 'run_id': run.id if run else False, 'output_type': output_type,
            'version': previous.version + 1 if previous else 1, 'supersedes_id': previous.id,
            'status': status, 'idempotency_key': key, 'payload': values, 'audit_input_hash': key,
            'attachment_id': attachment.id, 'artifact_sha256': artifact_hash, 'artifact_mimetype': mimetype,
        }, 'output_generation')
        if previous:
            self.env.cr.execute(
                'UPDATE logistics_idp_output SET status = %s WHERE id = %s AND status = %s',
                ['superseded', previous.id, 'generated'],
            )
            previous.invalidate_recordset(['status'])
        return output


class LogisticsPolicySource(models.Model):
    _name = 'logistics.idp.policy.source'
    _description = 'Effective-dated Demo Policy Source'
    _inherit = 'logistics.idp.immutable.snapshot'

    code = fields.Char(required=True, readonly=True, index=True)
    version = fields.Char(required=True, readonly=True)
    company_id = fields.Many2one('res.company', readonly=True, index=True)
    jurisdiction = fields.Char(required=True, default='VN', readonly=True, index=True)
    regime = fields.Char(required=True, default='ALL', readonly=True, index=True)
    source_tier = fields.Selection([
        ('demo', 'Demo'),
        ('customer_reference', 'Customer Reference'),
        ('authoritative_tier_1', 'Authoritative Tier 1'),
        ('authoritative_tier_2', 'Authoritative Tier 2'),
        ('approved_provider_tier_3', 'Approved Provider Tier 3'),
        ('early_warning_tier_4', 'Early Warning Tier 4'),
    ], required=True, default='demo', readonly=True)
    citation = fields.Char(required=True, readonly=True)
    provenance = fields.Text(readonly=True)
    trade02_binding = fields.Text(readonly=True)
    effective_from = fields.Date(required=True, readonly=True)
    effective_to = fields.Date(readonly=True)
    recorded_at = fields.Datetime(required=True, readonly=True, index=True, default=fields.Datetime.now)
    state = fields.Selection([('draft', 'Draft'), ('active', 'Active'), ('retired', 'Retired')], required=True, default='draft', readonly=True)
    supersedes_id = fields.Many2one('logistics.idp.policy.source', readonly=True, ondelete='restrict')
    expected_refresh_hours = fields.Float(required=True, default=24, readonly=True)
    stale_after_hours = fields.Float(required=True, default=48, readonly=True)
    last_successful_sync_at = fields.Datetime(readonly=True)
    stale_disposition = fields.Selection([
        ('warning', 'Warning'), ('review', 'Review'), ('block', 'Block'),
    ], required=True, default='review', readonly=True)
    freshness_status = fields.Selection([
        ('fresh', 'Fresh'), ('warning', 'Warning'), ('review', 'Review'), ('block', 'Block'),
    ], compute='_compute_freshness_status')
    verification_status = fields.Selection([
        ('verified', 'Verified'), ('unverified', 'Unverified'),
    ], compute='_compute_governance_status')
    legal_authority = fields.Boolean(compute='_compute_governance_status')
    activation_eligible = fields.Boolean(compute='_compute_governance_status')
    activation_blocker = fields.Char(compute='_compute_governance_status')
    impacted_case_count = fields.Integer(compute='_compute_impacted_case_count')

    @api.model
    def _authority_registry(self):
        policy = json.loads((Path(__file__).resolve().parents[1] / 'config' / 'compliance_policy.json').read_text())
        registry = policy.get('authority_registry', {})
        if tools.config['test_enable'] and self.env.context.get('logistics_idp_test_authority_registry'):
            registry = {**registry, **self.env.context['logistics_idp_test_authority_registry']}
        return registry

    @api.model
    def _citation_is_governed(self, tier, citation):
        entry = self._authority_registry().get(tier, {}).get(citation)
        return bool(entry and entry.get('approved') is True and entry.get('source_tier') == tier)

    def _trade02_binding_is_verified(self):
        self.ensure_one()
        try:
            trade02 = json.loads(self.trade02_binding or '')
        except (TypeError, ValueError):
            return False
        if not isinstance(trade02, dict):
            return False
        identities = ('reviewer_identity', 'legal_owner_identity', 'maker_identity', 'checker_identity')
        return (
            trade02.get('status') == 'verified'
            and trade02.get('activation_allowed') is True
            and trade02.get('approval_status') == 'approved'
            and isinstance(trade02.get('policy_sha256'), str)
            and trade02['policy_sha256'] == self.payload_hash
            and all(isinstance(trade02.get(name), str) and re.fullmatch(r'[0-9a-f]{64}', trade02[name])
                    for name in ('canonical_oracle_sha256', 'receipt_sha256', 'pack_binding_sha256'))
            and all(isinstance(trade02.get(name), str) and trade02[name]
                    for name in identities)
            and len({trade02[name] for name in identities}) == len(identities)
        )

    def _compute_governance_status(self):
        for source in self:
            try:
                provenance = json.loads(source.provenance or '{}')
            except (TypeError, ValueError):
                provenance = {}
            if not isinstance(provenance, dict):
                provenance = {}
            source.verification_status = 'verified' if provenance.get('verification') == 'verified' else 'unverified'
            source.legal_authority = AUTHORITY_TIERS.get(source.source_tier, False)
            governed = source._citation_is_governed(source.source_tier, source.citation)
            trade02_verified = not source.code.startswith('TRADE_COMPLIANCE_') or source._trade02_binding_is_verified()
            source.activation_eligible = bool(
                source.state == 'draft' and source.legal_authority and governed
                and source.verification_status == 'verified' and trade02_verified
                and source.freshness_status == 'fresh')
            if source.state != 'draft':
                source.activation_blocker = 'Only draft candidates may be activated.'
            elif not source.legal_authority:
                source.activation_blocker = 'Source tier has no legal activation authority.'
            elif source.verification_status != 'verified':
                source.activation_blocker = 'REVIEW: source is unverified; activation forbidden.'
            elif not governed:
                source.activation_blocker = 'REVIEW: citation is not in the governed authority registry.'
            elif not trade02_verified:
                source.activation_blocker = 'REVIEW: TRADE-02 verified independent-oracle binding is required for activation.'
            elif source.freshness_status != 'fresh':
                source.activation_blocker = 'REVIEW: source freshness is missing, stale, or not eligible for activation.'
            else:
                source.activation_blocker = False

    _policy_freshness_thresholds = models.Constraint(
        'check(expected_refresh_hours > 0 AND stale_after_hours > 0 AND stale_after_hours >= expected_refresh_hours)',
        'Freshness cadence and threshold must be positive; stale threshold must not precede cadence.')

    def _compute_freshness_status(self):
        as_of = fields.Datetime.now().replace(tzinfo=timezone.utc)
        for source in self:
            source.freshness_status = source._freshness_values(as_of)['status']

    def _freshness_values(self, as_of):
        self.ensure_one()
        if not isinstance(as_of, datetime) or as_of.tzinfo is None or as_of.utcoffset() is None:
            raise ValidationError('Freshness as_of must be timezone-aware.')
        as_of = as_of.astimezone(timezone.utc)
        last_sync = self.last_successful_sync_at
        if not last_sync:
            return {'status': self.stale_disposition, 'age_hours': None,
                    'reason': 'Missing last successful synchronization; failed closed.'}
        last_sync = last_sync.replace(tzinfo=timezone.utc)
        age_hours = (as_of - last_sync).total_seconds() / 3600
        if age_hours < 0:
            raise ValidationError('Freshness as_of must not precede last successful synchronization.')
        stale = age_hours >= self.stale_after_hours
        return {'status': self.stale_disposition if stale else 'fresh', 'age_hours': age_hours,
                'reason': 'Age reached stale threshold.' if stale else 'Age remains below stale threshold.'}

    def evaluate_freshness(self, case, as_of):
        self.ensure_one()
        if not case.exists() or case.company_id != (self.company_id or self.env.company):
            raise ValidationError('Policy source and case must belong to the same allowed company.')
        values = self._freshness_values(as_of)
        canonical_as_of = as_of.astimezone(timezone.utc).isoformat().replace('+00:00', 'Z')
        payload = {
            'source_id': self.id, 'source_version': self.version, 'source_hash': self.payload_hash,
            'last_successful_sync_at': fields.Datetime.to_string(self.last_successful_sync_at),
            'expected_refresh_hours': self.expected_refresh_hours, 'stale_after_hours': self.stale_after_hours,
            'as_of': canonical_as_of, 'age_hours': values['age_hours'],
            'disposition': values['status'], 'reason': values['reason'],
        }
        audit_input_hash = hashlib.sha256(_canonical_json(payload).encode()).hexdigest()
        result_model = self.env['logistics.idp.check.result']
        existing = result_model.search([('audit_input_hash', '=', audit_input_hash)], limit=1)
        if existing:
            return existing
        verdict = 'pass' if values['status'] == 'fresh' else 'block' if values['status'] == 'block' else 'review'
        try:
            with self.env.cr.savepoint():
                return result_model._create_from_check_runner({
                    'case_id': case.id, 'policy_source_id': self.id, 'code': 'REGULATORY_SOURCE_FRESHNESS',
                    'required': True, 'expected': 'age_hours < stale_after_hours',
                    'actual': values['status'], 'tolerance': str(self.stale_after_hours),
                    'verdict': verdict, 'rationale': values['reason'], 'citation': self.citation,
                    'payload': payload, 'audit_input_hash': audit_input_hash,
                }, 'regulatory_source_freshness')
        except UniqueViolation:
            return result_model.search([('audit_input_hash', '=', audit_input_hash)], limit=1)

    def _impacted_case_domain(self):
        self.ensure_one()
        return [('company_id', '=', self.company_id.id or self.env.company.id),
                ('effective_date', '>=', self.effective_from),
                ('effective_date', '<=', self.effective_to or fields.Date.today())]

    def _compute_impacted_case_count(self):
        for policy in self:
            policy.impacted_case_count = self.env['logistics.idp.case'].search_count(policy._impacted_case_domain())

    def action_open_impacted_cases(self):
        self.ensure_one()
        return {'type': 'ir.actions.act_window', 'name': 'Impacted Cases',
                'res_model': 'logistics.idp.case', 'view_mode': 'list,form',
                'domain': self._impacted_case_domain(), 'context': {'create': False}}

    _policy_effective_range = models.Constraint(
        'check(effective_to IS NULL OR effective_to >= effective_from)',
        'Effective To must not be earlier than Effective From.')
    _policy_company_version_unique = models.Constraint(
        'unique nulls not distinct(company_id, code, version)',
        'This company or global policy version already exists.')

    @api.model
    def _controlled_create(self, values, service):
        if (self.env.context.get('_logistics_policy_write_token') is not _INTERNAL_POLICY_WRITE_TOKEN
                and self.env.context.get('_logistics_policy_loader_token') is not _INTERNAL_POLICY_LOADER_TOKEN
                and not (tools.config['test_enable'] and service == 'test_fixture')):
            raise AccessError('Policy sources may only be created by a trusted policy service.')
        return super()._controlled_create(values, service)

    @api.model_create_multi
    def create(self, vals_list):
        if (self.env.context.get('_logistics_policy_write_token') is not _INTERNAL_POLICY_WRITE_TOKEN
                and self.env.context.get('_logistics_policy_loader_token') is not _INTERNAL_POLICY_LOADER_TOKEN):
            raise AccessError('Policy sources may only be created by a trusted policy service.')
        for vals in vals_list:
            vals['recorded_at'] = fields.Datetime.now()
        return super().create(vals_list)

    def write(self, vals):
        if 'recorded_at' in vals:
            raise UserError('Policy recorded_at is immutable and server-assigned.')
        if (self.env.context.get('_logistics_policy_write_token') is not _INTERNAL_POLICY_WRITE_TOKEN
                and self.env.context.get('_logistics_policy_loader_token') is not _INTERNAL_POLICY_LOADER_TOKEN
                and self.env.context.get('_logistics_snapshot_token') is not _INTERNAL_SNAPSHOT_TOKEN):
            raise AccessError('Policy sources may only be changed by a trusted policy service.')
        return super().write(vals)

    @api.model
    def validate_compliance_context(self, context, case=None, authority=None):
        try:
            context = json.loads(context) if isinstance(context, str) else dict(context)
        except (TypeError, ValueError, json.JSONDecodeError) as exc:
            raise ValidationError('Compliance context must be valid JSON.') from exc
        if context.get('schema_version') != '1.0':
            raise ValidationError('Unsupported compliance context schema version.')
        if set(context) - {'schema_version', 'company_id', 'authority', 'transaction_time',
                           'effective_time', 'recorded_time', *COMPLIANCE_CONTEXT_FIELDS}:
            raise ValidationError('Compliance context contains unknown fields.')
        company_id = context.get('company_id')
        if (not isinstance(company_id, int) or company_id not in self.env.companies.ids
                or case and (not case.exists() or case.company_id.id != company_id)):
            raise ValidationError('Compliance context company is not allowed or does not match the case.')
        tier = authority or context.get('authority')
        if tier not in AUTHORITY_TIERS:
            raise ValidationError('Unknown compliance authority tier.')
        for field_name in COMPLIANCE_CONTEXT_FIELDS:
            context.setdefault(field_name, None)
        provenance_keys = {'document_id', 'field', 'source_kind', 'source_hash', 'extraction_run_id', 'authority'}
        evidence = context.get('evidence')
        evidence_valid = isinstance(evidence, list) and bool(evidence)
        for item in evidence or []:
            if not isinstance(item, dict) or set(item) != provenance_keys:
                evidence_valid = False
                continue
            document_id, run_id = item['document_id'], item['extraction_run_id']
            if (not isinstance(document_id, int) or isinstance(document_id, bool) or document_id <= 0
                    or not isinstance(run_id, int) or isinstance(run_id, bool) or run_id <= 0):
                raise ValidationError('Compliance evidence requires real document and extraction run IDs.')
            document = self.env['logistics.idp.document'].browse(document_id).exists()
            run = self.env['logistics.idp.extraction.run'].browse(run_id).exists()
            if not document or not run:
                raise ValidationError('Compliance evidence document or extraction run does not exist.')
            if run.document_id != document:
                raise ValidationError('Compliance evidence extraction run does not belong to the document.')
            if (document.company_id.id != company_id or run.company_id.id != company_id
                    or case and (document.case_id != case or run.case_id != case)):
                raise ValidationError('Compliance evidence ownership does not match the case or company.')
            if (item['authority'] not in AUTHORITY_TIERS
                    or AUTHORITY_RANK[item['authority']] > AUTHORITY_RANK[tier]):
                raise ValidationError('Compliance evidence authority exceeds policy authority.')
            if (not isinstance(item['source_hash'], str)
                    or not re.fullmatch(r'[0-9a-f]{64}', item['source_hash'])):
                evidence_valid = False
            elif item['source_kind'] == 'document' and item['source_hash'] != document.content_hash:
                raise ValidationError('Compliance evidence source hash does not match its referenced document.')
            elif item['source_kind'] == 'extraction_run' and item['source_hash'] != run.payload_hash:
                raise ValidationError('Compliance evidence source hash does not match its referenced extraction run.')
            elif item['source_kind'] not in ('document', 'extraction_run'):
                raise ValidationError('Compliance evidence source kind is invalid.')
        missing = [name for name in COMPLIANCE_CONTEXT_CRITICAL
                   if not context.get(name) or name == 'evidence' and not evidence_valid]
        parsed_times = {}
        for time_field in ('transaction_time', 'effective_time', 'recorded_time'):
            try:
                value = context.get(time_field)
                parsed_times[time_field] = datetime.fromisoformat(value.replace('Z', '+00:00'))
                if parsed_times[time_field].tzinfo is None:
                    raise ValueError
                context[time_field] = parsed_times[time_field].astimezone(timezone.utc).isoformat().replace('+00:00', 'Z')
            except (AttributeError, TypeError, ValueError):
                raise ValidationError('Compliance timestamps must be canonical timezone-aware ISO 8601 values.')
        if not (parsed_times['transaction_time'] <= parsed_times['effective_time'] <= parsed_times['recorded_time']):
            raise ValidationError('Compliance timestamps must satisfy transaction <= effective <= recorded ordering.')
        mappings = {}
        schemas = {
            'what': {'hs': 'hs_code'},
            'where': {'origin_country': 'origin_country_code', 'export_country': 'export_country_code',
                      'import_country': 'import_country_code'},
        }
        for section, aliases in schemas.items():
            value = context.get(section)
            if value is None:
                continue
            if not isinstance(value, dict):
                raise ValidationError('Compliance context %s must be an object.' % section)
            normalized = {}
            for key, item in value.items():
                target = aliases.get(key, key)
                if target not in set(aliases.values()):
                    raise ValidationError('Compliance context %s contains unknown fields.' % section)
                if target in normalized and normalized[target] != item:
                    raise ValidationError('Compliance context alias conflicts.')
                normalized[target] = item
                if key != target:
                    mappings['%s.%s' % (section, key)] = '%s.%s' % (section, target)
            context[section] = normalized
        transport = context.get('transport')
        if transport is not None:
            if not isinstance(transport, dict) or set(transport) - {'mode', 'extensions'}:
                raise ValidationError('Compliance context transport contains unknown fields.')
            extensions = transport.get('extensions', {})
            if not isinstance(extensions, dict) or any(not key.startswith('x-demo:') for key in extensions):
                raise ValidationError('Compliance context transport extensions are invalid.')
            if 'x-demo:mode' in extensions and extensions['x-demo:mode'] != transport.get('mode'):
                raise ValidationError('Compliance context transport extensions conflict.')
        canonical = _canonical_json(context)
        mapping_evidence = {'mappings': mappings, 'mapping_hash': hashlib.sha256(
            _canonical_json(mappings).encode()).hexdigest()}
        return {
            'schema_version': '1.0', 'context': context,
            'context_hash': hashlib.sha256(canonical.encode()).hexdigest(), 'mapping_evidence': mapping_evidence,
            'verdict': 'review' if missing or not AUTHORITY_TIERS[tier] else 'pass',
            'missing_critical': sorted(set(missing)), 'legal_authority': AUTHORITY_TIERS[tier],
        }

    @api.model
    def validate_policy_payload(self, payload, domain=None, pack=None, jurisdiction=None, regime=None,
                                require_policy_metadata=False):
        try:
            payload = json.loads(payload) if isinstance(payload, str) else payload
            if payload['schema_version'] != '1':
                raise KeyError('schema_version')
            horizontal = payload['horizontal']
            vertical = payload['vertical']
            overlays = payload['overlays']
            communication = payload['supplier_communication']
            selected_pack = pack or vertical['pack']
            if selected_pack not in vertical['packs']:
                raise KeyError('pack')
            config = vertical['packs'][selected_pack]
            if domain and vertical['domain'] != domain:
                raise KeyError('domain')
            if jurisdiction and overlays['jurisdiction']['code'] != jurisdiction:
                raise KeyError('jurisdiction')
            if regime and regime != 'ALL' and regime not in overlays['regimes']:
                raise KeyError('regime')
            if not all(key in horizontal for key in (
                    'document_lifecycle', 'comparison_lifecycle', 'classification', 'iap_charging',
                    'security_quality_gates')):
                raise KeyError('horizontal')
            if not all(key in config for key in (
                    'batch', 'attachment_categories', 'compare_stages', 'invoice_field_schema',
                    'extraction_aliases', 'purchase_order_lifecycle', 'reconciliation', 'counterpart_rules')):
                raise KeyError('pack')
            thresholds = horizontal['classification']['thresholds']
            if (set(thresholds) != {'trusted', 'strong_filename', 'ambiguous', 'iap_cutoff'}
                    or any(isinstance(value, bool) or not isinstance(value, (int, float)) or not 0 <= value <= 1
                           for value in thresholds.values())):
                raise KeyError('classification.thresholds')
            extraction = horizontal.get('extraction', {})
            if extraction.get('line_field_types') != {
                    'quantity': ['number', 'string', 'null'],
                    'unit_price': ['number', 'string', 'null'],
                    'price_per': ['number', 'string', 'null'],
                    'value': ['number', 'string', 'null']}:
                raise KeyError('extraction.line_field_types')
            acceptance = extraction.get('acceptance_threshold')
            if isinstance(acceptance, bool) or not isinstance(acceptance, (int, float)) or not 0 <= acceptance <= 1:
                raise KeyError('extraction.acceptance_threshold')
            templates = extraction.get('templates', {})
            if not isinstance(templates, dict):
                raise KeyError('extraction.templates')
            for role, template in templates.items():
                if (role != 'default' and role not in DOCUMENT_TYPES) or not isinstance(template, dict):
                    raise KeyError('extraction.templates')
                if set(template) - {'version', 'schema_version', 'prompt'}:
                    raise KeyError('extraction.templates')
                if (not isinstance(template.get('version'), str) or not template['version'].strip()
                        or not isinstance(template.get('schema_version'), str) or not template['schema_version'].strip()
                        or not isinstance(template.get('prompt'), str) or '{document_type}' not in template['prompt']
                        or '{schema}' not in template['prompt']):
                    raise KeyError('extraction.templates')
            for role, threshold in extraction.get('review_thresholds', {}).items():
                if (role not in DOCUMENT_TYPES or isinstance(threshold, bool)
                        or not isinstance(threshold, (int, float)) or not 0 <= threshold <= 1):
                    raise KeyError('extraction.review_thresholds')
            allowed_critical = set().union(*DOCUMENT_TYPES.values(), ('currency', 'document_date'))
            allowed_line_critical = {'lines[].%s' % field for field in (
                'description', 'material_code', 'quantity', 'unit_price', 'value', 'uom', 'currency',
                'price_per', 'quantity_1', 'quantity_2', 'remaining_quantity', 'hs_code')}
            for role, critical_fields in extraction.get('critical_fields', {}).items():
                if (role not in DOCUMENT_TYPES or not isinstance(critical_fields, list)
                        or any(field not in allowed_critical | allowed_line_critical for field in critical_fields)):
                    raise KeyError('extraction.critical_fields')
            families = set(config['batch']['expected_roles'])
            aliases = config['batch'].get('lifecycle_aliases', {})
            if (not families or any(source not in DOCUMENT_TYPES or not isinstance(target, str)
                                    or target not in families for source, target in aliases.items())):
                raise KeyError('batch')
            classification_aliases = config.get('classification', {}).get('aliases', {})
            for target, patterns in classification_aliases.items():
                if target not in DOCUMENT_TYPES or not isinstance(patterns, list) or not patterns:
                    raise KeyError('classification.aliases')
                for pattern in patterns:
                    if (not isinstance(pattern, str) or len(pattern) > 500 or re.search(r'\\[1-9]|\(\?<?[=!]|\([^)]*[+*][^)]*\)[+*]', pattern)):
                        raise KeyError('classification.aliases')
                    try:
                        re.compile(pattern)
                    except re.error as exc:
                        raise KeyError('classification.aliases') from exc
            schema_fields = config['invoice_field_schema'].get('fields', {})
            if any(not isinstance(spec, dict) or spec.get('type') not in ('string', 'date', 'object', 'array', 'number', 'integer', 'boolean')
                   for spec in schema_fields.values()):
                raise KeyError('invoice_field_schema')
            allowed_extraction = set().union(*DOCUMENT_TYPES.values(), ('supplier_address', 'buyer_name', 'total_gross',
                'total_quantity', 'currency', 'document_date', 'description', 'material_code', 'quantity', 'unit_price',
                'value', 'uom', 'price_per', 'quantity_1', 'quantity_2', 'remaining_quantity', 'hs_code',
                'invoice_serial', 'invoice_number', 'buyer_name', 'total_gross'))
            if any(not isinstance(source, str) or target not in allowed_extraction
                   for source, target in config['extraction_aliases'].items()):
                raise KeyError('extraction_aliases')
            codes = set()
            for rule in config['counterpart_rules']:
                trace_fields = ('code', 'source', 'version', 'input', 'outcome', 'test_ids')
                if (any(not rule.get(field) for field in trace_fields) or rule['code'] in codes
                        or not isinstance(rule['test_ids'], list) or rule['calendar_basis'] != 'calendar_days'
                        or not isinstance(rule['deadline_days'], int) or rule['deadline_days'] <= 0
                        or not rule['start_roles'] or not rule['expected_roles']
                        or rule['start_date']['document_field'] != 'document_date'
                        or rule['start_date']['multiple_triggers'] not in ('earliest_unresolved', 'explicit')):
                    raise KeyError('counterpart_rules')
                codes.add(rule['code'])
            if config['extensions']['certificate_of_origin']['enabled']:
                raise KeyError('certificate_of_origin')
            if not all(key in communication for key in (
                    'orchestration', 'recipient_role', 'warning_template_ref', 'overdue_template_ref',
                    'cadence_days', 'max_reminders', 'no_send_states')):
                raise KeyError('supplier_communication')
            if require_policy_metadata:
                metadata = payload.get('policy_metadata')
                effect = metadata.get('severity_decision_effect', {}) if isinstance(metadata, dict) else {}
                publication_date = metadata.get('source_publication_date') if isinstance(metadata, dict) else None
                try:
                    fields.Date.to_date(publication_date)
                except (TypeError, ValueError):
                    publication_date = None
                if (not publication_date or not isinstance(metadata.get('applicability_predicate'), dict)
                        or not metadata['applicability_predicate'] or not isinstance(effect, dict)
                        or effect.get('severity') not in ('low', 'medium', 'high', 'critical')
                        or effect.get('decision_effect') not in ('pass', 'review', 'block')
                        or not isinstance(metadata.get('evidence_requirements'), list)
                        or not metadata['evidence_requirements']
                        or any(not isinstance(item, str) or not item.strip()
                               for item in metadata['evidence_requirements'])):
                    raise ValidationError('policy_metadata is invalid.')
                predicate = metadata['applicability_predicate']
                supported = {'jurisdiction', 'regimes', 'parties', 'hs_codes'}
                unknown = set(predicate) - supported
                if unknown:
                    raise ValidationError(
                        'policy_metadata.applicability_predicate: unsupported selectors: %s.' %
                        ', '.join(sorted(unknown)))
                if 'jurisdiction' in predicate and (
                        not isinstance(predicate['jurisdiction'], str)
                        or not predicate['jurisdiction'].strip()):
                    raise ValidationError(
                        'policy_metadata.applicability_predicate.jurisdiction must be a non-empty string.')
                for selector in ('regimes', 'parties', 'hs_codes'):
                    if selector in predicate and (
                            not isinstance(predicate[selector], list) or not predicate[selector]
                            or any(not isinstance(item, str) or not item.strip()
                                   for item in predicate[selector])):
                        raise ValidationError(
                            'policy_metadata.applicability_predicate.%s must be a non-empty list of strings.'
                            % selector)
            if overlays.get('supplier_overrides') or overlays.get('company_overrides'):
                raise KeyError('overrides')
            export_controls = overlays.get('export_controls')
            if export_controls is not None:
                allowed = {'enabled', 'jurisdiction', 'required_fields', 'selectors', 'rules', 'missing_context_verdict', 'unknown_rule_verdict', 'no_pack_verdict'}
                if not isinstance(export_controls, dict) or set(export_controls) != allowed:
                    raise KeyError('export_controls')
                if export_controls['enabled'] not in (True, False) or not isinstance(export_controls['jurisdiction'], str):
                    raise KeyError('export_controls')
                if not isinstance(export_controls['required_fields'], list) or any(field not in CANONICAL_EXPORT_FIELDS for field in export_controls['required_fields']):
                    raise KeyError('export_controls.required_fields')
                selectors = export_controls['selectors']
                if not isinstance(selectors, dict) or set(selectors) != {'destination_countries', 'export_countries', 'origin_countries', 'transit_countries', 'transshipment_countries'} or any(not isinstance(value, list) or any(not isinstance(item, str) for item in value) for value in selectors.values()):
                    raise KeyError('export_controls.selectors')
                for rule in export_controls['rules']:
                    if not isinstance(rule, dict) or set(rule) - {'code', 'required_fields'} or not isinstance(rule.get('code'), str) or not isinstance(rule.get('required_fields'), list):
                        raise KeyError('export_controls.rules')
                for key in ('missing_context_verdict', 'unknown_rule_verdict', 'no_pack_verdict'):
                    if export_controls[key] not in ('PASS', 'REVIEW', 'BLOCK', 'NOT_APPLICABLE'):
                        raise KeyError('export_controls.%s' % key)
                if export_controls.get('activation_allowed') is True or export_controls.get('legal_authority') is True:
                    raise KeyError('export_controls.authority')
            restricted_parties = overlays.get('restricted_parties', [])
            if restricted_parties:
                if (not isinstance(restricted_parties, list) or any(
                        not isinstance(item, dict) or not item.get('list_name')
                        or not isinstance(item.get('entries'), list) or not item['entries']
                        for item in restricted_parties)):
                    raise KeyError('restricted_parties')
                for item in restricted_parties:
                    for entry in item['entries']:
                        if (not isinstance(entry, dict)
                                or not isinstance(entry.get('name'), str) or not entry['name'].strip()
                                or not isinstance(entry.get('aliases', []), list)
                                or any(not isinstance(alias, str) for alias in entry.get('aliases', []))):
                            raise KeyError('restricted_parties')
        except (KeyError, TypeError, json.JSONDecodeError) as exc:
            raise ValidationError('Invalid horizontal/vertical/overlays policy payload: %s.' % exc.args[0]) from exc
        return payload

    @api.model
    def evaluate_applicability(self, payload, case):
        """Deterministic applicability selector evaluation (SRS 31.8).

        Supported selectors: jurisdiction (policy-level), regimes, parties, hs_codes.
        Origin/destination/FTA selectors are rejected at import until entity fields exist.
        Selectors with no available case facts return 'indeterminate' (fail-closed).
        """
        metadata = payload.get('policy_metadata') or {}
        predicate = metadata.get('applicability_predicate')
        if not predicate:
            return {'verdict': 'matched', 'reasons': ['no_applicability_predicate']}
        reasons = []
        jurisdiction = predicate.get('jurisdiction')
        if jurisdiction:
            if payload['overlays']['jurisdiction']['code'] != jurisdiction:
                return {'verdict': 'not_matched', 'reasons': ['jurisdiction_mismatch']}
            reasons.append('jurisdiction_matched')
        regimes = predicate.get('regimes') or []
        if regimes:
            case_regimes = set()
            for document in case.document_ids:
                try:
                    metadata_doc = json.loads(
                        document.current_run_id.payload or '{}').get('metadata') or {}
                    for key in ('customs_regime', 'regime'):
                        if metadata_doc.get(key):
                            case_regimes.add(str(metadata_doc[key]))
                except (TypeError, ValueError):
                    continue
            if not case_regimes:
                return {'verdict': 'indeterminate', 'reasons': ['regime_facts_unavailable']}
            if not case_regimes.intersection(regimes):
                return {'verdict': 'not_matched', 'reasons': ['regime_not_matched']}
            reasons.append('regimes_matched')
        parties = predicate.get('parties') or []
        if parties:
            party = normalize_party(case.supplier_reference)
            if not party:
                return {'verdict': 'indeterminate', 'reasons': ['party_facts_unavailable']}
            if not any(normalize_party(expected) in party for expected in parties):
                return {'verdict': 'not_matched', 'reasons': ['party_not_matched']}
            reasons.append('parties_matched')
        hs_codes = predicate.get('hs_codes') or []
        if hs_codes:
            case_hs = set()
            snapshots = self.env['logistics.idp.po.snapshot'].sudo().search([
                ('company_id', '=', case.company_id.id),
                ('po_reference', '=', case.po_reference)], limit=1000)
            for snapshot in snapshots:
                try:
                    lines = json.loads(snapshot.payload or '{}').get('lines') or []
                except (TypeError, ValueError):
                    lines = []
                for line in lines:
                    code = str(line.get('hs_code') or '').replace('.', '').strip()
                    if code:
                        case_hs.add(code)
            if not case_hs:
                return {'verdict': 'indeterminate', 'reasons': ['hs_facts_unavailable']}
            expected = [str(code).replace('.', '').strip() for code in hs_codes]
            if not any(item and any(hs.startswith(item) for hs in case_hs) for item in expected):
                return {'verdict': 'not_matched', 'reasons': ['hs_not_matched']}
            reasons.append('hs_codes_matched')
        return {'verdict': 'matched', 'reasons': reasons}

    @api.model
    def evaluate_export_controls(self, payload, case=None):
        overlay = (payload or {}).get('overlays', {}).get('export_controls')
        base = {'policy_identity': hashlib.sha256(_canonical_json(overlay or {}).encode()).hexdigest(), 'matched_rules': [], 'required_evidence': [], 'reason_codes': []}
        if not overlay or not overlay.get('enabled'):
            return {**base, 'verdict': 'NOT_APPLICABLE', 'reason_codes': ['NO_EXPORT_CONTROLS_PACK']}
        if not isinstance(overlay.get('rules'), list) or not isinstance(overlay.get('selectors'), dict):
            return {**base, 'verdict': 'REVIEW', 'reason_codes': ['INVALID_EXPORT_CONTROLS_CONTEXT']}
        facts = payload.get('document') if isinstance(payload.get('document'), dict) else payload
        destination = facts.get('destination_country')
        if not isinstance(destination, str) or not destination.strip():
            return {**base, 'verdict': 'REVIEW', 'reason_codes': ['MISSING_DESTINATION_COUNTRY'], 'required_evidence': ['destination_country']}
        selectors = overlay['selectors']
        selected = selectors.get('destination_countries', [])
        if selected and destination not in selected:
            return {**base, 'verdict': 'NOT_APPLICABLE', 'reason_codes': ['SELECTOR_MISMATCH']}
        required = sorted(set(overlay.get('required_fields', [])))
        missing = [field for field in required if facts.get(field) in (None, '', [])]
        if missing:
            return {**base, 'verdict': 'REVIEW', 'reason_codes': ['MISSING_REQUIRED_EVIDENCE'], 'required_evidence': missing}
        base.update({'matched_rules': [rule['code'] for rule in overlay['rules']], 'audit_input_hash': hashlib.sha256(_canonical_json({'payload': payload, 'case_id': getattr(case, 'id', None)}).encode()).hexdigest()})
        return {**base, 'verdict': 'REVIEW', 'reason_codes': ['NON_AUTHORITATIVE_SYNTHETIC_OVERLAY']}

    @api.model
    def compute_resilience_metrics(self, company, as_of=None):
        """Deterministic trade-resilience KPIs from persisted governance snapshots (SRS 31.9).

        ponytail: measurable subset only - source freshness, detection/activation
        latency, open impact assessments and unresolved stale exposures.
        Time-to-clear, avoided re-screening and lane-coverage metrics wait for
        lane/country entity fields and case-resolution timestamps.
        """
        as_of = as_of or fields.Datetime.now().replace(tzinfo=timezone.utc)
        if not isinstance(as_of, datetime) or as_of.tzinfo is None or as_of.utcoffset() is None:
            raise ValidationError('as_of must be timezone-aware.')
        sources = self.sudo().search([('company_id', '=', company.id)])
        freshness = {'total': len(sources), 'fresh': 0, 'stale': 0, 'unknown': 0}
        exposures = []
        detection_days, activation_days = [], []
        for source in sources:
            try:
                values = source._freshness_values(as_of)
            except (ValidationError, ValueError):
                values = None
            if values is None or values['age_hours'] is None:
                freshness['unknown'] += 1
            elif values['status'] == 'fresh':
                freshness['fresh'] += 1
            else:
                freshness['stale'] += 1
                exposures.append({
                    'code': source.code, 'version': source.version,
                    'age_hours': round(values['age_hours'], 2),
                })
            metadata = {}
            try:
                metadata = (json.loads(source.payload or '{}') or {}).get('policy_metadata') or {}
            except (TypeError, ValueError):
                pass
            publication = metadata.get('source_publication_date')
            try:
                published = fields.Date.from_string(str(publication))
            except (TypeError, ValueError):
                published = None
            if published:
                if source.recorded_at:
                    detection_days.append(max((source.recorded_at.date() - published).days, 0))
                if source.effective_from:
                    activation_days.append(
                        max((fields.Date.from_string(source.effective_from) - published).days, 0))
        decisions = self.env['logistics.idp.policy.decision'].sudo().search(
            [('case_id.company_id', '=', company.id)])
        open_decisions = decisions.filtered(lambda record: not record.effective_to)

        def summarize(days):
            if not days:
                return {'count': 0, 'mean': None, 'max': None, 'min': None}
            return {'count': len(days), 'mean': round(sum(days) / len(days), 2),
                    'max': max(days), 'min': min(days)}

        return {
            'schema_version': '1.0',
            'generated_at': as_of.astimezone(timezone.utc).isoformat().replace('+00:00', 'Z'),
            'sources': freshness,
            'detection_latency_days': summarize(detection_days),
            'activation_latency_days': summarize(activation_days),
            'open_impact_assessments': {'total': len(decisions), 'open': len(open_decisions)},
            'unresolved_stale_exposures': exposures,
        }

    @api.model
    def fetch_external_candidates(self, registry_key, query, method='GET'):
        if (not self.env.user.has_group('insilos_logistics_idp.group_logistics_manager')
                and not self.env.user.has_group('insilos_logistics_idp.group_logistics_admin')
                or self.env.user.has_group('insilos_logistics_idp.group_logistics_integration')):
            raise UserError('Only Logistics Managers or Administrators may fetch policy candidates.')
        if method != 'GET':
            raise ValidationError('Candidate API supports GET only.')
        return fetch_candidates(registry_key, query)

    @api.model
    def import_candidate(self, values):
        if (not self.env.user.has_group('insilos_logistics_idp.group_logistics_manager')
                and not self.env.user.has_group('insilos_logistics_idp.group_logistics_admin')
                or self.env.user.has_group('insilos_logistics_idp.group_logistics_integration')):
            raise UserError('Only Logistics Managers or Administrators may import policy candidates.')
        if not isinstance(values, dict):
            raise ValidationError('Policy candidate must be a JSON object.')
        allowed = {'code', 'version', 'company_id', 'jurisdiction', 'regime', 'source_tier',
                   'citation', 'provenance', 'effective_from', 'effective_to', 'payload',
                   'compliance_context', 'expected_refresh_hours', 'stale_after_hours',
                   'last_successful_sync_at', 'stale_disposition'}
        try:
            raw = json.dumps(values, ensure_ascii=False, separators=(',', ':'), sort_keys=True, default=fields.Date.to_string)
        except (TypeError, ValueError) as exc:
            raise ValidationError('Policy candidate must contain JSON-compatible values.') from exc
        if len(raw.encode()) > MAX_POLICY_IMPORT_BYTES:
            raise ValidationError('Policy candidate exceeds the 1 MiB limit.')
        if set(values) - allowed:
            raise ValidationError('Policy candidate contains unknown fields.')
        required = allowed - {'company_id', 'effective_to', 'expected_refresh_hours', 'stale_after_hours',
                              'last_successful_sync_at', 'stale_disposition'}
        if any(values.get(name) in (None, '', []) for name in required):
            raise ValidationError('Policy candidate is missing required provenance or policy fields.')
        company_id = values.get('company_id')
        if company_id is not None and (not isinstance(company_id, int) or company_id not in self.env.companies.ids):
            raise ValidationError('Policy candidate company is not allowed.')
        if values['source_tier'] not in AUTHORITY_TIERS:
            raise ValidationError('Unknown compliance authority tier.')
        citation = values['citation']
        if not isinstance(citation, str) or len(citation) > 2048 or re.search(r'[\x00-\x1f\x7f]', citation):
            raise ValidationError('Policy candidate citation is invalid.')
        citation = citation.strip()
        if citation.startswith(('http:', 'https:')):
            parsed = urlsplit(citation)
            if (parsed.scheme != 'https' or not parsed.hostname or parsed.username or parsed.password
                    or parsed.query or parsed.fragment):
                raise ValidationError('Policy candidate citation URL must be clean HTTPS.')
            citation = parsed.geturl()
        elif not re.fullmatch(r'(?:LAW|REG|DECREE|CIRCULAR|STANDARD):[A-Z0-9][A-Z0-9._/-]{0,127}', citation):
            raise ValidationError('Policy candidate citation authority reference is invalid.')
        provenance = values['provenance']
        if not isinstance(provenance, dict) or not provenance.get('source') or not provenance.get('captured_at'):
            raise ValidationError('Policy candidate provenance is invalid.')
        try:
            captured_at = datetime.fromisoformat(provenance['captured_at'].replace('Z', '+00:00'))
        except (AttributeError, TypeError, ValueError) as exc:
            raise ValidationError('Policy candidate provenance captured_at must be timezone-aware ISO 8601.') from exc
        if captured_at.tzinfo is None or captured_at.utcoffset() is None:
            raise ValidationError('Policy candidate provenance captured_at must be timezone-aware ISO 8601.')
        provenance = {**provenance, 'captured_at': captured_at.astimezone(timezone.utc).isoformat().replace('+00:00', 'Z')}
        effective_from = fields.Date.to_date(values['effective_from'])
        effective_to = fields.Date.to_date(values.get('effective_to'))
        if effective_to and effective_to < effective_from:
            raise ValidationError('Effective To must not be earlier than Effective From.')
        if effective_to and effective_to < fields.Date.today():
            raise ValidationError('Stale policy candidates cannot be imported.')
        payload = self.validate_policy_payload(values['payload'], require_policy_metadata=True)
        normalized_context = self.validate_compliance_context(
            values['compliance_context'], authority=values['source_tier'])
        canonical_payload = _canonical_json(payload)
        payload_hash = hashlib.sha256(canonical_payload.encode()).hexdigest()
        provenance_json = _canonical_json(provenance)
        audit_input_hash = hashlib.sha256(_canonical_json({
            'payload_hash': payload_hash, 'context_hash': normalized_context['context_hash'],
            'citation': citation, 'provenance_hash': hashlib.sha256(provenance_json.encode()).hexdigest(),
            'import': {**{key: values.get(key) for key in (
                'code', 'version', 'company_id', 'jurisdiction', 'regime', 'source_tier')},
                'effective_from': fields.Date.to_string(effective_from),
                'effective_to': fields.Date.to_string(effective_to)},
        }).encode()).hexdigest()
        identity = [('company_id', '=', company_id or False), ('code', '=', values['code']),
                    ('version', '=', values['version'])]
        existing = self.search(identity, limit=1)
        if existing:
            if existing.payload_hash == payload_hash and existing.audit_input_hash == audit_input_hash:
                return existing
            raise ValidationError('This policy identity already exists with different content or governance metadata.')
        create_values = {key: values[key] for key in (
            'code', 'version', 'jurisdiction', 'regime', 'source_tier', 'effective_from', 'payload')}
        create_values.update(company_id=company_id, effective_to=effective_to, state='draft',
                             citation=citation, provenance=provenance_json,
                             audit_input_hash=audit_input_hash)
        try:
            with self.env.cr.savepoint():
                return self.with_context(
                    _logistics_policy_write_token=_INTERNAL_POLICY_WRITE_TOKEN,
                )._controlled_create(create_values, 'policy_candidate_import')
        except UniqueViolation:
            existing = self.search(identity, limit=1)
            if existing and existing.payload_hash == payload_hash and existing.audit_input_hash == audit_input_hash:
                return existing
            raise ValidationError('This policy identity already exists with different content or governance metadata.')

    @api.model
    def preview_impact(self, candidate, cutoff, horizon, record_limit, compliance_context):
        candidate.ensure_one()
        if candidate.state != 'draft':
            raise ValidationError('Impact preview requires a draft policy candidate.')
        try:
            cutoff = fields.Date.to_date(cutoff)
        except (TypeError, ValueError) as exc:
            raise ValidationError('Impact preview cutoff must be a valid explicit date.') from exc
        if not cutoff:
            raise ValidationError('Impact preview requires an explicit cutoff date.')
        if cutoff < candidate.effective_from:
            raise ValidationError('Impact preview cutoff must not precede candidate Effective From.')
        if candidate.effective_to and cutoff > candidate.effective_to:
            raise ValidationError('Impact preview cutoff must not exceed candidate Effective To.')
        try:
            horizon = fields.Date.to_date(horizon)
        except (TypeError, ValueError) as exc:
            raise ValidationError('Impact preview horizon must be a valid explicit date.') from exc
        if not horizon or horizon < cutoff:
            raise ValidationError('Impact preview horizon must be explicit and not precede cutoff.')
        if candidate.effective_to and horizon > candidate.effective_to:
            raise ValidationError('Impact preview horizon must not exceed candidate Effective To.')
        if not isinstance(record_limit, int) or isinstance(record_limit, bool) or not 1 <= record_limit <= 10000:
            raise ValidationError('Impact preview record limit must be between 1 and 10000.')
        # ponytail: pack comes from the candidate payload (pack-family driven); jurisdiction stays VN until multi-jurisdiction sources exist.
        candidate_payload = self.validate_policy_payload(
            candidate.payload, 'trade_compliance', None, 'VN', candidate.regime)
        pack = candidate_payload['vertical']['pack']
        normalized = self.validate_compliance_context(
            compliance_context, authority=candidate.source_tier)
        company = candidate.company_id or self.env.company
        if normalized['context']['company_id'] != company.id:
            raise ValidationError('Compliance context company does not match the candidate.')
        recorded_as_of = datetime.fromisoformat(
            normalized['context']['recorded_time'].replace('Z', '+00:00')).astimezone(timezone.utc).replace(tzinfo=None)
        active = self.select_effective(
            company, 'VN', candidate.regime, cutoff, code=candidate.code,
            recorded_as_of=recorded_as_of)
        if len(active) != 1:
            raise ValidationError('Impact preview requires exactly one active policy.')
        active_payload = self.validate_policy_payload(
            active.payload, 'trade_compliance', pack, 'VN', candidate.regime)

        def rules(payload):
            selected = payload['vertical']['packs'][pack]
            return selected, {rule['code']: rule for rule in selected['counterpart_rules']}

        candidate_pack, candidate_rules = rules(candidate_payload)
        active_pack, active_rules = rules(active_payload)
        domain_version = 'trade-compliance-impact-v2'
        domain = [
            ('company_id', '=', company.id),
            ('state', 'not in', ('completed', 'closed_duplicate', 'closed_other')),
            ('effective_date', '<=', horizon),
        ]
        cases = self.env['logistics.idp.case'].search(domain, order='id', limit=record_limit + 1)
        if len(cases) > record_limit:
            raise ValidationError('Impact preview record limit exceeded; narrow the horizon.')
        results = []
        manifest_cases = []
        for case in cases:
            evidence = [{'id': item.id, 'hash': item.payload_hash, 'status': item.status}
                        for item in case.evidence_ids.sorted('id')]
            documents = [{
                'id': item.id,
                'document_type': item.document_type,
                'content_hash': item.content_hash,
                'write_date': fields.Datetime.to_string(item.write_date),
                'current_run_id': item.current_run_id.id or False,
                'current_run_write_date': fields.Datetime.to_string(item.current_run_id.write_date)
                    if item.current_run_id else False,
                'current_run': {
                    'id': item.current_run_id.id,
                    'write_date': fields.Datetime.to_string(item.current_run_id.write_date),
                } if item.current_run_id else {'id': False, 'write_date': False},
                'current_run_payload_hash': hashlib.sha256(
                    (item.current_run_id.payload or '').encode()).hexdigest(),
            } for item in case.document_ids.sorted('id')]
            manifest_cases.append({
                'id': case.id,
                'company_id': case.company_id.id,
                'identity': f'{case._name}:{case.id}',
                'source_identity': [case.source_system, case.source_key, case.source_version],
                'state': case.state,
                'write_date': fields.Datetime.to_string(case.write_date),
                'effective_date': fields.Date.to_string(case.effective_date),
                'documents': documents,
                'evidence': evidence,
            })
            evidence_references = [{'model': 'logistics.idp.evidence', **item} for item in evidence]

            def evaluate(pack, configured_rules):
                aliases = pack['batch']['lifecycle_aliases']
                documents = [(document, aliases.get(document.document_type, document.document_type))
                             for document in case.document_ids.sorted('id')]
                roles = {role for _document, role in documents}
                outputs = []
                missing = []
                references = list(evidence_references)
                for code, rule in sorted(configured_rules.items()):
                    triggers = [(document, role) for document, role in documents
                                if role in set(rule['start_roles'])]
                    if not triggers:
                        continue
                    expected_present = bool(roles.intersection(rule['expected_roles']))
                    dates = []
                    for document, _role in triggers:
                        references.append({'model': 'logistics.idp.document', 'id': document.id,
                                           'hash': document.content_hash})
                        try:
                            value = json.loads(document.current_run_id.payload).get('payload', {}).get(
                                rule['start_date']['document_field'])
                            date = fields.Date.to_date(value)
                        except (TypeError, ValueError):
                            date = None
                        if date:
                            dates.append(date)
                    if not dates and not expected_present:
                        missing.append('%s.when' % code)
                        outputs.append((code, 'insufficient_context'))
                    else:
                        start = min(dates) if dates else case.effective_date
                        status = self.env['logistics.idp.case'].evaluate_counterpart_deadline(
                            start, cutoff, rule, expected_present)['status']
                        outputs.append((code, status))
                return outputs, sorted(set(missing)), sorted(
                    references, key=lambda item: (item['model'], item['id'], item['hash']))

            applicability = self.evaluate_applicability(candidate_payload, case)
            if applicability['verdict'] != 'matched':
                if applicability['verdict'] == 'not_matched':
                    outcome, reasons = 'not_affected', ['applicability_not_matched'] + applicability['reasons']
                else:
                    outcome, reasons = 'insufficient_context', applicability['reasons']
                before, after, missing = [], [], []
                before_refs, after_refs = [], []
            else:
                before, before_missing, before_refs = evaluate(active_pack, active_rules)
                after, after_missing, after_refs = evaluate(candidate_pack, candidate_rules)
                missing = sorted(set(before_missing + after_missing))
                if missing:
                    outcome, reasons = 'insufficient_context', missing
                elif before != after:
                    outcome, reasons = 'affected', ['counterpart_rule_outcome_changed']
                elif candidate.payload_hash == active.payload_hash:
                    outcome, reasons = 'not_affected', ['candidate_matches_active_hash']
                elif not before and not after:
                    outcome, reasons = 'not_affected', ['no_counterpart_rule_trigger']
                else:
                    outcome, reasons = 'possibly_affected', ['policy_hash_changed_outcome_unchanged']
            references = [json.loads(item) for item in sorted(
                {_canonical_json(reference) for reference in before_refs + after_refs})]
            results.append({'case_id': case.id, 'outcome': outcome, 'reasons': reasons,
                            'active_rule_outcomes': before, 'candidate_rule_outcomes': after,
                            'evidence_references': references,
                            'dimensions': {
                                'open_cases': {'outcome': outcome, 'reasons': reasons,
                                               'evidence_references': references},
                                'products_materials_and_hs_codes': {
                                    'outcome': 'insufficient_context',
                                    'reasons': ['entity_model_unavailable'], 'evidence_references': []},
                            }})
        counts = {name: 0 for name in (
            'affected', 'possibly_affected', 'not_affected', 'insufficient_context')}
        for result in results:
            counts[result['outcome']] += 1
        output = {'counts': counts, 'results': results}
        manifest = {
            'algorithm_version': '2', 'schema_version': '1', 'domain_version': domain_version,
            'candidate_id': candidate.id, 'candidate_hash': candidate.payload_hash,
            'candidate_effective_from': fields.Date.to_string(candidate.effective_from),
            'candidate_effective_to': fields.Date.to_string(candidate.effective_to) if candidate.effective_to else False,
            'active_id': active.id, 'active_hash': active.payload_hash,
            'recorded_as_of': fields.Datetime.to_string(recorded_as_of),
            'cutoff': fields.Date.to_string(cutoff), 'horizon': fields.Date.to_string(horizon),
            'record_limit': record_limit, 'context_hash': normalized['context_hash'],
            'cases': manifest_cases,
            'outcomes': [{'case_id': item['case_id'], 'outcome': item['outcome'],
                          'reasons': item['reasons']} for item in results],
        }
        return {**manifest, 'case_ids': cases.ids,
                'input_hash': hashlib.sha256(_canonical_json(manifest).encode()).hexdigest(),
                'output_hash': hashlib.sha256(_canonical_json(output).encode()).hexdigest(), **output}

    @api.model
    def audit_effective(self, company, jurisdiction, regime, legal_effective_date, recorded_as_of, code=None):
        if not legal_effective_date:
            raise ValidationError('Audit requires an explicit legal effective date.')
        policies = self.select_effective(company, jurisdiction, regime, legal_effective_date,
                                         code=code, recorded_as_of=recorded_as_of)
        return {'policy_ids': policies.ids, 'legal_effective_date': fields.Date.to_string(fields.Date.to_date(legal_effective_date)),
                'recorded_as_of': fields.Datetime.to_datetime(recorded_as_of)}

    @api.model
    def simulate_scenario(self, candidate, cutoff, horizon, record_limit, compliance_context, hypothetical):
        if not tools.config['test_enable']:
            raise ValidationError('Simulation is non-production only.')
        if hypothetical != {'id': candidate.id, 'payload_hash': candidate.payload_hash}:
            raise ValidationError('Simulation requires an explicit hypothetical candidate.')
        result = self.preview_impact(candidate, cutoff, horizon, record_limit, compliance_context)
        result['simulation'] = {'mode': 'non_production', 'read_only': True, 'legal_authority': False,
                                'disclaimer': 'Hypothetical simulation; no legal authority.'}
        return result

    @api.model
    def diff_candidate(self, candidate):
        candidate.ensure_one()
        active = self.select_effective(
            candidate.company_id or self.env.company, candidate.jurisdiction, candidate.regime,
            candidate.effective_from, code=candidate.code)
        def flatten(value, path='$'):
            if isinstance(value, dict):
                return {item: leaf for key in sorted(value) for item, leaf in flatten(value[key], '%s.%s' % (path, key)).items()}
            if isinstance(value, list):
                return {item: leaf for index, child in enumerate(value) for item, leaf in flatten(child, '%s[%s]' % (path, index)).items()}
            return {path: value}
        candidate_payload = json.loads(candidate.payload)
        active_payload = json.loads(active.payload) if active else {}
        left, right = flatten(active_payload), flatten(candidate_payload)
        changes = [{'path': path, 'active': left.get(path), 'candidate': right.get(path)}
                   for path in sorted(set(left) | set(right)) if left.get(path) != right.get(path)]
        return {'candidate_hash': candidate.payload_hash, 'active_hash': active.payload_hash if active else False,
                'active_id': active.id or False, 'changes': changes}

    @api.model
    def select_effective_pack(self, company, domain, pack, jurisdiction, regime, effective_date):
        policies = self.select_effective(company, jurisdiction, regime, effective_date)
        if not policies:
            return policies
        if len(policies) != 1:
            raise ValidationError('Overlapping active policy packs; selection failed closed.')
        self.validate_policy_payload(policies.payload, domain, pack, jurisdiction, regime)
        return policies

    @api.model
    def select_effective(self, company, jurisdiction, regime, effective_date, code=None,
                         recorded_as_of=None):
        recorded_as_of = fields.Datetime.to_datetime(recorded_as_of or fields.Datetime.now())
        domain = [
            ('state', '=', 'active'), ('company_id', 'in', (False, company.id)),
            ('jurisdiction', '=', jurisdiction), ('regime', 'in', ('ALL', regime)),
            ('effective_from', '<=', effective_date),
            '|', ('effective_to', '=', False), ('effective_to', '>=', effective_date),
            ('recorded_at', '<=', recorded_as_of),
        ]
        if code:
            domain.append(('code', '=', code))
        policies = self.search(domain, order='effective_from desc, version desc, id desc')
        exact = policies.filtered(lambda item: item.company_id == company and item.regime == regime)
        return exact or policies.filtered(lambda item: item.company_id == company) or policies.filtered(
            lambda item: item.regime == regime) or policies

    @api.model
    def _upgrade_effective_policy_history(self):
        if self.env.context.get('_logistics_policy_loader_token') is not _INTERNAL_POLICY_LOADER_TOKEN:
            raise AccessError('Policy history upgrade requires the trusted policy loader.')
        models_data = [
            {
                'name': 'Gemini 2.5 Flash',
                'model_id': 'google/gemini-2.5-flash',
                'supports_vision': True,
                'supports_json': True,
                'supports_pdf_input': True,
                'capability_evidence_source': 'successful_logistics_corpus; runs 6844/6846/6850; complete output requires 8192 tokens',
                'extraction_max_tokens': 8192,
            },
            {
                'name': 'Gemini 3.5 Flash Lite',
                'model_id': 'google/gemini-3.5-flash-lite',
                'supports_vision': True,
                'supports_json': True,
                'supports_pdf_input': True,
                'capability_evidence_source': 'successful_logistics_corpus; runs 6844/6846/6850; complete output requires 8192 tokens',
                'extraction_max_tokens': 8192,
            },
        ]
        for mdata in models_data:
            existing = self.env['openrouter.model'].search([('model_id', '=', mdata['model_id'])], limit=1)
            if existing:
                existing.write(mdata)
            else:
                self.env['openrouter.model'].create(mdata)

        proven_models = self.env['openrouter.model'].search([
            ('model_id', 'in', ['google/gemini-2.5-flash', 'google/gemini-3.5-flash-lite']),
        ])
        domain = self.env['openrouter.domain'].search([('code', '=', 'logistics_ocr')], limit=1)
        if domain:
            primary = proven_models.filtered(lambda m: m.model_id == 'google/gemini-2.5-flash')[:1]
            if primary:
                domain.primary_model_id = primary
            domain.fallback_model_ids = proven_models.filtered(lambda model: model != domain.primary_model_id)
        self.env.cr.execute("""
            WITH numbered AS (
                SELECT id, ROW_NUMBER() OVER (PARTITION BY document_id ORDER BY started_at, id) AS attempt
                FROM logistics_idp_extraction_run
            )
            UPDATE logistics_idp_extraction_run AS run
               SET attempt_number = numbered.attempt
              FROM numbered
             WHERE run.id = numbered.id AND run.attempt_number IS NULL
        """)
        self.env.cr.execute(
            'ALTER TABLE logistics_idp_extraction_run ALTER COLUMN attempt_number SET NOT NULL')
        baseline = self.env.ref(
            'insilos_logistics_idp.policy_trade_compliance_vn_reference',
            raise_if_not_found=False,
        )
        successor = self.env.ref(
            'insilos_logistics_idp.policy_trade_compliance_vn_reference_2026_2',
            raise_if_not_found=False,
        )
        if not baseline or not successor:
            return True
        payload = json.loads(successor.payload)
        extraction = payload['horizontal'].setdefault('extraction', {})
        extraction['line_field_types'] = {
            'quantity': ['number', 'string', 'null'],
            'unit_price': ['number', 'string', 'null'],
            'price_per': ['number', 'string', 'null'],
            'value': ['number', 'string', 'null'],
        }
        extraction['critical_fields'] = {
            document_type: list(required_fields) + (['lines[].quantity'] if 'lines' in required_fields else [])
            for document_type, required_fields in DOCUMENT_TYPES.items()
        }
        extraction['critical_fields']['import_declaration'] += ['regime', 'lines[].hs_code']
        payload['provenance']['version'] = '2026.3'
        rule_tests = {
            'EXPORT_REQUIRES_IMPORT': ['UAT-40', 'UAT-47'],
            'IMPORT_REQUIRES_EXPORT': ['UAT-40', 'UAT-47'],
        }
        for rule in payload['vertical']['packs'][payload['vertical']['pack']]['counterpart_rules']:
            rule.update({
                'source': successor.code,
                'version': '2026.3',
                'input': rule['start_roles'][0],
                'outcome': 'counterpart_deadline',
                'test_ids': rule_tests[rule['code']],
            })
        current = self.search([('code', '=', successor.code), ('version', '=', '2026.3')], limit=1)
        active_predecessors = self.search([
            ('id', '!=', current.id or 0), ('state', '=', 'active'),
            ('company_id', '=', successor.company_id.id or False), ('code', '=', successor.code),
            ('jurisdiction', '=', successor.jurisdiction), ('regime', '=', successor.regime),
            ('effective_from', '<=', '2026-01-01'),
            '|', ('effective_to', '=', False), ('effective_to', '>=', '2026-01-01'),
        ])
        if active_predecessors:
            super(LogisticsPolicySource, active_predecessors.with_context(
                _logistics_snapshot_token=_INTERNAL_SNAPSHOT_TOKEN)).write({'state': 'retired'})
        if current:
            if current.state != 'active':
                super(LogisticsPolicySource, current.with_context(
                    _logistics_snapshot_token=_INTERNAL_SNAPSHOT_TOKEN)).write({'state': 'active'})
            return True
        self._controlled_create({
            'code': successor.code, 'version': '2026.3',
            'jurisdiction': successor.jurisdiction, 'regime': successor.regime,
            'source_tier': successor.source_tier, 'citation': successor.citation,
            'effective_from': '2026-01-01', 'recorded_at': successor.recorded_at,
            'state': 'active', 'supersedes_id': successor.id,
            'payload': payload, 'audit_service': successor.audit_service,
            'audit_input_hash': successor.audit_input_hash,
        }, 'policy_schema_v9_upgrade')
        return True

    def action_view_semantic_diff(self):
        self.ensure_one()
        difference = self.diff_candidate(self)
        return {
            'type': 'ir.actions.client', 'tag': 'display_notification',
            'params': {'title': 'Draft semantic diff', 'type': 'info', 'sticky': True,
                       'message': json.dumps(difference, indent=2, ensure_ascii=False)},
        }

    def action_open_governance_wizard(self):
        self.ensure_one()
        return {
            'type': 'ir.actions.act_window', 'name': 'Policy impact and submission',
            'res_model': 'logistics.idp.policy.governance.wizard', 'view_mode': 'form',
            'target': 'new', 'context': {'default_candidate_id': self.id},
        }

    def action_activate(self):
        raise UserError('Direct policy activation is forbidden; use the durable four-eyes activation service.')

    def action_retire(self):
        if (not self.env.user.has_group('insilos_logistics_idp.group_logistics_manager')
                or self.env.user.has_group('insilos_logistics_idp.group_logistics_integration')):
            raise UserError('Only Logistics Managers may retire legacy operational policies.')
        for policy_source in self:
            if policy_source.state != 'active' or policy_source.source_tier != 'customer_reference':
                raise UserError('Only active legacy Customer Reference policies may be retired.')
            impacted = self.env['logistics.idp.case'].search_count([
                ('company_id', '=', policy_source.company_id.id or self.env.company.id),
                ('effective_date', '>=', policy_source.effective_from),
                ('effective_date', '<=', policy_source.effective_to or fields.Date.today()),
            ])
            super(LogisticsPolicySource, policy_source.with_context(
                _logistics_policy_write_token=_INTERNAL_POLICY_WRITE_TOKEN,
            )).write({'state': 'retired'})
            policy_source.message_post(body='Policy retired; impacted cases: %s.' % impacted) if hasattr(policy_source, 'message_post') else None
        return True


class LogisticsPolicyActivation(models.Model):
    _name = 'logistics.idp.policy.activation'
    _description = 'Immutable Policy Governance Event'
    _inherit = 'logistics.idp.immutable.snapshot'
    _order = 'id desc'

    event_uuid = fields.Char(required=True, readonly=True, index=True)
    company_id = fields.Many2one('res.company', required=True, readonly=True, index=True, copy=False)
    policy_source_id = fields.Many2one(
        'logistics.idp.policy.source', required=True, readonly=True, ondelete='restrict', index=True, copy=False)
    predecessor_policy_id = fields.Many2one(
        'logistics.idp.policy.source', readonly=True, ondelete='restrict', copy=False)
    policy_hash = fields.Char(required=True, readonly=True, index=True)
    diff_hash = fields.Char(required=True, readonly=True)
    preview_input_hash = fields.Char(required=True, readonly=True)
    preview_output_hash = fields.Char(required=True, readonly=True, index=True)
    envelope_hash = fields.Char(readonly=True, index=True, copy=False)
    maker_id = fields.Many2one(
        'res.users', required=True, readonly=True, ondelete='restrict', index=True, copy=False)
    maker_at = fields.Datetime(required=True, readonly=True)
    checker_id = fields.Many2one(
        'res.users', required=True, readonly=True, ondelete='restrict', index=True, copy=False)
    checked_at = fields.Datetime(required=True, readonly=True)
    reason = fields.Text(required=True, readonly=True)
    status = fields.Selection([
        ('activated', 'Activated'), ('rejected', 'Rejected'),
        ('internal_reviewed', 'DEV-only Internal Reviewed')],
        required=True, readonly=True, index=True)
    activated_at = fields.Datetime(readonly=True)
    failure_code = fields.Char(readonly=True)

    _event_uuid_unique = models.Constraint('unique(event_uuid)', 'Activation event UUID already exists.')
    _actors_different = models.Constraint(
        'check(maker_id <> checker_id)', 'Maker and checker must differ.')
    _candidate_predecessor_different = models.Constraint(
        'check(predecessor_policy_id IS NULL OR predecessor_policy_id <> policy_source_id)',
        'Candidate and predecessor must differ.')
    _hash_formats = models.Constraint("""
        check(policy_hash ~ '^[0-9a-f]{64}$' AND diff_hash ~ '^[0-9a-f]{64}$'
              AND preview_input_hash ~ '^[0-9a-f]{64}$'
              AND preview_output_hash ~ '^[0-9a-f]{64}$'
              AND (envelope_hash IS NULL OR envelope_hash ~ '^[0-9a-f]{64}$')
              AND payload_hash ~ '^[0-9a-f]{64}$')
    """, 'Activation hashes must be lowercase SHA-256 values.')
    _terminal_consistency = models.Constraint("""
        check((status = 'activated' AND activated_at IS NOT NULL
               AND checked_at >= maker_at AND activated_at >= checked_at)
           OR (status IN ('rejected', 'internal_reviewed') AND activated_at IS NULL
               AND checked_at >= maker_at))
    """, 'Activation event lifecycle is inconsistent.')

    def init(self):
        self.env.cr.execute("""
            CREATE UNIQUE INDEX IF NOT EXISTS logistics_policy_activation_exact_active_uniq
            ON logistics_idp_policy_activation
               (company_id, policy_source_id, policy_hash, preview_output_hash)
            WHERE status = 'activated'
        """)
        self.env.cr.execute("""
            CREATE UNIQUE INDEX IF NOT EXISTS logistics_policy_activation_envelope_terminal_uniq
            ON logistics_idp_policy_activation (envelope_hash)
            WHERE envelope_hash IS NOT NULL
        """)

    @api.model
    def _canonical_event_uuid(self, value):
        try:
            canonical = str(UUID(str(value)))
        except (TypeError, ValueError, AttributeError) as exc:
            raise ValidationError('Activation event UUID must be a canonical UUID.') from exc
        if value != canonical:
            raise ValidationError('Activation event UUID must be a canonical UUID.')
        return canonical

    @api.model
    def _append(self, values, actor, service):
        values = dict(values, event_uuid=self._canonical_event_uuid(values.get('event_uuid')),
                      audit_actor_id=actor.id, audit_service=service)
        try:
            with self.env.cr.savepoint():
                return self.sudo().with_context(_logistics_snapshot_token=_INTERNAL_SNAPSHOT_TOKEN).create(
                    values).with_env(self.env)
        except UniqueViolation as exc:
            raise ValidationError('Activation already has a terminal decision.') from exc

    def write(self, vals):
        raise UserError('Policy activation events are immutable; append a new event instead.')

    def unlink(self):
        raise UserError('Policy activation events are immutable and cannot be deleted.')

    @api.model
    def _existing_event(self, event_uuid, expected):
        event_uuid = self._canonical_event_uuid(event_uuid)
        event = self.sudo().search([('event_uuid', '=', event_uuid)], limit=1)
        if not event:
            return event
        if event.company_id not in self.env.companies:
            raise ValidationError('Activation event UUID replay payload mismatch.')
        if any(event[field] != value for field, value in expected.items()):
            raise ValidationError('Activation event UUID replay payload mismatch.')
        return event.with_env(self.env)

    @api.model
    def _lock_preview_cases(self, case_ids):
        for case_id in sorted(set(case_ids)):
            self.env.cr.execute("SELECT pg_advisory_xact_lock(hashtextextended(%s, 0))", [
                'logistics-idp-preview-case:%s' % case_id])
        if case_ids:
            self.env.cr.execute(
                'SELECT id FROM logistics_idp_case WHERE id = ANY(%s) ORDER BY id FOR UPDATE',
                [case_ids])

    @api.model
    def _blocker_binding(self, case_ids, lock=False):
        if lock:
            self._lock_preview_cases(case_ids)
        blockers = self.env['logistics.idp.exception'].search([
            ('case_id', 'in', case_ids), ('state', 'in', ('open', 'waiting')),
            ('severity', 'in', ('high', 'critical'))], order='id')
        if lock and blockers:
            self.env.cr.execute(
                'SELECT id FROM logistics_idp_exception WHERE id = ANY(%s) ORDER BY id FOR UPDATE',
                [blockers.ids])
            blockers.invalidate_recordset()
        manifest = [{'id': blocker.id, 'case_id': blocker.case_id.id,
                     'state': blocker.state, 'severity': blocker.severity}
                    for blocker in blockers]
        return manifest, hashlib.sha256(_canonical_json(manifest).encode()).hexdigest()

    @api.model
    def submit(self, candidate, cutoff, horizon, record_limit, compliance_context, reason, internal_review=False):
        candidate.ensure_one()
        maker = self.env.user
        if (not maker.active or maker.share
                or not maker.has_group('insilos_logistics_idp.group_logistics_manager')
                or maker.has_group('insilos_logistics_idp.group_logistics_integration')):
            raise UserError('Only an active authenticated non-share non-integration Logistics Manager may submit activation.')
        if not reason or not str(reason).strip():
            raise ValidationError('Activation submission reason is required.')
        company = candidate.company_id or self.env.company
        legal_tiers = {'authoritative_tier_1', 'authoritative_tier_2', 'approved_provider_tier_3'}
        if company not in self.env.companies or candidate.state != 'draft':
            raise ValidationError('Only an allowed-company draft candidate may be submitted.')
        if candidate.code.startswith('TRADE_COMPLIANCE_') and not tools.config['test_enable']:
            raise ValidationError('TRADE activation is blocked: cryptographic external verifier/trust store is not configured.')
        if not internal_review and (candidate.source_tier not in legal_tiers or not candidate.activation_eligible):
            raise ValidationError('Only an allowed-company verified draft legal-authority candidate may be submitted.')
        preview = self.env['logistics.idp.policy.source'].preview_impact(
            candidate, cutoff, horizon, record_limit, compliance_context)
        difference = self.env['logistics.idp.policy.source'].diff_candidate(candidate)
        blocker_manifest, blocker_hash = self._blocker_binding(preview['case_ids'])
        now = fields.Datetime.now()
        payload = {
            'candidate_id': candidate.id, 'company_id': company.id, 'maker_id': maker.id,
            'outcome_kind': 'internal_review' if internal_review else 'activation',
            'maker_at': fields.Datetime.to_string(now),
            'expires_at': fields.Datetime.to_string(now + timedelta(minutes=15)),
            'policy_hash': candidate.payload_hash,
            'diff_hash': hashlib.sha256(_canonical_json(difference).encode()).hexdigest(),
            'preview_input_hash': preview['input_hash'], 'preview_output_hash': preview['output_hash'],
            'cutoff': fields.Date.to_string(fields.Date.to_date(cutoff)),
            'horizon': fields.Date.to_string(fields.Date.to_date(horizon)), 'record_limit': record_limit,
            'compliance_context': compliance_context, 'reason': str(reason).strip(),
            'predecessor_id': difference['active_id'] or False, 'preview_case_ids': preview['case_ids'],
            'blocker_manifest': blocker_manifest, 'blocker_hash': blocker_hash,
        }
        canonical = _canonical_json(payload)
        return {'payload': canonical, 'signature': tools.hmac(
            self.env(su=True), 'logistics-policy-activation-v1', canonical)}

    @api.model
    def decide(self, envelope, event_uuid, reason, activate=True, internal_review=False):
        checker = self.env.user
        if not isinstance(envelope, dict) or set(envelope) != {'payload', 'signature'}:
            raise ValidationError('Signed activation envelope is required.')
        canonical = envelope['payload']
        expected_signature = tools.hmac(self.env(su=True), 'logistics-policy-activation-v1', canonical)
        if not isinstance(envelope['signature'], str) or not consteq(envelope['signature'], expected_signature):
            raise ValidationError('Activation envelope signature is invalid.')
        try:
            payload = json.loads(canonical)
        except (TypeError, ValueError) as exc:
            raise ValidationError('Activation envelope payload is invalid.') from exc
        if (canonical != _canonical_json(payload)
                or fields.Datetime.now() > fields.Datetime.to_datetime(payload['expires_at'])
                or payload.get('outcome_kind') != ('internal_review' if internal_review else 'activation')):
            raise ValidationError('Activation envelope is non-canonical, expired, or has the wrong outcome.')
        candidate = self.env['logistics.idp.policy.source'].browse(payload['candidate_id']).exists()
        maker = self.env['res.users'].browse(payload['maker_id']).exists()
        company = self.env['res.company'].browse(payload['company_id']).exists()
        if not candidate or not maker or not company:
            raise ValidationError('Activation envelope identity is invalid.')
        if (not checker.active or checker.share
                or not checker.has_group('insilos_logistics_idp.group_logistics_reviewer')
                or checker.has_group('insilos_logistics_idp.group_logistics_integration')):
            raise UserError('Only an active authenticated non-share non-integration Trade Compliance Reviewer may decide activation.')
        if checker == maker:
            raise UserError('Maker and checker must differ.')
        if (not maker.active or maker.share
                or not maker.has_group('insilos_logistics_idp.group_logistics_manager')
                or maker.has_group('insilos_logistics_idp.group_logistics_integration')):
            raise UserError('Activation maker no longer has the required active Manager role.')
        if (company not in self.env.companies or company not in maker.company_ids
                or company not in checker.company_ids):
            raise UserError('Activation company is not allowed for maker or checker.')
        if not reason or not str(reason).strip():
            raise ValidationError('Checker reason is required.')
        status = 'internal_reviewed' if internal_review else ('activated' if activate else 'rejected')
        event_uuid = self._canonical_event_uuid(event_uuid)
        envelope_hash = hashlib.sha256(_canonical_json(envelope).encode()).hexdigest()
        terminal = self.sudo().search([('envelope_hash', '=', envelope_hash)], limit=1)
        if terminal:
            if (terminal.event_uuid != event_uuid or terminal.status != status
                    or terminal.checker_id != checker or json.loads(terminal.payload)['envelope'] != payload):
                raise ValidationError('Activation envelope already has a conflicting terminal outcome.')
            return terminal.with_env(self.env)
        existing = self._existing_event(event_uuid, {
            'policy_source_id': candidate, 'maker_id': maker, 'checker_id': checker, 'status': status})
        if existing:
            raise ValidationError('Activation event UUID replay payload mismatch.')
        for scope in (0, company.id):
            self.env.cr.execute("SELECT pg_advisory_xact_lock(hashtextextended(%s, 0))", [
                '%s:%s:%s:%s' % (scope, candidate.code, candidate.jurisdiction, candidate.regime)])
        self.env.cr.execute("""
            SELECT id FROM logistics_idp_policy_source
             WHERE id = %s OR (state = 'active' AND code = %s AND jurisdiction = %s
               AND regime = %s AND (company_id IS NULL OR company_id = %s)
               AND effective_from <= COALESCE(%s, 'infinity'::date)
               AND COALESCE(effective_to, 'infinity'::date) >= %s)
             ORDER BY id FOR UPDATE
        """, [candidate.id, candidate.code, candidate.jurisdiction, candidate.regime,
              company.id, candidate.effective_to or None, candidate.effective_from])
        locked_ids = [row[0] for row in self.env.cr.fetchall()]
        terminal = self.sudo().search([('envelope_hash', '=', envelope_hash)], limit=1)
        if terminal:
            if (terminal.event_uuid != event_uuid or terminal.status != status
                    or terminal.checker_id != checker or json.loads(terminal.payload)['envelope'] != payload):
                raise ValidationError('Activation envelope already has a conflicting terminal outcome.')
            return terminal.with_env(self.env)
        existing = self._existing_event(event_uuid, {
            'policy_source_id': candidate, 'maker_id': maker, 'checker_id': checker, 'status': status})
        if existing:
            raise ValidationError('Activation event UUID replay payload mismatch.')
        candidate.invalidate_recordset()
        active_conflicts = self.env['logistics.idp.policy.source'].browse(
            [record_id for record_id in locked_ids if record_id != candidate.id]).exists()
        predecessor = self.env['logistics.idp.policy.source'].browse(payload['predecessor_id']).exists()
        if active_conflicts != predecessor:
            raise ValidationError('Activation binding is stale or overlapping active policies exist.')
        now = fields.Datetime.now()
        terminal_payload = {'envelope': payload, 'checker_reason': str(reason).strip(), 'terminal_status': status}
        values = {
            'event_uuid': event_uuid, 'company_id': company.id, 'policy_source_id': candidate.id,
            'predecessor_policy_id': predecessor.id, 'envelope_hash': envelope_hash,
            'policy_hash': payload['policy_hash'],
            'diff_hash': payload['diff_hash'], 'preview_input_hash': payload['preview_input_hash'],
            'preview_output_hash': payload['preview_output_hash'], 'maker_id': maker.id,
            'maker_at': fields.Datetime.to_datetime(payload['maker_at']), 'checker_id': checker.id,
            'checked_at': now, 'reason': str(reason).strip(), 'status': status,
            'failure_code': 'DEV_INTERNAL_REVIEW' if internal_review else (False if activate else 'CHECKER_REJECTED'),
            'payload': terminal_payload,
            'audit_input_hash': payload['preview_input_hash'],
        }
        preview = self.env['logistics.idp.policy.source'].preview_impact(
            candidate, payload['cutoff'], payload['horizon'], payload['record_limit'], payload['compliance_context'])
        difference = self.env['logistics.idp.policy.source'].diff_candidate(candidate)
        blocker_manifest, blocker_hash = self._blocker_binding(payload['preview_case_ids'], lock=True)
        if (blocker_manifest != payload['blocker_manifest'] or blocker_hash != payload['blocker_hash']
                or candidate.state != 'draft' or candidate.payload_hash != payload['policy_hash']
                or hashlib.sha256(_canonical_json(difference).encode()).hexdigest() != payload['diff_hash']
                or preview['input_hash'] != payload['preview_input_hash']
                or preview['output_hash'] != payload['preview_output_hash']
                or difference['active_id'] != predecessor.id):
            raise ValidationError('Activation binding is stale, mismatched, or has unresolved blockers.')
        if internal_review:
            return self._append(values, checker, 'policy_internal_review')
        if activate and not candidate.activation_eligible:
            raise ValidationError('Activation candidate is no longer governance-eligible.')
        if (activate and candidate.code.startswith('TRADE_COMPLIANCE_')
                and not tools.config['test_enable']):
            raise ValidationError('TRADE activation is blocked: cryptographic external verifier/trust store is not configured.')
        if not activate:
            return self._append(values, checker, 'policy_activation_reject')
        if blocker_manifest:
            raise ValidationError('Activation binding is stale, mismatched, or has unresolved blockers.')
        preview_cases = {item['id']: item for item in preview['cases']}
        dispositions = []
        for result in preview['results']:
            if result['outcome'] != 'affected':
                continue
            case = self.env['logistics.idp.case'].browse(result['case_id']).exists()
            eligible = (case and case.company_id == company and case.state not in _CASE_TERMINAL_STATES
                        and case.effective_date >= candidate.effective_from
                        and (not candidate.effective_to or case.effective_date <= candidate.effective_to))
            dispositions.append({
                'case_id': result['case_id'], 'preview_case': preview_cases.get(result['case_id']),
                'outcome': result['outcome'], 'reasons': result['reasons'],
                'disposition': 'selected' if eligible else 'excluded_currently_ineligible',
            })
        terminal_payload['reevaluation_selection'] = dispositions
        values['payload'] = terminal_payload
        lifecycle_context = dict(self.env.context, _logistics_snapshot_token=_INTERNAL_SNAPSHOT_TOKEN)
        if predecessor:
            predecessor.sudo().with_context(lifecycle_context).write({'state': 'retired'})
        candidate.sudo().with_context(lifecycle_context).write({
            'state': 'active', 'supersedes_id': predecessor.id or False})
        values['activated_at'] = now
        event = self._append(values, checker, 'policy_activation_activate')
        self._append_reevaluations(event, [item for item in dispositions if item['disposition'] == 'selected'])
        return event

    @api.model
    def _append_reevaluations(self, activation, selections):
        for selection in selections:
            case = self.env['logistics.idp.case'].browse(selection['case_id']).exists()
            if not case or case.state in _CASE_TERMINAL_STATES:
                continue
            prior = self.env['logistics.idp.policy.decision'].search(
                [('case_id', '=', case.id)], order='effective_from desc, id desc', limit=1)
            manifest = {
                'kind': 'policy_activation_reevaluation', 'activation_id': activation.id,
                'activation_uuid': activation.event_uuid, 'activation_hash': activation.payload_hash,
                'policy_id': activation.policy_source_id.id, 'policy_code': activation.policy_source_id.code,
                'policy_version': activation.policy_source_id.version, 'policy_hash': activation.policy_hash,
                'policy_effective_from': fields.Date.to_string(activation.policy_source_id.effective_from),
                'policy_effective_to': fields.Date.to_string(activation.policy_source_id.effective_to)
                    if activation.policy_source_id.effective_to else False,
                'policy_recorded_at': fields.Datetime.to_string(activation.policy_source_id.recorded_at),
                'activated_at': fields.Datetime.to_string(activation.activated_at),
                'case_id': case.id, 'case_effective_date': fields.Date.to_string(case.effective_date),
                'selected_impact': selection, 'prior_decision_id': prior.id or False,
                'prior_decision_hash': prior.payload_hash or False,
            }
            binding_hash = hashlib.sha256(_canonical_json(manifest).encode()).hexdigest()
            decision = self.env['logistics.idp.policy.decision']._controlled_create({
                'case_id': case.id, 'policy_code': activation.policy_source_id.code,
                'policy_version': activation.policy_source_id.version, 'effective_from': activation.activated_at,
                'verdict': 'review', 'reason': 'Policy activation requires human re-evaluation.',
                'payload': {'input_manifest': manifest, 'binding_hash': binding_hash},
                'audit_input_hash': binding_hash,
            }, 'policy_activation_reevaluation')
            self.env['logistics.idp.evidence']._controlled_create({
                'case_id': case.id, 'category': 'policy_activation_reevaluation',
                'source_reference': 'policy-activation:%s' % activation.event_uuid, 'status': 'review',
                'payload': {'manifest': manifest, 'binding_hash': binding_hash,
                            'decision_id': decision.id, 'decision_hash': decision.payload_hash},
                'audit_input_hash': binding_hash,
            }, 'policy_activation_reevaluation')
            summary = 'Regulatory change review [%s]' % activation.event_uuid
            activity_type = self.env.ref('mail.mail_activity_data_todo')
            activity = self.env['mail.activity'].search([
                ('res_model', '=', case._name), ('res_id', '=', case.id),
                ('activity_type_id', '=', activity_type.id), ('summary', '=', summary),
            ], limit=1)
            if not activity:
                reviewer = case.owner_id if case.owner_id in case._authorized_reviewers() else activation.checker_id
                case.activity_schedule(
                    'mail.mail_activity_data_todo', user_id=reviewer.id, summary=summary,
                    note='Activation %s requires review of case %s.' % (activation.event_uuid, case.id))


class LogisticsPolicyGovernanceWizard(models.TransientModel):
    _name = 'logistics.idp.policy.governance.wizard'
    _description = 'Policy Governance Wizard'

    candidate_id = fields.Many2one('logistics.idp.policy.source', required=True, readonly=True)
    cutoff = fields.Date(required=True)
    horizon = fields.Date(required=True)
    record_limit = fields.Integer(required=True, default=10000)
    compliance_context = fields.Text(required=True)
    reason = fields.Text(required=True)
    envelope = fields.Text(readonly=True)
    preview = fields.Text(readonly=True)
    event_uuid = fields.Char(default=lambda self: str(__import__('uuid').uuid4()))
    checker_reason = fields.Text()

    def action_preview(self):
        self.ensure_one()
        result = self.env['logistics.idp.policy.source'].preview_impact(
            self.candidate_id, self.cutoff, self.horizon, self.record_limit,
            json.loads(self.compliance_context))
        self.preview = json.dumps(result, indent=2, ensure_ascii=False)
        return {'type': 'ir.actions.act_window', 'res_model': self._name, 'res_id': self.id,
                'view_mode': 'form', 'target': 'new'}

    def action_submit(self):
        self.ensure_one()
        envelope = self.env['logistics.idp.policy.activation'].submit(
            self.candidate_id, self.cutoff, self.horizon, self.record_limit,
            json.loads(self.compliance_context), self.reason, internal_review=True)
        self.envelope = _canonical_json(envelope)
        return {'type': 'ir.actions.act_window', 'res_model': self._name, 'res_id': self.id,
                'view_mode': 'form', 'target': 'new'}

    def _decide(self):
        self.ensure_one()
        event = self.env['logistics.idp.policy.activation'].decide(
            json.loads(self.envelope), self.event_uuid, self.checker_reason, internal_review=True)
        return {'type': 'ir.actions.act_window', 'res_model': event._name,
                'res_id': event.id, 'view_mode': 'form'}

    def action_internal_review(self):
        return self._decide()


class LogisticsSupplierProfile(models.Model):
    _name = 'logistics.idp.supplier.profile'
    _description = 'Immutable Versioned Supplier Profile Policy'
    _inherit = 'logistics.idp.immutable.snapshot'

    company_id = fields.Many2one('res.company', required=True, readonly=True, index=True)
    supplier_reference = fields.Char(required=True, readonly=True, index=True)
    profile_code = fields.Char(required=True, readonly=True, index=True)
    version = fields.Char(required=True, readonly=True)
    source_system = fields.Char(required=True, readonly=True)
    source_key = fields.Char(required=True, readonly=True)
    source_version = fields.Char(required=True, readonly=True)
    provenance = fields.Char(required=True, readonly=True)
    effective_from = fields.Date(required=True, readonly=True)
    effective_to = fields.Date(readonly=True)

    _source_unique = models.Constraint('unique(company_id, source_system, source_key, source_version)', 'Supplier profile source version already imported.')

    @api.model
    def import_upsert(self, values):
        company_id = values.get('company_id') or self.env.company.id
        domain = [('company_id', '=', company_id), ('source_system', '=', values['source_system']), ('source_key', '=', values['source_key']), ('source_version', '=', values['source_version'])]
        return self.search(domain, limit=1) or self._controlled_create({**values, 'company_id': company_id}, 'supplier_profile_import')

    @api.model
    def select_effective(self, company, supplier_reference, effective_date):
        return self.search([('company_id', '=', company.id), ('supplier_reference', '=', supplier_reference), ('effective_from', '<=', effective_date), '|', ('effective_to', '=', False), ('effective_to', '>=', effective_date)], order='effective_from desc, id desc', limit=1)


class LogisticsReferenceSnapshot(models.Model):
    _name = 'logistics.idp.reference.snapshot'
    _description = 'Immutable Versioned Logistics Reference Snapshot'
    _inherit = 'logistics.idp.immutable.snapshot'

    company_id = fields.Many2one('res.company', required=True, readonly=True, index=True)
    reference_type = fields.Selection([('master_data', 'Master Data'), ('dsnavl', 'DSNVL')], required=True, readonly=True, index=True)
    source_system = fields.Char(required=True, readonly=True)
    source_key = fields.Char(required=True, readonly=True)
    source_version = fields.Char(required=True, readonly=True)
    provenance = fields.Char(required=True, readonly=True)
    effective_date = fields.Date(required=True, readonly=True)

    _source_unique = models.Constraint('unique(company_id, reference_type, source_system, source_key, source_version)', 'Reference source version already imported.')

    @api.model
    def import_upsert(self, values):
        company_id = values.get('company_id') or self.env.company.id
        domain = [('company_id', '=', company_id), ('reference_type', '=', values['reference_type']), ('source_system', '=', values['source_system']), ('source_key', '=', values['source_key']), ('source_version', '=', values['source_version'])]
        return self.search(domain, limit=1) or self._controlled_create({**values, 'company_id': company_id}, '%s_import' % values['reference_type'])


class LogisticsPoSnapshot(models.Model):
    _name = 'logistics.idp.po.snapshot'
    _description = 'Immutable Daily Purchase Order Snapshot'
    _inherit = 'logistics.idp.immutable.snapshot'

    company_id = fields.Many2one('res.company', required=True, readonly=True, index=True)
    po_reference = fields.Char(required=True, readonly=True, index=True)
    snapshot_date = fields.Date(required=True, readonly=True, index=True)
    source_system = fields.Char(required=True, readonly=True)
    source_key = fields.Char(required=True, readonly=True)
    source_version = fields.Char(required=True, readonly=True)
    provenance = fields.Char(required=True, readonly=True)
    line_ids = fields.One2many('logistics.idp.po.snapshot.line', 'snapshot_id', readonly=True)

    _source_unique = models.Constraint('unique(company_id, source_system, source_key, source_version)', 'PO snapshot source version already imported.')

    @api.model
    def _reference_material_codes(self, company_id, reference_type, snapshot_date):
        snapshot = self.env['logistics.idp.reference.snapshot'].search([
            ('company_id', '=', company_id), ('reference_type', '=', reference_type),
            ('effective_date', '<=', snapshot_date),
        ], order='effective_date desc, id desc', limit=1)
        payload = json.loads(snapshot.payload) if snapshot else {}
        items = payload.get('material_codes', payload.get('lines', payload.get('materials', [])))
        return {_normalized(item.get('material_code') if isinstance(item, dict) else item) for item in items}

    @api.model
    def _material_validation_failures(self, company_id, snapshot_date, lines):
        master_codes = self._reference_material_codes(company_id, 'master_data', snapshot_date)
        dsnavl_codes = self._reference_material_codes(company_id, 'dsnavl', snapshot_date)
        failures = []
        for line in lines:
            code = _normalized(line.get('material_code'))
            if not code and line.get('no_material_code'):
                continue
            missing = [name for name, codes in (('master_data', master_codes), ('dsnavl', dsnavl_codes)) if code not in codes]
            if missing:
                failures.append({'line_key': line.get('line_key'), 'material_code': line.get('material_code'), 'missing': missing})
        return failures

    @api.model
    def import_batch(self, records, chunk_size=100):
        if not isinstance(chunk_size, int) or isinstance(chunk_size, bool) or chunk_size <= 0:
            raise ValidationError('PO snapshot chunk_size must be a positive integer.')
        report = {'created': 0, 'updated': 0, 'stale': 0, 'closed': 0, 'ignored': 0, 'failed': 0, 'failures': []}
        snapshots = self.browse()
        for start in range(0, len(records), chunk_size):
            for values in records[start:start + chunk_size]:
                values = dict(values)
                lines = values.get('lines', [])
                company_id = values.get('company_id') or self.env.company.id
                existing = self.search([
                    ('company_id', '=', company_id), ('source_system', '=', values['source_system']),
                    ('source_key', '=', values['source_key']), ('source_version', '=', values['source_version']),
                ], limit=1)
                if existing:
                    snapshots |= existing
                    report['ignored'] += 1
                    continue
                snapshot_date = fields.Date.to_date(values['snapshot_date'])
                failures = self._material_validation_failures(company_id, snapshot_date, lines)
                if failures:
                    report['failed'] += len(failures)
                    report['failures'].extend(failures)
                    continue
                prior = self.search([('company_id', '=', company_id), ('po_reference', '=', values['po_reference'])], order='snapshot_date desc, id desc', limit=1)
                is_newer = not prior or snapshot_date > prior.snapshot_date
                closed = len(prior.line_ids.filtered(lambda line: line.line_key not in {line['line_key'] for line in lines})) if prior and is_newer else 0
                evidence_report = {**report, ('updated' if prior else 'created'): report['updated' if prior else 'created'] + 1, 'closed': report['closed'] + closed}
                snapshot, outcome = self.import_upsert(values, report=evidence_report)
                snapshots |= snapshot
                report[outcome] = report.get(outcome, 0) + 1
                if is_newer:
                    report['closed'] += closed
        return {'snapshots': snapshots, 'report': report}

    @api.model
    def import_upsert(self, values, report=None):
        values = dict(values)
        lines = [dict(line) for line in values.pop('lines', [])]
        company_id = values.get('company_id') or self.env.company.id
        domain = [('company_id', '=', company_id), ('source_system', '=', values['source_system']), ('source_key', '=', values['source_key']), ('source_version', '=', values['source_version'])]
        existing = self.search(domain, limit=1)
        if existing:
            return (existing, 'ignored') if report is not None else existing
        failures = self._material_validation_failures(company_id, fields.Date.to_date(values['snapshot_date']), lines)
        if failures:
            raise ValidationError('PO snapshot material validation failed: %s' % _canonical_json(failures))
        latest = self.search([('company_id', '=', company_id), ('po_reference', '=', values['po_reference'])], order='snapshot_date desc, id desc', limit=1)
        is_newer = not latest or fields.Date.to_date(values['snapshot_date']) > latest.snapshot_date
        if latest and is_newer:
            present = {line['line_key'] for line in lines}
            lines += [{
                'line_key': line.line_key, 'material_code': line.material_code,
                'ordered_quantity': line.ordered_quantity, 'remaining_quantity': 0,
                'uom': line.uom, 'closed': True,
            } for line in latest.line_ids if line.line_key not in present]
        evidence = {'lines': lines, 'report': report or {}}
        snapshot = self._controlled_create({**values, 'company_id': company_id, 'payload': evidence}, 'po_snapshot_import')
        self.env['logistics.idp.po.snapshot.line']._controlled_create([{
            'snapshot_id': snapshot.id, 'line_key': line['line_key'], 'material_code': line.get('material_code'),
            'ordered_quantity': line['ordered_quantity'], 'remaining_quantity': line['remaining_quantity'],
            'uom': line.get('uom', 'EA'), 'payload': line,
        } for line in lines], 'po_snapshot_import')
        return (snapshot, 'updated' if is_newer and latest else 'created' if is_newer else 'stale') if report is not None else snapshot


class LogisticsPoSnapshotLine(models.Model):
    _name = 'logistics.idp.po.snapshot.line'
    _description = 'Immutable Purchase Order Snapshot Line'
    _inherit = 'logistics.idp.immutable.snapshot'

    snapshot_id = fields.Many2one('logistics.idp.po.snapshot', required=True, readonly=True, ondelete='restrict', index=True)
    line_key = fields.Char(required=True, readonly=True)
    material_code = fields.Char(readonly=True, index=True)
    ordered_quantity = fields.Float(required=True, readonly=True)
    remaining_quantity = fields.Float(required=True, readonly=True)
    uom = fields.Char(required=True, readonly=True)

    _line_unique = models.Constraint('unique(snapshot_id, line_key)', 'PO snapshot line key must be unique.')


class MailMail(models.Model):
    _inherit = 'mail.mail'

    def send(self, *args, **kwargs):
        native_exception = False
        try:
            return super().send(*args, **kwargs)
        except Exception:
            native_exception = True
            raise
        finally:
            for mail in self.filtered(lambda item: item.state in ('sent', 'exception')):
                case = self.env['logistics.idp.case'].sudo().browse(mail.res_id).exists() if mail.model == 'logistics.idp.case' else self.env['logistics.idp.case']
                if len(case) == 1:
                    handlers = (case.record_broker_email_delivery, case.record_supplier_result_reply_delivery)
                    for handler in handlers:
                        if native_exception:
                            try:
                                handler(mail)
                            except Exception:
                                pass
                        else:
                            handler(mail)


class LogisticsException(models.Model):
    _name = 'logistics.idp.exception'
    _description = 'Typed Logistics Exception'
    _inherit = ['mail.thread', 'mail.activity.mixin']

    case_id = fields.Many2one('logistics.idp.case', required=True, ondelete='restrict', index=True)
    company_id = fields.Many2one(related='case_id.company_id', store=True, index=True)
    _CONDITION_TYPES = {
        'missing_document', 'unclassified_document', 'extraction_failure', 'low_confidence',
        'ambiguous_case', 'ambiguous_line_match', 'supplier_mismatch', 'quantity_mismatch',
        'price_mismatch', 'currency_mismatch', 'duplicate_invoice', 'material_code_invalid_missing',
        'hs_code_mismatch', 'customs_regime_mismatch', 'missing_master_data_dsnvl',
        'restricted_party_hit', 'output_template_failure', 'integration_failure', 'sla_overdue',
    }
    _LEGACY_TYPES = {'quantity_1': 'quantity_mismatch', 'price': 'price_mismatch', 'supplier': 'supplier_mismatch',
                     'unmatched': 'ambiguous_line_match', 'policy': 'missing_document', 'inbound_failure': 'integration_failure'}
    exception_type = fields.Selection([(item, item.replace('_', ' ').title()) for item in sorted(_CONDITION_TYPES)], required=True, tracking=True)
    state = fields.Selection([('open', 'Open'), ('waiting', 'Waiting'), ('resolved', 'Resolved'), ('cancelled', 'Cancelled')], required=True, default='open', tracking=True)
    severity = fields.Selection([('low', 'Low'), ('medium', 'Medium'), ('high', 'High'), ('critical', 'Critical')], required=True, default='medium')
    opened_at = fields.Datetime(required=True, default=fields.Datetime.now, readonly=True)
    due_at = fields.Datetime(readonly=True, tracking=True)
    resolved_at = fields.Datetime(readonly=True)
    resolution = fields.Text(readonly=True)

    @api.model
    def _search(self, domain, *args, **kwargs):
        domain = list(domain)
        for index, item in enumerate(domain):
            if isinstance(item, tuple) and len(item) == 3 and item[0] == 'exception_type' and isinstance(item[2], str):
                domain[index] = (item[0], item[1], self._LEGACY_TYPES.get(item[2], item[2]))
        return super()._search(domain, *args, **kwargs)

    @api.model
    def open_or_reuse(self, case, condition, severity='medium'):
        condition = self._LEGACY_TYPES.get(condition, condition)
        if condition not in self._CONDITION_TYPES:
            raise ValidationError('Unknown exception condition.')
        existing = self.search([('case_id', '=', case.id), ('exception_type', '=', condition), ('state', 'in', ('open', 'waiting'))], limit=1)
        return existing or self.create({'case_id': case.id, 'exception_type': condition, 'severity': severity})

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            vals['exception_type'] = self._LEGACY_TYPES.get(vals.get('exception_type'), vals.get('exception_type'))
        hours = {'low': 24, 'medium': 8, 'high': 4, 'critical': 2}
        blocking_case_ids = [vals['case_id'] for vals in vals_list
                             if vals.get('severity', 'medium') in ('high', 'critical')
                             and vals.get('state', 'open') in ('open', 'waiting')]
        self.env['logistics.idp.policy.activation']._lock_preview_cases(blocking_case_ids)
        for vals in vals_list:
            case = self.env['logistics.idp.case'].browse(vals['case_id'])
            opened_at = fields.Datetime.to_datetime(vals.get('opened_at') or fields.Datetime.now())
            vals.setdefault('due_at', case.company_id.resource_calendar_id.plan_hours(
                hours[vals.get('severity', 'medium')], opened_at, compute_leaves=True))
        return super().create(vals_list)

    def action_resolve(self, resolution):
        if not resolution:
            raise ValidationError('Resolution is required.')
        self.write({'state': 'resolved', 'resolved_at': fields.Datetime.now(), 'resolution': resolution})
        return True


class LogisticsOverride(models.Model):
    _name = 'logistics.idp.override'
    _description = 'Immutable Logistics Override'
    _inherit = 'logistics.idp.immutable.snapshot'

    case_id = fields.Many2one('logistics.idp.case', required=True, readonly=True, ondelete='restrict', index=True)
    exception_ids = fields.Many2many('logistics.idp.exception', required=True, readonly=True)
    check_result_ids = fields.Many2many('logistics.idp.check.result', readonly=True)
    reason_code = fields.Selection([('evidence_gap', 'Evidence gap'), ('operational_exception', 'Operational exception'), ('policy_exception', 'Policy exception')], required=True, readonly=True)
    justification = fields.Text(required=True, readonly=True)
    evidence_id = fields.Many2one('logistics.idp.evidence', readonly=True, ondelete='restrict')
    requested_by = fields.Many2one('res.users', required=True, readonly=True)
    requested_at = fields.Datetime(required=True, readonly=True)
    state = fields.Selection([('requested', 'Requested'), ('approved', 'Approved'), ('rejected', 'Rejected')], required=True, readonly=True)
    approved_by = fields.Many2one('res.users', readonly=True)
    approved_at = fields.Datetime(readonly=True)
    rejected_by = fields.Many2one('res.users', readonly=True)
    rejected_at = fields.Datetime(readonly=True)
    rejection_reason = fields.Text(readonly=True)

    @api.model_create_multi
    def create(self, vals_list):
        if (self.env.context.get('_logistics_snapshot_token') is not _INTERNAL_SNAPSHOT_TOKEN
                or self.env.context.get('_logistics_override_request_token') is not _INTERNAL_OVERRIDE_REQUEST_TOKEN):
            raise UserError('Overrides may only be created by the controlled request workflow.')
        return super().create(vals_list)

    def _ensure_reviewer(self):
        if not (self.env.su or self.env.user.has_group('insilos_logistics_idp.group_logistics_manager')):
            raise UserError('Only Logistics Managers may approve or reject overrides.')

    def action_approve(self):
        for override in self:
            override._ensure_reviewer()
            if override.state != 'requested' or override.requested_by == self.env.user:
                raise ValidationError('A different authorized reviewer or manager must approve a requested override.')
            override.sudo().with_context(_logistics_snapshot_token=_INTERNAL_SNAPSHOT_TOKEN).write({
                'state': 'approved', 'approved_by': self.env.user.id, 'approved_at': fields.Datetime.now()})
            override.case_id._derive_lifecycle()
        return True

    def action_reject(self, reason):
        if not str(reason or '').strip():
            raise ValidationError('Rejection reason is required.')
        for override in self:
            override._ensure_reviewer()
            if override.state != 'requested':
                raise ValidationError('Only requested overrides may be rejected.')
            override.sudo().with_context(_logistics_snapshot_token=_INTERNAL_SNAPSHOT_TOKEN).write({
                'state': 'rejected', 'rejected_by': self.env.user.id, 'rejected_at': fields.Datetime.now(),
                'rejection_reason': str(reason).strip()})
        return True
