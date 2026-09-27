import ast
import base64
import hashlib
import importlib.util
import io
import json
import re
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
REPOSITORY = Path(__file__).resolve().parents[4]
POLICY = json.loads((ROOT / "config/compliance_policy.json").read_text())
VN_LEGAL_PACK = json.loads((ROOT / "config/vn_legal_candidate_pack.json").read_text())
CASES = json.loads((Path(__file__).parent / "golden/sanitized_cases.json").read_text())
ONBOARDING = json.loads((ROOT / "config/onboarding_profile.json").read_text())
RELEASE = json.loads((ROOT / "config/release_package.json").read_text())
CLASSIFIER_PATH = Path(__file__).parent / "golden/build_sanitized_corpus.py"
SPEC = importlib.util.spec_from_file_location("build_sanitized_corpus", CLASSIFIER_PATH)
CLASSIFIER = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(CLASSIFIER)


class HardeningContractTest(unittest.TestCase):
    def test_corpus_fail_to_review(self):
        self.assertFalse(CASES["metadata"]["formal_golden_acceptance"])
        for case in CASES["cases"]:
            self.assertIn(case["oracle"]["verdict"], ("pass", "allowed_failure"))
            self.assertTrue(case["oracle"]["checks"])

    def test_rule_traceability_mapping_and_coverage_gate(self):
        self.assertEqual({case["trace_category"] for case in CASES["cases"]}, CLASSIFIER.TRACE_CATEGORIES)
        CLASSIFIER.verify(CASES)
        original = CLASSIFIER.load_workbook
        class BrokenWorkbook:
            def __getitem__(self, name):
                return type("Sheet", (), {"cell": lambda self, row, column: type("Cell", (), {"value": "broken"})()})()
        CLASSIFIER.load_workbook = lambda *args, **kwargs: BrokenWorkbook() if not kwargs.get("data_only") else original(*args, **kwargs)
        try:
            with self.assertRaisesRegex(ValueError, "Requirement Coverage formula differs"):
                CLASSIFIER.workbook_contract()
        finally:
            CLASSIFIER.load_workbook = original

    def test_authoritative_srs_inventory_and_false_pass_gate(self):
        trace = CASES["authoritative_srs_traceability"]
        rebuilt = CLASSIFIER.authoritative_srs_inventory(CASES["cases"], CASES["executor_evidence"]["test_ids"])
        self.assertEqual(trace, rebuilt)
        self.assertEqual(trace["authority_sha256"], hashlib.sha256(CLASSIFIER.SRS.read_bytes()).hexdigest())
        self.assertGreater(trace["coverage_counts"]["explicitly_numbered"], 150)
        self.assertGreater(trace["coverage_counts"]["stable_unnumbered"], 50)
        expected_count = sum(item["coverage"] == "covered" for item in trace["items"])
        self.assertEqual(trace["coverage_counts"]["capability_specific_covered"], expected_count)
        self.assertEqual(trace["coverage_counts"]["achieved"], expected_count)
        workspaces = [item for item in trace["items"] if item["srs_id"] in CLASSIFIER.IMPLEMENTATION_EVIDENCE and item["srs_id"] not in CLASSIFIER.CAPABILITY_EXECUTORS]
        self.assertEqual({item["evidence_status"] for item in workspaces}, {"partial"})
        self.assertTrue(all(item["implementation_symbols"] and not item["executable_test_ids"] for item in workspaces))
        explicit = {item["srs_id"] for item in trace["items"] if CLASSIFIER.SRS_EXPLICIT_ID.match(item["srs_id"])}
        source_explicit = set(re.findall(r"(?m)^#{2,4} ((?:FR|AI|RULE|EXC|DUP|SLA|RPT|INT|SEC|NFR|UAT)-\d+[A-Z]?|VA-\d+)\b", CLASSIFIER.SRS.read_text()))
        self.assertEqual(explicit, source_explicit)
        generic = next(item for item in trace["items"] if item["corpus_ids"] and not item["executable_test_ids"])
        self.assertEqual((generic["evidence_status"], generic["coverage"], generic["executable_test_ids"]),
                         ("partial", "not_covered", []))
        forged = dict(generic, evidence_status="achieved", coverage="covered")
        self.assertFalse(forged["executable_test_ids"], "generic case must not become capability PASS")

    def test_capability_executor_registry_rejects_invalid_missing_and_generic_mappings(self):
        validated = CLASSIFIER.validate_capability_executors()
        self.assertEqual(len(validated), len(CLASSIFIER.CAPABILITY_EXECUTORS))
        original = dict(CLASSIFIER.CAPABILITY_EXECUTORS)
        invalid = dict(original)
        invalid.pop("SEC-003")
        with self.assertRaisesRegex(ValueError, "executor allowlist must contain exactly the reviewed"):
            CLASSIFIER.validate_capability_executors(invalid)
        generic = dict(original)
        generic["SEC-003"] = ("development_uat_harness.*", "models/logistics_idp.py", ["ImmutableSnapshot"], "aggregate")
        with self.assertRaisesRegex(ValueError, "generic, wildcard"):
            CLASSIFIER.validate_capability_executors(generic)
        missing = dict(original)
        missing["SEC-003"] = ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_missing", "models/logistics_idp.py", ["ImmutableSnapshot"], "missing")
        with self.assertRaisesRegex(ValueError, "class/method missing"):
            CLASSIFIER.validate_capability_executors(missing)
        missing_symbol = dict(original)
        missing_symbol["SEC-003"] = (original["SEC-003"][0], "models/logistics_idp.py", ["MissingSymbol"], "missing")
        with self.assertRaisesRegex(ValueError, "implementation symbol missing"):
            CLASSIFIER.validate_capability_executors(missing_symbol)
        implementation = dict(CLASSIFIER.IMPLEMENTATION_EVIDENCE)
        invalid_implementation = dict(implementation)
        invalid_implementation["SRS-SECTION-7-3-LOGISTICS-CASE-WORKSPACE"] = (
            "models/logistics_idp.py", ["LogisticsCase.*"], "generic")
        with self.assertRaisesRegex(ValueError, "wildcard or method-level"):
            CLASSIFIER.validate_implementation_evidence(invalid_implementation)
        invalid_implementation["SRS-SECTION-7-3-LOGISTICS-CASE-WORKSPACE"] = (
            "models/logistics_idp.py", ["MissingSymbol"], "missing")
        with self.assertRaisesRegex(ValueError, "implementation evidence symbol missing"):
            CLASSIFIER.validate_implementation_evidence(invalid_implementation)

    def test_source_identity_is_idempotent_and_version_sensitive(self):
        item = CASES["cases"][0]["artifact"]
        content = base64.b64decode(item["content"])
        self.assertEqual(hashlib.sha256(content).hexdigest(), item["sha256"])
        self.assertNotEqual(item["sha256"], hashlib.sha256(content + b"-V2").hexdigest())

    def test_audit_and_completion_contract(self):
        for case in CASES["cases"]:
            self.assertTrue(case["expected_result"])
            self.assertTrue(case["preconditions"])
            self.assertTrue(case["artifact"]["sha256"])
            self.assertTrue(case["oracle"]["lineage"])

    def test_upgrade_lifecycle_has_single_versioned_authority(self):
        manifest = (ROOT / '__manifest__.py').read_text()
        module_init = (ROOT / '__init__.py').read_text()
        migration = (ROOT / 'migrations/19.0.1.1.0/post-migrate.py').read_text()
        self.assertNotIn('logistics_idp_policy_upgrade.xml', manifest)
        self.assertFalse((ROOT / 'data/logistics_idp_policy_upgrade.xml').exists())
        self.assertIn("'post_init_hook': 'post_init_hook'", manifest)
        self.assertIn('_logistics_policy_loader_token=_INTERNAL_POLICY_LOADER_TOKEN', module_init)
        self.assertIn('._upgrade_effective_policy_history()', module_init)
        self.assertIn("env['logistics.idp.case']._upgrade_migration_service_principal()", module_init)
        self.assertNotIn('_backfill_review_workflow', module_init)
        for method in ('_upgrade_effective_policy_history', '_backfill_review_workflow',
                       '_upgrade_migration_service_principal'):
            self.assertEqual(migration.count(method), 1)

    def test_notification_policy_has_no_sender_path(self):
        sources = '\n'.join(path.read_text() for path in (ROOT / 'models').glob('*.py'))
        demo_policies = '\n'.join(path.read_text() for path in (ROOT / 'data').glob('logistics_idp_policy*.xml'))
        self.assertIn('"send_enabled":false', demo_policies)
        self.assertNotIn('"send_enabled":true', demo_policies)
        self.assertNotRegex(sources, r'\.send_mail\s*\([^)]*force_send\s*=\s*True|env\[[\'\"]mail\.mail[\'\"]\]\.create\s*\(')

    def test_retry_and_generic_erp_contract_fail_safe(self):
        retry = POLICY["retry_contract"]
        self.assertEqual(retry["identity_fields"], ["source_system", "source_key", "source_version"])
        self.assertEqual(retry["duplicate_result"], "existing_snapshot")
        self.assertEqual(retry["exhausted_result"], "REVIEW")
        self.assertEqual(retry["max_attempts"], 3)
        self.assertEqual(retry["backoff_minutes"], [1, 2, 4])
        self.assertEqual(set(retry["observable_fields"]), {"state", "attempt_count", "last_attempt_at", "next_attempt_at", "last_error"})
        self.assertEqual(POLICY["erp_adapter"]["operations"], ["import"])

    def test_onboarding_contract_is_profile_driven(self):
        self.assertEqual(set(POLICY["onboarding_required_fields"]), {"profile_code", "source_system", "adapter_type"})

    def test_onboarding_and_release_package_are_fail_safe(self):
        self.assertEqual(ONBOARDING["allowed_operations"], ["import"])
        self.assertEqual(ONBOARDING["validation_mode"], "fail_to_review")
        self.assertTrue(RELEASE["synthetic_development_corpus_included"])
        self.assertFalse(RELEASE["acceptance_corpus_included"])
        self.assertFalse(RELEASE["formal_release_gate_passed"])
        self.assertIn("formal_golden_uat_nfr_acceptance", RELEASE["deferred_artifacts"])

    def test_release_metadata_consistency_gate_fails_closed(self):
        manifest_path = ROOT / "__manifest__.py"
        manifest = ast.literal_eval(manifest_path.read_text())
        governance = RELEASE["release_governance"]
        self.assertEqual(RELEASE["module_version"], manifest["version"])
        self.assertEqual(RELEASE["module_version"], "19.0.1.3.12")
        orphan_repair = ROOT / "migrations/19.0.1.3.11/pre-migrate.py"
        self.assertTrue(orphan_repair.is_file())
        self.assertIn("logistics_idp_migration_orphan_archive", orphan_repair.read_text())
        self.assertEqual(governance["manifest_sha256"], hashlib.sha256(manifest_path.read_bytes()).hexdigest())
        self.assertEqual(governance["source_sha256"]["__manifest__.py"], governance["manifest_sha256"])
        self.assertRegex(governance["manifest_sha256"], r"^[0-9a-f]{64}$")
        self.assertIn(governance["provenance_commit"], (None,))
        self.assertFalse(governance["immutable_commit_sha_claimed"])
        self.assertIsNone(governance["bundle_sha256"])
        self.assertEqual(governance["bundle_sha256_status"], "not_generated")
        self.assertEqual(governance["manifest_sha256_status"], "verified")
        for field in ("verifier_evidence_status", "restore_evidence_status", "rollback_evidence_status"):
            self.assertEqual(governance[field], "not_run")
        self.assertFalse(governance["promotion_ready"] or RELEASE["formal_release_gate_passed"])

    def test_mes_boundary_is_read_only(self):
        self.assertEqual(POLICY["mes_mode"], "inbound_read_only")
        self.assertEqual(set(POLICY["mes_forbidden_operations"]), {"author", "mutate", "write_back", "plan", "execute"})

    def test_trade_compliance_wp1_contract_is_fail_closed(self):
        taxonomy = POLICY['authority_taxonomy']
        self.assertEqual(set(taxonomy), {'demo', 'customer_reference', 'early_warning_tier_4',
                                         'authoritative_tier_1', 'authoritative_tier_2',
                                         'approved_provider_tier_3'})
        self.assertFalse(taxonomy['demo']['legal_authority'])
        self.assertFalse(taxonomy['customer_reference']['legal_authority'])
        self.assertFalse(taxonomy['early_warning_tier_4']['legal_authority'])
        external = POLICY['source_classes']['external_api_candidate']
        option = VN_LEGAL_PACK['source_options']['external_api']
        required = {'provider_id', 'endpoint_config_reference', 'source_key', 'source_version',
                    'retrieved_at_utc', 'response_sha256', 'content_sha256', 'request_fingerprint',
                    'citation_clean_url', 'license_terms', 'retention_terms', 'content_owner',
                    'independent_oracle_status'}
        self.assertEqual(set(external['required_metadata']), required)
        self.assertEqual(set(option['required_metadata']), required)
        self.assertEqual(external['default_source_tier'], 'early_warning_tier_4')
        self.assertEqual(option['source_tier'], 'early_warning_tier_4')
        self.assertFalse(external['provider_tier_selection_allowed'] or external['activation_allowed'])
        self.assertFalse(option['provider_may_select_tier'] or option['activation'] or option['legal_authority'])
        self.assertEqual((external['verification'], external['missing_or_invalid_metadata_verdict']),
                         ('unverified', 'REVIEW'))
        self.assertEqual(option['expected_terminal_status'], 'review')
        self.assertEqual(external['provider_registry'], 'closed_allowlist')
        self.assertIn('no_raw_arbitrary_url', external['endpoint_policy'])
        self.assertIn('redirect_without_full_revalidation', external['network_denied'])
        self.assertEqual(external['credential_policy'], 'config_or_IAP_only_never_payload_or_log')
        self.assertEqual(external['approval_policy'], 'no_auto_activation_four_eyes')
        evidence = {source.get('classification'): source for source in VN_LEGAL_PACK['sources']}
        self.assertEqual(set(evidence), {None, 'business_acceptance', 'synthetic_expected', 'customer_presentation'})
        self.assertTrue(all(not source['claims_legal_truth'] and not source['independent']
                            for source in VN_LEGAL_PACK['sources']))
        schema = POLICY['compliance_context_schema']
        self.assertEqual(schema['schema_version'], '1.0')
        self.assertEqual(schema['missing_or_unknown_verdict'], 'REVIEW')
        self.assertEqual(set(schema['critical_fields']), {
            'who', 'when', 'evidence', 'transaction_time', 'effective_time', 'recorded_time'})
        demo = (ROOT / 'data/logistics_idp_policy_demo.xml').read_text()
        self.assertNotIn('>official<', demo)
        self.assertNotIn('"source_tier":"official"', demo)

    def test_partial_verify_is_conditional_gate_pass_without_legal_activation(self):
        policy = POLICY["gate_policy"]["PARTIAL_VERIFY"]
        candidate = VN_LEGAL_PACK["gate_policy"]
        self.assertEqual(policy["derived_gate_status"], "PASS_WITH_REVIEW")
        self.assertEqual(candidate["raw_status"], "PARTIAL_VERIFY")
        self.assertEqual(candidate["derived_gate_status"], "PASS_WITH_REVIEW")
        self.assertEqual(CLASSIFIER.gate_status("PARTIAL_VERIFY"), "PASS_WITH_REVIEW")
        self.assertEqual(CLASSIFIER.gate_status("PARTIAL"), "PARTIAL")
        for contract in (policy, candidate):
            self.assertEqual(contract["legal_verdict"], "REVIEW")
            self.assertEqual(contract["content_status"], "unverified")
            self.assertFalse(contract["activation_allowed"])
            self.assertEqual(contract["activation_blocker"], "BLOCKED_APPROVED_POLICY_DATASET")
            self.assertEqual(set(contract["forbidden_activation_tiers"]), {"early_warning_tier_4", "customer_reference"})
        self.assertFalse(candidate["automatic_activation"])
        self.assertIn("four_eyes", candidate["approval"])

    def test_traluat_documents_candidate_is_get_only_unverified_and_inactive(self):
        provider = POLICY["provider_registry"]["traluat_documents"]
        option = VN_LEGAL_PACK["source_options"]["traluat_documents"]
        self.assertEqual(provider["base_endpoint"], "https://traluat.com/api/documents")
        self.assertEqual(provider["allowed_methods"], ["GET"])
        self.assertEqual(provider["source_tier"], "early_warning_tier_4")
        self.assertFalse(provider["legal_authority"] or provider["auto_activation"])
        self.assertEqual((provider["verification"], provider["license_status"], provider["content_owner_status"]),
                         ("pending_independent_verification", "pending", "pending"))
        self.assertFalse(option["enabled"] or option["activation"] or option["automatic_activation"])

    def test_traluat_probe_taxonomy_and_seed_metadata_are_fail_closed(self):
        probe = VN_LEGAL_PACK["traluat_probe_contract"]
        expected = {"core_customs", "procedures_declarations", "tax_tariff", "hs_classification",
                    "customs_valuation", "origin_co_fta", "prohibited_restricted_goods",
                    "licenses_quotas", "specialized_inspection", "trade_remedies_antidumping",
                    "strategic_trade_control", "sanctions_low_confidence"}
        families = {item["id"]: item for item in probe["query_families"]}
        self.assertEqual(set(families), expected)
        self.assertEqual(families["sanctions_low_confidence"]["confidence"], "low")
        for family in families.values():
            self.assertTrue(family["keywords"] and family["expected_domains"] and
                            family["expected_agencies"] and family["expected_types"])
            self.assertTrue(family["accept"] and family["reject"])
        capture = probe["common_capture"]
        self.assertIn("response meta", capture["pagination"])
        self.assertEqual(capture["dedupe"], ["id", "docNumber+issueDate"])
        self.assertEqual(capture["immutability"], ["raw_response_sha256", "content_sha256"])
        self.assertEqual(set(capture["required_metadata"]),
                         {"status", "effective_date", "citation_slug", "retrieved_at_utc", "freshness"})
        self.assertEqual({item["doc_number"] for item in probe["seed_candidates"]}, {
            "95/VBHN-VPQH", "54/2014/QH13", "05/2017/QH14", "94/VBHN-VPQH",
            "107/2016/QH13", "08/2015/NĐ-CP", "39/2015/TT-BTC", "31/2018/NĐ-CP",
            "21/VBHN-BCT", "10/2018/NĐ-CP", "85/2019/NĐ-CP", "259/2025/NĐ-CP"})
        self.assertTrue(all(item["id"] is None for item in probe["seed_candidates"]))
        self.assertEqual(probe["seed_status"], "unverified_metadata_only_not_legal_truth")
        limitations = " ".join(probe["provider_limitations"])
        self.assertIn("provider claim, not independently verified", limitations)
        self.assertIn("Search ranking is noisy", limitations)
        self.assertIn("relationship and version review", limitations)

    def test_authenticated_reference_governance(self):
        tier = POLICY["source_classes"]["restricted_authenticated_product_reference"]
        self.assertEqual({key for key, value in tier.items() if value is False}, {"compliance_authority", "automated_ingestion", "raw_release_allowed"})
        self.assertTrue(tier["synthetic_derivation_allowed"])
        self.assertEqual(set(tier["required_metadata"]), {"ownership_authorization_reference", "capture_actor", "capture_time", "source_locator", "content_sha256", "retention_policy", "approver"})
        self.assertEqual(set(RELEASE["corpus_classes"]), {"raw_reference_evidence", "sanitized_acceptance_corpus", "synthetic_release_corpus"})
        self.assertTrue(all(value is False for value in RELEASE["raw_reference_claim_restrictions"].values()))

    def test_reference_document_traceability_is_complete_and_content_free(self):
        trace = RELEASE["reference_document_traceability"]
        self.assertFalse(trace["content_embedded"] or trace["pii_embedded"])
        sources = {item["id"]: item for item in trace["sources"]}
        self.assertEqual(set(sources), {"meeting_workbook", "uat_workbook", "solution_description", "product_intent"})
        for source in sources.values():
            path = REPOSITORY / source["path"]
            self.assertRegex(source["sha256"], r"^[0-9a-f]{64}$")
            if path.exists():
                self.assertEqual(source["sha256"], hashlib.sha256(path.read_bytes()).hexdigest())
            self.assertTrue(source["coverage"]["complete"])
        for source_id, expected_count in (("meeting_workbook", 16), ("uat_workbook", 9)):
            source = sources[source_id]
            path = REPOSITORY / source["path"]
            if path.exists():
                self.assertEqual(source["coverage"]["items"], load_workbook(path, read_only=True).sheetnames)
            self.assertEqual(len(source["coverage"]["items"]), expected_count)
        for source_id in ("solution_description", "product_intent"):
            source = sources[source_id]
            path = REPOSITORY / source["path"]
            page_count = source["coverage"]["page_count"]
            if path.exists():
                self.assertEqual(page_count, len(re.findall(rb'/Type\s*/Page(?!s)\b', path.read_bytes())))
            covered = {page for start, end in source["coverage"]["ranges"] for page in range(start, end + 1)}
            self.assertEqual(covered, set(range(1, page_count + 1)))
        self.assertEqual([item["rank"] for item in trace["precedence"]], list(range(1, 6)))
        self.assertEqual({item["source"] for item in trace["requirement_mappings"]}, set(sources))
        self.assertTrue(all(item["evidence"] and item["requirements"] for item in trace["requirement_mappings"]))

    def test_offline_classifier_rejects_secrets_without_returning_payload(self):
        content = b"session_token=synthetic-secret-value"
        result = CLASSIFIER.classify_raw_file(content, "text/plain")
        self.assertEqual(result["classification"], "rejected_secret_or_unsupported")
        self.assertIn("secret_or_session", result["findings"])
        self.assertEqual(result["sha256"], hashlib.sha256(content).hexdigest())
        self.assertNotIn(content.decode(), json.dumps(result))

    def test_offline_classifier_inspects_pdf_xlsx_and_image_metadata(self):
        pdf = CLASSIFIER.classify_raw_file(b"%PDF-1.4 /Author(x) /URI(url) /EmbeddedFile BT %%EOF", "application/pdf")
        self.assertEqual(pdf["classification"], "synthetic_derivation_only")
        self.assertTrue({"pdf_metadata", "pdf_links", "pdf_embedded_files", "pdf_ocr_text_layer"} <= set(pdf["findings"]))
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w") as archive:
            archive.writestr("xl/workbook.xml", '<sheet state="hidden"/><externalLink/>')
            archive.writestr("xl/worksheets/sheet1.xml", "<f>SUM(A1)</f>")
            archive.writestr("xl/comments1.xml", "<comments/>")
            archive.writestr("docProps/core.xml", "<lastModifiedBy>x</lastModifiedBy>")
        xlsx = CLASSIFIER.classify_raw_file(stream.getvalue(), CLASSIFIER.XLSX_MIME)
        self.assertEqual(xlsx["classification"], "synthetic_derivation_only")
        self.assertTrue({"xlsx_hidden_content", "xlsx_comments", "xlsx_formulas", "xlsx_external_links", "xlsx_revision_metadata"} <= set(xlsx["findings"]))
        image = CLASSIFIER.classify_raw_file(b"\xff\xd8\xffExif\x00\x00<x:xmpmeta/>", "image/jpeg")
        self.assertIn("image_exif_xmp", image["findings"])

    def test_image_barcode_is_decoded_or_explicitly_rejected(self):
        image = CLASSIFIER.classify_raw_file(b"\xff\xd8\xffSYNTHETIC\xff\xd9", "image/jpeg")
        self.assertIn(image["classification"], ("synthetic_derivation_only", "rejected_secret_or_unsupported"))
        if image["classification"] == "rejected_secret_or_unsupported":
            self.assertIn("qr_barcode_decode_unsupported", image["findings"])

    def test_restricted_observation_derivation_is_deterministic_and_gated(self):
        observations = {
            "taxonomy": ["purchase_order"],
            "field_types": ["string", "decimal"],
            "workflow_states": ["received", "review"],
            "validation_categories": ["schema", "cross_document"],
            "approval_status": "pending",
        }
        first = CLASSIFIER.derive_synthetic_fixtures(observations)
        self.assertEqual(first, CLASSIFIER.derive_synthetic_fixtures(observations))
        self.assertEqual(first[0]["fixture_id"], "SYN-DERIVED-001")
        self.assertTrue(first[0]["contact"].endswith(".invalid"))
        self.assertTrue(first[0]["metadata"]["derived_from_restricted_reference"])
        self.assertNotIn("derived_fixtures", CLASSIFIER.build(observations))
        observations["approval_status"] = "approved"
        self.assertEqual(len(CLASSIFIER.build(observations)["derived_fixtures"]), 1)
        with self.assertRaises(ValueError):
            CLASSIFIER.derive_synthetic_fixtures({**observations, "raw_payload": "forbidden"})

    def test_non_reidentification_rejects_exact_and_fuzzy_overlap(self):
        with self.assertRaisesRegex(ValueError, "exact"):
            CLASSIFIER.assert_non_reidentifying({"value": "customer-12345"}, ["customer-12345"])
        with self.assertRaisesRegex(ValueError, "fuzzy"):
            CLASSIFIER.assert_non_reidentifying({"value": "customer-12346"}, ["customer-12345"])

    def test_reference_pipeline_requires_external_quarantine_and_independent_oracle(self):
        with self.assertRaisesRegex(ValueError, "ngoài Git"):
            CLASSIFIER.reference_evidence(Path(__file__))
        self.assertEqual(CLASSIFIER._approved_oracles(None), {})
        pending = {"oracles": [{"source_hash": "a" * 64, "review_status": "candidate_requires_independent_ui_review", "reviewer_independent_of_exporter": False}]}
        path = Path("/tmp/idp-pending-oracle-self-check.json")
        path.write_text(json.dumps(pending))
        try:
            self.assertEqual(CLASSIFIER._approved_oracles(path), {})
        finally:
            path.unlink()

    def test_reference_counts_keep_record_uniqueness_separate_from_attachment_rejection(self):
        inventory = Path("/tmp/idp-inventory-self-check.jsonl")
        errors = Path("/tmp/idp-errors-self-check.jsonl")
        inventory.write_text(json.dumps({"source_id": "accepted", "source_attachment_count": 0, "attachments": []}) + "\n")
        errors.write_text(json.dumps({"source_id": "rejected", "category": "http_404"}) + "\n")
        try:
            evidence = CLASSIFIER.reference_evidence(inventory, error_inventory_path=errors, combined=True)
            self.assertEqual((evidence["raw_status"], evidence["derived_gate_status"]), ("PARTIAL_VERIFY", "PASS_WITH_REVIEW"))
            self.assertEqual(evidence["counts"]["run"]["detailed_golden"], {"detailed_pass": 198, "detailed_fail": 0, "golden_pass": 46, "golden_fail": 0, "unmapped": 1, "golden_partial": 2})
            self.assertEqual((evidence["claims"]["legal_verdict"], evidence["claims"]["content_status"]), ("REVIEW", "unverified"))
            self.assertFalse(evidence["claims"]["activation_allowed"])
            self.assertEqual(evidence["claims"]["activation_blocker"], "BLOCKED_APPROVED_POLICY_DATASET")
            self.assertEqual(evidence["counts"]["run"]["source_records"], {"discovered": 2, "fetched": 2, "unique": 2, "failed": 1, "skipped": 0})
            self.assertEqual(evidence["counts"]["run"]["attachments"]["rejected"], 1)
        finally:
            inventory.unlink()
            errors.unlink()

    def test_nfr_invokes_dashboard_service_contract(self):
        runner = (Path(__file__).parent / "run_local_nfr.py").read_text()
        self.assertIn("case_model.get_dashboard_data(filters)", runner)
        self.assertNotIn("case_model.read_group([], ['id:count'], ['state'])", runner)
        self.assertIn("'sudo': False", runner)

    def test_reconciliation_evidence_native_views_expose_payload_and_case_link(self):
        views = (ROOT / "views/logistics_idp_views.xml").read_text()
        self.assertIn('id="view_logistics_evidence_form"', views)
        self.assertIn('string="Persisted Reconciliation Result"', views)
        self.assertIn('<field name="payload" nolabel="1"/>', views)
        self.assertIn('<field name="case_id"/>', views)
        self.assertIn('<field name="evidence_ids" readonly="1"><list><field name="category"/><field name="source_reference"/><field name="status" widget="badge"/><field name="payload"/>', views)

    def test_internal_screening_ui_claims_are_neutral(self):
        views = (ROOT / "views/logistics_idp_views.xml").read_text()
        onboarding = (ROOT / "views/is_logistics_idp_onboarding_views.xml").read_text()
        report = (ROOT / "reports/is_logistics_idp_report_templates.xml").read_text()
        dashboard = (ROOT / "static/src/dashboard/dashboard.xml").read_text()
        self.assertIn('Internal Screening Review Center', views)
        self.assertNotIn('Document Clearance', views)
        self.assertIn('hồ sơ NSW để người dùng rà soát', onboarding)
        self.assertIn('Kết quả kiểm tra nội bộ (Trade):', report)
        self.assertNotIn('text-success', report[report.index('Kết quả kiểm tra nội bộ (Trade):'):report.index('Kết quả sàng lọc nội bộ tự động:')])
        self.assertIn('Internal screening status', dashboard)

    def test_dashboard_client_contract(self):
        manifest = (ROOT / "__manifest__.py").read_text()
        views = (ROOT / "views/logistics_idp_views.xml").read_text()
        js = (ROOT / "static/src/dashboard/dashboard.js").read_text()
        template = (ROOT / "static/src/dashboard/dashboard.xml").read_text()
        self.assertIn("web.assets_backend", manifest)
        self.assertIn('id="action_logistics_overview"', views)
        self.assertLess(views.index('id="menu_logistics_overview"'), views.index('id="menu_logistics_cases"'))
        self.assertIn('registry.category("actions").add("insilos_logistics_idp.dashboard"', js)
        self.assertIn('"get_dashboard_data", [this.cleanFilters()]', js)
        self.assertIn('"dashboard_drilldown", [kind, this.cleanFilters(), value]', js)
        self.assertRegex(template, r't-if="state\.busyAction"[\s\S]*t-elif="state\.loading"')
        self.assertIn('t-if="state.error"', template)
        self.assertIn('No records match these filters.', template)
        self.assertIn('if (!available) return _t("Not Available")', js)
        self.assertIn('metric.available !== false', template)
        self.assertNotRegex("\n".join((js, template)), r'class="[^"]*\bo[_-]')

    def test_synthetic_development_corpus_is_not_formal_acceptance(self):
        self.assertEqual(len(CASES["cases"]), 198)
        self.assertEqual(len({golden for case in CASES["cases"] for golden in case["golden_ids"]}), 48)
        self.assertTrue(CASES["metadata"]["non_production"])
        self.assertFalse(CASES["metadata"]["formal_golden_acceptance"])
        self.assertNotIn("Swarovski", json.dumps(CASES))


if __name__ == "__main__":
    unittest.main()
