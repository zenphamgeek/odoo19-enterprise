import argparse
import ast
import base64
import csv
import hashlib
import io
import json
import re
import zipfile
from difflib import SequenceMatcher
from datetime import datetime, timezone
from email.message import EmailMessage
from pathlib import Path

from openpyxl import Workbook, load_workbook

HERE = Path(__file__).resolve().parent
ROOT = Path(__file__).resolve().parents[5]
WORKBOOK = ROOT / "docs/industries/Insilos_IDP_Logistics_UAT_v1.4.xlsx"
SRS = ROOT / "docs/industries/Insilos_Vertical_IDP_Logistics_SRS_v1.4_Trade_Compliance_MES_Boundary.md"
OUTPUT = HERE / "sanitized_cases.json"
CROSSWALK = ROOT / "insilos/apps/insilos_logistics_idp/docs/tests/srs_canonical_crosswalk_v1.json"
DEVELOPMENT_UAT_RESULT = ROOT / "insilos/apps/insilos_logistics_idp/docs/tests/development_uat_result.json"
MIMES = ("application/pdf", "application/pdf", "image/png", "image/jpeg", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "text/csv", "message/rfc822", "application/json")
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TEXT_FINDINGS = {
    "pii": re.compile(rb"(?:[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}|\+?\d[\d ().-]{7,}\d|(?:address|tax[ _-]?id)\s*[:=])", re.I),
    "business_identifier": re.compile(rb"(?:PO|invoice|B/L|declaration|account|tenant)[ _#:/-]?[A-Z0-9-]{3,}", re.I),
    "name_or_company": re.compile(rb"(?:customer|company|supplier|contact)[ _-]?name\s*[:=]", re.I),
    "secret_or_session": re.compile(rb"(?:BEGIN (?:RSA |OPENSSH )?PRIVATE KEY|(?:password|secret|bearer|csrf|session|cookie|token)\s*[:=])", re.I),
    "signed_url": re.compile(rb"https?://[^\s<>\"]+[?&](?:x-amz-signature|signature|sig|token)=", re.I),
    "qr_or_barcode": re.compile(rb"(?:qr|barcode)[ _-]?(?:payload|value)\s*[:=]", re.I),
    "reidentification_risk": re.compile(rb"(?:date.of.birth|license.plate|exact.location|unique.combination)\s*[:=]", re.I),
}


def canonical(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()


def digest(value):
    return hashlib.sha256(value if isinstance(value, bytes) else canonical(value)).hexdigest()


def rows(name, header_row=1):
    sheet = load_workbook(WORKBOOK, data_only=True, read_only=True)[name]
    values = sheet.iter_rows(values_only=True)
    for _ in range(header_row - 1):
        next(values)
    header = next(values)
    return [dict(zip(header, row)) for row in values if any(value is not None for value in row)]


def workbook_contract():
    coverage = rows("Requirement Coverage")
    formula_sheet = load_workbook(WORKBOOK, data_only=False, read_only=True)["Requirement Coverage"]
    for row in range(2, len(coverage) + 2):
        expected = (f'=COUNTIF(\'UAT Cases\'!$C$2:$C$199,"*"&A{row}&"*")',
                    f'=COUNTIF(\'Golden 48\'!$F$2:$F$49,"*"&A{row}&"*")',
                    f'=IF(C{row}>0,"Covered","Review")')
        if tuple(formula_sheet.cell(row, column).value for column in range(3, 6)) != expected:
            raise ValueError("Requirement Coverage formula differs at row %d" % row)
    detailed = rows("UAT Cases")
    golden = rows("Golden 48")
    requirement_ids = {str(item["Requirement ID"]) for item in coverage}
    srs_sections = {match.group(1) for match in re.finditer(r"^#{1,4}\\s+(\\d+[A-Z]?(?:\\.\\d+[A-Z]?)*)\\b", SRS.read_text(), re.M)}
    detailed_requirements = {}
    for row in detailed:
        value = str(row.get("Requirement") or "")
        explicit = set(REQUIREMENT_ID.findall(value))
        sections = {match.group(1) for match in re.finditer(r"§(\\d+[A-Z]?(?:\\.\\d+[A-Z]?)*)\\b", value)}
        detailed_requirements[row["Test ID"]] = (explicit | (sections & requirement_ids & srs_sections)) & requirement_ids
    golden_map = {}
    for row in golden:
        for detailed_id in str(row.get("Mapped Detailed Tests") or "").split(","):
            if detailed_id.strip() in detailed_requirements:
                golden_map.setdefault(detailed_id.strip(), set()).add(row["Golden ID"])
    for item in coverage:
        requirement_id = str(item["Requirement ID"])
        mapped_detailed = {test_id for test_id, ids in detailed_requirements.items() if requirement_id in ids}
        mapped_golden = {golden_id for test_id in mapped_detailed for golden_id in golden_map.get(test_id, ())}
        item.update({"Detailed Test Count": len(mapped_detailed), "Golden Mapping Count": len(mapped_golden),
                     "Coverage Status": "Covered" if mapped_detailed else "Review"})
    return {
        "requirement_coverage": coverage,
        "source_traceability": rows("Source Traceability"),
        "acceptance_metrics": rows("Acceptance Metrics"),
        "test_data": rows("Test Data", 3),
    }


def payload(case_id, index):
    value = f"{10 + index / 100:.2f}"
    return {"document_type": "purchase_order", "confidence": 0.99, "source_spans": [{"page": 1, "text": case_id}], "payload": {"supplier": f"Synthetic Supplier {index:03d}", "supplier_address": f"{index} Synthetic Street", "supplier_number": f"SUP-{index:03d}", "po_reference": case_id, "total_value": value, "document_date": f"2026-04-{index % 28 + 1:02d}", "currency": "USD", "lines": [{"material_code": f"SYN-{index:04d}", "description": f"Synthetic item {index:03d}", "quantity": index + 1, "unit_price": value, "uom": "EA", "value": value, "custom_code": f"PO-{index:04d}"}]}}


def _minimal_pdf(data, page_count):
    payload = base64.b64encode(canonical(data))
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [" + b" ".join(f"{3 + page * 2} 0 R".encode() for page in range(page_count)) + b"] /Count " + str(page_count).encode() + b" >>",
    ]
    for page in range(page_count):
        content_id = 4 + page * 2
        stream = b"% synthetic payload " + payload + b"\n"
        objects.extend((
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents " + str(content_id).encode() + b" 0 R >>",
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"endstream",
        ))
    output, offsets = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n"), []
    for object_id, object_body in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f"{object_id} 0 obj\n".encode() + object_body + b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    output.extend(b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets))
    output.extend(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return bytes(output)


def binary(mimetype, data, multipage=False):
    encoded = canonical(data)
    if mimetype == "application/json":
        return encoded
    if mimetype == "text/csv":
        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=("document_type", "confidence", "supplier", "supplier_address", "supplier_number", "po_reference", "total_value", "document_date", "currency", "lines"))
        writer.writeheader(); writer.writerow({"document_type": data["document_type"], "confidence": data["confidence"], **data["payload"], "lines": json.dumps(data["payload"]["lines"], separators=(",", ":"))})
        return output.getvalue().encode()
    if mimetype == "message/rfc822":
        message = EmailMessage(); message["From"] = "synthetic-sender@example.invalid"; message["To"] = "synthetic-receiver@example.invalid"; message["Subject"] = "Synthetic logistics fixture"; message.set_content(encoded.decode())
        return message.as_bytes()
    if mimetype == "application/pdf":
        return _minimal_pdf(data, 2 if multipage else 1)
    if mimetype == "image/png":
        return b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR" + b"\x00" * 16 + encoded
    if mimetype == "image/jpeg":
        return b"\xff\xd8\xff\xe0SYNTHETIC" + encoded + b"\xff\xd9"
    workbook = Workbook()
    workbook.properties.created = datetime(2026, 1, 1, tzinfo=timezone.utc)
    sheet = workbook.active
    sheet.title = "Normalized Document"
    sheet.append(("document_type", "confidence", "supplier", "supplier_address", "supplier_number", "po_reference", "total_value", "document_date", "currency"))
    sheet.append((data["document_type"], data["confidence"], data["payload"]["supplier"], data["payload"]["supplier_address"], data["payload"]["supplier_number"],
                  data["payload"]["po_reference"], data["payload"]["total_value"], data["payload"]["document_date"], data["payload"]["currency"]))
    sheet.append(())
    sheet.append(("material_code", "description", "quantity", "unit_price", "uom", "value", "custom_code"))
    for line in data["payload"]["lines"]:
        sheet.append(tuple(line.get(field) for field in ("material_code", "description", "quantity", "unit_price", "uom", "value", "custom_code")))
    generated = io.BytesIO()
    workbook.save(generated)
    output = io.BytesIO()
    with zipfile.ZipFile(generated) as source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            value = source.read(name)
            if name == "docProps/core.xml":
                value = re.sub(br"<dcterms:modified[^>]*>.*?</dcterms:modified>", b"", value)
            info = zipfile.ZipInfo(name, (2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, value)
    return output.getvalue()


def _detected_mime(content):
    if content.startswith(b"%PDF-"):
        return "application/pdf"
    if content.startswith(b"PK\x03\x04"):
        return XLSX_MIME
    if content.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png"
    if content.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    if b"\x00" not in content[:1024]:
        return "text/plain"
    return "application/octet-stream"


def _barcode_payloads(content, detected_mime):
    if detected_mime not in {"image/png", "image/jpeg"}:
        return (), False
    try:
        from PIL import Image
        from pyzbar.pyzbar import decode
    except (ImportError, OSError):
        return (), False
    try:
        return tuple(item.data for item in decode(Image.open(io.BytesIO(content)))), True
    except (OSError, ValueError):
        return (), True


def classify_raw_file(content, declared_mime=None):
    """Return metadata and finding categories only; never retain raw payload."""
    detected_mime = _detected_mime(content)
    findings = {name for name, pattern in TEXT_FINDINGS.items() if pattern.search(content)}
    barcode_payloads, barcode_decoder_available = _barcode_payloads(content, detected_mime)
    if barcode_payloads:
        findings.add("qr_or_barcode")
    elif detected_mime in {"image/png", "image/jpeg"} and not barcode_decoder_available:
        findings.add("qr_barcode_decode_unsupported")
    if declared_mime and declared_mime != detected_mime:
        findings.add("mime_mismatch")
    supported = {"application/pdf", XLSX_MIME, "image/png", "image/jpeg", "text/plain"}
    if detected_mime == "application/pdf":
        pdf_markers = {
            "pdf_metadata": (b"/Author", b"/Creator", b"/Producer", b"/Metadata"),
            "pdf_links": (b"/URI",),
            "pdf_embedded_files": (b"/EmbeddedFile", b"/Filespec"),
            "pdf_ocr_text_layer": (b"BT", b"/ToUnicode"),
        }
        findings.update(name for name, markers in pdf_markers.items() if any(marker in content for marker in markers))
    elif detected_mime == XLSX_MIME:
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                names = archive.namelist()
                parts = {name: archive.read(name) for name in names}
            joined = b"\n".join(parts.values())
            xlsx_markers = {
                "xlsx_hidden_content": (b'state="hidden"', b'state="veryHidden"', b' hidden="1"'),
                "xlsx_comments": (b"<comments", b"/comments"),
                "xlsx_formulas": (b"<f",),
                "xlsx_external_links": (b"externalLink", b"externalLinks"),
                "xlsx_revision_metadata": (b"revision", b"lastModifiedBy"),
            }
            findings.update(name for name, markers in xlsx_markers.items() if any(marker in joined for marker in markers))
        except zipfile.BadZipFile:
            findings.add("malformed_container")
    elif detected_mime in {"image/png", "image/jpeg"}:
        if any(marker in content for marker in (b"Exif\x00\x00", b"http://ns.adobe.com/xap/1.0/", b"<x:xmpmeta")):
            findings.add("image_exif_xmp")
    rejected = findings & {"secret_or_session", "signed_url", "mime_mismatch", "malformed_container", "qr_barcode_decode_unsupported"}
    derivation_only = findings & {"business_identifier", "name_or_company", "qr_or_barcode", "reidentification_risk", "pdf_embedded_files", "pdf_ocr_text_layer", "xlsx_hidden_content", "xlsx_comments", "xlsx_formulas", "xlsx_external_links", "xlsx_revision_metadata", "image_exif_xmp"}
    if detected_mime not in supported or rejected:
        classification = "rejected_secret_or_unsupported"
    elif derivation_only:
        classification = "synthetic_derivation_only"
    elif findings:
        classification = "sanitization_candidate"
    else:
        classification = "raw_restricted"
    return {"sha256": digest(content), "size": len(content), "declared_mime": declared_mime, "detected_mime": detected_mime, "classification": classification, "findings": sorted(findings)}


DERIVATION_KEYS = {"taxonomy", "field_types", "workflow_states", "validation_categories", "approval_status"}
TRANSFORMATION_VERSION = "restricted-observations-v1"
ORACLE_VERSION = "synthetic-business-oracle-v1"
REQUIREMENT_ID = re.compile(r"(?:4A\.\d+|[A-Z]{2,5}-\d+[A-Z]?)")
TRACE_CATEGORIES = {
    "test-only", "extraction/schema", "policy config", "deterministic rule",
    "workflow/UI", "external/non-runnable",
}
SRS_EXPLICIT_ID = re.compile(r"^(?:FR|AI|RULE|EXC|DUP|SLA|RPT|INT|SEC|NFR|UAT)-\d+[A-Z]?$|^VA-\d+$")
SRS_STABLE_GROUPS = {
    "ROLE": "5.1", "DOC": "9", "PHASE": "26", "DOD": "27", "SECTION": None,
}
TEST_PACKAGES = {
    "idp": ("odoo.addons.insilos_logistics_idp.tests", HERE.parent),
    "kg": ("odoo.addons.insilos_knowledge_graph.tests", ROOT / "insilos/apps/insilos_knowledge_graph/tests"),
}
CAPABILITY_EXECUTORS = {
    "AI-003": ("test_ws2_ocr_quality.TestWs2OcrQuality.test_iap_strict_ocr_persists_canonical_run_and_rejects_invalid_geometry", "services/document_processor.py", ["IAPDocumentProcessor.process", "extraction_response_schema"], "IAP extraction requests strict typed JSON schema and rejects schema-invalid payloads before use"),
    "AI-005": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_extraction_low_confidence_routes_review", "models/logistics_idp.py", ["LogisticsCase.add_extraction"], "below-threshold extraction routes to review with payload retained"),
    "AI-006": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_processed_viewer_exposes_current_payload_and_attempt_history", "models/logistics_idp.py", ["LogisticsDocument.process", "LogisticsExtractionRun.create"], "processing/reprocess persist provider/model/schema/prompt/template versions, timestamps, and immutable retry_count sequence"),
    "AI-007": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_ai007_local_logical_documents_are_stable_and_malformed_safe", ("models/logistics_idp.py", "services/document_processor.py"), ["LogisticsDocument.process", "LocalDocumentProcessor.process"], "local synthetic logical_documents validate each typed item, create attachment/hash-linked deterministic logical children and immutable runs, replay idempotently, and fail malformed input without child business documents"),
    "EXC-001": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_exc001_canonical_types_are_company_scoped_and_idempotent", "models/logistics_idp.py", ["LogisticsException.open_or_reuse"], "canonical exception taxonomy preserves company/case scope and reuses one open exception per deterministic condition"),
    "DUP-001": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_dossier_exact_duplicate_disposition_is_durable_and_company_safe", "models/logistics_idp.py", ["LogisticsDocument.intake_content"], "same-content cross-case intake reuses the canonical document and persists its SHA-256 exact-duplicate disposition"),
    "DUP-003": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_dossier_exact_duplicate_disposition_is_durable_and_company_safe", "models/logistics_idp.py", ["LogisticsDocument.intake_content"], "same-content cross-case intake links the canonical document through one durable excluded-from-aggregation duplicate disposition"),
    "DUP-004": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_semantic_duplicate_links_canonical_case", "models/logistics_idp.py", ["LogisticsCase.mark_semantic_duplicate"], "manager closure sets Closed-Duplicated with its canonical case reference and immutable paired semantic-duplicate evidence"),
    "FR-102": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr102_upload_case_inbox_dedup_and_boundaries", "models/logistics_idp.py", ["LogisticsCase.action_open_upload_wizard", "LogisticsDocument.upload_intake", "LogisticsDocumentUploadWizard.action_upload"], "operator upload queues unknown documents to an explicit case or a collecting manual inbox case without synchronous provider processing; same-case exact hashes reuse the canonical document and boundaries reject terminal, empty, and unsupported MIME input"),
    "FR-104": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr104_supplier_email_filename_is_audit_only_until_content_classification", "models/logistics_idp.py", ["LogisticsInboundJob._process_one", "LogisticsDocument.process"], "queued supplier-email filenames and subjects remain audit metadata; intake persists unknown/0/review and only content processing selects the canonical type"),
    "FR-105": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_intake_exact_idempotency_and_deterministic_correlation", "models/logistics_idp.py", ["LogisticsCase.intake"], "exact identity reuse and PO correlation"),
    "FR-106": (("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr106_same_normalized_subject_keeps_latest_active_history_and_replay", "test_logistics_idp_security.TestLogisticsIdpSecurity.test_fr106_thread_history_is_company_scoped_and_intake_controlled"), "models/logistics_idp.py", ["LogisticsCase._record_same_subject_thread_entry", "LogisticsEmailThreadEntry"], "same-company normalized same-subject intake retains immutable history with exactly one newest active entry; canonical replay is idempotent and history remains ACL/company controlled"),
    "FR-107":  (("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr107_different_subject_requires_idempotent_review_and_explicit_active_thread_selection", "test_logistics_idp_security.TestLogisticsIdpSecurity.test_fr107_active_thread_selection_denies_unauthorized_and_cross_company"), "models/logistics_idp.py", ["LogisticsCase._schedule_different_subject_review", "LogisticsCase.action_select_active_thread"], "different-subject ambiguity creates one review activity without mutating case/thread, subject, documents, or state; reviewer/manager explicitly selects only a same-company active thread"),
    "FR-108": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_case_fails_to_review_and_source_identity_is_unique", "models/logistics_idp.py", ["LogisticsCase"], "unique case source identity constraint"),
    "FR-201": ("test_logistics_idp_phase_1_3.TestLogisticsIdpPurePhase13.test_fr201_purchase_order_missing_invalid_and_low_confidence_review", ("services/document_schemas.py", "services/document_processor.py"), ["document_schemas.validate", "document_processor.extraction_response_schema"], "synthetic processor missing, invalid, and low-confidence review paths cover Supplier Name, Supplier Address, Supplier Number, Total Value, PO Date, Currency, Material Code, Description, Unit Price, Quantity, UoM, Value, and Custom Code; runtime critical-check positive path is not evidenced"),
    "FR-202": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr202_batch_chunks_larger_than_100_and_reports_created", "models/logistics_idp.py", ["LogisticsPoSnapshot.import_batch"], "configurable default-100 batching safely imports 101 PO snapshots and reports creations"),
    "FR-203": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr203_latest_snapshot_closes_absent_material_without_mutating_prior", "models/logistics_idp.py", ["LogisticsPoSnapshot.import_upsert"], "latest PO snapshot represents absent material lines as remaining zero while prior snapshot remains immutable"),
    "FR-204": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr204_snapshot_retry_is_ignored_without_duplicates", "models/logistics_idp.py", ["LogisticsPoSnapshot.import_batch"], "source/version retry is ignored without duplicate snapshot and report exposes counts"),
    "FR-205": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr205_material_reference_failures_and_no_material_line", "models/logistics_idp.py", ["LogisticsPoSnapshot.import_batch"], "material codes require latest master-data and DSNVL references; explicit no-material lines are allowed"),
    "FR-301": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr301_supplier_draft_invoice_policy_modes_effectivity_isolation_and_replay", ("models/logistics_idp.py", "services/reconciliation.py"), ["LogisticsCase.reconcile_documents", "validate_draft_invoice_policy", "draft_invoice_mode_check"], "local deterministic supplier draft-invoice policy only; missing, invalid, cross-company, out-of-window, or ambiguous profiles fail closed to review"),
    "FR-302": ("test_logistics_idp_phase_1_3.TestLogisticsIdpPurePhase13.test_fr302_draft_identity_final_number_policy_is_scoped", "services/reconciliation.py", ["validate_draft_invoice_policy", "reconcile_documents"], "local deterministic reconciliation only: final number is preserved as observed evidence and returns REVIEW when disallowed; no legal, provider, or external authority assertion"),
    "FR-303": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_draft_po_multiline_identity_unmatched_evidence_is_persisted", "models/logistics_idp.py", ["LogisticsCase.reconcile_documents"], "draft invoice PO reconciliation matches supplier and lines by configured identity cascade"),
    "FR-304": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr304_draft_result_persists_differences_warnings_and_evidence", "models/logistics_idp.py", ["LogisticsCase.reconcile_documents"], "draft result persists pass/fail/warning verdicts, matched and unmatched lines, quantity/price/supplier differences, regime warnings, and linked evidence hashes"),
    "FR-305": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr305_supplier_result_reply_is_configured_threaded_authorized_and_idempotent", "models/logistics_idp.py", ["LogisticsInboundJob.intake_supplier_email", "LogisticsCase.queue_supplier_result_reply", "MailMail.send"], "manager-only native queued reply accepts only an actual supplier-email intake draft-VAT-invoice document/evidence with matching hash and immutable intake_supplier_email audit markers; forged controlled evidence has no side effect; native send appends exactly-one sent/exception delivery evidence without raw failure reason, retains force_send=False and exception-mail reuse"),
    "FR-306": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_po_invoice_reconciliation_and_recheck_preserve_runs", "models/logistics_idp.py", ["LogisticsCase.recheck"], "re-check appends immutable reconciliation evidence"),
    "FR-403": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr403_missing_regime_persists_review_and_replay_is_side_effect_free", ("models/logistics_idp.py", "services/reconciliation.py"), ["LogisticsCase.reconcile_documents", "line_regime_decisions", "reconcile_documents"], "missing customs regime persists review with payload retained and replay is side-effect-free; no legal or provider inference"),
    "FR-404": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr404_runtime_declaration_groups_bind_manifest_and_replay", ("models/logistics_idp.py", "services/reconciliation.py"), ["declaration_group_check", "LogisticsCase.reconcile_documents"], "effective declaration-group count passes only exact normalized groups; mismatch blocks; absent, out-of-window, ambiguous rules review fail-closed; manifest-bound replay has no side effect"),
    "FR-405": ("test_logistics_idp_phase_1_3.TestLogisticsIdpPurePhase13.test_fr405_line_regimes_use_explicit_po_and_invoice_lines", "services/reconciliation.py", ["line_regime_decisions", "reconcile_documents"], "explicit PO/invoice line comparison blocks customs-regime mismatches, routes missing or conflicting regimes to review; no legal inference"),
    "FR-402": ("test_logistics_idp_phase_1_3.TestLogisticsIdpPurePhase13.test_fr402_no_material_default_is_policy_driven_not_core", "services/reconciliation.py", ["line_regime_decisions", "reconcile_documents"], "no-material-code lines without regime follow the effective no_material_behavior rule: custom_code with one allowed_regimes value assigns that regime and passes the regime check into declaration-group validation; review default or absent/multi-value rule keeps missing-customs-regime review; no provider/IAP calls"),
    "FR-406": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr406_final_export_hs_reconciliation_is_immutable_idempotent_and_scoped", ("models/logistics_idp.py", "services/reconciliation.py"), ["customs_checks", "LogisticsCase._final_export_hs_checks", "LogisticsCase.reconcile_documents", "LogisticsException.open_or_reuse"], "final export HS is observed-only against DSNVL/master; match, mismatch and missing inputs append immutable checks/evidence, mismatch reuses one case/company exception, master remains unchanged"),
    "FR-501": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr501_main_invoice_policy_effectivity_isolation_and_replay", ("models/logistics_idp.py", "services/reconciliation.py"), ["LogisticsCase.reconcile_documents", "main_invoice_check", "validate_main_invoice_policy", "reconcile_documents"], "local deterministic policy defaults to numbered main VAT invoices; exact profile substitutes only; invalid, ambiguous, out-of-window, or foreign profiles review fail-closed; replay reuses matching reconciliation evidence"),
    "FR-502": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr502_main_invoice_requires_exact_approved_draft_snapshot_and_replay", ("models/logistics_idp.py", "services/reconciliation.py"), ["LogisticsCase.reconcile_documents", "reconcile_draft_main"], "same-case approved draft VAT selection is exact and deterministic; manifest snapshots selected draft/current run; supplier, identity, quantity, unit price and line total mismatch blocks; missing, ambiguous or unapproved draft reviews; replay reuses immutable semantic evidence"),
    "FR-503": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr503_skipped_profile_reconciles_main_directly_and_never_compares_draft", ("models/logistics_idp.py", "services/reconciliation.py"), ["LogisticsCase.reconcile_documents", "reconcile_draft_main"], "effective same-company skipped draft policy validates a numbered main VAT invoice directly against PO/reference; evidence/manifest draft_main is false, valid draft input does not invoke FR-502 comparison, critical PO mismatch blocks, and identical replay reuses immutable evidence"),
    "FR-504": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr504_same_type_logical_invoices_preserve_each_payload_and_replay_runs", ("models/logistics_idp.py", "services/document_processor.py"), ["LogisticsDocument.process", "LocalDocumentProcessor.process"], "local synthetic logical-document splitting preserves two same-type invoice payloads independently, appends one run per stable child on replay, and does not invoke provider/IAP; no legal or provider claim"),
    "FR-505": ("test_logistics_idp_phase_1_3.TestLogisticsIdpPurePhase13.test_smart_filename_classifier_and_batch_structure", "services/document_processor.py", ["batch_structure"], "draft/main invoice lifecycle ordering"),
    "FR-506": ("test_logistics_idp_phase_1_3.TestLogisticsIdpPurePhase13.test_smart_filename_classifier_and_batch_structure", "services/document_processor.py", ["batch_structure"], "duplicate semantic document roles"),
    "FR-507": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr507_variance_blocks_before_artifact", "models/logistics_idp.py", ["LogisticsCase.generate_sap_erp_output"], "invoice-to-SAP total variance blocks artifact creation"),
    "FR-601": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr601_template_positions_sheet_and_style_preserved", "models/logistics_idp.py", ["LogisticsCase.generate_sap_erp_output"], "SAP template sheet, positions, and template style are preserved"),
    "FR-602": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr602_field_level_configured_default_sources", "models/logistics_idp.py", ["LogisticsCase.generate_sap_erp_output"], "SRS field defaults: invoice quantity/UoM/amount/number/date, PO terms/items, DSNVL/master fields; config cannot override source class"),
    "FR-603": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr603_missing_source_blocks_without_output", "models/logistics_idp.py", ["LogisticsCase.generate_sap_erp_output"], "missing authoritative source blocks SAP output creation"),
    "FR-604": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr604_blank_invoice_and_po_material_use_configured_po_code_template", "models/logistics_idp.py", ["LogisticsCase.generate_sap_erp_output"], "blank material custom code uses configured PO number/item template"),
    "FR-604A": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr604a_supplemental_final_customs_upload_reuses_case_and_preserves_history", "models/logistics_idp.py", ["LogisticsCase.action_open_final_customs_upload_wizard", "LogisticsDocument.upload_final_customs_document", "LogisticsDocument.action_reclassify"], "authorized same-company nonterminal supplemental upload reuses the case, preserves immutable history and provenance, classifies final declaration, and processes only the new document"),
    "FR-605A": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr605a_configured_rounding_is_persisted_with_versioned_evidence", "models/logistics_idp.py", ["LogisticsCase._sap_erp_amount", "LogisticsCase._sap_erp_rounding", "LogisticsCase.generate_sap_erp_output"], "deterministic supplier/profile/currency/field/template rule selection, fallback, version and rounded-result evidence"),
    "FR-606": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr606_stable_po_order_and_idempotency", "models/logistics_idp.py", ["LogisticsCase.generate_sap_erp_output"], "stable PO ordering and idempotent SAP output"),
    "FR-607": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr607_configured_custom_code_format", "models/logistics_idp.py", ["LogisticsCase.generate_sap_erp_output"], "configured SAP custom-code format and price normalization"),
    "FR-608": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr608_source_mapping_version_and_provenance", "models/logistics_idp.py", ["LogisticsCase.generate_sap_erp_output"], "SAP field-source mapping provenance and version evidence"),
    "FR-609": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr609_policy_blocks_missing_ambiguous_and_unbound_before_evidence", "models/logistics_idp.py", ["LogisticsCase._authoritative_reference", "LogisticsCase.generate_sap_erp_output", "LogisticsCase.generate_customs_output", "LogisticsCase.reconcile_import_declaration"], "versioned stage/field exact-or-wildcard output policy binds SAP downstream and E13/E15; invalid, unbound, or stale profiles fail closed before reconciliation evidence"),
    "FR-706": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_vietnam_outputs_import_reference_shipping_and_gate_pass", "models/logistics_idp.py", ["LogisticsCase.reconcile_import_declaration"], "customs quantity and UoM comparison"),
    "FR-801": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr801_broker_package_missing_requirement_creates_no_output", "models/logistics_idp.py", ["LogisticsCase.generate_broker_package"], "configured required categories block package creation when unavailable"),
    "FR-802": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr802_broker_package_archives_qdtq_bytes_and_generated_output", "models/logistics_idp.py", ["LogisticsCase.generate_broker_package", "LogisticsOutput._generate_broker_package"], "policy-gated broker ZIP preserves configured document bytes and generated spreadsheet output"),
    "FR-804": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr804_native_broker_email_queue_is_configured_idempotent_and_auditable", "models/logistics_idp.py", ["LogisticsCase.queue_broker_email"], "manager-only native template queue uses effective profile recipients, current broker package attachment, immutable request/queued evidence, timeline, and deterministic duplicate reuse"),
    "NFR-009": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr804_native_broker_email_queue_is_configured_idempotent_and_auditable", "models/logistics_idp.py", ["LogisticsCase.queue_broker_email", "LogisticsCase.record_broker_email_delivery"], "native queued mail retains deterministic immutable request, queued receipt, optional terminal delivery evidence, and never direct-sends"),
    "FR-901": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr902_full_import_declaration_canonical_evidence_and_mismatches", "models/logistics_idp.py", ["LogisticsCase.reconcile_import_declaration"], "latest import declaration target uses E13/E15 canonical outputs"),
    "FR-902": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr902_full_import_declaration_canonical_evidence_and_mismatches", "models/logistics_idp.py", ["LogisticsCase.reconcile_import_declaration"], "canonical import declaration reconciliation covers declaration identity, parties, currency, totals, codes, HS, regime, quantity, UoM, line totals, aggregation and immutable evidence"),
    "FR-903": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr902_full_import_declaration_canonical_evidence_and_mismatches", "models/logistics_idp.py", ["LogisticsCase.reconcile_import_declaration"], "aggregate customs quantity by material/custom code and regime"),
    "FR-904": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_derived_stage_requires_terminal_checks_and_no_processing_before_completion", "models/logistics_idp.py", ["LogisticsCase._derive_lifecycle", "LogisticsCase.request_override", "LogisticsCase.action_complete"], "required IMPORT_DECLARATION review/block rejects completion; pass or manager-controlled override reaches ready before manager completion"),
    "FR-1001": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr1000_pre_customs_generation", "models/logistics_idp.py", ["LogisticsCase.generate_shipping_plan"], "Shipping Plan generates after SAP/ERP without declaration evidence"),
    "FR-1002": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr1002_full_mapping_template_values",  "models/logistics_idp.py", ["LogisticsCase.generate_shipping_plan"], "Shipping Plan maps all required default template fields and preserves source provenance"),
    "FR-1003": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr1003_trusted_provenance_late_declaration_supersession_idempotency", "models/logistics_idp.py", ["LogisticsCase.generate_shipping_plan"], "Shipping Plan maps trusted sources, supersedes after late declaration evidence, and remains idempotent"),
    "FR-1101": ("test_logistics_idp_configuration.TestLogisticsIdpConfiguration.test_legacy_customer_reference_is_read_only_operational_runtime_configuration", "models/logistics_idp.py", ["LogisticsPolicySource.select_effective_pack"], "profile effective enablement"),
    "FR-1102": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr1102_controlled_manager_approval_required_checks_current_override", "models/logistics_idp.py", ["LogisticsCase.generate_gate_pass"], "Gate Pass requires an enabled current manager-approved Shipping Plan and valid current required checks"),
    "FR-1104": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr1104_native_gate_pass_distribution_is_manager_scoped_configured_and_versioned", "models/logistics_idp.py", ["LogisticsCase.queue_gate_pass_distribution"], "manager-only native Gate Pass queue resolves one same-company profile, persists immutable intent evidence, reuses replay, and queues a fresh intent for a superseding output"),
    "FR-1103": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr1103_gate_template_semantic_fields", "models/logistics_idp.py", ["LogisticsCase.generate_gate_pass"], "Gate Pass maps required semantic fields from the approved Shipping Plan with template positions and provenance"),
    "SEC-001": ("test_logistics_idp_security.TestLogisticsIdpSecurity.test_unassigned_operator_cannot_read_case", "models/logistics_idp.py", ["LogisticsCase"], "least-privilege record visibility"),
    "SEC-003": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_sec003_completed_run_and_output_are_immutable_at_orm_sql_and_ui_boundaries", "models/logistics_idp.py", ["LogisticsExtractionRun", "LogisticsOutput"], "completed extraction run/output ORM and SQL mutation rejection; snapshot/artifact preservation; static read-only UI contract"),
    "SLA-004": ("test_logistics_idp_configuration.TestLogisticsIdpConfiguration.test_sla004_counterpart_alert_intents_are_idempotent_cadenced_and_suppressed", "models/logistics_idp.py", ["LogisticsCase._record_counterpart_alert_intent", "LogisticsCase.evaluate_compliance_deadlines"], "counterpart deadline alert intents retain deterministic immutable receipts, replay without duplication, create configured cadence successors, cap reminders, suppress no-send states, and never send mail"),
    "NFR-002": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_generic_inbound_attachment_m2m_is_rejected_on_create_and_write", "models/logistics_idp.py", ["LogisticsInboundJob.create", "LogisticsInboundJob.write", "LogisticsInboundJob.intake_supplier_email"], "generic inbound attachment M2M creation and write reject direct injection, including legacy context; controlled supplier-email intake alone receives the private linking token"),
    "NFR-003": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_nfr003_critical_inputs_are_run_bound_and_never_silent_pass", "models/logistics_idp.py", ["LogisticsDocument.process", "LogisticsDocument._create_extraction_critical_check", "LogisticsCase._derive_lifecycle"], "every completed extraction run creates an immutable run-bound critical-input check; configured canonical missing/invalid fields, warnings, low confidence, unknown or ambiguous classification require review and cannot silently pass"),
    "SRS-SECTION-31-2": ("kg:test_logistics_projection.TestLogisticsProjection.test_312_lineage_projection_is_deterministic_provenance_bounded_and_company_isolated", ("kg:services/logistics_projection.py", "kg:services/graph_projection.py", "kg:services/graph_service.py", "kg:services/graph_query.py"), ["LogisticsProjection.project_case", "GraphProjection.project_edge", "GraphService.neighbors", "GraphQuery.paths"], "case-centric relational KG projection is deterministic, provenance-bound, evidence-readable, bounded, company-isolated, and fail-closed for unverified trust"),
    "SRS-SECTION-4A-1-CORE-PRINCIPLE-POLICY-AS-DATA": ("test_logistics_idp_configuration.TestLogisticsIdpConfiguration.test_policy_candidate_import_is_draft_idempotent_and_diff_is_read_only", "models/logistics_idp.py", ["LogisticsPolicySource.validate_policy_payload", "LogisticsPolicySource.import_candidate", "LogisticsPolicySource.select_effective_pack"], "partial local import-bound metadata/hash/effective selection; full policy-as-data remains unproven"),
    "SRS-SECTION-4A-2-TRADE-COMPLIANCE-CONTEXT": ("test_logistics_idp_configuration.TestLogisticsIdpConfiguration.test_trade_compliance_authority_and_context_contract", "models/logistics_idp.py", ["LogisticsPolicySource.validate_compliance_context"], "canonical compliance context requires material evidence, rejects unknown authority, and fails missing context to review"),
    "SRS-SECTION-4A-3-SEPARATE-THREE-CONCEPTS": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_operational_exposure_is_separate_candidate_evidence_and_idempotent", "models/logistics_idp.py", ["LogisticsCase.record_operational_exposure"], "partial separate tier4/internal signal no mutating legal/document states, no external ingestion/authority"),
    "SRS-SECTION-4A-4-REGULATORY-SOURCE-TRUST-TIERS": ("test_logistics_idp_configuration.TestLogisticsIdpConfiguration.test_trade_compliance_authority_and_context_contract", "models/logistics_idp.py", ["LogisticsPolicySource._compute_governance_status"], "authority tiers distinguish authoritative evidence from customer-reference evidence, which routes to review"),
    "SRS-SECTION-4A-5-REGULATORY-CHANGE-LIFECYCLE": ("test_logistics_idp_configuration.TestLogisticsIdpConfiguration.test_srs_4a5_activation_reevaluates_only_bound_affected_cases", "models/logistics_idp.py", ["LogisticsPolicyActivation.decide", "LogisticsPolicyActivation._append_reevaluations", "LogisticsPolicySource.preview_impact"], "governed activation re-evaluates only exact affected, open, same-company cases in the preview window; decision/evidence lineage is immutable and replay is idempotent; approved legal source/full regulatory lifecycle, notifications, and broad applicability/impact facets remain unproven"),
    "SRS-SECTION-4A-8-SCENARIO-SHOCK-SIMULATION": ("test_logistics_idp_configuration.TestLogisticsIdpConfiguration.test_srs_4a8_simulation_is_test_only_deterministic_read_only_and_non_legal", "models/logistics_idp.py", ["LogisticsPolicySource.simulate_scenario"], "local test-only hypothetical simulation is deterministic, read-only, non-production, and carries no legal authority"),
    "SRS-SECTION-4A-9-STANDARDS-ALIGNMENT": ("test_logistics_idp_configuration.TestLogisticsIdpConfiguration.test_srs_4a9_cross_border_context_is_canonical_deterministic_and_fail_closed", "models/logistics_idp.py", ["LogisticsPolicySource.validate_compliance_context"], "local canonical aliases and hash-bound mapping evidence are deterministic; unknown/conflicting fields and extensions fail closed; incomplete context routes review; no standards or legal claim"),
    "SRS-SECTION-8-3-DERIVED-STAGE-RULE":  ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_derived_stage_requires_terminal_checks_and_no_processing_before_completion", "models/logistics_idp.py", ["LogisticsCase._derive_lifecycle", "LogisticsCase.action_complete"], "required review/block prevents ready/completion; pass/not_applicable aggregate reaches ready only without processing; manager completion has no processing documents"),
    "SRS-SECTION-7-1-LOGISTICS-IDP-CONTROL-TOWER": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_srs_section_7_1_control_tower_available_kpis_and_required_drilldowns", "models/logistics_idp.py", ["LogisticsCase.get_dashboard_data", "LogisticsCase.dashboard_drilldown"], "all supported Control Tower KPIs, every real required drilldown, company-scoped domains/results, and explicit unavailable KPI semantics"),
    "SRS-SECTION-7-2-DOCUMENT-INBOX": (("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_srs_section_7_2_open_source_email_is_persisted_same_company_read_only", "test_logistics_idp_security.TestLogisticsIdpSecurity.test_open_source_email_denies_missing_ambiguous_cross_company_non_email_and_unauthorized"), "models/logistics_idp.py", ["LogisticsDocument.action_open_source_email"], "authorized source-email documents resolve exactly one persisted same-company incoming mail.message by stored reference, return a read-only form action, and deny missing, ambiguous, cross-company, non-email, and unauthorized access without mutation or external mail"),
    "SRS-SECTION-7-5-EXCEPTION-AUDIT-CENTER": (("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_srs_section_7_5_exception_audit_center_projects_persisted_categories_read_only", "test_logistics_idp_security.TestLogisticsIdpSecurity.test_exception_audit_center_excludes_foreign_company_and_exposes_no_write_context"), "models/logistics_idp.py", ["LogisticsCase.exception_audit_domain", "LogisticsCase.action_open_exception_audit_center"], "authorized users receive a deterministic, company-scoped, read-only case projection of persisted exceptions, review/block checks, extraction/correlation errors, SLA breaches, overrides, and reprocessing evidence"),
    "SRS-SECTION-11-2-DESCRIPTION-MATCHING": ("test_logistics_idp_phase_1_3.TestLogisticsIdpPurePhase13.test_srs_11_2_description_similarity_is_candidate_only_never_pass", "services/reconciliation.py", ["description_similarity_candidates", "reconcile_documents"], "normalized-description similarity yields candidate suggestions only, never a standalone pass; below-threshold, disabled, or absent configuration produces no candidates and leaves the verdict unchanged"),
    "UAT-40": ("test_logistics_idp_configuration.TestLogisticsIdpConfiguration.test_policy_impact_preview_is_deterministic_read_only_and_fails_closed", "models/logistics_idp.py", ["LogisticsPolicySource.preview_impact"], "effective-dated candidate impact preview is deterministic, read-only, and identifies affected cases"),
    "UAT-42": (("test_logistics_idp_configuration.TestLogisticsIdpConfiguration.test_policy_impact_preview_is_deterministic_read_only_and_fails_closed", "test_logistics_idp_configuration.TestLogisticsIdpConfiguration.test_impact_preview_applicability_selectors_gate_cases"), "models/logistics_idp.py", ["LogisticsPolicySource.preview_impact", "LogisticsPolicySource.evaluate_applicability"], "deterministic country/HS applicability identifies affected cases and excludes unrelated or insufficient-context cases"),
    "UAT-43": ("test_restricted_party.TestRestrictedPartyRuntime.test_screening_hit_routes_review_and_is_idempotent", ("models/logistics_idp.py", "services/restricted_party.py"), ["LogisticsCase._restricted_party_screening", "screen_parties"], "ambiguous synthetic fuzzy hit creates an idempotent review with policy/list evidence; never auto-blocks"),
    "UAT-45": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_operational_exposure_is_separate_candidate_evidence_and_idempotent", "models/logistics_idp.py", ["LogisticsCase.record_operational_exposure"], "news-like operational exposure remains candidate evidence and does not mutate compliance verdict"),
    "UAT-47": ("test_logistics_idp_configuration.TestLogisticsIdpConfiguration.test_regulatory_source_freshness_fails_safe", "models/logistics_idp.py", ["LogisticsPolicySource.evaluate_freshness"], "stale mandatory regulatory source records review rather than silently passing"),
    "UAT-48": ("test_logistics_idp_configuration.TestLogisticsIdpConfiguration.test_srs_4a8_simulation_is_test_only_deterministic_read_only_and_non_legal", "models/logistics_idp.py", ["LogisticsPolicySource.simulate_scenario"], "hypothetical impact simulation is deterministic, read-only, non-production, and has no legal authority"),
    "SEC-006": ("test_logistics_idp_configuration.TestLogisticsIdpConfiguration.test_g5_submit_activate_and_reject_terminal_lifecycle", "models/logistics_idp.py", ["LogisticsPolicyActivation.submit", "LogisticsPolicyActivation.decide"], "versioned maker submission and reviewer terminal approval/rejection govern policy activation"),
    "SEC-007": ("test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_sec007_decision_manifest_binds_inputs_and_recheck_appends", "models/logistics_idp.py", ["LogisticsCase.reconcile_documents", "LogisticsPolicyDecision.create"], "immutable decision manifest binds exact reconciliation inputs and rechecks append decisions"),
    "SEC-008": ("test_logistics_idp_configuration.TestLogisticsIdpConfiguration.test_g5_replay_tamper_self_approval_and_stale_binding", "models/logistics_idp.py", ["LogisticsPolicyActivation.decide"], "maker/checker separation rejects self-approval before blocking-rule activation"),
    "SRS-SECTION-31-2-TRADE-COMPLIANCE-KNOWLEDGE-GRAPH-LOGICAL-GRAPH-NOT-MANDATORY-NEW-GRAPH-DATABASE": ("kg:test_logistics_projection.TestLogisticsProjection.test_projects_case_checks_and_exceptions", "kg:services/logistics_projection.py", ["LogisticsProjection.project_case"], "case-centric logical KG projects deterministic case, checks, policy sources, decisions, supplier profiles, PO material/HS lines, and exceptions; unresolved exception evidence remains candidate and trade/legal retrieval stays disabled"),
    "SRS-SECTION-31-4-ORIGIN-FTA-C-O": ("test_logistics_idp_configuration.TestLogisticsIdpConfiguration.test_origin_evidence_evaluator_is_deterministic_and_fail_closed", "services/reconciliation.py", ["evaluate_origin_evidence"], "structural Origin/FTA/C/O evidence intake is deterministic and review-only; legal qualification, approved origin rule packs, and authoritative policy remain blockers"),
    "SRS-SECTION-31-3-CUSTOMS-COMPLIANCE": ("test_logistics_idp_configuration.TestLogisticsIdpConfiguration.test_customs_compliance_capabilities_are_deterministic_and_fail_closed", ("models/logistics_idp.py", "services/reconciliation.py"), ["LogisticsCase.evaluate_taxable_value", "LogisticsCase.evaluate_compliance_deadlines", "customs_checks", "compare_line"], "customs regime and HS comparison, explicit-effective-FX taxable valuation, declaration line consistency and counterpart deadlines are deterministic and fail closed; licenses/permits/certificates, safeguard/remedy duty rates and post-entry review remain unproven"),
    "SRS-SECTION-31-5-RESTRICTED-PARTY-SANCTIONS-COMPLIANCE": (("test_restricted_party.TestRestrictedPartyPure.test_canonical_party_roles_and_identifiers_are_deterministic", "test_restricted_party.TestRestrictedPartyRuntime.test_screening_hit_routes_review_and_is_idempotent"), ("models/logistics_idp.py", "services/restricted_party.py"), ["canonical_party_evidence", "LogisticsCase._restricted_party_screening", "screen_parties"], "provider-neutral canonical party roles and identifiers are deterministic; synthetic restricted-party lists fuzzy-screen active same-company overlays; matches route to review with idempotent check results and a restricted_party_hit exception, never block, and foreign-company lists are ignored"),
    "SRS-SECTION-31-7-TARIFF-TRADE-REMEDIES-AND-ECONOMIC-SECURITY-MEASURES": ("test_logistics_idp_configuration.TestLogisticsIdpConfiguration.test_tariff_remedy_security_evaluator_is_deterministic_and_fail_closed", "services/reconciliation.py", ["evaluate_tariff_remedy_security"], "deterministic structural tariff/remedy/security intake is review-only and fail-closed; approved legal datasets, real schedules, rates, and legal PASS/BLOCK remain unproven"),
    "SRS-SECTION-31-6-EXPORT-CONTROLS": ("test_logistics_idp_configuration.TestLogisticsIdpConfiguration.test_export_controls_contract_and_fail_closed_evaluator", "models/logistics_idp.py", ["LogisticsPolicySource.validate_policy_payload", "LogisticsPolicySource.evaluate_export_controls"], "schema-ready optional export-controls overlay and deterministic fail-closed review-only evaluator; legal dataset blocker remains"),
    "SRS-SECTION-31-8-REGULATORY-IMPACT-INTELLIGENCE": ("test_logistics_idp_configuration.TestLogisticsIdpConfiguration.test_regulatory_change_notification_intent_is_deterministic_and_review_only", ("models/logistics_idp.py", "services/reconciliation.py"), ["LogisticsPolicySource.evaluate_applicability", "LogisticsPolicySource.preview_impact", "build_regulatory_change_notification"], "deterministic applicability and structural regulatory-change notification intent remain pending review, hash-bound, no-send, no auto activation; unsupported selectors remain deferred"),
    "SRS-SECTION-31-9-GEOPOLITICAL-TRADE-RESILIENCE": ("test_logistics_idp_configuration.TestLogisticsIdpConfiguration.test_resilience_metrics_are_deterministic_and_scoped", "models/logistics_idp.py", ["LogisticsPolicySource.compute_resilience_metrics"], "deterministic resilience metrics over persisted governance snapshots: source freshness, publication-to-ingestion and publication-to-activation latency, open impact assessments and unresolved stale-policy exposures; time-to-clear and lane-coverage KPIs deferred until lane/country entities exist"),
    "SRS-SECTION-31-1-FROM-IDP-TO-TRADE-COMPLIANCE-INTELLIGENCE": ("test_logistics_idp_configuration.TestLogisticsIdpConfiguration.test_idp_to_trade_compliance_intelligence_evolution_chain", ("models/logistics_idp.py", "services/reconciliation.py"), ["LogisticsPolicySource.validate_policy_payload", "LogisticsPolicySource.preview_impact", "LogisticsPolicySource.compute_resilience_metrics", "customs_checks", "compare_line"], "evolution chain holds end to end: cross-document control, customs compliance, governed trade-compliance packs, candidate impact preview and resilience metrics all sit above the evidence acquisition layer; full global legal-content coverage stays phased"),
    "SRS-SECTION-31-10-AI-ROLE": ("test_logistics_idp_configuration.TestLogisticsIdpConfiguration.test_ai_role_guardrails_hold", "models/logistics_idp.py", ["LogisticsPolicySource.import_candidate", "LogisticsPolicyActivation.submit", "LogisticsPolicyActivation.decide", "LogisticsPolicySource.write", "LogisticsPolicySource.select_effective_pack"], "AI-drafted candidates stay draft, self-approval of blocking rules is rejected, authoritative dates are immutable, and conflicting active packs fail closed instead of being silently resolved; summarization/explanation surfaces remain unproven"),
}
ACHIEVED_CAPABILITY_EXECUTORS = set(CAPABILITY_EXECUTORS) - {'NFR-009', 'SRS-SECTION-4A-1-CORE-PRINCIPLE-POLICY-AS-DATA', 'SRS-SECTION-4A-3-SEPARATE-THREE-CONCEPTS', 'SRS-SECTION-4A-5-REGULATORY-CHANGE-LIFECYCLE'}

IMPLEMENTATION_EVIDENCE = {
    "SRS-SECTION-7-1-LOGISTICS-IDP-CONTROL-TOWER": ("models/logistics_idp.py", ["LogisticsCase"], "native case control-tower model/views implemented; full KPI/drill-down acceptance not proven"),
    "SRS-SECTION-7-2-DOCUMENT-INBOX": ("models/logistics_idp.py", ["LogisticsDocument"], "native inbox includes a role-gated, same-company, uniquely persisted incoming-email opener returning a read-only mail form; duplicate marking and other inbox acceptance remain partial"),
    "SRS-SECTION-7-3-LOGISTICS-CASE-WORKSPACE": ("models/logistics_idp.py", ["LogisticsCase"], "native case workspace model/views implemented; full cockpit acceptance not proven"),
}

WP5_SCENARIO_MATRIX = {
    "duplicate_scan": {"source_document": "TỰ ĐỘNG KIỂM SOÁT CHỨNG TỪ XUẤT NHẬP KHẨU.pdf", "source_section": "Kiểm soát hàng hoá trùng lặp", "requirements": ["DUP-001", "DUP-003"], "decision": "current", "owner": "Logistics", "implementation_symbol": "LogisticsDocument.intake_content", "tests": ["test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_dossier_exact_duplicate_disposition_is_durable_and_company_safe"], "status": "achieved", "assertion": "same-content dossier intake preserves one canonical document and durable excluded duplicate disposition"},
    "repeated_subject": {"source_document": "TỰ ĐỘNG KIỂM SOÁT CHỨNG TỪ XUẤT NHẬP KHẨU.pdf", "source_section": "Tiếp nhận email cùng/khác subject", "requirements": ["FR-106", "FR-107"], "decision": "current", "owner": "Logistics", "implementation_symbol": "LogisticsCase._record_same_subject_thread_entry", "tests": ["test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr106_same_normalized_subject_keeps_latest_active_history_and_replay", "test_logistics_idp_security.TestLogisticsIdpSecurity.test_fr106_thread_history_is_company_scoped_and_intake_controlled", "test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr107_different_subject_requires_idempotent_review_and_explicit_active_thread_selection", "test_logistics_idp_security.TestLogisticsIdpSecurity.test_fr107_active_thread_selection_denies_unauthorized_and_cross_company"], "status": "achieved", "assertion": "same-company normalized same-subject intake retains immutable history with exactly one newest active entry and idempotent canonical replay; different subject preserves the existing case and creates one idempotent review activity until reviewer/manager explicitly selects a related same-company active thread"},
    "grouping": {"source_document": "TỰ ĐỘNG KIỂM SOÁT CHỨNG TỪ XUẤT NHẬP KHẨU.pdf", "source_section": "Gộp dòng hàng", "requirements": ["UAT-07", "FR-303"], "decision": "current", "owner": "Logistics", "implementation_symbol": "LogisticsCase.reconcile_documents", "tests": ["test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_two_po_lines_group_to_one_invoice_line_by_identity"], "status": "achieved", "assertion": "two PO lines with deterministic shared identity aggregate remaining quantity against one supplier line and preserve both PO references"},
    "uom_per_1000_pricing": {"source_document": "Tài_liệu_mô_tả_giải_pháp_đối_chiếu_chứng_từ_tự_động.pdf", "source_section": "Unit Price sau hệ số Per", "requirements": ["RULE-003", "UAT-08"], "decision": "current", "owner": "Finance", "implementation_symbol": "normalize_price", "tests": ["test_logistics_idp_phase_1_3.TestLogisticsIdpPurePhase13.test_currency_uom_per_rounding_and_customs_fail_safe"], "status": "achieved", "assertion": "1000 G to 1 KG plus price_per=1000 and deterministic FX normalizes to pass"},
    "ordering": {"source_document": "Tài_liệu_mô_tả_giải_pháp_đối_chiếu_chứng_từ_tự_động.pdf", "source_section": "Sequence Warning", "requirements": ["FR-606", "UAT-28", "UAT-29"], "decision": "no-action", "owner": "Logistics", "implementation_symbol": "LogisticsCase.reconcile_documents", "tests": ["test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_profile_sequence_warn_and_block_preserve_identity_mapping", "test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_reconciliation_po_order_grouping_currency_and_regime_are_deterministic"], "status": "achieved", "assertion": "deterministic identity mapping retains PO order across grouped lines; sequence_policy warn routes review and block rejects; historical QQ position issue remains no-action"},
    "currency": {"source_document": "Tài_liệu_mô_tả_giải_pháp_đối_chiếu_chứng_từ_tự_động.pdf", "source_section": "Đa tiền tệ", "requirements": ["FR-705", "UAT-19"], "decision": "current", "owner": "Finance", "implementation_symbol": "normalize_currency", "tests": ["test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_reconciliation_po_order_grouping_currency_and_regime_are_deterministic", "test_logistics_idp_configuration.TestLogisticsIdpConfiguration.test_usd_vnd_taxable_value_requires_explicit_effective_fx"], "status": "achieved", "assertion": "currency mismatch routes to review; explicit effective USD/VND FX produces deterministic invoice and taxable values; missing or reversed FX routes review"},
    "missing_master": {"source_document": "Tài_liệu_mô_tả_giải_pháp_đối_chiếu_chứng_từ_tự_động.pdf", "source_section": "Cảnh báo thiếu Master Data/DSNVL", "requirements": ["FR-205", "UAT-18"], "decision": "current", "owner": "Logistics", "implementation_symbol": "LogisticsPoSnapshot.import_batch", "tests": ["test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr205_material_reference_failures_and_no_material_line"], "status": "achieved", "assertion": "missing required master/DSNVL fails; explicit no-material lines remain allowed"},
    "customs_regime_changes": {"source_document": "TỰ ĐỘNG KIỂM SOÁT CHỨNG TỪ XUẤT NHẬP KHẨU.pdf", "source_section": "Kiểm soát quy định thay đổi", "requirements": ["FR-404", "UAT-17", "SRS-SECTION-4A-5-REGULATORY-CHANGE-LIFECYCLE"], "decision": "current", "owner": "Logistics", "implementation_symbol": "LogisticsCase.reconcile_documents", "tests": ["test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_reconciliation_po_order_grouping_currency_and_regime_are_deterministic", "test_logistics_idp_configuration.TestLogisticsIdpConfiguration.test_lidp03_effective_regime_policy_selection_stays_non_authoritative", "test_logistics_idp_configuration.TestLogisticsIdpConfiguration.test_policy_impact_preview_is_deterministic_read_only_and_fails_closed", "test_logistics_idp_configuration.TestLogisticsIdpConfiguration.test_srs_4a5_activation_reevaluates_only_bound_affected_cases"], "status": "partial", "assertion": "deterministic reconciliation routes mixed E11/E15 lines to review; effective-dated selection chooses the exact non-authoritative customer-reference policy; impact preview is deterministic and read-only; synthetic governed activation re-evaluates affected open cases only and creates one idempotent internal notification; partial solely because approved legal authority/dataset/oracle and real legal activation remain unproven if authority requires them"},
    "import_declaration_timing": {"source_document": "Tài_liệu_mô_tả_giải_pháp_đối_chiếu_chứng_từ_tự_động.pdf", "source_section": "Cảnh báo 10 ngày", "requirements": ["SLA-001", "FR-1001", "FR-1003", "UAT-38"], "decision": "superseded", "owner": "Logistics", "implementation_symbol": "LogisticsCase.generate_shipping_plan", "tests": ["test_logistics_idp_runtime.TestLogisticsIdpRuntime.test_fr1003_trusted_provenance_late_declaration_supersession_idempotency", "test_logistics_idp_configuration.TestLogisticsIdpConfiguration.test_counterpart_deadline_boundaries_and_configuration"], "status": "achieved", "assertion": "configured deadline is pending before warning offset, warning on the pre-deadline boundary, overdue exactly on day 10, satisfied when counterpart exists; late declaration supersedes Shipping Plan idempotently"},
}


def _stable_slug(value):
    return re.sub(r"[^A-Z0-9]+", "-", value.upper()).strip("-")


def validate_capability_executors(registry=CAPABILITY_EXECUTORS):
    expected_ids = set(CAPABILITY_EXECUTORS)
    if set(registry) != expected_ids:
        raise ValueError("executor allowlist must contain exactly the reviewed %d requirements" % len(expected_ids))
    validated = {}
    for requirement_id, (relative_ids, source_paths, symbols, assertion_scope) in registry.items():
        relative_ids = (relative_ids,) if isinstance(relative_ids, str) else relative_ids
        source_paths = (source_paths,) if isinstance(source_paths, str) else source_paths
        resolved_ids = []
        for relative_id in relative_ids:
            package_key, separator, test_id = relative_id.partition(":")
            package_key, test_id = (package_key, test_id) if separator else ("idp", relative_id)
            if package_key not in TEST_PACKAGES or "*" in test_id or "harness" in test_id.lower() or len(test_id.split(".")) != 3 or not test_id.split(".")[-1].startswith("test_"):
                raise ValueError("generic, wildcard, or class-level executor ID rejected: %s" % relative_id)
            module_name, class_name, method_name = test_id.split(".")
            test_path = TEST_PACKAGES[package_key][1] / (module_name + ".py")
            if not test_path.is_file():
                raise ValueError("executor test module missing: %s" % test_path)
            classes = {node.name: node for node in ast.parse(test_path.read_text()).body if isinstance(node, ast.ClassDef)}
            methods = {node.name for node in classes[class_name].body if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef))} if class_name in classes else set()
            if method_name not in methods:
                raise ValueError("executor test class/method missing: %s" % relative_id)
            resolved_ids.append(TEST_PACKAGES[package_key][0] + "." + test_id)
        implementation = {}
        for source_path in source_paths:
            package_key, separator, path = source_path.partition(":")
            package_key, path = (package_key, path) if separator else ("idp", source_path)
            implementation_path = (HERE.parents[1] if package_key == "idp" else ROOT / "insilos/apps/insilos_knowledge_graph") / path
            if package_key not in TEST_PACKAGES or not implementation_path.is_file():
                raise ValueError("implementation source missing: %s" % implementation_path)
            implementation[source_path] = {node.name: node for node in ast.parse(implementation_path.read_text()).body if isinstance(node, (ast.ClassDef, ast.FunctionDef, ast.AsyncFunctionDef))}
        for symbol in symbols:
            parts = symbol.split(".")
            if not any((parts[0] in top and (len(parts) == 1 or parts[1] in {child.name for child in top[parts[0]].body if isinstance(child, (ast.FunctionDef, ast.AsyncFunctionDef))})) or (Path(source_path).stem == parts[0] and parts[1] in top) for source_path, top in implementation.items()):
                raise ValueError("implementation symbol missing: %s" % symbol)
        if not assertion_scope.strip():
            raise ValueError("assertion_scope missing: %s" % requirement_id)
        validated[requirement_id] = {
            "unittest_ids": resolved_ids,
            "implementation_symbols": symbols,
            "assertion_scope": assertion_scope,
        }
    return validated


def validate_wp5_scenario_matrix(matrix=WP5_SCENARIO_MATRIX):
    accepted = {"duplicate_scan", "repeated_subject", "grouping", "uom_per_1000_pricing", "ordering", "currency", "missing_master", "customs_regime_changes", "import_declaration_timing"}
    if set(matrix) != accepted:
        raise ValueError("two-PDF matrix has duplicate or unmapped accepted scenarios")
    source_rows = [(item.get("source_document"), item.get("source_section")) for item in matrix.values()]
    if len(source_rows) != len(set(source_rows)) or {row[0] for row in source_rows} != {"Tài_liệu_mô_tả_giải_pháp_đối_chiếu_chứng_từ_tự_động.pdf", "TỰ ĐỘNG KIỂM SOÁT CHỨNG TỪ XUẤT NHẬP KHẨU.pdf"}:
        raise ValueError("two-PDF source rows must be unique and mapped")
    for scenario, item in matrix.items():
        required = {"source_document", "source_section", "requirements", "decision", "owner", "implementation_symbol", "tests", "status", "assertion"}
        if set(item) != required or item["decision"] not in {"current", "superseded", "no-action"} or item["owner"] not in {"Finance", "Logistics"} or item["status"] not in {"achieved", "partial", "missing"} or not item["requirements"] or not item["tests"] or not item["assertion"]:
            raise ValueError("invalid two-PDF scenario mapping: %s" % scenario)
        for test in item["tests"]:
            module_name, class_name, method_name = test.split(".")
            tree = ast.parse((HERE.parent / (module_name + ".py")).read_text())
            classes = {node.name: node for node in tree.body if isinstance(node, ast.ClassDef)}
            methods = {node.name for node in classes[class_name].body if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef))}
            if method_name not in methods:
                raise ValueError("two-PDF executor missing: %s" % test)
    if {item["decision"] for item in matrix.values()} != {"current", "superseded", "no-action"}:
        raise ValueError("two-PDF matrix must preserve current, superseded and no-action decisions")
    return matrix


def validate_implementation_evidence(registry=IMPLEMENTATION_EVIDENCE):
    if set(registry) != {"SRS-SECTION-7-1-LOGISTICS-IDP-CONTROL-TOWER", "SRS-SECTION-7-2-DOCUMENT-INBOX", "SRS-SECTION-7-3-LOGISTICS-CASE-WORKSPACE"}:
        raise ValueError("implementation evidence allowlist must contain exactly workspace sections 7.1/7.2/7.3")
    validated = {}
    for requirement_id, (source_path, symbols, assertion_scope) in registry.items():
        if any("*" in symbol or "." in symbol for symbol in symbols):
            raise ValueError("wildcard or method-level implementation evidence rejected: %s" % requirement_id)
        implementation_path = HERE.parents[1] / source_path
        tree = ast.parse(implementation_path.read_text())
        known = {node.name for node in tree.body if isinstance(node, (ast.ClassDef, ast.FunctionDef, ast.AsyncFunctionDef))}
        if not symbols or any(symbol not in known for symbol in symbols):
            raise ValueError("implementation evidence symbol missing: %s" % requirement_id)
        validated[requirement_id] = {"implementation_symbols": symbols, "assertion_scope": assertion_scope}
    return validated


def canonical_srs_crosswalk(raw_inventory):
    data = json.loads(CROSSWALK.read_text())
    if data.get("schema_version") != 1 or data.get("source_authority") != str(SRS.relative_to(ROOT)):
        raise ValueError("canonical SRS crosswalk schema")
    if data.get("source_sha256") != digest(SRS.read_bytes()):
        raise ValueError("stale canonical SRS crosswalk source")
    exclusions = data.get("exclusions")
    if not isinstance(exclusions, list) or not exclusions:
        raise ValueError("canonical SRS crosswalk exclusions missing")
    raw = {item["srs_id"]: item["title"] for item in raw_inventory}
    excluded = [item.get("source_id") for item in exclusions]
    if len(excluded) != len(set(excluded)) or any(item not in raw for item in excluded):
        raise ValueError("canonical SRS crosswalk unknown or duplicate source ID")
    if any(raw[item.get("source_id")] != item.get("source_heading") for item in exclusions):
        raise ValueError("stale canonical SRS crosswalk source heading")
    if any(item.get("disposition") != "non-atomic operating profile" or not item.get("rationale") for item in exclusions):
        raise ValueError("canonical SRS crosswalk disposition")
    governed = set(raw) - set(excluded)
    if data.get("raw_heading_count") != len(raw) or data.get("governed_requirement_count") != len(governed):
        raise ValueError("canonical SRS crosswalk counts")
    return data, governed


def authoritative_srs_inventory(cases, passed_test_ids=()):
    text = SRS.read_text()
    executors = validate_capability_executors()
    implementations = validate_implementation_evidence()
    passed_test_ids = set(passed_test_ids)
    headings = [(len(level), title.strip()) for level, title in re.findall(r"^(#{1,4})\s+(.+?)\s*$", text, re.M)]
    explicit = []
    for _level, title in headings:
        match = re.match(r"((?:FR|AI|RULE|EXC|DUP|SLA|RPT|INT|SEC|NFR|UAT)-\d+[A-Z]?|VA-\d+)\b", title)
        if match:
            explicit.append((match.group(1), title))
    stable = []
    section = None
    for level, title in headings:
        plain = re.sub(r"[`*_]", "", title)
        if level == 1:
            section = re.match(r"(\d+[A-Z]?)\.", plain)
            section = section.group(1) if section else None
        kind = None
        if section == "5" and level == 3:
            kind = "ROLE"
        elif section == "9" and level == 1:
            kind = "DOC"
        elif section == "26" and level == 2 and plain.startswith("Phase "):
            kind = "PHASE"
        elif section == "27":
            kind = "DOD"
        elif level in (1, 2) and not SRS_EXPLICIT_ID.match(plain.split(" ", 1)[0]):
            kind = "SECTION"
        if kind:
            stable.append((f"SRS-{kind}-{_stable_slug(plain)}", plain))
    mapped = {}
    for case in cases:
        for requirement_id in case["requirement_ids"]:
            mapped.setdefault(requirement_id, []).append(case["scenario_id"])
    inventory = []
    for requirement_id, title in sorted(dict(explicit + stable).items()):
        corpus_ids = sorted(set(mapped.get(requirement_id, [])))
        executor = executors.get(requirement_id)
        implementation = implementations.get(requirement_id)
        passed = executor and all(test_id in passed_test_ids for test_id in executor["unittest_ids"])
        partial = bool(implementation or corpus_ids)
        status = "achieved" if passed else "partial" if partial else "missing"
        inventory.append({
            "srs_id": requirement_id, "title": title,
            "product_spec_ids": [requirement_id] if requirement_id in mapped else [],
            "implementation_symbols": executor["implementation_symbols"] if executor else implementation["implementation_symbols"] if implementation else [],
            "executable_test_ids": executor["unittest_ids"] if executor else [],
            "assertion_scope": executor["assertion_scope"] if executor else implementation["assertion_scope"] if implementation else None,
            "corpus_ids": corpus_ids, "evidence_status": status,
            "coverage": "covered" if passed else "not_covered" if partial else "missing",
            "reason": "focused exact unittest PASS and static guards passed" if passed else "implementation evidence only; full acceptance not proven" if implementation else "executor not reconciled to current focused PASS evidence" if executor else "generic corpus case has no capability-specific executor/assertion" if corpus_ids else "no executable mapping",
        })
    counts = {status: sum(item["evidence_status"] == status for item in inventory)
              for status in ("achieved", "partial", "missing", "out_of_scope", "deferred")}
    counts.update(total=len(inventory), explicitly_numbered=len(explicit), stable_unnumbered=len(stable),
                  capability_specific_covered=sum(item["coverage"] == "covered" for item in inventory))
    crosswalk, governed_ids = canonical_srs_crosswalk(inventory)
    return {"schema_version": 1, "authority": str(SRS.relative_to(ROOT)), "authority_sha256": digest(SRS.read_bytes()),
            "items": inventory, "coverage_counts": counts,
            "canonical_crosswalk": {"path": str(CROSSWALK.relative_to(ROOT)), "sha256": digest(CROSSWALK.read_bytes()),
                                     "raw_heading_count": len(inventory), "governed_requirement_count": len(governed_ids),
                                     "excluded_source_ids": sorted(set(item["srs_id"] for item in inventory) - governed_ids)}}


def trace_category(row):
    suite = str(row.get("Suite") or "").lower()
    text = " ".join(str(row.get(key) or "") for key in ("Suite", "Type", "Title", "Steps", "Expected Result")).lower()
    if suite == "architecture & scope":
        return "test-only"
    if suite == "operations, ui & reporting":
        return "workflow/UI"
    if suite == "trade compliance & resilience" and any(term in text for term in ("policy", "effective", "version", "regulatory")):
        return "policy config"
    if any(term in text for term in ("external", "connector", "performance", "penetration", "outlook", "sap", "trigger.dev")):
        return "external/non-runnable"
    if any(term in text for term in ("extract", "schema", "ocr", "field accuracy", "classif")):
        return "extraction/schema"
    if any(term in text for term in ("rule engine", "deadline", "duplicate", "tolerance", "reconcil", "counterpart")):
        return "deterministic rule"
    if any(term in text for term in ("dashboard", "screen", "report", "workflow", "activity", "review queue", "workspace")):
        return "workflow/UI"
    if any(term in text for term in ("policy", "configuration", "profile", "effective-date", "version")):
        return "policy config"
    return "test-only"


def assert_non_reidentifying(fixtures, sensitive_values, fuzzy_threshold=0.88):
    generated = json.dumps(fixtures, ensure_ascii=False, sort_keys=True).casefold()
    for value in sensitive_values:
        sensitive = str(value).strip().casefold()
        if not sensitive:
            continue
        if sensitive in generated:
            raise ValueError("exact sensitive overlap")
        for token in re.findall(r"[\w@.+-]{5,}", generated):
            if SequenceMatcher(None, sensitive, token).ratio() >= fuzzy_threshold:
                raise ValueError("fuzzy sensitive overlap")
    return True


def derive_synthetic_fixtures(observations, sensitive_values=()):
    if not isinstance(observations, dict) or set(observations) - DERIVATION_KEYS:
        raise ValueError("derivation accepts abstract observations only")
    required = DERIVATION_KEYS - {"approval_status"}
    if not required <= observations.keys() or any(not isinstance(observations[key], (list, tuple)) for key in required):
        raise ValueError("invalid abstract observations")
    fixtures = [{
        "fixture_id": f"SYN-DERIVED-{index:03d}",
        "taxonomy": str(taxonomy),
        "field_types": [str(value) for value in observations["field_types"]],
        "workflow_states": [str(value) for value in observations["workflow_states"]],
        "validation_categories": [str(value) for value in observations["validation_categories"]],
        "contact": f"derived-{index:03d}@example.invalid",
        "metadata": {
            "derived_from_restricted_reference": True,
            "transformation_version": TRANSFORMATION_VERSION,
            "approval_status": observations.get("approval_status", "pending"),
        },
    } for index, taxonomy in enumerate(observations["taxonomy"], 1)]
    assert_non_reidentifying(fixtures, sensitive_values)
    return fixtures


def sanitized(value):
    return re.sub(r"Swarovski", "reference-customer", str(value or ""), flags=re.I)


def fixture(row, index):
    case_id = row["Test ID"]
    text = " ".join(str(row.get(key) or "") for key in ("Suite", "Type", "Title", "Preconditions", "Test Data", "Expected Result")).lower()
    mode = "success"
    if any(word in text for word in ("timeout", "unavailable", "retry", "dead-letter")):
        mode = "timeout"
    elif any(word in text for word in ("malformed", "corrupt", "wrong mime", "unsupported file", "invalid file")):
        mode = "malformed"
    mimetype = MIMES[(index - 1) % len(MIMES)]
    data = payload(case_id, index)
    content = binary(mimetype, data, multipage=(index % 11 == 0))
    metadata = {"scenario_id": case_id}
    if mimetype != "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        metadata["provider_response"] = data
    expected_error = None
    if mode == "timeout":
        metadata["fixture"] = "timeout"; expected_error = "TimeoutError"
    elif mode == "malformed":
        mimetype = "application/pdf"; content = b"%PDF-MALFORMED-SYNTHETIC"; expected_error = "ProcessingError"
    body = {"artifact_id": f"{case_id}-INPUT", "media_type": mimetype, "encoding": "base64", "content": base64.b64encode(content).decode(), "metadata": metadata}
    body["sha256"] = digest(content)
    normalized = json.loads(json.dumps(data["payload"]))
    for line in normalized["lines"]:
        line["quantity"], line["unit_price"] = str(line["quantity"]), str(line["unit_price"])
    oracle = {"classification": "purchase_order" if not expected_error else None, "normalized_fields": normalized if not expected_error else {}, "correlations": {"case_id": case_id}, "checks": ["binary_validation", "schema_validation", "deterministic_hash"], "verdict": "allowed_failure" if expected_error else "pass", "outputs": [] if expected_error else ["normalized_document"], "lineage": [body["artifact_id"]], "allowed_failures": [expected_error] if expected_error else []}
    oracle["sha256"] = digest({k: v for k, v in oracle.items() if k != "sha256"})
    return body, oracle


def build(derived_observations=None, sensitive_values=()):
    executors = validate_capability_executors()
    wp5_matrix = validate_wp5_scenario_matrix()
    passed_test_ids = {test_id for requirement_id in ACHIEVED_CAPABILITY_EXECUTORS for test_id in executors[requirement_id]["unittest_ids"]}
    detailed = rows("UAT Cases")
    golden_rows = rows("Golden 48")
    golden_map = {}
    for row in golden_rows:
        detailed_ids = [item.strip() for item in str(row["Mapped Detailed Tests"] or "").split(",") if item.strip()]
        for detailed_id in detailed_ids:
            golden_map.setdefault(detailed_id, []).append(row["Golden ID"])
    cases = []
    for index, row in enumerate(detailed, 1):
        artifact, oracle = fixture(row, index)
        requirement_ids = sorted(set(REQUIREMENT_ID.findall(str(row["Requirement"] or ""))))
        cases.append({"scenario_id": row["Test ID"], "golden_ids": golden_map.get(row["Test ID"], []), "trace_category": trace_category(row), "suite": sanitized(row["Suite"]), "requirement": sanitized(row["Requirement"]), "requirement_ids": requirement_ids, "provenance": {"type": "synthetic", "source_hash": artifact["sha256"], "sanitization_class": "fully_synthetic", "oracle_version": ORACLE_VERSION}, "terminal_status": "error" if oracle["allowed_failures"] else "pass", "phase": sanitized(row["Phase"]), "priority": sanitized(row["Priority"]), "test_type": sanitized(row["Type"]), "title": sanitized(row["Title"]), "preconditions": sanitized(row["Preconditions"]), "test_data": sanitized(row["Test Data"]), "steps": sanitized(row["Steps"]), "expected_result": sanitized(row["Expected Result"]), "evidence_required": sanitized(row["Evidence Required"]), "source_traceability": sanitized(row["Source / Traceability"]), "release_gate": row["Release Gate"], "artifact": artifact, "oracle": oracle})
    corpus = {"schema_version": 6, "corpus_id": "insilos-logistics-idp-198-synthetic-development-v6", "metadata": {"synthetic": True, "non_production": True, "deterministic": True, "formal_golden_acceptance": False, "independent_compliance_certification": False}, "workbook": {"path": WORKBOOK.name, "sha256": digest(WORKBOOK.read_bytes())}, "inventory": {"detailed": 198, "golden": 48}, "engine_capability_tests": [{"test_id": "ENGINE-SYNTHETIC-MULTI-ATTACHMENT-001", "provenance_type": "synthetic", "claim_scope": "engine_capability_only_not_observed_reference_stratum", "attachment_count": 2, "attachment_hashes": [cases[0]["artifact"]["sha256"], cases[1]["artifact"]["sha256"]], "expected": {"accepted": 2, "terminal_status": "pass"}}], "workbook_contract": workbook_contract(), "cases": cases}
    corpus["authoritative_srs_traceability"] = authoritative_srs_inventory(cases, passed_test_ids)
    corpus["wp5_source_requirement_test_matrix"] = wp5_matrix
    corpus["executor_evidence"] = {"scope": "focused exact unittest run", "result": "PASS", "test_ids": sorted(passed_test_ids)}
    if derived_observations:
        derived = derive_synthetic_fixtures(derived_observations, sensitive_values)
        if all(item["metadata"]["approval_status"] == "approved" for item in derived):
            corpus["derived_fixtures"] = derived
    corpus["manifest_sha256"] = digest({k: v for k, v in corpus.items() if k != "manifest_sha256"})
    return corpus


def verify(corpus):
    rebuilt = build()
    assert corpus == rebuilt, "corpus differs from deterministic rebuild"
    assert len(corpus["cases"]) == 198 and len({item["scenario_id"] for item in corpus["cases"]}) == 198
    assert len({golden for item in corpus["cases"] for golden in item["golden_ids"]}) == 48
    trace = corpus["authoritative_srs_traceability"]
    crosswalk, governed_ids = canonical_srs_crosswalk(trace["items"])
    assert trace["coverage_counts"]["total"] == crosswalk["raw_heading_count"] == 311
    assert len(governed_ids) == crosswalk["governed_requirement_count"] == 310
    capability = corpus["engine_capability_tests"][0]
    assert corpus["wp5_source_requirement_test_matrix"] == validate_wp5_scenario_matrix()
    assert {item["status"] for item in corpus["wp5_source_requirement_test_matrix"].values()} == {"achieved", "partial"}
    assert capability["provenance_type"] == "synthetic" and capability["claim_scope"] == "engine_capability_only_not_observed_reference_stratum"
    assert capability["attachment_count"] == len(set(capability["attachment_hashes"])) == 2
    forbidden = re.compile(r"Swarovski|BEGIN (?:RSA |OPENSSH )?PRIVATE KEY|(?:password|secret|token)\s*[:=]", re.I)
    executable_requirements = {requirement for case in corpus["cases"] for requirement in case["requirement_ids"]}
    for coverage in corpus["workbook_contract"]["requirement_coverage"]:
        if str(coverage["Coverage Status"]).lower() == "covered":
            assert coverage["Requirement ID"] in executable_requirements, f"covered requirement without executable mapping: {coverage['Requirement ID']}"
    deferred_requirements = {"4A.5", "4A.6"}
    deferred_golden = {
        golden_id
        for case in corpus["cases"]
        if deferred_requirements.intersection(case["requirement_ids"])
        for golden_id in case["golden_ids"]
    }
    assert deferred_golden == {"UAT-40", "UAT-41"}, "deferred trade Golden mapping drifted"
    for case in corpus["cases"]:
        content = base64.b64decode(case["artifact"]["content"], validate=True)
        assert digest(content) == case["artifact"]["sha256"]
        assert case["provenance"] == {"type": "synthetic", "source_hash": case["artifact"]["sha256"], "sanitization_class": "fully_synthetic", "oracle_version": ORACLE_VERSION}
        assert case["requirement_ids"] == sorted(set(REQUIREMENT_ID.findall(case["requirement"])))
        assert case["terminal_status"] in {"pass", "review", "error", "quarantined"}
        assert case["trace_category"] in TRACE_CATEGORIES
        assert digest({k: v for k, v in case["oracle"].items() if k != "sha256"}) == case["oracle"]["sha256"]
        assert not forbidden.search(json.dumps(case, ensure_ascii=False))
    print(f"Synthetic corpus: VALID detailed=198 golden=48 manifest={corpus['manifest_sha256']}")


def _zero_counts():
    return {
        "source_records": {"discovered": 0, "fetched": 0, "unique": 0, "failed": 0, "skipped": 0},
        "documents": {"accepted": 0, "pass": 0, "review": 0, "error": 0, "quarantined": 0},
        "attachments": {"discovered": 0, "downloaded": 0, "rejected": 0},
        "pages": {"discovered": 0, "processed": 0, "failed": 0},
        "extraction_runs": {"attempted": 0, "succeeded": 0, "review": 0, "error": 0},
        "detailed_golden": {"detailed_pass": 0, "detailed_fail": 0, "golden_pass": 0, "golden_fail": 0, "unmapped": 0},
        "iap_ledger": {"holds": 0, "consumes": 0, "releases": 0, "refunds": 0},
        "trigger": {"attempts": 0, "retries": 0, "dead_letter": 0},
        "parity": {"pass": 0, "fail": 0, "mismatch": 0},
    }


def gate_status(raw_status):
    return "PASS_WITH_REVIEW" if raw_status == "PARTIAL_VERIFY" else raw_status


def _development_acceptance_counts():
    result = json.loads(DEVELOPMENT_UAT_RESULT.read_text())
    detailed = result["detailed_case_counts"]
    golden = result["golden_counts"]
    assert golden == {"passed": 46, "partial": 2, "failed": 0, "not_run": 0, "skipped": 0}
    return detailed, golden


def _approved_oracles(path):
    if not path:
        return {}
    payload = json.loads(Path(path).read_text())
    return {str(item["source_hash"]): item for item in payload.get("oracles", [])
            if item.get("review_status") == "independently_approved"
            and item.get("reviewer_independent_of_exporter") is True}


def reference_evidence(inventory_path, oracle_path=None, limit=3, combined=False, source_summary=None, error_inventory_path=None):
    inventory_path = Path(inventory_path).resolve()
    if ROOT in inventory_path.parents:
        raise ValueError("raw/quarantine inventory phải nằm ngoài Git worktree")
    approved = _approved_oracles(oracle_path)
    records, source_hashes, strata = [], set(), {"no_attachment": 0, "exactly_one": 0, "multi_attachment": 0}
    pages = attachments = attachments_downloaded = 0
    with inventory_path.open() as stream:
        for line in stream:
            item = json.loads(line)
            source_hash = digest(str(item["source_id"]).encode())
            source_hashes.add(source_hash)
            downloaded = len(item.get("attachments") or [])
            count = int(item.get("source_attachment_count", downloaded))
            stratum = "no_attachment" if count == 0 else "exactly_one" if count == 1 else "multi_attachment"
            strata[stratum] += 1
            attachments += count
            attachments_downloaded += downloaded
            page_count = len(re.findall(r"--- Page \d+ / \d+ ---", str(item.get("raw_data") or "")))
            pages += page_count
            sampled_strata = {record["stratum"] for record in records}
            if len(records) < limit and stratum in {"no_attachment", "exactly_one"} and stratum not in sampled_strata:
                oracle = approved.get(source_hash)
                records.append({
                    "source_hash": source_hash, "stratum": stratum, "attachment_count": count,
                    "attachment_hashes": sorted(value["sha256"] for value in item.get("attachments") or []),
                    "pages": page_count, "provenance_type": "raw_quarantined",
                    "sanitization_class": "metadata_only_non_reidentifying",
                    "transformation_version": TRANSFORMATION_VERSION,
                    "oracle_status": "independently_approved" if oracle else "candidate_requires_independent_ui_review",
                    "terminal_status": oracle.get("expected_terminal_status", "review") if oracle else "review",
                    "requirement_ids": oracle.get("requirement_ids", ["IDP-CORPUS-001", "IDP-CORPUS-003"]) if oracle else ["IDP-CORPUS-001", "IDP-CORPUS-003"],
                })
    if source_summary:
        summary = json.loads(Path(source_summary).read_text())
        expected = summary["strata"]
        if expected["multi_attachment"] != 0:
            raise ValueError("source summary multi_attachment trái known source limit")
        if expected["no_attachment"] != strata["no_attachment"]:
            raise ValueError("source summary no_attachment mismatch")
        strata = expected
        attachments = expected["exactly_one"]
    rejected_source_hashes = set()
    if error_inventory_path:
        error_inventory_path = Path(error_inventory_path).resolve()
        if ROOT in error_inventory_path.parents:
            raise ValueError("raw/quarantine errors phải nằm ngoài Git worktree")
        with error_inventory_path.open() as stream:
            rejected_source_hashes = {digest(str(json.loads(line)["source_id"]).encode()) for line in stream if line.strip()}
    counts = _zero_counts()
    fetched = source_hashes | rejected_source_hashes
    total = len(fetched) if error_inventory_path else sum(strata.values())
    counts["source_records"].update(discovered=total, fetched=total, unique=total, failed=len(rejected_source_hashes), skipped=0)
    counts["documents"].update({"accepted": len(records), "review": sum(r["terminal_status"] == "review" for r in records), "pass": sum(r["terminal_status"] == "pass" for r in records)})
    counts["attachments"].update(discovered=attachments + len(rejected_source_hashes), downloaded=attachments_downloaded, rejected=len(rejected_source_hashes))
    counts["pages"].update(discovered=pages)
    if combined:
        detailed, golden = _development_acceptance_counts()
        counts["detailed_golden"].update(
            detailed_pass=detailed["passed"], detailed_fail=detailed["failed"],
            golden_pass=golden["passed"], golden_fail=golden["failed"],
            unmapped=sum(r["oracle_status"] != "independently_approved" for r in records),
        )
        counts["detailed_golden"]["golden_partial"] = golden["partial"]
    else:
        counts["detailed_golden"]["unmapped"] = sum(r["oracle_status"] != "independently_approved" for r in records)
    counts["parity"].update({"pass": 1, "mismatch": 0})
    limitations = {"multi_attachment": "not_observed/source_limit" if strata["multi_attachment"] == 0 else "observed"}
    raw_status = "PASS" if records and all(r["oracle_status"] == "independently_approved" for r in records) and strata["multi_attachment"] > 0 else "PARTIAL_VERIFY" if records else "REVIEW"
    gate = gate_status(raw_status)
    evidence = {
        "schema_version": 1, "raw_status": raw_status, "derived_gate_status": gate, "gate": gate, "scope": "combined" if combined else "small_reference",
        "environment": "local_deterministic_metadata_only", "generated_at": datetime.now(timezone.utc).isoformat(),
        "source_inventory": {"records": total, "accepted_records": len(source_hashes), "attachment_rejected_records": len(rejected_source_hashes), "strata": strata, "limitations": limitations},
        "records": records, "counts": {"run": counts, "cumulative": json.loads(json.dumps(counts))},
        "hashes": {"inventory_sha256": digest(inventory_path.read_bytes()), "error_inventory_sha256": digest(error_inventory_path.read_bytes()) if error_inventory_path else None, "corpus_manifest_sha256": build()["manifest_sha256"], "oracle_file_sha256": digest(Path(oracle_path).read_bytes()) if oracle_path else None},
        "oracle_policy": {"independent_approval_required": True, "candidate_self_approval_allowed": False},
        "failures": [] if gate in {"PASS", "PASS_WITH_REVIEW"} else ["no candidate records available"],
        "review_requirements": ["independent oracle and observed multi-attachment stratum unavailable"] if gate == "PASS_WITH_REVIEW" else [],
        "claims": {"full_pass": gate == "PASS", "provider_bulk_calls": False, "legal_verdict": "REVIEW", "content_status": "unverified", "activation_allowed": False, "activation_blocker": "BLOCKED_APPROVED_POLICY_DATASET"},
    }
    assert counts["source_records"]["discovered"] == len(source_hashes | rejected_source_hashes)
    assert not (source_hashes & rejected_source_hashes)
    assert counts["attachments"]["rejected"] == len(rejected_source_hashes)
    assert counts["attachments"]["discovered"] >= strata["exactly_one"] + 2 * strata["multi_attachment"]
    assert all("raw_data" not in record for record in records)
    return evidence


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--verify", action="store_true")
    parser.add_argument("--reference-inventory")
    parser.add_argument("--oracle")
    parser.add_argument("--source-summary")
    parser.add_argument("--error-inventory")
    parser.add_argument("--evidence-output")
    parser.add_argument("--limit", type=int, default=3)
    parser.add_argument("--combined", action="store_true")
    options = parser.parse_args()
    if options.reference_inventory:
        evidence = reference_evidence(options.reference_inventory, options.oracle, options.limit, options.combined, options.source_summary, options.error_inventory)
        if options.evidence_output:
            Path(options.evidence_output).write_text(json.dumps(evidence, ensure_ascii=False, indent=2) + "\n")
        print(f"Reference corpus gate: {evidence['gate']} records={evidence['source_inventory']['records']} strata={evidence['source_inventory']['strata']}")
        return
    corpus = json.loads(OUTPUT.read_text()) if options.verify else build()
    if not options.verify:
        OUTPUT.write_text(json.dumps(corpus, ensure_ascii=False, indent=2) + "\n")
    verify(corpus)


if __name__ == "__main__":
    main()
