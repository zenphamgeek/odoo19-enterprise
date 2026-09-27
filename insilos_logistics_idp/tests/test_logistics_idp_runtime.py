import ast
import base64
import hashlib
import io
import json
import zipfile
from email import policy
from email.parser import BytesParser
from datetime import timedelta
from pathlib import Path
from unittest.mock import MagicMock, patch

from odoo import fields
from ..models.logistics_idp import _INTERNAL_CASE_LIFECYCLE_TOKEN
from odoo.exceptions import AccessError, UserError, ValidationError
from odoo.tests.common import TransactionCase, new_test_user as _new_test_user, tagged
from openpyxl import load_workbook

from .common import assert_unique_constraint, new_logistics_test_user, unique_fixture
from .development_uat_harness import _assert_case_actual, _target_passes, run_orm_corpus
from ..services.document_processor import MAX_BYTES
from ..services.reconciliation import reconcile_documents


def new_test_user(env, **values):
    return new_logistics_test_user(_new_test_user, env, **values)


@tagged('post_install', '-at_install', 'logistics_idp')
class TestLogisticsIdpRuntime(TransactionCase):
    def test_stale_processing_job_is_recovered_for_retry(self):
        job = self.env['logistics.idp.inbound.job'].create({
            'profile_code': 'fixture', 'source_system': 'fixture',
            'source_key': unique_fixture('stale'), 'source_version': '1',
            'payload': json.dumps({'name': 'fixture', 'provenance': 'fixture',
                                   'effective_date': '2026-01-01'}),
            'state': 'processing',
            'last_attempt_at': fields.Datetime.subtract(fields.Datetime.now(), hours=1),
        })
        self.assertEqual(self.env['logistics.idp.inbound.job']._recover_stale_processing(), 1)
        self.assertEqual((job.state, job.last_error), ('retry', 'INBOUND_FAILURE:StaleProcessing'))
        self.assertTrue(job.message_ids.filtered(
            lambda message: 'automatically recovered from stale processing' in message.body))

    def test_uat_evidence_fails_wrong_actual_and_unmet_target(self):
        case = {"oracle": {"classification": "purchase_order", "normalized_fields": {"supplier": "A", "lines": []}}}
        with self.assertRaisesRegex(AssertionError, "classification"):
            _assert_case_actual(case, {"classification": "invoice", "normalized_fields": {"supplier": "A", "lines": []}})
        self.assertFalse(_target_passes(">= 99%", 0.98))

    def test_srs_section_7_5_exception_audit_center_projects_persisted_categories_read_only(self):
        case = self.env['logistics.idp.case'].create({
            'name': 'AUDIT-' + unique_fixture('7-5'), 'source_system': 'fixture',
            'source_key': unique_fixture('audit'), 'source_version': 'v1', 'provenance': 'fixture',
            'effective_date': '2026-01-01', 'sla_deadline': fields.Datetime.subtract(fields.Datetime.now(), hours=1),
        })
        document = self.env['logistics.idp.document'].intake_content(case, unique_fixture('audit-doc').encode(), 'application/pdf', process=False)
        document.write({'status': 'review'})
        run_values = {'case_id': case.id, 'document_id': document.id, 'provider': 'fixture',
                      'model_version': 'v1', 'schema_version': 'v1', 'prompt_version': 'v1',
                      'started_at': fields.Datetime.now(), 'completed_at': fields.Datetime.now(),
                      'status': 'error', 'payload': {'fixture': 'audit'}}
        self.env['logistics.idp.extraction.run']._controlled_create(run_values, 'audit_center')
        reprocessed = self.env['logistics.idp.extraction.run']._controlled_create(run_values, 'audit_center')
        check = self.env['logistics.idp.check.result']._controlled_create({
            'case_id': case.id, 'run_id': reprocessed.id, 'code': 'AUDIT', 'verdict': 'block',
            'rationale': 'fixture', 'payload': {'fixture': 'audit'},
        }, 'audit_center')
        exception = self.env['logistics.idp.exception'].create({
            'case_id': case.id, 'exception_type': 'extraction_failure', 'state': 'waiting'})
        self.env['logistics.idp.override']._controlled_create({
            'case_id': case.id, 'exception_ids': [(6, 0, exception.ids)], 'reason_code': 'evidence_gap',
            'justification': 'fixture', 'requested_by': self.env.user.id, 'requested_at': fields.Datetime.now(),
            'state': 'approved', 'approved_by': self.env.user.id, 'approved_at': fields.Datetime.now(),
            'payload': {'fixture': 'audit'},
        }, 'audit_center')
        model = self.env['logistics.idp.case']
        for category in ('exceptions', 'checks', 'extraction', 'correlations', 'sla', 'overrides', 'reprocessing'):
            self.assertIn(case, model.search(model.exception_audit_domain(category)))
        action = model.action_open_exception_audit_center()
        self.assertEqual((action['res_model'], action['context']), ('logistics.idp.case', {'create': False, 'edit': False, 'delete': False}))
        self.assertIn(case, model.search(action['domain']))
        self.assertEqual(check.verdict, 'block')

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.po = {'supplier': 'SUP-1', 'regime': 'E13', 'lines': [
            {'material_code': 'MAT-1', 'unit_price': '1000', 'price_per': '1000', 'remaining_quantity': '10', 'custom_code': 'PO-1-10'},
        ]}
        cls.invoice = {'supplier': 'SUP-1', 'regime': 'E13', 'lines': [
            {'material_code': 'mat-1', 'unit_price': '1', 'quantity': '4', 'custom_code': 'PO-1-10'},
        ]}
        cls.mail_count = cls.env['mail.mail'].sudo().search_count([])
        cls.operator = new_test_user(
            cls.env, login='logistics-runtime-operator', context={'no_reset_password': True},
            groups='insilos_logistics_idp.group_logistics_operator',
        )

    def test_vendor_bill_creation_fails_closed_without_accounting_control(self):
        reviewer = new_test_user(
            self.env, login='external-reviewer-' + unique_fixture('reviewer'),
            context={'no_reset_password': True}, groups='insilos_logistics_idp.group_logistics_reviewer')
        document = self.env['logistics.idp.document']._intake_synthetic_content(
            self.case, b'SYNTHETIC VENDOR INVOICE', 'text/plain', process=False)
        document.write({'status': 'valid', 'extracted_doc_number': unique_fixture('INV')})
        before = self.env['account.move'].search_count([('move_type', '=', 'in_invoice')])
        for user in (self.operator, reviewer):
            with self.subTest(user=user.login), self.assertRaisesRegex(UserError, 'immutable independent approval'):
                document.with_user(user).with_context(
                    approved_logistics_action=True,
                    logistics_correlation_id=unique_fixture('vendor'),
                ).action_create_vendor_bill()
        self.assertEqual(self.env['account.move'].search_count([('move_type', '=', 'in_invoice')]), before)

    def test_external_nsw_dossier_fails_closed_without_governed_submission_control(self):
        reviewer = new_test_user(
            self.env, login='nsw-reviewer-' + unique_fixture('reviewer'), context={'no_reset_password': True},
            groups='insilos_logistics_idp.group_logistics_reviewer,insilos_chemical_trade_compliance.group_chemical_compliance_user')
        document = self.env['logistics.idp.document']._intake_synthetic_content(
            self.case, b'SYNTHETIC NSW DOSSIER', 'text/plain', process=False)
        document.write({'status': 'valid', 'extracted_doc_number': unique_fixture('NSW')})
        before = self.env['is.chemical.compliance.dossier'].search_count([])
        for user in (self.operator, reviewer):
            with self.subTest(user=user.login), self.assertRaisesRegex(UserError, 'immutable independent approval'):
                document.with_user(user).with_context(
                    approved_logistics_action=True,
                    logistics_correlation_id=unique_fixture('nsw'),
                ).action_create_customs_dossier()
        self.assertEqual(self.env['is.chemical.compliance.dossier'].search_count([]), before)

    def test_external_action_partial_failure_rolls_back_without_iap_or_mes_write(self):
        reviewer = new_test_user(
            self.env, login='failure-reviewer-' + unique_fixture('reviewer'),
            context={'no_reset_password': True},
            groups='insilos_logistics_idp.group_logistics_reviewer,account.group_account_invoice')
        document = self.env['logistics.idp.document']._intake_synthetic_content(
            self.case, b'SYNTHETIC PARTIAL FAILURE', 'text/plain', process=False)
        document.write({'status': 'valid', 'extracted_doc_number': unique_fixture('FAIL')})
        before = {
            'moves': self.env['account.move'].search_count([('move_type', '=', 'in_invoice')]),
            'charges': self.env['iap.charge'].search_count([]),
            'ledger': self.env['iap.ledger'].search_count([]),
            'mes': self.env['logistics.idp.mes.reference'].search_count([]),
        }
        with patch.object(type(document), '_audit_external_action', side_effect=ValidationError('audit failed')), \
                self.assertRaisesRegex(ValidationError, 'audit failed'), self.env.cr.savepoint():
            document.with_user(reviewer).with_context(
                approved_logistics_action=True,
                logistics_correlation_id=unique_fixture('failure'),
            ).action_create_vendor_bill()
        self.assertEqual({
            'moves': self.env['account.move'].search_count([('move_type', '=', 'in_invoice')]),
            'charges': self.env['iap.charge'].search_count([]),
            'ledger': self.env['iap.ledger'].search_count([]),
            'mes': self.env['logistics.idp.mes.reference'].search_count([]),
        }, before)

    def setUp(self):
        super().setUp()
        suffix = unique_fixture(self._testMethodName)
        self.case = self.env['logistics.idp.case'].create({
            'name': 'SYNTHETIC-CASE-1-' + suffix, 'source_system': 'synthetic-erp',
            'source_key': 'CASE-1-' + suffix, 'source_version': 'v1',
            'provenance': 'fixture:development', 'effective_date': '2026-01-01',
            'po_reference': 'PO-1', 'supplier_reference': 'SUP-1', 'email_subject': 'Shipment PO-1',
        })
        stream = io.BytesIO()
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.title = 'SAP'
        workbook.active['A1'] = 'template'
        workbook.save(stream)
        self.sap_template = self.env['ir.attachment'].create({
            'name': 'sap-template.xlsx', 'raw': stream.getvalue(),
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            'company_id': self.case.company_id.id,
        })
        self.sap_profile = self.env['logistics.idp.supplier.profile'].import_upsert({
            'supplier_reference': 'SUP-1', 'profile_code': 'DEFAULT', 'version': '1',
            'source_system': 'synthetic-erp', 'source_key': 'SUP-1-' + suffix, 'source_version': 'v1',
            'provenance': 'fixture:development', 'effective_from': '2026-01-01', 'payload': {
                'draft_invoice': {'mode': 'skipped', 'final_number_allowed': False},
                'authoritative_reference_policy': {'policy_version': 'fr609-v1', 'rules': [
                    {'id': 'sap-downstream', 'stage': str(self.case.task_id.stage_id.id), 'field': 'sap_erp_output', 'source_type': 'output', 'output_types': ['sap_erp_output'], 'downstream_field': 'import_declaration'},
                    {'id': 'sap-wildcard', 'stage': str(self.case.task_id.stage_id.id), 'field': '*', 'source_type': 'output', 'output_types': ['sap_erp_output']},
                    {'id': 'e13-downstream', 'stage': str(self.case.task_id.stage_id.id), 'field': 'e13', 'source_type': 'output', 'output_types': ['e13']},
                    {'id': 'e15-downstream', 'stage': str(self.case.task_id.stage_id.id), 'field': 'e15', 'source_type': 'output', 'output_types': ['e15']},
                    {'id': 'declaration-reference', 'stage': str(self.case.task_id.stage_id.id), 'field': 'import_declaration', 'source_type': 'output', 'output_types': ['e13', 'e15']},
                ]},
                'broker_package': {
                    'qdtq_document_types': ['customs_declaration'],
                    'exclude_document_types_without_qdtq': True,
                    'required_categories': {'customs': {'output_types': ['e13']}},
                },
                'sap_erp_output': {'stages': {str(self.case.task_id.stage_id.id): {
                    'mapping_version': 'sap-v1', 'source_precedence': ['po_snapshot', 'invoice_evidence', 'dsnavl', 'master_data'],
                    'template_attachment_id': self.sap_template.id, 'sheet': 'SAP', 'start_row': 2, 'sort_keys': ['line_key'],
                    'line_identity_keys': ['line_key'], 'price_source': 'line.unit_price',
                    'amount_source': 'line.amount', 'custom_code_template': 'SAP-{material_code}',
                    'amount_column': 'amount',
                    'columns': [
                        {'name': 'material', 'source': 'po_line.material_code', 'position': 3},
                        {'name': 'custom', 'source': 'line.custom_code', 'position': 1},
                        {'name': 'amount', 'source': 'line.amount', 'position': 2},
                    ],
                }}},
            },
        })

    def _write_sap_config(self, config):
        payload = json.dumps(config, sort_keys=True, separators=(',', ':'))
        self.cr.execute('UPDATE logistics_idp_supplier_profile SET payload = %s, payload_hash = %s WHERE id = %s',
                        [payload, hashlib.sha256(payload.encode()).hexdigest(), self.sap_profile.id])
        self.sap_profile.invalidate_recordset(['payload', 'payload_hash'])

    def _retag_authoritative_policy_rules(self):
        config = json.loads(self.sap_profile.payload)
        for rule in config['authoritative_reference_policy']['rules']:
            rule['stage'] = str(self.case.task_id.stage_id.id)
        self._write_sap_config(config)

    def _sap_sources(self, po_lines, invoice_lines):
        suffix = unique_fixture(self._testMethodName)
        po = self.env['logistics.idp.po.snapshot']._controlled_create({
            'company_id': self.case.company_id.id, 'po_reference': self.case.po_reference,
            'snapshot_date': self.case.effective_date, 'source_system': 'synthetic-erp',
            'source_key': 'PO-' + suffix, 'source_version': 'v1', 'provenance': 'fixture:development',
            'payload': {'lines': po_lines},
        }, 'runtime_test')
        invoice = self.env['logistics.idp.evidence']._controlled_create({
            'case_id': self.case.id, 'category': 'invoice', 'source_reference': 'INV-' + suffix,
            'status': 'valid', 'payload': {
                'invoice_number': 'INV-LOCAL', 'bill_number': 'BILL-LOCAL',
                'goods_description': 'Invoice-described goods', 'package_count': '3',
                'gross_weight': '12.5', 'net_weight': '10', 'carrier': 'Trusted Carrier',
                'note': 'Invoice note', 'lines': invoice_lines,
            },
        }, 'runtime_test')
        return po, invoice

    def test_fr406_final_export_hs_reconciliation_is_immutable_idempotent_and_scoped(self):
        def document(document_type, hs_code):
            item = self.env['logistics.idp.document'].intake_content(
                self.case, unique_fixture('fr406-' + document_type + hs_code).encode(), 'application/json', process=False)
            item.sudo().write({'document_type': document_type})
            run = self.env['logistics.idp.extraction.run']._controlled_create({
                'case_id': self.case.id, 'document_id': item.id, 'provider': 'fixture', 'model_version': 'v1',
                'schema_version': 'v1', 'prompt_version': 'v1', 'started_at': fields.Datetime.now(),
                'completed_at': fields.Datetime.now(), 'status': 'valid',
                'payload': {'payload': {'lines': [{'material_code': 'MAT-406', 'hs_code': hs_code}]}},
            }, 'fr406_test')
            item.with_context(_logistics_document_type_token=True).write({'current_run_id': run.id})
            return item

        master = document('dsnavl', '100100')
        document('export_declaration_final', '100100')
        po, invoice = {'supplier': 'SUP-1', 'lines': []}, {'supplier': 'SUP-1', 'lines': []}
        first = self.case.reconcile_documents(po, invoice)
        self.assertEqual(self.case.check_result_ids.filtered(lambda item: item.code == 'FINAL_EXPORT_HS').verdict, 'pass')
        self.assertEqual(json.loads(master.current_run_id.payload)['payload']['lines'][0]['hs_code'], '100100')
        self.assertEqual(first, self.case.reconcile_documents(po, invoice))
        document('export_declaration_final', '999999')
        self.case.reconcile_documents(po, invoice)
        checks = self.case.check_result_ids.filtered(lambda item: item.code == 'FINAL_EXPORT_HS')
        self.assertEqual(checks.sorted('id').mapped('verdict'), ['pass', 'block'])
        self.assertEqual(self.env['logistics.idp.exception'].search_count([
            ('case_id', '=', self.case.id), ('company_id', '=', self.case.company_id.id), ('exception_type', '=', 'hs_code_mismatch')]), 1)
        self.assertEqual(json.loads(master.current_run_id.payload)['payload']['lines'][0]['hs_code'], '100100')
        self.case.document_ids.filtered(lambda item: item.document_type == 'dsnavl').sudo().write({'document_type': 'unknown'})
        document('export_declaration_final', '888888')
        self.case.reconcile_documents(po, invoice)
        self.assertEqual(self.case.check_result_ids.filtered(lambda item: item.code == 'FINAL_EXPORT_HS').sorted('id')[-1].verdict, 'review')

    def test_fr404_runtime_declaration_groups_bind_manifest_and_replay(self):
        config = json.loads(self.sap_profile.payload)
        config['declaration_group_rules'] = [{
            'rule_id': 'e11-e15', 'regimes': ['E11', 'E15'], 'declaration_count': 2,
            'effective_from': '2026-01-01', 'effective_to': '2026-12-31',
        }]
        self._write_sap_config(config)
        po = {'supplier': 'SUP-1', 'lines': [{'material_code': 'M', 'remaining_quantity': 2, 'unit_price': 1}]}
        invoice = {'supplier': 'SUP-1', 'lines': [
            {'material_code': 'M', 'regime': 'E11', 'quantity': 1, 'unit_price': 1},
            {'material_code': 'M', 'regime': 'E15', 'quantity': 1, 'unit_price': 1},
        ]}
        correct = self.case.reconcile_documents(po, invoice, reconciliation_inputs={'declaration_groups': ['TK-2', 'TK-1']})
        self.assertEqual((correct.status, json.loads(correct.payload)['declaration_group_check']['actual_count']), ('valid', 2))
        decision = self.env['logistics.idp.policy.decision'].search([('case_id', '=', self.case.id)], order='id desc', limit=1)
        self.assertEqual(json.loads(decision.payload)['input_manifest']['references']['declaration_groups'], ['TK-1', 'TK-2'])
        before = self.env['logistics.idp.policy.decision'].search_count([('case_id', '=', self.case.id)])
        self.assertEqual(self.case.reconcile_documents(po, invoice, reconciliation_inputs={'declaration_groups': ['TK-1', 'TK-2']}), correct)
        self.assertEqual(self.env['logistics.idp.policy.decision'].search_count([('case_id', '=', self.case.id)]), before)
        self.assertEqual(self.case.reconcile_documents(po, invoice, reconciliation_inputs={'declaration_groups': ['TK-1']}).status, 'invalid')
        self.assertEqual(self.case.reconcile_documents(po, invoice).status, 'review')
        config['declaration_group_rules'][0]['effective_to'] = '2025-12-31'
        self._write_sap_config(config)
        self.assertEqual(self.case.reconcile_documents(po, invoice, reconciliation_inputs={'declaration_groups': ['TK-1', 'TK-2']}).status, 'review')
        config['declaration_group_rules'][0]['effective_to'] = '2026-12-31'
        config['declaration_group_rules'].append({**config['declaration_group_rules'][0], 'rule_id': 'duplicate'})
        self._write_sap_config(config)
        self.assertEqual(self.case.reconcile_documents(po, invoice, reconciliation_inputs={'declaration_groups': ['TK-1', 'TK-2']}).status, 'review')

    def test_fr604a_supplemental_final_customs_upload_reuses_case_and_preserves_history(self):
        old_check = self.env['logistics.idp.check.result']._controlled_create({
            'case_id': self.case.id, 'code': 'FINAL_CUSTOMS_DOCUMENT', 'required': True,
            'verdict': 'review', 'rationale': 'Final customs document missing.', 'payload': {'missing': True},
        }, 'runtime_test')
        old_hash, old_payload = old_check.payload_hash, old_check.payload
        old_evidence = self.env['logistics.idp.evidence']._controlled_create({
            'case_id': self.case.id, 'category': 'extraction', 'source_reference': 'missing-final',
            'status': 'review', 'payload': {'missing': 'export_declaration_final'},
        }, 'runtime_test')
        old_evidence_hash = old_evidence.payload_hash
        payload = {'document_type': 'export_declaration_final', 'confidence': 1, 'payload': {
            'declaration_number': 'QDTQ-604A', 'regime': 'E13', 'lines': [{'item': '1'}]}, 'source_spans': []}
        content = json.dumps(payload).encode()
        documents_before = self.env['logistics.idp.document'].search_count([('case_id', '=', self.case.id)])
        cases_before = self.env['logistics.idp.case'].search_count([])
        document = self.env['logistics.idp.document'].with_user(self.operator).upload_final_customs_document(
            self.case, content, 'application/json', 'Tokhaihq7X_QDTQ.json')
        self.assertEqual((document.case_id, document.document_type, document.content_hash),
                         (self.case, 'export_declaration_final', hashlib.sha256(content).hexdigest()))
        self.assertEqual((self.env['logistics.idp.case'].search_count([]),
                          self.env['logistics.idp.document'].search_count([('case_id', '=', self.case.id)])),
                         (cases_before, documents_before + 1))
        self.assertEqual((document.source_channel, document.create_uid, document.version), ('upload', self.operator, 1))
        self.assertTrue(document.create_date and document.attachment_id and document.document_id)
        self.assertEqual((old_check.verdict, old_check.payload_hash, old_check.payload), ('review', old_hash, old_payload))
        self.assertEqual((old_evidence.status, old_evidence.payload_hash), ('review', old_evidence_hash))
        self.assertEqual(document.extraction_run_ids.document_id, document)
        self.assertEqual(self.env['logistics.idp.check.result'].search([('run_id', '=', document.current_run_id.id)]).mapped('run_id'), document.current_run_id)
        outsider = new_test_user(self.env, login='fr604a-outsider-' + unique_fixture('fr604a'), context={'no_reset_password': True})
        with self.assertRaises(UserError):
            self.env['logistics.idp.document'].with_user(outsider).upload_final_customs_document(
                self.case, content + b'x', 'application/json', 'denied.json')
        foreign_company = self.env['res.company'].create({'name': 'FR604A Foreign'})
        foreign_case = self.env['logistics.idp.case'].with_company(foreign_company).create({
            'name': 'FR604A Foreign', 'source_system': 'synthetic', 'source_key': unique_fixture('fr604a-foreign'),
            'source_version': '1', 'provenance': 'fixture', 'effective_date': '2026-01-01', 'company_id': foreign_company.id})
        with self.assertRaises(ValidationError):
            self.env['logistics.idp.document'].with_user(self.operator).upload_final_customs_document(
                foreign_case, content + b' ', 'application/json', 'foreign.json')
        self.case.write({'state': 'completed'})
        with self.assertRaisesRegex(ValidationError, 'terminal'):
            self.env['logistics.idp.document'].with_user(self.operator).upload_final_customs_document(
                self.case, content + b'  ', 'application/json', 'terminal.json')

    def test_sap_erp_output_uses_authoritative_sources_and_persists_provenance(self):
        po, invoice = self._sap_sources(
            [{'line_key': '2', 'material_code': 'MAT-2'}, {'line_key': '1', 'material_code': 'MAT-1'}],
            [{'line_key': '2', 'material_code': 'mat-2', 'quantity': '2', 'unit_price': '3'}, {'line_key': '1', 'material_code': 'mat-1', 'quantity': '1', 'unit_price': '4'}])
        output = self.case.generate_sap_erp_output()
        payload = json.loads(output.payload)
        self.assertEqual(payload['source_records'], [{'id': po.id, 'hash': po.payload_hash}, {'id': invoice.id, 'hash': invoice.payload_hash}])
        self.assertEqual(payload['rows'][0]['sources']['material'], {
            'mapping_version': 'sap-v1', 'source': 'po_line.material_code',
            'record_type': 'logistics.idp.po.snapshot', 'record_id': po.id, 'record_hash': po.payload_hash,
        })
        self.assertEqual([row['values']['material'] for row in payload['rows']], ['MAT-1', 'MAT-2'])
        self.assertEqual(payload['invoice_total'], '10')
        binding = payload['authoritative_policy_binding']
        self.assertEqual((binding['stage'], binding['field'], binding['rule_id'], binding['policy_version'], binding['output_types']),
                         (str(self.case.task_id.stage_id.id), 'sap_erp_output', 'sap-downstream', 'fr609-v1', ['sap_erp_output']))
        self.assertEqual(binding['downstream_authoritative_binding']['rule_id'], 'declaration-reference')
        with self.assertRaisesRegex(ValidationError, 'no caller-provided'):
            self.case.generate_sap_erp_output({'lines': []}, {'lines': []})

    def test_sap_erp_blank_material_custom_code_and_price_normalization(self):
        self._sap_sources(
            [{'line_key': 'L1', 'material_code': 'MAT-1'}],
            [{'line_key': 'L1', 'quantity': '4', 'unit_price': '15', 'price_per': '10'}])
        output = self.case.generate_sap_erp_output()
        values = json.loads(output.payload)['rows'][0]['values']
        self.assertEqual((values['material'], values['custom'], values['amount']), ('MAT-1', 'SAP-MAT-1', '6'))

    def test_sap_erp_blank_invoice_custom_code_uses_po_code(self):
        self._sap_sources(
            [{'line_key': 'L1', 'material_code': 'MAT-1', 'custom_code': 'PO-CODE'}],
            [{'line_key': 'L1', 'quantity': '1', 'unit_price': '2', 'custom_code': ''}])
        values = json.loads(self.case.generate_sap_erp_output().payload)['rows'][0]['values']
        self.assertEqual(values['custom'], 'PO-CODE')

    def test_sap_erp_blocks_invalid_price_per_and_foreign_sources(self):
        self._sap_sources([{'line_key': 'L1', 'material_code': 'MAT-1'}], [{'line_key': 'L1', 'quantity': '1', 'unit_price': '2', 'price_per': '0'}])
        with self.assertRaisesRegex(ValidationError, 'price_per'):
            self.case.generate_sap_erp_output()
        self.assertEqual(self.case._sap_erp_amount({'quantity': '4', 'value': '15', 'price_per': '10'}, 'line.value'), 6)
        foreign_case = self.env['logistics.idp.case'].create({
            'name': 'Foreign-' + unique_fixture('sap'), 'source_system': 'synthetic-erp',
            'source_key': unique_fixture('foreign'), 'source_version': 'v1', 'provenance': 'fixture:development',
            'effective_date': '2026-01-01',
        })
        self.env['logistics.idp.evidence']._controlled_create({
            'case_id': foreign_case.id, 'category': 'invoice', 'source_reference': 'foreign', 'status': 'valid',
            'payload': {'lines': [{'line_key': 'L1', 'quantity': '1', 'unit_price': '9'}]},
        }, 'runtime_test')
        with self.assertRaisesRegex(ValidationError, 'price_per'):
            self.case.generate_sap_erp_output()

    def test_derived_stage_requires_terminal_checks_and_no_processing_before_completion(self):
        manager = new_test_user(
            self.env, login='logistics-runtime-manager-' + unique_fixture('lifecycle'),
            context={'no_reset_password': True}, groups='insilos_logistics_idp.group_logistics_manager')
        reviewer = new_test_user(
            self.env, login='logistics-runtime-reviewer-' + unique_fixture('lifecycle'),
            context={'no_reset_password': True}, groups='insilos_logistics_idp.group_logistics_reviewer')

        def case_with_import_check(verdict):
            case = self.env['logistics.idp.case'].create({
                'name': 'SYNTHETIC-LIFECYCLE-' + verdict, 'source_system': 'synthetic-erp',
                'source_key': unique_fixture('lifecycle-' + verdict), 'source_version': 'v1',
                'provenance': 'fixture:development', 'effective_date': '2026-01-01',
                'document_status': 'pass', 'reconciliation_status': 'pass',
                'compliance_status': 'pass', 'output_status': 'not_applicable',
            })
            check = self.env['logistics.idp.check.result']._controlled_create({
                'case_id': case.id, 'code': 'IMPORT_DECLARATION', 'required': True, 'verdict': verdict,
                'rationale': 'synthetic lifecycle fixture', 'payload': {},
            }, 'runtime_lifecycle_test')
            return case, check

        for verdict, state in (('review', 'review'), ('block', 'blocked')):
            case, _check = case_with_import_check(verdict)
            case._derive_lifecycle()
            self.assertEqual(case.state, state)
            with self.assertRaises(ValidationError):
                case.with_user(manager).action_complete('synthetic completion')

        case, check = case_with_import_check('pass')
        processing = self.env['logistics.idp.document'].create({
            'case_id': case.id, 'content_hash': unique_fixture('lifecycle-processing'),
            'mimetype': 'application/pdf', 'status': 'processing',
        })
        case._derive_lifecycle()
        self.assertEqual(case.state, 'processing')
        processing.write({'status': 'valid'})
        case._derive_lifecycle()
        self.assertEqual((case.state, case.verdict), ('ready', 'pass'))
        self.assertTrue(case.with_user(manager).action_complete('synthetic completion'))
        self.assertEqual(case.state, 'completed')
        self.assertFalse(case.document_ids.filtered(lambda document: document.status == 'processing'))

        case, check = case_with_import_check('review')
        exception = self.env['logistics.idp.exception'].create({
            'case_id': case.id, 'exception_type': 'missing_document', 'severity': 'medium'})
        override = case.with_user(reviewer).request_override(
            exception, 'evidence_gap', 'reviewer accepted declaration evidence', checks=check)
        self.assertEqual(override.state, 'requested')
        self.assertEqual(case.state, 'review')
        override.with_user(manager).action_approve()
        self.assertEqual(check.verdict, 'review')
        self.assertEqual(case.state, 'ready')
        self.assertTrue(case.with_user(manager).action_complete('controlled override completion'))
        self.assertEqual(case.state, 'completed')

    def test_case_fails_to_review_and_source_identity_is_unique(self):
        self.assertEqual(self.case.verdict, 'review')
        assert_unique_constraint(self, lambda: self.env['logistics.idp.case'].create({
            'name': 'Duplicate', 'company_id': self.case.company_id.id,
            'source_system': self.case.source_system, 'source_key': self.case.source_key,
            'source_version': self.case.source_version, 'provenance': 'fixture:development', 'effective_date': '2026-01-01',
        }))

    def test_intake_exact_idempotency_and_deterministic_correlation(self):
        values = {
            'name': 'Repeat', 'source_system': self.case.source_system, 'source_key': self.case.source_key,
            'source_version': self.case.source_version, 'provenance': 'fixture:development', 'effective_date': '2026-01-01',
        }
        self.assertEqual(self.env['logistics.idp.case'].intake(values), self.case)
        correlated = self.env['logistics.idp.case'].intake({**values, 'source_key': 'ATT-2', 'po_reference': 'PO-1'})
        self.assertEqual(correlated, self.case)

    def test_ai007_local_logical_documents_are_stable_and_malformed_safe(self):
        content = json.dumps({'logical_documents': [
            {'document_type': 'draft_vat_invoice', 'confidence': 1,
             'payload': {'supplier': 'SUP', 'invoice_number': 'VAT-1', 'lines': [{}]}},
            {'document_type': 'packing_list', 'confidence': 1,
             'payload': {'supplier': 'SUP', 'lines': [{}]}},
        ]}, sort_keys=True).encode()
        document = self.env['logistics.idp.document'].with_user(self.operator)._intake_synthetic_content(
            self.case, content, 'application/json')
        children = self.case.document_ids.filtered(lambda item: item.logical_parent_id == document).sorted('logical_index')
        self.assertEqual([(item.logical_index, item.document_type) for item in children],
                         [(1, 'draft_vat_invoice'), (2, 'packing_list')])
        self.assertTrue(all(item.content_hash == document.content_hash and item.attachment_id == document.attachment_id
                            and item.current_run_id for item in children))
        identities = [(item.logical_index, item.logical_source_identity, item.current_run_id.id) for item in children]
        document.with_user(self.operator)._reprocess_synthetic()
        replay = self.case.document_ids.filtered(lambda item: item.logical_parent_id == document).sorted('logical_index')
        self.assertEqual([(item.logical_index, item.logical_source_identity) for item in replay],
                         [(index, identity) for index, identity, run_id in identities])
        self.assertEqual([len(item.extraction_run_ids) for item in replay], [2, 2])
        malformed = self.env['logistics.idp.document'].with_user(self.operator)._intake_synthetic_content(
            self.case, json.dumps({'logical_documents': [{'document_type': 'packing_list', 'confidence': 1,
            'payload': {'supplier': 'SUP'}}]}).encode(), 'application/json')
        self.assertEqual(malformed.status, 'error')
        self.assertFalse(self.case.document_ids.filtered(lambda item: item.logical_parent_id == malformed))

    def test_fr504_same_type_logical_invoices_preserve_each_payload_and_replay_runs(self):
        content = json.dumps({'logical_documents': [
            {'document_type': 'main_vat_invoice', 'confidence': 1,
             'payload': {'supplier': 'SUP', 'invoice_number': 'VAT-504-A', 'lines': [{}]}},
            {'document_type': 'main_vat_invoice', 'confidence': 1,
             'payload': {'supplier': 'SUP', 'invoice_number': 'VAT-504-B', 'lines': [{}]}},
        ]}, sort_keys=True).encode()
        with patch('odoo.addons.insilos_logistics_idp.services.document_processor.IAPDocumentProcessor.process') as iap_process:
            document = self.env['logistics.idp.document'].with_user(self.operator)._intake_synthetic_content(
                self.case, content, 'application/json')
            children = self.case.document_ids.filtered(lambda item: item.logical_parent_id == document).sorted('logical_index')
            self.assertEqual([(item.logical_index, item.document_type) for item in children],
                             [(1, 'main_vat_invoice'), (2, 'main_vat_invoice')])
            self.assertEqual([json.loads(item.normalized_fields)['invoice_number'] for item in children],
                             ['VAT-504-A', 'VAT-504-B'])
            self.assertEqual([len(item.extraction_run_ids) for item in children], [1, 1])
            document.with_user(self.operator)._reprocess_synthetic()
        replay = self.case.document_ids.filtered(lambda item: item.logical_parent_id == document).sorted('logical_index')
        self.assertEqual(replay, children)
        self.assertEqual([json.loads(item.normalized_fields)['invoice_number'] for item in replay],
                         ['VAT-504-A', 'VAT-504-B'])
        self.assertEqual([len(item.extraction_run_ids) for item in replay], [2, 2])
        iap_process.assert_not_called()

    def test_extraction_low_confidence_routes_review(self):
        extraction = self.case.add_extraction('SYNTHETIC-ATT', {'invoice_number': 'INV-1'}, confidence=.79, threshold=.8)
        self.assertEqual(extraction.status, 'review')
        self.assertEqual(json.loads(extraction.payload)['data']['invoice_number'], 'INV-1')

    def test_nfr003_critical_inputs_are_run_bound_and_never_silent_pass(self):
        policy_source = self.env['logistics.idp.policy.source'].select_effective_pack(
            self.case.company_id, 'trade_compliance', 'vn_fdi_fta', 'VN', 'ALL', self.case.effective_date)
        policy = self.env['logistics.idp.policy.source'].validate_policy_payload(
            policy_source.payload, 'trade_compliance', 'vn_fdi_fta', 'VN', 'ALL')
        extraction = policy['horizontal']['extraction']
        fields_config = extraction['critical_fields']['main_vat_invoice']
        header = next(path for path in fields_config if not path.startswith('lines[].'))
        line = next(path for path in fields_config if path.startswith('lines[].'))

        def run_for(payload, confidence=1, classified=1, document_type='main_vat_invoice'):
            document = self.env['logistics.idp.document'].intake_content(
                self.case, unique_fixture('nfr003').encode(), 'application/json', process=False)
            run = self.env['logistics.idp.extraction.run']._controlled_create({
                'case_id': self.case.id, 'document_id': document.id, 'provider': 'fixture',
                'model_version': 'v1', 'schema_version': 'v1', 'prompt_version': 'v1',
                'started_at': fields.Datetime.now(), 'completed_at': fields.Datetime.now(), 'status': 'review',
                'payload': {'document_type': document_type, 'payload': payload, 'confidence': confidence,
                            'classification_confidence': classified, 'validation_warnings': [], 'effective_policy': policy},
            }, 'nfr003_test')
            document._create_extraction_critical_check(run)
            return run, self.case.check_result_ids.filtered(lambda check: check.run_id == run and check.code == 'EXTRACTION_CRITICAL_INPUTS')

        complete = {'lines': [{}]}
        for path in fields_config:
            if path.startswith('lines[].'):
                complete['lines'][0][path.split('.', 1)[1]] = 'x'
            elif path != 'lines':
                complete[path] = 'x'
        _, header_check = run_for({key: value for key, value in complete.items() if key != header})
        _, line_check = run_for({**complete, 'lines': [{line.split('.', 1)[1]: None}]})
        self.assertEqual((header_check.verdict, line_check.verdict), ('review', 'review'))
        valid_run, valid_check = run_for(complete)
        self.assertEqual(valid_check.verdict, 'pass')
        self.case._derive_lifecycle()
        self.assertNotEqual(self.case.state, 'ready')
        _, ambiguous_check = run_for(complete, classified=0)
        self.assertEqual(ambiguous_check.verdict, 'review')
        document = valid_run.document_id
        document._create_extraction_critical_check(valid_run)
        self.assertEqual(len(self.case.check_result_ids.filtered(lambda check: check.run_id == valid_run and check.code == 'EXTRACTION_CRITICAL_INPUTS')), 1)
        document.current_run_id = ambiguous_check.run_id
        self.case.document_status = 'pass'
        self.case._derive_lifecycle()
        self.assertEqual(self.case.state, 'review')
        reprocessed = self.env['logistics.idp.extraction.run']._controlled_create({
            'case_id': self.case.id, 'document_id': document.id, 'provider': 'fixture',
            'model_version': 'v1', 'schema_version': 'v1', 'prompt_version': 'v1',
            'started_at': fields.Datetime.now(), 'completed_at': fields.Datetime.now(), 'status': 'review',
            'payload': {'document_type': 'main_vat_invoice', 'payload': complete, 'confidence': 1,
                        'classification_confidence': 1, 'validation_warnings': [], 'effective_policy': policy},
        }, 'nfr003_reprocess_test')
        document._create_extraction_critical_check(reprocessed)
        document.current_run_id = reprocessed
        reprocessed_check = self.case.check_result_ids.filtered(
            lambda check: check.run_id == reprocessed and check.code == 'EXTRACTION_CRITICAL_INPUTS')
        (self.case.document_ids - document).current_run_id = False

        no_policy = self.env['logistics.idp.extraction.run']._controlled_create({
            'case_id': self.case.id, 'document_id': document.id, 'provider': 'fixture',
            'model_version': 'v1', 'schema_version': 'v1', 'prompt_version': 'v1',
            'started_at': fields.Datetime.now(), 'completed_at': fields.Datetime.now(), 'status': 'review',
            'payload': {'document_type': 'purchase_order', 'payload': {'supplier': 'SUP', 'po_reference': 'PO', 'lines': [{}]},
                        'confidence': 1, 'classification_confidence': 1, 'validation_warnings': []},
        }, 'nfr003_no_policy_test')
        document._create_extraction_critical_check(no_policy)
        no_policy_check = self.case.check_result_ids.filtered(lambda check: check.run_id == no_policy)
        self.assertEqual((no_policy_check.verdict, no_policy_check.payload and json.loads(no_policy_check.payload)['policy_missing']), ('review', True))

        self.case.document_ids.write({'status': 'valid'})
        self.case.write({
            'document_status': 'pass', 'reconciliation_status': 'pass',
            'compliance_status': 'pass', 'output_status': 'pass'})
        self.env['logistics.idp.check.result']._controlled_create({
            'case_id': self.case.id, 'code': 'IMPORT_DECLARATION', 'required': True,
            'verdict': 'pass', 'rationale': 'synthetic lifecycle fixture', 'payload': {},
        }, 'nfr003_import_ready')
        self.case._derive_lifecycle()
        self.assertEqual((ambiguous_check.verdict, reprocessed_check.verdict, document.current_run_id, self.case.state),
                         ('review', 'pass', reprocessed, 'ready'),
                         self.case.check_result_ids.filtered(lambda item: item.required).mapped(lambda item: (item.code, item.verdict, item.run_id.id, item.run_id.document_id.current_run_id.id)))
        self.assertEqual((valid_run.attempt_number, reprocessed.attempt_number), (1, 2))

        for document_type, payload in (
                ('import_declaration', {'declaration_number': 'D', 'lines': [{'hs_code': 'HS'}]}),
                ('purchase_order', {'supplier': 'SUP', 'lines': [{}]}),
                ('main_vat_invoice', {'supplier': 'SUP', 'lines': [{}]})):
            _, check = run_for(payload, document_type=document_type)
            self.assertEqual(check.verdict, 'review')

    def test_native_ocr_attachment_binding_and_immutable_correction_evidence(self):
        document = self.env['logistics.idp.document'].intake_content(
            self.case, b'{"document_type":"invoice","confidence":1,"payload":{}}',
            'application/json', {'filename': 'native-viewer.json'}, process=False)
        self.assertEqual(document.message_main_attachment_id, document.attachment_id)
        self.assertEqual(document.extract_attachment_id, document.attachment_id)
        run_count = len(document.extraction_run_ids)
        run_payload = document.current_run_id.payload
        run_hash = document.current_run_id.payload_hash
        document.write({'corrected_invoice_number': 'INV-CORRECTED'})
        evidence = self.case.evidence_ids.filtered(lambda item: item.category == 'ocr_manual_correction')
        self.assertEqual(len(evidence), 1)
        correction = json.loads(evidence.payload)
        self.assertEqual(correction, {
            'after': {'corrected_invoice_number': 'INV-CORRECTED'},
            'before': {'corrected_invoice_number': False},
            'document_id': document.id,
            'run_hash': run_hash,
            'run_id': document.current_run_id.id,
        })
        document.write({'corrected_invoice_number': 'INV-CORRECTED'})
        document.write({'page_count': 2})
        self.assertEqual(len(self.case.evidence_ids.filtered(
            lambda item: item.category == 'ocr_manual_correction')), 1)
        self.assertEqual(len(document.extraction_run_ids), run_count)
        self.assertEqual((document.current_run_id.payload, document.current_run_id.payload_hash),
                         (run_payload, run_hash))
        with self.assertRaises(UserError):
            evidence.write({'payload': '{}'})

    def test_processed_viewer_exposes_current_payload_and_attempt_history(self):
        metadata = {'document_type': 'invoice', 'confidence': .95, 'payload': {
            'supplier': 'SUP-1', 'invoice_number': 'INV-1', 'lines': [{'quantity': 1, 'unit_price': 1}],
        }}
        document = self.env['logistics.idp.document'].with_user(self.operator)._intake_synthetic_content(
            self.case, b'%PDF-1.4\n1 0 obj <</Type /Page>> endobj\n%%EOF',
            'application/pdf', metadata)
        document.with_user(self.operator)._reprocess_synthetic(metadata)
        runs = document.extraction_run_ids.sorted('attempt_number')
        self.assertEqual([(run.attempt_number, run.retry_count) for run in runs], [(1, 0), (2, 1)])
        for run in runs:
            self.assertTrue(run.provider and run.model_version and run.schema_version and run.prompt_version)
            self.assertTrue(run.started_at and run.completed_at and run.template_version)
        with self.assertRaisesRegex(ValidationError, 'retry count is allocated'):
            self.env['logistics.idp.extraction.run']._controlled_create({
                'case_id': self.case.id, 'document_id': document.id, 'provider': 'fixture',
                'model_version': 'model-v1', 'schema_version': 'schema-v1', 'prompt_version': 'prompt-v1',
                'started_at': fields.Datetime.now(), 'status': 'pending', 'retry_count': 99,
                'payload': {'pending': True},
            }, 'viewer_test')
        with self.assertRaises(UserError):
            runs[0].write({'retry_count': 99})
        self.assertEqual(json.loads(document.normalized_fields)['invoice_number'], 'INV-1')
        self.assertEqual(document.extraction_template_version, runs[-1].template_version)
        with patch('requests.sessions.Session.request') as external_request:
            document.read(['extraction_payload', 'normalized_fields', 'source_spans'])
        external_request.assert_not_called()

    def test_srs_section_7_2_open_source_email_is_persisted_same_company_read_only(self):
        reference = '<source-email-%s@example.test>' % unique_fixture('source-email')
        document = self.env['logistics.idp.document'].with_user(self.operator).intake_content(
            self.case, b'source-email', 'application/pdf', {'source_message_reference': reference},
            source_channel='email', process=False)
        message = self.env['mail.message'].create({
            'message_id': reference, 'message_type': 'email', 'record_company_id': self.env.company.id,
        })
        before = (document.case_id.id, self.case.evidence_ids.ids, message.write_date)
        with patch('requests.sessions.Session.request') as external_request:
            action = document.with_user(self.operator).action_open_source_email()
        external_request.assert_not_called()
        self.assertEqual((action['res_model'], action['res_id'], action['view_mode']), ('mail.message', message.id, 'form'))
        self.assertEqual(action['context'], {'create': False, 'edit': False, 'delete': False})
        self.assertEqual(before, (document.case_id.id, self.case.evidence_ids.ids, message.write_date))

    def test_srs_section_7_2_inbox_projects_extracted_supplier_unique_case_candidate_without_mutation(self):
        reference = unique_fixture('inbox-po')
        target = self.env['logistics.idp.case'].create({
            'name': 'INBOX TARGET', 'source_system': 'fixture', 'source_key': unique_fixture('inbox-target'),
            'source_version': 'v1', 'provenance': 'fixture', 'effective_date': '2026-01-01',
            'po_reference': reference, 'owner_id': self.operator.id,
        })
        document = self.env['logistics.idp.document'].with_user(self.operator)._intake_synthetic_content(
            self.case, b'%PDF-1.4\n1 0 obj <</Type /Page>> endobj\n%%EOF', 'application/pdf', {
                'document_type': 'purchase_order', 'confidence': 1, 'allow_review': True,
                'payload': {'supplier': 'SUP-1', 'po_reference': reference, 'lines': [{}]},
            })
        before = (document.case_id, self.env['logistics.idp.case'].search_count([]),
                  self.env['logistics.idp.document'].search_count([]))
        self.env.invalidate_all()
        document = self.env['logistics.idp.document'].browse(document.id)
        self.assertFalse(document.current_run_id.error)
        self.assertEqual((document.predicted_supplier, document.related_case_candidate_id), ('SUP-1', target))
        self.assertFalse(document.inbox_action_required)
        self.assertEqual(before, (document.case_id, self.env['logistics.idp.case'].search_count([]),
                                  self.env['logistics.idp.document'].search_count([])))

    def test_srs_section_7_2_inbox_ambiguous_or_foreign_candidate_requires_review_without_mutation(self):
        reference = unique_fixture('inbox-po')
        for index in range(2):
            self.env['logistics.idp.case'].create({
                'name': 'INBOX AMBIGUOUS %s' % index, 'source_system': 'fixture',
                'source_key': unique_fixture('inbox-ambiguous'), 'source_version': 'v1', 'provenance': 'fixture',
                'effective_date': '2026-01-01', 'po_reference': reference, 'owner_id': self.operator.id,
            })
        document = self.env['logistics.idp.document'].with_user(self.operator)._intake_synthetic_content(
            self.case, b'%PDF-1.4\n1 0 obj <</Type /Page>> endobj\n%%EOF', 'application/pdf', {
                'document_type': 'purchase_order', 'confidence': 1,
                'payload': {'supplier': 'SUP-1', 'po_reference': reference, 'lines': [{}]},
            })
        foreign_reference = unique_fixture('inbox-foreign-po')
        foreign_company = self.env['res.company'].create({'name': unique_fixture('inbox-foreign')})
        self.env['logistics.idp.case'].with_company(foreign_company).create({
            'name': 'INBOX FOREIGN', 'company_id': foreign_company.id, 'source_system': 'fixture',
            'source_key': unique_fixture('inbox-foreign'), 'source_version': 'v1', 'provenance': 'fixture',
            'effective_date': '2026-01-01', 'po_reference': foreign_reference,
        })
        before = (document.case_id, self.env['logistics.idp.case'].search_count([]),
                  self.env['logistics.idp.document'].search_count([]))
        self.assertFalse(document.related_case_candidate_id)
        self.assertIn('Review / assign case', document.inbox_action_required)
        self.assertEqual(before, (document.case_id, self.env['logistics.idp.case'].search_count([]),
                                  self.env['logistics.idp.document'].search_count([])))
        foreign_only = self.env['logistics.idp.document'].with_user(self.operator)._intake_synthetic_content(
            self.case, b'%PDF-1.4\n1 0 obj <</Type /Page>> endobj\n%%EOF foreign', 'application/pdf', {
                'document_type': 'purchase_order', 'confidence': 1,
                'payload': {'supplier': 'SUP-1', 'po_reference': foreign_reference, 'lines': [{}]},
            })
        self.assertFalse(foreign_only.related_case_candidate_id)
        self.assertIn('Review / assign case', foreign_only.inbox_action_required)

    def test_intake_binds_documents_and_attachment_to_case_company(self):
        company = self.env['res.company'].create({'name': 'IDP Isolated Company'})
        outsider = new_test_user(
            self.env, login='logistics-documents-outsider', company_id=self.env.company.id,
            company_ids=[(6, 0, [self.env.company.id])], groups='documents.group_documents_user')
        case = self.env['logistics.idp.case'].with_company(company).create({
            'name': 'ISOLATED', 'company_id': company.id, 'owner_id': self.operator.id,
            'source_system': 'fixture', 'source_key': unique_fixture('isolated'),
            'source_version': '1', 'provenance': 'fixture', 'effective_date': '2026-01-01',
        })
        document = self.env['logistics.idp.document'].intake_content(
            case, unique_fixture('isolated-content').encode(), 'application/pdf', process=False)
        bound, attachment = document.document_id, document.attachment_id
        self.assertIn(attachment, document._get_mail_thread_data_attachments())
        self.assertEqual((bound.company_id, bound.owner_id), (company, self.operator))
        self.assertEqual(attachment.company_id, company)
        self.assertFalse(bound.folder_id)
        self.assertEqual(bound.access_internal, 'none')
        outsider_env = self.env(user=outsider)
        self.assertFalse(outsider_env['documents.document'].search([('id', '=', bound.id)]))
        self.assertFalse(outsider_env['ir.attachment'].search([('id', '=', attachment.id)]))
        with self.assertRaises(AccessError):
            outsider_env['ir.attachment'].browse(attachment.id).read(['raw'])
        with self.assertRaises(AccessError):
            outsider_env['ir.attachment'].get_file_viewer_data([attachment.id])

    def test_file_viewer_data_is_read_only_and_respects_acl(self):
        document = self.env['logistics.idp.document'].intake_content(
            self.case, unique_fixture('viewer-content').encode(), 'application/pdf',
            {'filename': 'viewer.pdf'}, process=False)
        with patch('odoo.addons.insilos_logistics_idp.services.document_processor.IAPDocumentProcessor.process') as process, \
             patch.object(type(self.env['openrouter.router']), 'complete', autospec=True) as complete, \
             patch.object(type(self.env['iap.charge']), 'hold', autospec=True) as hold, \
             patch.object(type(self.env['iap.charge']), 'consume', autospec=True) as consume:
            data = self.env['ir.attachment'].get_file_viewer_data([document.attachment_id.id])
            document.read(['extraction_payload', 'normalized_fields', 'source_spans'])
        self.assertEqual(len(data), 1)
        self.assertEqual(data[0]['id'], document.attachment_id.id)
        self.assertEqual((data[0]['name'], data[0]['mimetype']), ('viewer.pdf', 'application/pdf'))
        process.assert_not_called()
        complete.assert_not_called()
        hold.assert_not_called()
        consume.assert_not_called()
        forbidden = new_test_user(
            self.env, login=unique_fixture('viewer-no-access'),
            groups='base.group_user')
        with self.assertRaises(AccessError):
            self.env['ir.attachment'].with_user(forbidden).get_file_viewer_data([document.attachment_id.id])

    def test_processed_viewer_handles_missing_attachment_and_run(self):
        document = self.env['logistics.idp.document'].create({
            'case_id': self.case.id, 'content_hash': unique_fixture('missing-attachment'),
            'mimetype': 'application/pdf', 'status': 'error',
        })
        self.assertFalse(document.attachment_id)
        self.assertEqual(json.loads(document.extraction_payload), {})
        self.assertEqual(json.loads(document.normalized_fields), {})
        self.assertEqual(json.loads(document.source_spans), [])

    def test_snapshot_is_canonical_deduplicated_and_immutable(self):
        reference = 'ATT-' + self._testMethodName
        payload = {'b': 2, 'a': 1, 'test': self._testMethodName}
        evidence = self.env['logistics.idp.evidence']._controlled_create({
            'case_id': self.case.id, 'category': 'invoice', 'source_reference': reference,
            'status': 'valid', 'payload': json.dumps(payload),
        }, 'runtime_test')
        self.assertEqual(json.loads(evidence.payload), {'a': 1, 'b': 2, 'test': self._testMethodName})
        with self.assertRaises(UserError):
            evidence.write({'payload': '{}'})
        with self.assertRaises(UserError):
            evidence.unlink()
        with self.assertRaisesRegex(ValidationError, 'Exact evidence replay'):
            self.env['logistics.idp.evidence']._controlled_create({
                'case_id': self.case.id, 'category': 'invoice', 'source_reference': reference + '-duplicate',
                'payload': json.dumps(payload),
            }, 'runtime_test')

    def test_sec003_completed_run_and_output_are_immutable_at_orm_sql_and_ui_boundaries(self):
        document = self.env['logistics.idp.document'].intake_content(
            self.case, b'{"invoice_number":"SEC-003"}', 'application/json',
            {'filename': 'sec003.json'}, process=False)
        run = self.env['logistics.idp.extraction.run']._controlled_create({
            'case_id': self.case.id, 'document_id': document.id, 'provider': 'fixture',
            'model_version': 'model-v1', 'schema_version': 'schema-v1', 'prompt_version': 'prompt-v1',
            'started_at': fields.Datetime.now(), 'status': 'pending',
            'company_id': self.env.company.id + 1, 'document_type': 'invoice', 'page_count': 999,
            'payload': {'payload': {'invoice_number': 'SEC-003'}},
        }, 'sec003_test')
        self.assertEqual((run.company_id, run.document_type, run.page_count),
                         (document.company_id, document.document_type, document.page_count))
        self.env.flush_all()
        self.env.cr.execute(
            "UPDATE logistics_idp_extraction_run SET completed_at = CURRENT_TIMESTAMP, status = 'valid' WHERE id = %s",
            [run.id])
        run.invalidate_recordset(['completed_at', 'status'])
        output = self.env['logistics.idp.output'].generate(
            self.case, 'report_email', {'subject': 'SEC-003'}, run=run)
        before = (run.payload, run.payload_hash, run.model_version, run.prompt_version,
                  output.payload, output.payload_hash, output.version, output.artifact_sha256,
                  output.attachment_id.raw)
        for record, values in ((run, {'payload': '{}'}), (output, {'payload': '{}'})):
            with self.assertRaises(UserError):
                record.write(values)
            with self.assertRaises(UserError):
                record.unlink()
        for field_name, value in (
            ('id', str(run.id + 1000000)), ('payload', "'{}'"), ('case_id', 'NULL'),
            ('company_id', 'NULL'), ('document_id', 'NULL'), ('document_type', "'invoice'"),
            ('page_count', '0'), ('attempt_number', '0'), ('confidence', '0'),
            ('duration_seconds', '0'), ('template_version', "'tampered'"),
            ('error', "'tampered'"), ('audit_actor_id', str(self.operator.id)),
            ('create_uid', str(self.operator.id)), ('write_uid', str(self.operator.id)),
            ('create_date', 'CURRENT_TIMESTAMP'), ('write_date', 'CURRENT_TIMESTAMP'),
            ('audit_service', "'tampered'"),
        ):
            with self.assertRaisesRegex(Exception, 'Completed extraction runs are immutable'), self.env.cr.savepoint():
                self.env.cr.execute('UPDATE logistics_idp_extraction_run SET %s = %s WHERE id = %%s' % (
                    field_name, value), [run.id])
        for field_name, value in (
            ('payload', "'{}'"), ('case_id', 'NULL'), ('run_id', 'NULL'),
            ('output_type', "'e11'"), ('idempotency_key', "'tampered'"),
            ('supersedes_id', str(output.id)), ('audit_service', "'tampered'"), ('status', "'review'"),
        ):
            with self.assertRaisesRegex(Exception, 'Generated outputs are immutable'), self.env.cr.savepoint():
                self.env.cr.execute('UPDATE logistics_idp_output SET %s = %s WHERE id = %%s' % (
                    field_name, value), [output.id])
        for field_name, value in (
            ('id', str(output.id + 1000000)), ('audit_actor_id', str(self.operator.id)),
            ('create_uid', str(self.operator.id)), ('write_uid', str(self.operator.id)),
            ('create_date', 'CURRENT_TIMESTAMP'), ('write_date', 'CURRENT_TIMESTAMP'),
        ):
            with self.assertRaisesRegex(Exception, 'Generated outputs are immutable'), self.env.cr.savepoint():
                self.env.cr.execute(
                    "UPDATE logistics_idp_output SET status = 'superseded', %s = %s WHERE id = %%s" % (
                        field_name, value), [output.id])
        self.env.cr.execute("UPDATE logistics_idp_output SET status = 'superseded' WHERE id = %s", [output.id])
        output.invalidate_recordset(['status'])
        self.assertEqual(output.status, 'superseded')
        with self.assertRaisesRegex(Exception, 'Generated output artifacts are immutable'), self.env.cr.savepoint():
            self.env.cr.execute("UPDATE ir_attachment SET db_datas = NULL WHERE id = %s", [output.attachment_id.id])
        with self.assertRaisesRegex(Exception, 'Generated output artifacts are immutable'), self.env.cr.savepoint():
            self.env.cr.execute("DELETE FROM ir_attachment WHERE id = %s", [output.attachment_id.id])
        unrelated = self.env['ir.attachment'].create({
            'name': 'unrelated-partner-image.svg', 'res_model': 'res.partner', 'res_id': self.operator.partner_id.id,
            'raw': b'<svg/>', 'mimetype': 'image/svg+xml',
        })
        from runpy import run_path
        migrate = run_path(str(Path(__file__).parents[1] / 'migrations/19.0.1.3.7/post-migrate.py'))['migrate']
        self.env.flush_all()
        for attempt in range(2):
            migrate(self.env.cr, '19.0.1.3.6')
            self.env.cr.execute("SELECT pg_get_functiondef('logistics_idp_output_artifact_immutable()'::regprocedure)")
            definition = self.env.cr.fetchone()[0]
            if attempt:
                self.assertEqual(definition, previous_definition)
            previous_definition = definition
            for statement in (
                "UPDATE ir_attachment SET name = 'tampered' WHERE id = %s",
                "DELETE FROM ir_attachment WHERE id = %s",
            ):
                with self.assertRaisesRegex(Exception, 'Generated output artifacts are immutable'), self.env.cr.savepoint():
                    self.env.cr.execute(statement, [output.attachment_id.id])
            name = 'ordinary-persisted-%s' % attempt
            self.env.cr.execute("UPDATE ir_attachment SET name = %s WHERE id = %s RETURNING name", [name, unrelated.id])
            self.assertEqual(self.env.cr.fetchone(), (name,))
            self.env.cr.execute("SELECT name FROM ir_attachment WHERE id = %s", [unrelated.id])
            self.assertEqual(self.env.cr.fetchone(), (name,))
        self.env.cr.execute("DELETE FROM ir_attachment WHERE id = %s RETURNING id", [unrelated.id])
        self.assertEqual(self.env.cr.fetchone(), (unrelated.id,))
        self.env.cr.execute("SELECT id FROM ir_attachment WHERE id = %s", [unrelated.id])
        self.assertIsNone(self.env.cr.fetchone())
        run.invalidate_recordset()
        output.invalidate_recordset()
        output.attachment_id.invalidate_recordset()
        self.assertEqual((run.payload, run.payload_hash, run.model_version, run.prompt_version,
                          output.payload, output.payload_hash, output.version, output.artifact_sha256,
                          output.attachment_id.raw), before)
        views = Path(__file__).parents[1] / 'views/logistics_idp_views.xml'
        source = views.read_text()
        self.assertIn('view_logistics_run_form', source)
        self.assertIn('<form create="0" edit="0" delete="0">', source)
        self.assertNotIn('name="unlink"', source)

    def test_fr301_supplier_draft_invoice_policy_modes_effectivity_isolation_and_replay(self):
        po = {'supplier': 'FR301', 'regime': 'E13', 'lines': [{'material_code': 'M', 'regime': 'E13', 'quantity': 1, 'unit_price': 1}]}
        base_invoice = {'supplier': 'FR301', 'regime': 'E13', 'lines': [{'material_code': 'M', 'regime': 'E13', 'quantity': 1, 'unit_price': 1}]}
        expectations = {
            'required': (None, 'review'), 'optional': (None, 'pass'), 'skipped': (None, 'pass'),
            'sales_invoice': ('sales_invoice', 'pass'), 'commercial_invoice': ('commercial_invoice', 'pass'),
        }
        for mode, (document_type, verdict) in expectations.items():
            result = reconcile_documents(po, {**base_invoice, **({'document_type': document_type} if document_type else {})},
                                         policy={'draft_invoice': {'mode': mode, 'final_number_allowed': False}})
            self.assertEqual(result['verdict'], verdict)
        case = self.env['logistics.idp.case'].create({
            'name': 'FR301-' + unique_fixture('case'), 'source_system': 'fixture', 'source_key': unique_fixture('FR301'),
            'source_version': '1', 'provenance': 'fixture:development', 'effective_date': '2026-06-01', 'supplier_reference': 'FR301',
        })
        profile = self.env['logistics.idp.supplier.profile'].import_upsert({
            'company_id': case.company_id.id, 'supplier_reference': 'FR301', 'profile_code': 'FR301-REQUIRED', 'version': 'v1',
            'source_system': 'fixture', 'source_key': unique_fixture('FR301-profile'), 'source_version': '1',
            'provenance': 'fixture:development', 'effective_from': '2026-01-01', 'effective_to': '2026-12-31',
            'payload': {'draft_invoice': {'mode': 'required', 'final_number_allowed': False}},
        })
        evidence = case.reconcile_documents(po, base_invoice)
        decision = case.decision_ids[0]
        binding = json.loads(decision.payload)['input_manifest']['supplier_profile']
        self.assertEqual((evidence.status, binding['profile_code'], binding['version'], binding['payload_hash']),
                         ('review', profile.profile_code, profile.version, profile.payload_hash))
        self.assertEqual(reconcile_documents(po, base_invoice, policy={'draft_invoice': {'mode': 'required', 'final_number_allowed': False}}),
                         reconcile_documents(po, base_invoice, policy={'draft_invoice': {'mode': 'required', 'final_number_allowed': False}}))
        other_company = self.env['res.company'].create({'name': 'FR301 Other Company'})
        self.env.user.company_ids |= other_company
        self.env['logistics.idp.supplier.profile'].with_company(other_company).import_upsert({
            'company_id': other_company.id, 'supplier_reference': 'FR301', 'profile_code': 'FR301-FOREIGN', 'version': 'v1',
            'source_system': 'fixture', 'source_key': unique_fixture('FR301-foreign'), 'source_version': '1',
            'provenance': 'fixture:development', 'effective_from': fields.Date.to_string(self.case.effective_date),
            'payload': {'draft_invoice': {'mode': 'optional', 'final_number_allowed': False}},
        })
        case.effective_date = '2027-01-01'
        self.assertEqual(case.reconcile_documents(po, base_invoice).status, 'review')

    def test_fr403_missing_regime_persists_review_and_replay_is_side_effect_free(self):
        self.case.supplier_reference = 'FR403'
        po = {'supplier': 'FR403', 'lines': [{'material_code': 'MAT-403', 'quantity': 1, 'unit_price': 1}]}
        invoice = {'supplier': 'FR403', 'lines': [{'material_code': 'MAT-403', 'quantity': 1, 'unit_price': 1}]}
        self.env['logistics.idp.supplier.profile'].import_upsert({
            'company_id': self.case.company_id.id, 'supplier_reference': 'FR403', 'profile_code': 'FR403', 'version': 'v1',
            'source_system': 'fixture', 'source_key': unique_fixture('FR403-profile'), 'source_version': '1',
            'provenance': 'fixture:development', 'effective_from': '2025-01-01',
            'payload': {'draft_invoice': {'mode': 'skipped', 'final_number_allowed': False}},
        })
        snapshots = self.env['logistics.idp.reference.snapshot']
        for reference_type in ('master_data', 'dsnavl'):
            snapshots.import_upsert({'reference_type': reference_type, 'source_system': 'fixture',
                'source_key': unique_fixture('FR403-' + reference_type), 'source_version': 'v1',
                'provenance': 'fixture:development', 'effective_date': '2025-01-01',
                'payload': {'lines': [{'material_code': 'MAT-403'}]}})
        master_before = [(item.id, item.payload_hash) for item in snapshots.search([('reference_type', 'in', ('master_data', 'dsnavl'))])]
        with patch('odoo.addons.insilos_logistics_idp.models.logistics_idp.IAPDocumentProcessor.process') as provider:
            evidence = self.case.reconcile_documents(po, invoice)
            decision = self.case.decision_ids[-1]
            self.assertEqual((evidence.status, self.case.verdict, decision.verdict), ('review', 'review', 'review'))
            self.assertEqual([(item['material_code'], item['regime'], item['reason'], item['result'])
                              for item in json.loads(evidence.payload)['results']],
                             [('MAT-403', '', 'missing customs regime', 'review')])
            self.assertEqual([(item['material_code'], item['regime'], item['reason'], item['result'])
                              for item in json.loads(decision.payload)['reconciliation']['results']],
                             [('MAT-403', '', 'missing customs regime', 'review')])
            self.assertEqual(self.case.reconcile_documents(po, invoice), evidence)
            provider.assert_not_called()
        self.assertEqual([(item.id, item.payload_hash) for item in snapshots.search([('reference_type', 'in', ('master_data', 'dsnavl'))])], master_before)
        self.assertEqual((len(self.case.evidence_ids.filtered(lambda item: item.category == 'reconciliation')),
                          len(self.case.decision_ids)), (1, 1))

    def test_fr501_main_invoice_policy_effectivity_isolation_and_replay(self):
        po = {'supplier': 'FR501', 'regime': 'E13', 'lines': [{'material_code': 'M', 'regime': 'E13', 'quantity': 1, 'unit_price': 1}]}
        invoice = {'supplier': 'FR501', 'regime': 'E13', 'document_type': 'main_vat_invoice', 'invoice_number': 'VAT-1',
                   'lines': [{'material_code': 'M', 'regime': 'E13', 'quantity': 1, 'unit_price': 1}]}
        policy = {'draft_invoice': {'mode': 'skipped', 'final_number_allowed': False}}
        self.assertEqual(reconcile_documents(po, invoice, policy=policy)['verdict'], 'pass')
        self.assertEqual(reconcile_documents(po, {**invoice, 'invoice_number': ''}, policy=policy)['verdict'], 'review')
        for document_type in ('sales_invoice', 'commercial_invoice'):
            self.assertEqual(reconcile_documents(po, {**invoice, 'document_type': document_type}, policy=policy)['verdict'], 'review')
        self.assertEqual(reconcile_documents(po, {**invoice, 'document_type': 'sales_invoice'}, policy={
            **policy, 'main_invoice': {'allowed_substitutes': ['sales_invoice']}})['verdict'], 'pass')
        case = self.env['logistics.idp.case'].create({
            'name': 'FR501-' + unique_fixture('case'), 'source_system': 'fixture', 'source_key': unique_fixture('FR501'),
            'source_version': '1', 'provenance': 'fixture:development', 'effective_date': '2026-06-01', 'supplier_reference': 'FR501',
        })
        profile = self.env['logistics.idp.supplier.profile'].import_upsert({
            'company_id': case.company_id.id, 'supplier_reference': 'FR501', 'profile_code': 'FR501-SALES', 'version': 'v1',
            'source_system': 'fixture', 'source_key': unique_fixture('FR501-profile'), 'source_version': '1',
            'provenance': 'fixture:development', 'effective_from': '2026-01-01', 'effective_to': '2026-12-31',
            'payload': {**policy, 'main_invoice': {'allowed_substitutes': ['sales_invoice']}},
        })
        sales = {**invoice, 'document_type': 'sales_invoice'}
        first = case.reconcile_documents(po, sales)
        self.assertEqual(first.status, 'valid')
        self.assertEqual(case.reconcile_documents(po, sales).id, first.id)
        self.assertEqual(len(case.evidence_ids.filtered(lambda item: item.category == 'reconciliation')), 1)
        other_company = self.env['res.company'].create({'name': 'FR501 Other Company'})
        self.env.user.company_ids |= other_company
        self.env['logistics.idp.supplier.profile'].with_company(other_company).import_upsert({
            'company_id': other_company.id, 'supplier_reference': 'FR501', 'profile_code': 'FR501-FOREIGN', 'version': 'v1',
            'source_system': 'fixture', 'source_key': unique_fixture('FR501-foreign'), 'source_version': '1',
            'provenance': 'fixture:development', 'effective_from': '2026-01-01', 'payload': {**policy, 'main_invoice': {'allowed_substitutes': ['sales_invoice']}},
        })
        case.effective_date = '2027-01-01'
        self.assertEqual(case.reconcile_documents(po, sales).status, 'review')
        self.cr.execute('UPDATE logistics_idp_supplier_profile SET payload = %s, payload_hash = %s WHERE id = %s',
                        [json.dumps({'draft_invoice': {'mode': 'skipped', 'final_number_allowed': False}, 'main_invoice': {'allowed_substitutes': ['invoice']}}), '0' * 64, profile.id])
        case.effective_date = '2026-06-01'
        self.assertEqual(case.reconcile_documents(po, sales).status, 'review')

    def test_fr502_main_invoice_requires_exact_approved_draft_snapshot_and_replay(self):
        self.case.supplier_reference = 'FR502'
        po = {'supplier': 'FR502', 'regime': 'E13', 'lines': [{'material_code': 'M', 'regime': 'E13', 'quantity': 1, 'unit_price': 2, 'line_total': 2}]}
        main = {'supplier': 'FR502', 'regime': 'E13', 'document_type': 'main_vat_invoice', 'invoice_number': 'VAT-502',
                'lines': [{'material_code': 'M', 'regime': 'E13', 'quantity': 1, 'unit_price': 2, 'line_total': 2}]}
        self.env['logistics.idp.supplier.profile'].import_upsert({
            'company_id': self.case.company_id.id, 'supplier_reference': 'FR502', 'profile_code': unique_fixture('FR502'),
            'version': 'v1', 'source_system': 'fixture', 'source_key': unique_fixture('FR502-profile'), 'source_version': '1',
            'provenance': 'fixture:development', 'effective_from': '2025-01-01',
            'payload': {'draft_invoice': {'mode': 'required', 'final_number_allowed': False}},
        })
        missing = self.case.reconcile_documents(po, main)
        self.assertEqual(missing.status, 'review')
        draft = self.env['logistics.idp.document'].with_user(self.operator).intake_content(
            self.case, b'%PDF-1.4\n%%EOF FR502', 'application/pdf',
            {'document_type_hint': 'draft_vat_invoice', 'status': 'valid'}, process=False)
        run = self.env['logistics.idp.extraction.run']._controlled_create({
            'case_id': self.case.id, 'document_id': draft.id, 'provider': 'fixture', 'model_version': 'v1',
            'schema_version': 'v1', 'prompt_version': 'v1', 'started_at': fields.Datetime.now(),
            'completed_at': fields.Datetime.now(), 'status': 'valid',
            'payload': {'payload': {'supplier': 'FR502', 'regime': 'E13', 'lines': [
                {'material_code': 'M', 'regime': 'E13', 'quantity': 1, 'unit_price': 2, 'line_total': 2}]}},
        }, 'fr502')
        draft.write({'current_run_id': run.id})
        passed = self.case.reconcile_documents(po, main)
        self.assertEqual(passed.status, 'valid', passed.payload)
        manifest = json.loads(self.env['logistics.idp.policy.decision'].search([('case_id', '=', self.case.id)], order='id desc', limit=1).payload)['input_manifest']
        self.assertEqual(manifest['draft_invoice']['selected']['id'], draft.id)
        before = (draft.current_run_id.payload, draft.current_run_id.payload_hash, json.loads(passed.payload))
        self.assertEqual(self.case.reconcile_documents(po, main).id, passed.id)
        self.assertEqual(before, (draft.current_run_id.payload, draft.current_run_id.payload_hash, json.loads(passed.payload)))
        mismatch = self.case.reconcile_documents(po, {**main, 'lines': [{**main['lines'][0], 'line_total': 3}]})
        self.assertEqual(mismatch.status, 'invalid')
        draft.write({'status': 'review'})
        self.assertEqual(self.case.reconcile_documents(po, main).status, 'review')
        draft.write({'status': 'valid'})
        duplicate = self.env['logistics.idp.document'].with_user(self.operator).intake_content(
            self.case, b'%PDF-1.4\n%%EOF FR502 duplicate', 'application/pdf',
            {'document_type_hint': 'draft_vat_invoice', 'status': 'valid'}, process=False)
        duplicate_run = self.env['logistics.idp.extraction.run']._controlled_create({
            'case_id': self.case.id, 'document_id': duplicate.id, 'provider': 'fixture', 'model_version': 'v1',
            'schema_version': 'v1', 'prompt_version': 'v1', 'started_at': fields.Datetime.now(),
            'completed_at': fields.Datetime.now(), 'status': 'valid',
            'payload': {'payload': {'supplier': 'FR502', 'regime': 'E13', 'lines': [
                {'material_code': 'M', 'regime': 'E13', 'quantity': 1, 'unit_price': 2}]}},
        }, 'fr502')
        duplicate.write({'current_run_id': duplicate_run.id})
        self.assertEqual(self.case.reconcile_documents(po, main).status, 'review')
        self.assertEqual(duplicate.company_id, self.case.company_id)

    def test_fr503_skipped_profile_reconciles_main_directly_and_never_compares_draft(self):
        self.case.supplier_reference = 'FR503'
        po = {'supplier': 'FR503', 'regime': 'E13', 'lines': [{'material_code': 'M', 'regime': 'E13', 'quantity': 1, 'unit_price': 2, 'line_total': 2}]}
        main = {'supplier': 'FR503', 'regime': 'E13', 'document_type': 'main_vat_invoice', 'invoice_number': 'VAT-503',
                'lines': [{'material_code': 'M', 'regime': 'E13', 'quantity': 1, 'unit_price': 2, 'line_total': 2}]}
        self.env['logistics.idp.supplier.profile'].import_upsert({
            'company_id': self.case.company_id.id, 'supplier_reference': 'FR503', 'profile_code': unique_fixture('FR503'),
            'version': 'v1', 'source_system': 'fixture', 'source_key': unique_fixture('FR503-profile'), 'source_version': '1',
            'provenance': 'fixture:development', 'effective_from': '2025-01-01',
            'payload': {'draft_invoice': {'mode': 'skipped', 'final_number_allowed': False}},
        })
        direct = self.case.reconcile_documents(po, main)
        self.assertEqual((direct.status, json.loads(direct.payload)['draft_main']), ('valid', False))
        direct_before = (direct.id, direct.payload, direct.payload_hash)
        self.assertEqual(self.case.reconcile_documents(po, main).id, direct.id)
        self.assertEqual((direct.id, direct.payload, direct.payload_hash), direct_before)
        draft = self.env['logistics.idp.document'].with_user(self.operator).intake_content(
            self.case, b'%PDF-1.4\n%%EOF FR503', 'application/pdf',
            {'document_type_hint': 'draft_vat_invoice', 'status': 'valid'}, process=False)
        run = self.env['logistics.idp.extraction.run']._controlled_create({
            'case_id': self.case.id, 'document_id': draft.id, 'provider': 'fixture', 'model_version': 'v1',
            'schema_version': 'v1', 'prompt_version': 'v1', 'started_at': fields.Datetime.now(),
            'completed_at': fields.Datetime.now(), 'status': 'valid', 'payload': {'payload': {'supplier': 'FR503', 'regime': 'E13', 'lines': [
                {'material_code': 'M', 'regime': 'E13', 'quantity': 1, 'unit_price': 999, 'line_total': 999}]}},
        }, 'fr503')
        draft.write({'current_run_id': run.id})
        with patch('odoo.addons.insilos_logistics_idp.models.logistics_idp.reconcile_draft_main') as draft_compare:
            passed = self.case.reconcile_documents(po, main)
            self.assertEqual((passed.status, json.loads(passed.payload)['draft_main']), ('valid', False))
            draft_compare.assert_not_called()
        self.assertEqual(self.case.reconcile_documents(po, main).id, passed.id)
        self.assertEqual(self.case.reconcile_documents(po, {**main, 'lines': [{**main['lines'][0], 'quantity': 3}]}).status, 'invalid')

    def test_sec007_decision_manifest_binds_inputs_and_recheck_appends(self):
        evidence = self.case.reconcile_documents(self.po, self.invoice, reconciliation_inputs={
            'po_reference': 'PO-1', 'invoice_reference': 'INV-1', 'declaration_reference': 'TK-1',
        })
        decision = self.env['logistics.idp.policy.decision'].search([('case_id', '=', self.case.id)], limit=1)
        payload = json.loads(decision.payload)
        manifest = payload['input_manifest']
        self.assertEqual(decision.audit_input_hash, payload['binding_hash'])
        self.assertEqual(payload['binding_hash'], hashlib.sha256(
            json.dumps(manifest, sort_keys=True, separators=(',', ':'), ensure_ascii=True).encode()).hexdigest())
        self.assertEqual(manifest['references'], {
            'po_reference': 'PO-1', 'invoice_reference': 'INV-1', 'declaration_reference': 'TK-1'})
        self.assertEqual(json.loads(evidence.payload)['decision_hash'], decision.payload_hash)
        before = (decision.payload, decision.payload_hash, decision.audit_input_hash)
        self.case.recheck(self.po, self.invoice)
        decisions = self.env['logistics.idp.policy.decision'].search([('case_id', '=', self.case.id)])
        self.assertEqual(len(decisions), 2)
        self.assertEqual((decision.payload, decision.payload_hash, decision.audit_input_hash), before)

    def test_reconciliation_without_effective_profile_serializes_false(self):
        self.case.effective_date = '2025-12-31'
        evidence = self.case.reconcile_documents(self.po, self.invoice)
        decision = self.env['logistics.idp.policy.decision'].search([('case_id', '=', self.case.id)], limit=1)
        self.assertEqual((evidence.status, self.case.verdict), ('review', 'review'))
        self.assertIs(json.loads(decision.payload)['input_manifest']['supplier_profile'], False)
        self.assertIsInstance(json.loads(evidence.payload), dict)

    def test_sec007_decision_sql_update_is_rejected(self):
        self.case.reconcile_documents(self.po, self.invoice)
        decision = self.env['logistics.idp.policy.decision'].search([('case_id', '=', self.case.id)], limit=1)
        with self.assertRaisesRegex(Exception, 'Policy decisions are immutable'), self.env.cr.savepoint():
            self.env.cr.execute("UPDATE logistics_idp_policy_decision SET reason = 'tampered' WHERE id = %s", [decision.id])

    def test_sec007_cross_case_manifest_binding_rejects(self):
        self.case.reconcile_documents(self.po, self.invoice)
        decision = self.env['logistics.idp.policy.decision'].search([('case_id', '=', self.case.id)], limit=1)
        other = self.env['logistics.idp.case'].create({
            'name': 'SEC-007-other', 'source_system': 'fixture', 'source_key': unique_fixture('SEC-007'),
            'source_version': '1', 'provenance': 'fixture:development', 'effective_date': '2026-01-01',
        })
        with self.assertRaisesRegex(ValidationError, 'manifest binding'):
            self.env['logistics.idp.policy.decision']._controlled_create({
                'case_id': other.id, 'supplier_profile_id': decision.supplier_profile_id.id,
                'policy_code': decision.policy_code, 'policy_version': decision.policy_version,
                'effective_from': fields.Datetime.now(), 'verdict': decision.verdict, 'reason': decision.reason,
                'payload': decision.payload, 'audit_input_hash': decision.audit_input_hash,
            }, 'runtime_test')

    def test_sec007_cross_company_supplier_profile_rejects(self):
        self.case.reconcile_documents(self.po, self.invoice)
        decision_model = self.env['logistics.idp.policy.decision']
        decision = decision_model.search([('case_id', '=', self.case.id)], limit=1)
        other_company = self.env['res.company'].create({'name': 'SEC-007 Other Company'})
        self.env.user.company_ids |= other_company
        profile = self.env['logistics.idp.supplier.profile'].with_company(other_company).import_upsert({
            'company_id': other_company.id, 'supplier_reference': 'SUP-1', 'profile_code': 'FOREIGN',
            'version': '1', 'source_system': 'fixture', 'source_key': unique_fixture('SEC-007-foreign'),
            'source_version': '1', 'provenance': 'fixture:development', 'effective_from': '2026-01-01', 'payload': {},
        })
        payload = json.loads(decision.payload)
        manifest = payload['input_manifest']
        manifest['supplier_profile'] = {
            'id': profile.id, 'profile_code': profile.profile_code, 'version': profile.version,
            'payload_hash': profile.payload_hash, 'effective_from': fields.Date.to_string(profile.effective_from), 'effective_to': False,
        }
        payload['binding_hash'] = hashlib.sha256(
            json.dumps(manifest, sort_keys=True, separators=(',', ':'), ensure_ascii=True).encode()).hexdigest()
        before = decision_model.search_count([('case_id', '=', self.case.id)])
        with self.assertRaisesRegex(ValidationError, 'manifest binding'):
            decision_model._controlled_create({
                'case_id': self.case.id, 'supplier_profile_id': profile.id,
                'policy_code': profile.profile_code, 'policy_version': profile.version,
                'effective_from': fields.Datetime.now(), 'verdict': decision.verdict, 'reason': decision.reason,
                'payload': payload, 'audit_input_hash': payload['binding_hash'],
            }, 'runtime_test')
        self.assertEqual(decision_model.search_count([('case_id', '=', self.case.id)]), before)

    def test_po_invoice_reconciliation_and_recheck_preserve_runs(self):
        failed = self.case.reconcile_documents(self.po, {**self.invoice, 'lines': [{**self.invoice['lines'][0], 'quantity': '11'}]})
        self.assertEqual(failed.status, 'invalid')
        passed = self.case.recheck(self.po, self.invoice)
        self.assertEqual(passed.status, 'valid')
        self.assertNotEqual(failed.payload_hash, passed.payload_hash)
        self.assertEqual(len(self.case.evidence_ids.filtered(lambda item: item.category == 'reconciliation')), 2)

    def test_matching_uses_identity_not_sequence_and_price_per(self):
        po = {'supplier': 'SUP-1', 'regime': 'E13', 'lines': [self.po['lines'][0], {
            'material_code': 'MAT-2', 'unit_price': '2', 'remaining_quantity': '3', 'custom_code': 'PO-1-20',
        }]}
        invoice = {'supplier': 'SUP-1', 'regime': 'E13', 'lines': [
            {'material_code': 'MAT-2', 'unit_price': '2', 'quantity': '1'}, self.invoice['lines'][0],
        ]}
        result = json.loads(self.case.reconcile_documents(po, invoice).payload)
        self.assertEqual([line['po_line']['material_code'] for line in result['results']], ['MAT-1', 'MAT-2'])
        self.assertTrue(all(line['result'] == 'pass' for line in result['results']))

    def test_draft_po_multiline_identity_unmatched_evidence_is_persisted(self):
        po = {'supplier': 'SUP-1', 'regime': 'E13', 'lines': [
            {'material_code': 'MAT-1', 'unit_price': '1', 'remaining_quantity': '4'},
            {'material_code': 'MAT-2', 'unit_price': '2', 'remaining_quantity': '3'},
        ]}
        invoice = {'supplier': 'SUP-1', 'regime': 'E13', 'lines': [
            {'material_code': 'MAT-2', 'unit_price': '2', 'quantity': '1'},
            {'material_code': 'MAT-3', 'unit_price': '3', 'quantity': '1'},
        ]}
        evidence = self.case.reconcile_documents(po, invoice)
        persisted = self.env['logistics.idp.evidence'].browse(evidence.id)
        result = json.loads(persisted.payload)
        self.assertEqual(persisted.category, 'reconciliation')
        self.assertEqual([line['identity'] for line in result['results']], ['MAT-1', 'MAT-2', 'MAT-3'])
        self.assertEqual([line['result'] for line in result['results']], ['review', 'pass', 'review'])
        self.assertEqual(self.case.verdict, 'review')

    def test_fr304_draft_result_persists_differences_warnings_and_evidence(self):
        po = {'supplier': 'SUP-1', 'lines': [
            {'material_code': 'MAT-1', 'unit_price': '1', 'remaining_quantity': '4'},
            {'material_code': 'MAT-2', 'unit_price': '2', 'remaining_quantity': '3'},
        ]}
        failed = self.case.reconcile_documents(po, {
            'supplier': 'SUP-2', 'regime': 'E13', 'lines': [
                {'material_code': 'MAT-1', 'unit_price': '2', 'quantity': '5'},
            ],
        })
        failed_result = json.loads(failed.payload)
        difference = failed_result['results'][0]
        self.assertEqual((failed.status, failed_result['verdict'], difference['result']), ('invalid', 'block', 'block'))
        self.assertEqual((difference['quantity_1_ok'], difference['price_ok'], difference['supplier_ok']),
                         (False, False, False))
        self.assertEqual(difference['regime'], 'E13')
        self.assertTrue(failed.payload_hash)
        self.assertEqual(failed.case_id, self.case)

        warning = self.case.recheck(po, {
            'supplier': 'SUP-1', 'regime': 'E15', 'lines': [
                {'material_code': 'MAT-2', 'unit_price': '2', 'quantity': '1'},
                {'material_code': 'MAT-3', 'unit_price': '3', 'quantity': '1'},
            ],
        })
        warning_result = json.loads(warning.payload)
        self.assertEqual((warning.status, warning_result['verdict']), ('review', 'review'))
        self.assertEqual([(line['identity'], line['regime'], line['reason']) for line in warning_result['results'] if line['result'] == 'review'],
                         [('MAT-1', 'E15', 'unmatched'), ('MAT-3', 'E15', 'unmatched')])
        self.assertEqual(len(self.case.evidence_ids.filtered(lambda item: item.category == 'reconciliation')), 2)
        self.assertNotEqual(failed.payload_hash, warning.payload_hash)

    def test_two_po_lines_group_to_one_invoice_line_by_identity(self):
        po = {'supplier': 'SUP-1', 'regime': 'E13', 'lines': [
            {'po_reference': 'PO-1', 'material_code': 'MAT-1', 'unit_price': '1', 'remaining_quantity': '2'},
            {'po_reference': 'PO-2', 'material_code': 'MAT-1', 'unit_price': '1', 'remaining_quantity': '3'},
        ]}
        invoice = {'supplier': 'SUP-1', 'regime': 'E13', 'lines': [
            {'material_code': 'MAT-1', 'unit_price': '1', 'quantity': '5'},
        ]}
        result = json.loads(self.case.reconcile_documents(po, invoice).payload)
        self.assertEqual((result['verdict'], result['results'][0]['cardinality']), ('pass', 'many-to-one'))
        self.assertEqual([line['po_reference'] for line in result['results'][0]['po_lines']], ['PO-1', 'PO-2'])

    def test_profile_sequence_warn_and_block_preserve_identity_mapping(self):
        po = {'supplier': 'SUP-1', 'lines': [
            {'material_code': 'MAT-1', 'unit_price': '1', 'remaining_quantity': '1'},
            {'material_code': 'MAT-2', 'unit_price': '2', 'remaining_quantity': '1'},
        ]}
        invoice = {'supplier': 'SUP-1', 'lines': [
            {'material_code': 'MAT-2', 'unit_price': '2', 'quantity': '1'},
            {'material_code': 'MAT-1', 'unit_price': '1', 'quantity': '1'},
        ]}
        config = json.loads(self.sap_profile.payload)
        config['sequence_policy'] = 'warn'
        self._write_sap_config(config)
        warned = json.loads(self.case.reconcile_documents(po, invoice).payload)
        self.assertEqual((warned['verdict'], warned['sequence']),
                         ('review', {'mismatch': True, 'policy': 'warn', 'result': 'review'}))
        self.assertEqual([line['identity'] for line in warned['results']], ['MAT-1', 'MAT-2'])

        config['sequence_policy'] = 'block'
        self._write_sap_config(config)
        blocked = json.loads(self.case.reconcile_documents(po, invoice).payload)
        self.assertEqual((blocked['verdict'], blocked['sequence']['result']), ('block', 'block'))
        self.assertEqual([line['identity'] for line in blocked['results']], ['MAT-1', 'MAT-2'])

    def test_reconciliation_po_order_grouping_currency_and_regime_are_deterministic(self):
        po = {'supplier': 'SUP-1', 'regime': 'E13', 'lines': [
            {'material_code': 'MAT-1', 'unit_price': '1', 'currency': 'USD', 'remaining_quantity': '4'},
            {'material_code': 'MAT-2', 'unit_price': '2', 'currency': 'USD', 'remaining_quantity': '3'},
        ]}
        invoice = {'supplier': 'SUP-1', 'regime': 'E13', 'lines': [
            {'material_code': 'MAT-2', 'unit_price': '2', 'currency': 'USD', 'quantity': '1'},
            {'material_code': 'MAT-1', 'unit_price': '1', 'currency': 'USD', 'quantity': '1'},
            {'material_code': 'MAT-1', 'unit_price': '1', 'currency': 'USD', 'quantity': '2'},
        ]}
        result = json.loads(self.case.reconcile_documents(po, invoice).payload)
        self.assertEqual([item['identity'] for item in result['results']], ['MAT-1', 'MAT-2'])
        self.assertEqual((result['results'][0]['cardinality'], result['verdict']), ('one-to-many', 'pass'))

        currency = json.loads(self.case.recheck(po, {**invoice, 'lines': [
            {**invoice['lines'][0], 'currency': 'VND'}, *invoice['lines'][1:],
        ]}).payload)
        self.assertEqual((currency['verdict'], currency['results'][1]['result']), ('review', 'review'))

        mixed = json.loads(self.case.recheck(po, {**invoice, 'lines': [
            {**invoice['lines'][0], 'regime': 'E13'},
            {**invoice['lines'][1], 'regime': 'E11'}, {**invoice['lines'][2], 'regime': 'E15'},
        ]}).payload)
        self.assertEqual((mixed['verdict'], mixed['results'][0]['reason']),
                         ('review', 'mixed customs regimes require explicit grouping policy'))

    def test_fr106_same_normalized_subject_keeps_latest_active_history_and_replay(self):
        base = {'name': 'Mail A', 'source_system': 'supplier_email', 'source_version': '1',
                'provenance': 'fixture:development', 'effective_date': '2026-01-01',
                'supplier_reference': 'SUP-FR106', 'email_subject': ' Shipment   Update ',
                'thread_reference': 'thread-1'}
        case = self.env['logistics.idp.case'].intake({**base, 'source_key': unique_fixture('fr106-a')})
        latest = self.env['logistics.idp.case'].intake({**base, 'name': 'Mail B',
            'source_key': unique_fixture('fr106-b'), 'email_subject': 'shipment update', 'thread_reference': 'thread-2'})
        self.assertEqual(latest, case)
        entries = self.env['logistics.idp.email.thread.entry'].search([('case_id', '=', case.id)], order='id')
        self.assertEqual((len(entries), entries.mapped('is_active'), case.thread_reference), (2, [False, True], 'thread-2'))
        self.assertEqual(self.env['logistics.idp.case'].intake({**base, 'name': 'Mail B',
            'source_key': latest.source_key, 'email_subject': 'shipment update', 'thread_reference': 'thread-2'}), case)
        entries = self.env['logistics.idp.email.thread.entry'].search([('case_id', '=', case.id)])
        self.assertEqual((len(entries), len(entries.filtered('is_active'))), (2, 1))

    def test_fr107_different_subject_requires_idempotent_review_and_explicit_active_thread_selection(self):
        base = {
            'name': 'Mail A', 'source_system': 'supplier_email', 'source_version': '1',
            'provenance': 'fixture:development', 'effective_date': '2026-01-01',
            'supplier_reference': 'SUP-SUBJECT', 'email_subject': 'Shipment update',
        }
        active = self.env['logistics.idp.case'].intake({**base, 'source_key': unique_fixture('subject-a'),
                                                        'thread_reference': 'thread-old'})
        ambiguous = self.env['logistics.idp.case'].intake({
            **base, 'name': 'Mail B', 'source_key': unique_fixture('subject-c'),
            'email_subject': 'Unresolved other subject', 'thread_reference': 'thread-other',
        })
        before = (active.thread_reference, ambiguous.email_subject, ambiguous.document_ids.ids, ambiguous.state)
        self.assertEqual(self.env['logistics.idp.case'].intake({
            **base, 'name': 'Mail B', 'source_key': ambiguous.source_key,
            'email_subject': ambiguous.email_subject, 'thread_reference': ambiguous.thread_reference,
        }), ambiguous)
        activities = ambiguous.activity_ids.filtered(lambda activity: activity.summary == 'FR-107: Select active supplier email thread')
        self.assertEqual((len(activities), before), (1, ('thread-old', 'Unresolved other subject', [], 'review')))
        reviewer = new_test_user(self.env, login='fr107-reviewer-' + unique_fixture('reviewer'),
                                 context={'no_reset_password': True}, groups='insilos_logistics_idp.group_logistics_reviewer')
        self.assertTrue(ambiguous.with_user(reviewer).action_select_active_thread(active))
        self.assertEqual((ambiguous.active_thread_case_id, ambiguous.thread_reference, ambiguous.email_subject,
                          ambiguous.document_ids.ids, ambiguous.state),
                         (active, 'thread-other', 'Unresolved other subject', [], 'review'))
        self.assertFalse(ambiguous.activity_ids.filtered(lambda activity: activity.summary == 'FR-107: Select active supplier email thread'))

    def test_semantic_duplicate_links_canonical_case(self):
        duplicate = self.env['logistics.idp.case'].create({
            'name': 'SYNTHETIC-DUP', 'source_system': 'upload', 'source_key': 'DUP-1', 'source_version': 'v1',
            'provenance': 'fixture:development', 'effective_date': '2026-01-01',
        })
        duplicate.mark_semantic_duplicate(self.case)
        self.assertEqual(duplicate.canonical_case_id, self.case)
        self.assertEqual(duplicate.state, 'closed_duplicate')

    def test_vietnam_outputs_import_reference_shipping_and_gate_pass(self):
        line = {**self.po['lines'][0], 'quantity': '4'}
        e13 = self.case.generate_customs_output('E13', [line], 'vn-2026')
        self.assertEqual(json.loads(e13.payload)['regime'], 'E13')
        self.case.generate_customs_output('E15', [line], 'vn-2026')
        self.assertEqual(self.case.generate_broker_package().status, 'generated')
        with self.assertRaisesRegex(ValidationError, 'no caller-provided'):
            self.case.generate_shipping_plan({'plan_key': 'PLAN-1'})
        with self.assertRaisesRegex(ValidationError, 'no caller-provided'):
            self.case.generate_gate_pass(e13)

    def test_fr801_broker_package_missing_requirement_creates_no_output(self):
        with self.assertRaisesRegex(ValidationError, 'missing required categories'):
            self.case.generate_broker_package()
        self.assertFalse(self.case.output_ids.filtered(lambda item: item.output_type == 'broker_package'))
        with self.assertRaisesRegex(ValidationError, 'does not accept caller-provided'):
            self.case.generate_broker_package(['customs_e13'])

    def test_exc002_controlled_override_workflow(self):
        reviewer = new_test_user(
            self.env, login='logistics-runtime-override-reviewer-' + unique_fixture('override'),
            context={'no_reset_password': True}, groups='insilos_logistics_idp.group_logistics_reviewer')
        manager = new_test_user(
            self.env, login='logistics-runtime-override-manager-' + unique_fixture('override'),
            context={'no_reset_password': True}, groups='insilos_logistics_idp.group_logistics_manager')
        check = self.env['logistics.idp.check.result']._controlled_create({
            'case_id': self.case.id, 'code': 'OVERRIDE_AUTHORIZATION', 'required': True, 'verdict': 'review',
            'rationale': 'synthetic override authorization', 'payload': {},
        }, 'runtime_override_authorization_test')
        exception = self.env['logistics.idp.exception'].create({
            'case_id': self.case.id, 'exception_type': 'missing_document', 'severity': 'medium'})
        with self.assertRaisesRegex(UserError, 'authorized reviewers'):
            self.case.with_user(self.operator).request_override(exception, 'evidence_gap', 'unauthorized', checks=check)
        override = self.case.with_user(reviewer).request_override(
            exception, 'evidence_gap', 'authorized justification', checks=check)
        self.assertEqual((override.state, override.requested_by, override.approved_by), ('requested', reviewer, False))
        self.assertEqual(check.verdict, 'review')
        with self.assertRaisesRegex(UserError, 'Logistics Managers'):
            override.with_user(reviewer).action_approve()
        self.assertTrue(override.with_user(manager).action_approve())
        self.assertEqual((override.state, override.approved_by), ('approved', manager))
        foreign = self.env['logistics.idp.case'].create({
            'name': 'Foreign override', 'source_system': 'synthetic', 'source_key': unique_fixture('override-foreign'),
            'source_version': 'v1', 'provenance': 'fixture', 'effective_date': '2026-01-01'})
        foreign_exception = self.env['logistics.idp.exception'].create({
            'case_id': foreign.id, 'exception_type': 'missing_document'})
        with self.assertRaisesRegex(ValidationError, 'same-case exceptions'):
            self.case.with_user(manager).request_override(foreign_exception, 'evidence_gap', 'cross case')
        with self.assertRaisesRegex(ValidationError, 'one or more same-case'):
            self.case.with_user(manager).request_override(self.env['logistics.idp.exception'], 'evidence_gap', 'empty')

    def test_broker_package_preserves_required_check_pass_after_override(self):
        manager = new_test_user(
            self.env, login='logistics-runtime-broker-manager-' + unique_fixture('broker'),
            context={'no_reset_password': True}, groups='insilos_logistics_idp.group_logistics_manager')
        config = json.loads(self.sap_profile.payload)
        config['broker_package']['required_checks'] = ['BROKER_CLEARANCE']
        self._write_sap_config(config)
        self.case.generate_customs_output('E13', [dict(self.po['lines'][0], quantity='1')])
        check = self.env['logistics.idp.check.result']._controlled_create({
            'case_id': self.case.id, 'code': 'BROKER_CLEARANCE', 'verdict': 'pass',
            'rationale': 'synthetic broker clearance', 'payload': {},
        }, 'runtime_broker_test')
        exception = self.env['logistics.idp.exception'].create({
            'case_id': self.case.id, 'exception_type': 'missing_document', 'severity': 'medium'})
        self.case.with_user(manager).request_override(
            exception, 'operational_exception', 'override does not replace clearance check', checks=check)
        package = self.case.generate_broker_package()
        self.assertEqual(package.status, 'generated')
        self.assertEqual(check.verdict, 'pass')

    def test_broker_package_generic_output_generation_is_rejected(self):
        manager = new_test_user(
            self.env, login='logistics-runtime-broker-output-manager-' + unique_fixture('broker'),
            context={'no_reset_password': True}, groups='insilos_logistics_idp.group_logistics_manager')
        values = {'entries': [], 'exclusions': []}
        output = self.env['logistics.idp.output']
        with self.assertRaisesRegex(ValidationError, 'case policy gate'):
            output.with_user(manager).generate(self.case, 'broker_package', values)
        with self.assertRaisesRegex(ValidationError, 'case policy gate'):
            output.sudo().generate(self.case, 'broker_package', values)
        self.assertFalse(self.case.output_ids.filtered(lambda item: item.output_type == 'broker_package'))
        self.case.generate_customs_output('E13', [dict(self.po['lines'][0], quantity='1')])
        self.assertEqual(self.case.generate_broker_package().status, 'generated')

    def test_fr804_native_broker_email_queue_is_configured_idempotent_and_auditable(self):
        manager = new_test_user(
            self.env, login='logistics-runtime-email-manager-' + unique_fixture('broker-email'),
            context={'no_reset_password': True}, groups='insilos_logistics_idp.group_logistics_manager')
        self.case.generate_customs_output('E13', [dict(self.po['lines'][0], quantity='1')])
        package = self.case.generate_broker_package()
        server = self.env['ir.mail_server'].create({
            'name': 'FR804 native sender', 'smtp_host': 'localhost', 'smtp_port': 25,
            'smtp_encryption': 'none', 'from_filter': 'sender@example.test',
        })
        template = self.env['mail.template'].create({
            'name': 'Broker Vietnamese package', 'model_id': self.env['ir.model']._get_id('logistics.idp.case'),
            'mail_server_id': server.id, 'use_default_to': False,
            'email_from': 'Người gửi <sender@example.test>',
            'subject': 'Hồ sơ môi giới {{ object.name }}',
            'body_html': '<p style="font-family: Arial,; font-size: 12pt;;">Xin chào<br/>Dòng tiếng Việt</p><table><tr><td><strong>Hồ sơ môi giới</strong></td></tr></table>',
        })
        config = json.loads(self.sap_profile.payload)
        config['broker_package']['email'] = {'template_id': template.id, 'to': ' Broker@Example.Test;broker@example.test ', 'cc': ['cc@example.test'], 'reply_to': 'reply@example.test'}
        self._write_sap_config(config)
        before = self.env['mail.mail'].sudo().search_count([])
        with self.assertRaisesRegex(UserError, 'Only Logistics Managers'):
            self.case.with_user(self.operator).queue_broker_email()
        self.assertEqual(self.env['mail.mail'].sudo().search_count([]), before)
        before_evidence = len(self.case.evidence_ids)
        before_snapshots = self.env['ir.attachment'].sudo().search_count([('mimetype', '=', 'message/rfc822')])
        with patch.object(type(server), '_build_email__', side_effect=RuntimeError('unexpected native MIME failure')):
            with self.assertRaisesRegex(RuntimeError, 'unexpected native MIME failure'):
                self.case.with_user(manager).queue_broker_email()
        self.assertEqual(self.env['mail.mail'].sudo().search_count([]), before)
        self.assertEqual(len(self.case.evidence_ids), before_evidence)
        self.assertEqual(self.env['ir.attachment'].sudo().search_count([('mimetype', '=', 'message/rfc822')]), before_snapshots)
        mail = self.case.with_user(manager).queue_broker_email()
        self.assertEqual((mail.state, mail.email_to, mail.email_cc, mail.reply_to, mail.auto_delete), ('outgoing', 'broker@example.test', 'cc@example.test', 'reply@example.test', False))
        self.assertIn(package.attachment_id, mail.attachment_ids)
        self.assertIn('<strong>Hồ sơ môi giới</strong>', mail.body_html)
        self.assertEqual(self.case.with_user(manager).queue_broker_email(), mail)
        request = self.case.evidence_ids.filtered(lambda item: item.category == 'broker_email_request')
        self.assertEqual(len(request), 1)
        payload = json.loads(request.payload)
        self.assertTrue(payload['mime']['available'])
        self.assertEqual(payload['mime']['charset'], 'utf-8')
        self.assertEqual(payload['mime']['attachment_sha256'], package.artifact_sha256)
        snapshot = self.env['ir.attachment'].browse(payload['mime']['attachment_id'])
        self.assertEqual((snapshot.res_model, snapshot.res_id, snapshot.mimetype), ('logistics.idp.case', self.case.id, 'message/rfc822'))
        self.assertNotIn(snapshot, mail.attachment_ids)
        self.assertEqual(set(mail.attachment_ids.mapped('mimetype')), {'application/zip'})
        raw = snapshot.raw
        parsed = BytesParser(policy=policy.default).parsebytes(raw)
        self.assertIn('=?utf-8?', raw.decode('ascii', 'ignore').lower())
        self.assertIn('Xin chào', parsed.get_body(('plain',)).get_content())
        self.assertIn('Dòng tiếng Việt', parsed.get_body(('html',)).get_content())
        self.assertEqual({'text/plain', 'text/html'}, {part.get_content_type() for part in parsed.walk() if part.get_content_type().startswith('text/')})
        self.assertEqual([part.get_content_type() for part in parsed.iter_attachments()], ['application/zip'])
        self.assertEqual(payload['mime']['parts'], ['%s/%s' % (part.get_content_maintype(), part.get_content_subtype()) for part in parsed.walk()])
        self.assertEqual(payload['mime']['sha256'], hashlib.sha256(raw).hexdigest())
        self.assertEqual(payload['mime']['size'], len(raw))
        self.assertFalse(self.case.with_user(manager).queue_broker_email().attachment_ids.filtered(lambda attachment: attachment.mimetype == 'message/rfc822'))
        self.assertEqual(len(self.case.evidence_ids.filtered(lambda item: item.category == 'broker_email_queued')), 1)
        self.assertTrue(self.case.message_ids.filtered(lambda message: 'Broker package email queued' in message.body))
        mail.sudo().write({'state': 'sent'})
        self.assertNotIn(snapshot, mail.attachment_ids)
        self.case.record_broker_email_delivery(mail)
        self.case.record_broker_email_delivery(mail)
        delivery = self.case.evidence_ids.filtered(lambda item: item.category == 'broker_email_delivery')
        self.assertEqual(len(delivery), 1)
        self.assertEqual(json.loads(delivery.payload)['state'], 'sent')
        secret = 'smtp://broker:injected-secret@internal.example.test recipient@example.test'
        mail.sudo().write({'state': 'exception', 'failure_type': 'mail_smtp', 'failure_reason': secret})
        self.case.record_broker_email_delivery(mail)
        self.case.record_broker_email_delivery(mail)
        failure = self.case.evidence_ids.filtered(lambda item: item.source_reference.endswith(':exception'))
        self.assertEqual(len(failure), 1)
        failure_payload = json.loads(failure.payload)
        self.assertEqual((failure_payload['failure_type'], failure_payload['failure_code']), ('mail_smtp', 'broker_delivery_failed'))
        self.assertNotIn('failure_reason', failure_payload)
        self.assertNotIn(secret, failure.payload)
        config['broker_package'].pop('email')
        self._write_sap_config(config)
        before = (self.env['mail.mail'].sudo().search_count([]), len(self.case.evidence_ids))
        with self.assertRaisesRegex(ValidationError, 'configuration is invalid'):
            self.case.with_user(manager).queue_broker_email()
        self.assertEqual((self.env['mail.mail'].sudo().search_count([]), len(self.case.evidence_ids)), before)

    def test_fr305_supplier_result_reply_is_configured_threaded_authorized_and_idempotent(self):
        manager = new_test_user(
            self.env, login='logistics-runtime-result-manager-' + unique_fixture('result-email'),
            context={'no_reset_password': True}, groups='insilos_logistics_idp.group_logistics_manager')
        template = self.env['mail.template'].create({
            'name': 'Draft result', 'model_id': self.env['ir.model']._get_id('logistics.idp.case'),
            'use_default_to': False, 'email_from': 'IDP <idp@example.test>',
            'subject': 'Draft result {{ object.name }}', 'body_html': '<p>Result</p>',
        })
        config = json.loads(self.sap_profile.payload)
        config['email_domains'] = ['example.test']
        config['draft_invoice'] = {'result_email': {'template_id': template.id, 'cc': ' Internal@Example.Test;internal@example.test '}}
        self._write_sap_config(config)
        before = (self.env['mail.mail'].sudo().search_count([]), len(self.case.evidence_ids))
        with self.assertRaisesRegex(UserError, 'Only Logistics Managers'):
            self.case.with_user(self.operator).queue_supplier_result_reply()
        with self.assertRaisesRegex(ValidationError, 'trusted supplier-email draft invoice source'):
            self.case.with_user(manager).queue_supplier_result_reply()
        self.assertEqual((self.env['mail.mail'].sudo().search_count([]), len(self.case.evidence_ids)), before)
        caller_document = self.env['logistics.idp.document'].intake_content(
            self.case, b'draft-result-untrusted', 'application/pdf', {
                'filename': 'untrusted.pdf', 'source_sender': 'Supplier@Example.Test',
                'source_message_reference': '<supplier-thread@example.test>',
            }, source_channel='queue', process=False)
        with self.assertRaisesRegex(ValidationError, 'trusted supplier-email draft invoice'):
            self.case.with_user(manager).queue_supplier_result_reply()
        self.assertFalse(self.case.evidence_ids.filtered(lambda item: item.category.startswith('supplier_result_reply_')))
        synthetic = self.env['logistics.idp.document']._intake_synthetic_content(
            self.case, b'draft-result-source', 'application/pdf', {
                'filename': 'draft.pdf', 'document_type': 'draft_vat_invoice',
                'source_sender': 'Supplier@Example.Test', 'source_message_reference': '<supplier-thread@example.test>',
            }, source_channel='queue', process=False)
        self.env['logistics.idp.evidence']._controlled_create({
            'case_id': self.case.id, 'category': 'email_attachment', 'source_reference': 'forged-draft-result-source',
            'status': 'valid', 'payload': {'content_hash': synthetic.content_hash,
                                            'audit_service': 'intake_supplier_email', 'source': 'intake_supplier_email'},
        }, 'supplier_email_intake')
        before = (self.env['mail.mail'].sudo().search_count([]), len(self.case.evidence_ids))
        with self.assertRaisesRegex(ValidationError, 'trusted supplier-email draft invoice'):
            self.case.with_user(manager).queue_supplier_result_reply()
        self.assertEqual((self.env['mail.mail'].sudo().search_count([]), len(self.case.evidence_ids)), before)
        job = self.env['logistics.idp.inbound.job'].intake_supplier_email({
            'message_id': '<supplier-thread@example.test>', 'source_sender': 'Supplier@Example.Test',
            'sender_domain': 'example.test', 'company_id': self.env.company.id,
        }, [{'filename': 'draft vat invoice.pdf', 'mimetype': 'application/pdf',
             'content': b'%PDF-1.4\n/Type /Page\n%%EOF'}])
        self.assertTrue(job._process_one())
        reply_case = job.case_id
        document = reply_case.document_ids
        self.assertEqual(document.document_type, 'unknown')
        with patch('odoo.addons.insilos_logistics_idp.models.logistics_idp.IAPDocumentProcessor.process',
                   return_value={'document_type': 'draft_vat_invoice', 'payload': {}, 'confidence': 1,
                                 'classification_confidence': 1, 'validation_warnings': [], 'source_spans': []}):
            document.with_user(self.operator).process(explicit=True)
        self.assertEqual(document.document_type, 'draft_vat_invoice')
        intake_evidence = reply_case.evidence_ids.filtered(lambda item: item.category == 'email_attachment')
        self.assertEqual(json.loads(intake_evidence.payload)['audit_service'], 'intake_supplier_email')
        mail = reply_case.with_user(manager).queue_supplier_result_reply()
        self.assertEqual((mail.state, mail.email_to, mail.email_cc, mail.auto_delete),
                         ('outgoing', 'supplier@example.test', 'internal@example.test', False))
        self.assertEqual(ast.literal_eval(mail.headers)['In-Reply-To'], '<supplier-thread@example.test>')
        self.assertEqual(reply_case.with_user(manager).queue_supplier_result_reply(), mail)
        server = self.env['ir.mail_server']
        with patch.object(type(server), '_disable_send', return_value=False), \
                patch.object(type(server), '_connect__', return_value=MagicMock()), \
                patch.object(type(server), 'send_email', return_value='<sent@example.test>'):
            mail.send()
        self.assertEqual(mail.state, 'sent')
        delivery = reply_case.evidence_ids.filtered(lambda item: item.category == 'supplier_result_reply_delivery')
        self.assertEqual((len(delivery), json.loads(delivery.payload)['state']), (1, 'sent'))
        secret = 'smtp://supplier:injected-secret@internal.example.test'
        mail.sudo().write({'state': 'outgoing'})
        with patch.object(type(server), '_disable_send', return_value=False), \
                patch.object(type(server), '_connect__', return_value=MagicMock()), \
                patch.object(type(server), 'send_email', side_effect=RuntimeError(secret)):
            mail.send()
        self.assertEqual(mail.state, 'exception')
        delivery = reply_case.evidence_ids.filtered(lambda item: item.category == 'supplier_result_reply_delivery')
        self.assertEqual(len(delivery), 2)
        failure = delivery.filtered(lambda item: item.source_reference.endswith(':exception'))
        self.assertEqual((len(failure), json.loads(failure.payload)['failure_type'], json.loads(failure.payload)['failure_code']),
                         (1, 'unknown', 'supplier_result_reply_delivery_failed'))
        self.assertNotIn('failure_reason', failure.payload)
        self.assertNotIn(secret, failure.payload)
        self.assertEqual(reply_case.with_user(manager).queue_supplier_result_reply(), mail)
        config['draft_invoice']['result_email']['cc'] = 'bad@example.test\r\nBcc: injected@example.test'
        self._write_sap_config(config)
        with self.assertRaisesRegex(ValidationError, 'safe valid list'):
            self.case.with_user(manager).queue_supplier_result_reply()
        config.pop('draft_invoice')
        self._write_sap_config(config)
        with self.assertRaisesRegex(ValidationError, 'configuration is invalid'):
            self.case.with_user(manager).queue_supplier_result_reply()

    def test_fr802_broker_package_archives_qdtq_bytes_and_generated_output(self):
        qdtq = self.env['logistics.idp.document']._intake_synthetic_content(
            self.case, b'qdtq-bytes', 'application/pdf', {'filename': 'qdtq.pdf', 'document_type': 'customs_declaration'}, process=False)
        non_qdtq = self.env['logistics.idp.document']._intake_synthetic_content(
            self.case, b'invoice-bytes', 'application/pdf', {'filename': 'invoice.pdf', 'document_type': 'invoice'}, process=False)
        e13 = self.case.generate_customs_output('E13', [dict(self.po['lines'][0], quantity='1')])
        package = self.case.generate_broker_package()
        with zipfile.ZipFile(io.BytesIO(package.attachment_id.raw)) as archive:
            manifest = json.loads(archive.read('manifest.json'))
            names = archive.namelist()
            self.assertEqual(archive.read(next(name for name in names if name.startswith('documents/%s-' % qdtq.id))), b'qdtq-bytes')
            self.assertEqual(archive.read(next(name for name in names if name.startswith('outputs/e13-'))), e13.attachment_id.raw)
        self.assertNotIn(non_qdtq.id, [entry['source_record_id'] for entry in manifest['entries']])
        self.assertEqual(manifest['profile'], {'id': self.sap_profile.id, 'version': '1', 'hash': self.sap_profile.payload_hash})
        self.assertEqual(manifest['exclusions'][0]['rule'], 'exclude_document_types_without_qdtq')

    def test_fr803_broker_package_idempotent_and_supersedes_changed_input(self):
        self.case.generate_customs_output('E13', [dict(self.po['lines'][0], quantity='1')])
        first = self.case.generate_broker_package()
        self.assertEqual(self.case.generate_broker_package(), first)
        self.env['logistics.idp.document']._intake_synthetic_content(
            self.case, b'new-qdtq-bytes', 'application/pdf', {'filename': 'new-qdtq.pdf', 'document_type': 'customs_declaration'}, process=False)
        second = self.case.generate_broker_package()
        self.assertEqual((second.version, second.supersedes_id), (2, first))
        self.assertEqual(first.status, 'superseded')
        profile = json.loads(self.sap_profile.payload)
        profile['broker_package']['qdtq_document_types'].append('invoice')
        payload = json.dumps(profile, sort_keys=True, separators=(',', ':'))
        self.cr.execute('UPDATE logistics_idp_supplier_profile SET payload = %s, payload_hash = %s WHERE id = %s',
                        [payload, hashlib.sha256(payload.encode()).hexdigest(), self.sap_profile.id])
        self.sap_profile.invalidate_recordset(['payload', 'payload_hash'])
        third = self.case.generate_broker_package()
        self.assertEqual((third.version, third.supersedes_id, second.status), (3, second, 'superseded'))

    def _configure_local_direct_outputs(self, gate_enabled=True, declaration_optional=True):
        stream = io.BytesIO()
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.title = 'Local'
        workbook.active['A1'] = 'template'
        for position in range(1, 5):
            workbook.active.cell(position, 1).number_format = '@'
        workbook.save(stream)
        template = self.env['ir.attachment'].create({
            'name': 'local-direct.xlsx', 'raw': stream.getvalue(),
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            'company_id': self.case.company_id.id,
        })
        payload = json.loads(self.sap_profile.payload)
        common = {'mapping_version': 'local-v1', 'template_attachment_id': template.id, 'sheet': 'Local',
                  'logical_key': 'case.source_key', 'declaration_optional_pre_customs': declaration_optional,
                  'fields': [
                      {'name': name, 'source': source, 'position': position,
                       'required': name != 'declaration_number' or not declaration_optional}
                      for position, (name, source) in enumerate((
                          ('sequence', 'case.source_key'), ('delivery_date', 'case.effective_date'),
                          ('declaration_number', 'declaration.declaration_number'), ('po_number', 'sap.po_reference'),
                          ('invoice_number', 'sap.invoice_number'), ('bill_number', 'sap.bill_number'),
                          ('supplier', 'sap.supplier_reference'), ('goods_description', 'invoice.goods_description'),
                          ('package_count', 'invoice.package_count'), ('gross_weight', 'invoice.gross_weight'),
                          ('net_weight', 'invoice.net_weight'), ('carrier', 'invoice.carrier'), ('note', 'invoice.note'),
                      ), 1)]}
        payload['shipping_plan'] = common
        payload['gate_pass'] = {**common, 'enabled': gate_enabled, 'requires_validated_declaration': True, 'required_checks': [],
                                'fields': [{'name': name, 'source': 'shipping_plan.logical_key', 'position': position, 'required': True}
                                           for position, name in enumerate(('number', 'validity', 'vehicle_or_carrier', 'approval_or_signature'), 1)]}
        canonical = json.dumps(payload, sort_keys=True, separators=(',', ':'))
        self.cr.execute('UPDATE logistics_idp_supplier_profile SET payload = %s, payload_hash = %s WHERE id = %s',
                        [canonical, hashlib.sha256(canonical.encode()).hexdigest(), self.sap_profile.id])
        self.sap_profile.invalidate_recordset(['payload', 'payload_hash'])
        self._sap_sources([{'line_key': 'L1', 'material_code': 'MAT-1'}],
                          [{'line_key': 'L1', 'quantity': '1', 'unit_price': '2'}])
        return template, self.case.generate_sap_erp_output()

    def _valid_local_declaration(self):
        self.case.sudo().owner_id = self.operator
        line = {'material_code': 'MAT-1', 'custom_code': 'PO-1-10', 'hs_code': '711311', 'quantity': '1',
                'uom': 'EA', 'line_total': '2', 'supplier': 'SUP-1', 'importer': 'IMP-1', 'currency': 'USD'}
        self._retag_authoritative_policy_rules()
        self.case.generate_customs_output('E13', [line], 'local-v1')
        self._retag_authoritative_policy_rules()
        self.case.generate_customs_output('E15', [line], 'local-v1')
        self._retag_authoritative_policy_rules()
        self.case.generate_customs_output('E13', [line], 'local-v2')
        self.case.generate_customs_output('E15', [line], 'local-v2')
        document = self.env['logistics.idp.document'].with_user(self.operator)._intake_synthetic_content(
            self.case, json.dumps({'document_type': 'import_declaration', 'confidence': 1, 'payload': {
                'declaration_number': 'TK-LOCAL', 'supplier': 'SUP-1', 'importer': 'IMP-1', 'currency': 'USD',
                'total_invoice_value': '4', 'lines': [{**line, 'regime': 'E13'}, {**line, 'regime': 'E15'}],
            }}, sort_keys=True).encode(), 'application/json',
            {'document_type': 'import_declaration', 'confidence': 1})
        self._retag_authoritative_policy_rules()
        self.case.generate_customs_output('E13', [line], 'local-v3')
        self.case.generate_customs_output('E15', [line], 'local-v3')
        declaration = self.case.reconcile_import_declaration(document)
        self.assertEqual(declaration.status, 'valid')
        reconciliation = self.env['logistics.idp.evidence']._controlled_create({
            'case_id': self.case.id, 'category': 'reconciliation', 'source_reference': 'TK-LOCAL', 'status': 'valid', 'payload': {},
        }, 'runtime_test')
        return self.env['logistics.idp.evidence']._controlled_create({
            'case_id': self.case.id, 'category': 'import_declaration', 'source_reference': 'TK-LOCAL', 'status': 'valid',
            'payload': {'declaration': {'declaration_number': 'TK-LOCAL'}, 'reconciliation_hash': reconciliation.payload_hash},
        }, 'runtime_test')

    def test_fr1000_pre_customs_generation(self):
        _template, sap = self._configure_local_direct_outputs()
        plan = self.case.generate_shipping_plan()
        self.assertEqual(json.loads(plan.payload)['source_records'][0], {
            'model': 'logistics.idp.output', 'id': sap.id, 'hash': sap.payload_hash,
        })

    def test_fr1002_full_mapping_template_values(self):
        template, sap = self._configure_local_direct_outputs()
        plan = self.case.generate_shipping_plan()
        payload = json.loads(plan.payload)
        names = ('sequence', 'delivery_date', 'declaration_number', 'po_number', 'invoice_number',
                 'bill_number', 'supplier', 'goods_description', 'package_count', 'gross_weight',
                 'net_weight', 'carrier', 'note')
        expected = {
            'sequence': self.case.source_key, 'delivery_date': '2026-01-01', 'declaration_number': None,
            'po_number': 'PO-1', 'invoice_number': 'INV-LOCAL', 'bill_number': 'BILL-LOCAL',
            'supplier': 'SUP-1', 'goods_description': 'Invoice-described goods', 'package_count': '3',
            'gross_weight': '12.5', 'net_weight': '10', 'carrier': 'Trusted Carrier', 'note': 'Invoice note',
        }
        self.assertEqual(payload['template_attachment_id'], template.id)
        self.assertEqual(payload['mapped_fields'], expected)
        self.assertEqual(payload['field_provenance']['goods_description']['record_type'], 'logistics.idp.evidence')
        workbook = load_workbook(io.BytesIO(plan.attachment_id.raw))
        self.assertEqual([workbook['Local'].cell(position, 1).value for position in range(1, 14)], list(expected.values()))
        self.assertEqual(payload['source_records'][0]['id'], sap.id)

    def test_fr1003_trusted_provenance_late_declaration_supersession_idempotency(self):
        self._configure_local_direct_outputs()
        config = json.loads(self.sap_profile.payload)
        config['shipping_plan']['fields'][2]['source'] = 'declaration.declaration_number'
        self._write_sap_config(config)
        first = self.case.generate_shipping_plan()
        declaration = self._valid_local_declaration()
        second = self.case.generate_shipping_plan()
        self.assertEqual((second.version, second.supersedes_id, first.status), (2, first, 'superseded'))
        payload = json.loads(second.payload)
        self.assertEqual(payload['mapped_fields']['declaration_number'], 'TK-LOCAL')
        self.assertEqual(payload['field_provenance']['declaration_number'], {
            'mapping_version': 'local-v1', 'source': 'declaration.declaration_number',
            'record_type': 'logistics.idp.evidence', 'record_id': declaration.id,
            'record_hash': declaration.payload_hash,
        })
        self.assertEqual(self.case.generate_shipping_plan(), second)

    def test_fr1102_controlled_manager_approval_required_checks_current_override(self):
        self._configure_local_direct_outputs()
        config = json.loads(self.sap_profile.payload)
        config['gate_pass']['required_checks'] = ['GATE_CLEARANCE']
        self._write_sap_config(config)
        self._retag_authoritative_policy_rules()
        plan = self.case.generate_shipping_plan()
        manager = new_test_user(
            self.env, login='logistics-runtime-gate-manager-' + unique_fixture('gate'),
            context={'no_reset_password': True}, groups='insilos_logistics_idp.group_logistics_manager')
        approval = plan.with_user(manager).action_approve_shipping_plan()
        self.assertEqual((approval.code, approval.verdict), ('shipping_plan_approval', 'pass'))
        replacement = self._valid_local_declaration()
        current = self.case.generate_shipping_plan()
        self.assertEqual((current.version, current.supersedes_id, plan.status), (2, plan, 'superseded'))
        with self.assertRaisesRegex(ValidationError, 'enabled current approved Shipping Plan'):
            self.case.generate_gate_pass()
        current.with_user(manager).action_approve_shipping_plan()
        self.assertEqual(self.case.generate_shipping_plan(), current)
        self.assertTrue(replacement)
        with self.assertRaisesRegex(ValidationError, 'valid required checks GATE_CLEARANCE'):
            self.case.generate_gate_pass()
        check = self.env['logistics.idp.check.result']._controlled_create({
            'case_id': self.case.id, 'code': 'GATE_CLEARANCE', 'verdict': 'pass',
            'rationale': 'synthetic gate clearance', 'payload': {},
        }, 'runtime_gate_test')
        self.assertTrue(self.case.generate_gate_pass())
        exception = self.env['logistics.idp.exception'].create({
            'case_id': self.case.id, 'exception_type': 'missing_document', 'severity': 'medium'})
        self.case.with_user(manager).request_override(
            exception, 'operational_exception', 'override does not replace gate clearance', checks=check)
        self.assertTrue(self.case.generate_gate_pass())

    def test_fr1103_gate_template_semantic_fields(self):
        _template, _sap = self._configure_local_direct_outputs()
        declaration = self._valid_local_declaration()
        plan = self.case.generate_shipping_plan()
        manager = new_test_user(
            self.env, login='logistics-runtime-gate-template-manager-' + unique_fixture('gate'),
            context={'no_reset_password': True}, groups='insilos_logistics_idp.group_logistics_manager')
        plan.with_user(manager).action_approve_shipping_plan()
        expected = {
            'number': 'PO-1', 'validity': '2026-01-01',
            'vehicle_or_carrier': 'Trusted Carrier', 'approval_or_signature': 'Invoice note',
        }
        config = json.loads(self.sap_profile.payload)
        config['gate_pass'].update({'mapping_version': 'gate-v1', 'fields': [
            {'name': name, 'source': source, 'position': position, 'required': True}
            for name, source, position in (
                ('number', 'shipping_plan.mapped_fields.po_number', 1),
                ('validity', 'shipping_plan.mapped_fields.delivery_date', 2),
                ('vehicle_or_carrier', 'shipping_plan.mapped_fields.carrier', 3),
                ('approval_or_signature', 'shipping_plan.mapped_fields.note', 4),
            )
        ]})
        self._write_sap_config(config)
        first = self.case.generate_gate_pass()
        payload = json.loads(first.payload)
        self.assertEqual(payload['mapped_fields'], expected)
        self.assertEqual(payload['source_records'][-1], {
            'model': 'logistics.idp.output', 'id': plan.id, 'hash': plan.payload_hash,
        })
        for name, source in ((
                'number', 'shipping_plan.mapped_fields.po_number'),
                ('validity', 'shipping_plan.mapped_fields.delivery_date'),
                ('vehicle_or_carrier', 'shipping_plan.mapped_fields.carrier'),
                ('approval_or_signature', 'shipping_plan.mapped_fields.note')):
            self.assertEqual(payload['field_provenance'][name], {
                'mapping_version': 'gate-v1', 'source': source,
                'record_type': 'logistics.idp.output', 'record_id': plan.id, 'record_hash': plan.payload_hash,
            })
        workbook = load_workbook(io.BytesIO(first.attachment_id.raw))
        self.assertEqual([workbook['Local'].cell(position, 1).value for position in range(1, 5)], list(expected.values()))
        self.assertEqual([workbook['Local'].cell(position, 1).number_format for position in range(1, 5)], ['@'] * 4)
        self.assertTrue(declaration)
        config['gate_pass'].update({'mapping_version': 'gate-v2', 'fields': [
            {**field, 'position': 5 - field['position']} for field in config['gate_pass']['fields']
        ]})
        self._write_sap_config(config)
        second = self.case.generate_gate_pass()
        self.assertEqual((second.version, second.supersedes_id, first.status), (2, first, 'superseded'))
        payload = json.loads(second.payload)
        self.assertEqual(payload['mapped_fields'], expected)
        self.assertEqual([field['position'] for field in payload['fields']], [4, 3, 2, 1])
        self.assertTrue(all(item['mapping_version'] == 'gate-v2' for item in payload['field_provenance'].values()))
        workbook = load_workbook(io.BytesIO(second.attachment_id.raw))
        self.assertEqual([workbook['Local'].cell(position, 1).value for position in range(1, 5)], list(reversed(list(expected.values()))))
        self.assertEqual([workbook['Local'].cell(position, 1).number_format for position in range(1, 5)], ['@'] * 4)
        config['gate_pass']['fields'].pop()
        self._write_sap_config(config)
        with self.assertRaisesRegex(ValidationError, 'Gate Pass semantic required fields'):
            self.case.generate_gate_pass()

    def test_fr1104_native_gate_pass_distribution_is_manager_scoped_configured_and_versioned(self):
        self._configure_local_direct_outputs()
        self._valid_local_declaration()
        manager = new_test_user(
            self.env, login='logistics-runtime-gate-distribution-manager-' + unique_fixture('gate'),
            context={'no_reset_password': True}, groups='insilos_logistics_idp.group_logistics_manager')
        plan = self.case.generate_shipping_plan()
        plan.with_user(manager).action_approve_shipping_plan()
        gate_pass = self.case.generate_gate_pass()
        template = self.env['mail.template'].create({
            'name': 'Gate Pass distribution', 'model_id': self.env['ir.model']._get_id('logistics.idp.case'),
            'use_default_to': False, 'email_from': 'Gate <gate@example.test>',
            'subject': 'Gate Pass {{ object.name }}', 'body_html': '<p>Gate Pass</p>',
        })
        config = json.loads(self.sap_profile.payload)
        config['gate_pass']['distribution'] = {
            'channel': 'native_mail', 'template_id': template.id,
            'to': ' Gate@Example.Test;gate@example.test ',
        }
        self._write_sap_config(config)
        before = self.env['mail.mail'].sudo().search_count([])
        with self.assertRaisesRegex(UserError, 'Only Logistics Managers'):
            self.case.with_user(self.operator).queue_gate_pass_distribution()
        self.assertEqual(self.env['mail.mail'].sudo().search_count([]), before)
        mail = self.case.with_user(manager).queue_gate_pass_distribution()
        self.assertEqual((mail.state, mail.email_to, mail.auto_delete), ('outgoing', 'gate@example.test', False))
        self.assertIn(gate_pass.attachment_id, mail.attachment_ids)
        self.assertEqual(self.case.with_user(manager).queue_gate_pass_distribution(), mail)
        request = self.case.evidence_ids.filtered(lambda item: item.category == 'gate_pass_distribution_request')
        self.assertEqual(len(request), 1)
        payload = json.loads(request.payload)
        self.assertEqual((payload['gate_pass'], payload['shipping_plan'], payload['recipients'], payload['channel']), (
            {'id': gate_pass.id, 'hash': gate_pass.artifact_sha256, 'version': gate_pass.version},
            {'id': plan.id, 'hash': plan.payload_hash}, 'gate@example.test', 'native_mail'))
        config['gate_pass']['distribution']['to'] = ''
        self._write_sap_config(config)
        with self.assertRaisesRegex(ValidationError, 'recipients'):
            self.case.with_user(manager).queue_gate_pass_distribution()
        config['gate_pass']['distribution']['to'] = 'gate@example.test'
        self._write_sap_config(config)
        duplicate = self.env['logistics.idp.supplier.profile'].import_upsert({
            'supplier_reference': self.case.supplier_reference, 'profile_code': 'GATE-DUP', 'version': '1',
            'source_system': 'fixture', 'source_key': unique_fixture('gate-duplicate'), 'source_version': 'v1',
            'provenance': 'fixture', 'effective_from': self.case.effective_date, 'payload': config,
        })
        with self.assertRaisesRegex(ValidationError, 'one effective same-company'):
            self.case.with_user(manager).queue_gate_pass_distribution()
        self.cr.execute('UPDATE logistics_idp_supplier_profile SET effective_to = %s WHERE id = %s',
                        ['2025-12-31', duplicate.id])
        config['gate_pass']['mapping_version'] = 'gate-distribution-v2'
        self._write_sap_config(config)
        newer = self.case.generate_gate_pass()
        self.assertEqual((newer.version, gate_pass.status), (2, 'superseded'))
        fresh = self.case.with_user(manager).queue_gate_pass_distribution()
        self.assertNotEqual(fresh, mail)
        self.assertEqual(len(self.case.evidence_ids.filtered(lambda item: item.category == 'gate_pass_distribution_queued')), 2)

    def test_shipping_contract_rejects_incomplete_unknown_source(self):
        self._configure_local_direct_outputs()
        with self.assertRaisesRegex(ValidationError, 'no caller-provided'):
            self.case.generate_shipping_plan({'source': 'caller'})
        config = json.loads(self.sap_profile.payload)
        config['shipping_plan']['fields'][0]['source'] = 'case.env'
        self._write_sap_config(config)
        with self.assertRaisesRegex(ValidationError, 'not allowed'):
            self.case.generate_shipping_plan()
        config['shipping_plan']['fields'][0]['source'] = 'case.source_key'
        config['shipping_plan']['fields'].pop()
        self._write_sap_config(config)
        with self.assertRaisesRegex(ValidationError, 'semantic required fields'):
            self.case.generate_shipping_plan()

    def test_unauthorized_normal_user_approval(self):
        self._configure_local_direct_outputs()
        plan = self.case.generate_shipping_plan()
        with self.assertRaisesRegex(UserError, 'Only Logistics Managers'):
            plan.with_user(self.operator).action_approve_shipping_plan()

    def test_reconciliation_is_deterministic(self):
        self.case.reconcile()
        self.assertEqual(self.case.verdict, 'review')
        self.env['logistics.idp.evidence']._controlled_create({
            'case_id': self.case.id, 'category': 'invoice', 'source_reference': 'ATT-3', 'status': 'valid', 'payload': '{}',
        }, 'runtime_test')
        self.case.reconcile()
        self.assertEqual(self.case.verdict, 'pass')

    def test_auditor_cannot_create_or_modify(self):
        auditor = new_test_user(self.env, login='logistics-auditor', groups='insilos_logistics_idp.group_logistics_auditor')
        with self.assertRaises(UserError):
            self.case.with_user(auditor).write({'name': 'Forbidden'})
        with self.assertRaises(UserError):
            self.env['logistics.idp.evidence'].with_user(auditor).create({
                'case_id': self.case.id, 'category': 'invoice', 'source_reference': 'ATT-4', 'payload': '{}',
            })

    def test_inbound_job_retries_then_dead_letters_with_observability(self):
        job = self.env['logistics.idp.inbound.job'].create({
            'profile_code': 'local-test', 'source_system': 'synthetic-erp', 'source_key': unique_fixture('INVALID-1'),
            'source_version': 'v1', 'payload': '{}', 'max_attempts': 2,
        })
        mail_count = self.env['mail.mail'].sudo().search_count([])
        template = type(self.env['mail.template'])
        mail = type(self.env['mail.mail'])
        with patch.object(template, 'send_mail', autospec=True) as send_mail, \
                patch.object(mail, 'create', autospec=True) as create_mail:
            self.assertFalse(job._process_one())
            self.assertEqual((job.state, job.attempt_count), ('retry', 1))
            self.assertTrue(job.last_attempt_at and job.next_attempt_at > job.last_attempt_at)
            self.assertEqual(job.last_error, 'INBOUND_FAILURE:ValidationError')
            self.assertFalse(job._process_one())
        send_mail.assert_not_called()
        create_mail.assert_not_called()
        self.assertEqual((job.state, job.attempt_count), ('dead', 2))
        self.assertEqual(self.env['mail.mail'].sudo().search_count([]), mail_count)

    def test_inbound_terminal_failure_is_safe_immutable_idempotent_and_quiet(self):
        secret = 'top-secret-inbound-payload'
        job = self.env['logistics.idp.inbound.job'].intake_supplier_email({
            'message_id': '<terminal-%s@example.test>' % unique_fixture('inbound'),
            'company_id': self.env.company.id,
        }, [
            {'filename': 'secret-0.pdf', 'mimetype': 'application/pdf', 'content': b'%PDF-1.4'},
            {'filename': 'secret-1.pdf', 'mimetype': 'application/pdf', 'content': b'%PDF-1.4'},
        ])
        case = self.env['logistics.idp.case'].create({
            'name': 'Failure case', 'provenance': 'fixture:development', 'effective_date': '2026-01-01',
            'source_system': job.source_system, 'source_key': job.source_key, 'source_version': job.source_version,
        })
        attachments = job.attachment_ids
        job.write({'payload': '{"secret":"%s"' % secret, 'case_id': case.id, 'max_attempts': 2})
        mail_count = self.env['mail.mail'].sudo().search_count([])
        template = type(self.env['mail.template'])
        mail = type(self.env['mail.mail'])
        with patch.object(template, 'send_mail', autospec=True) as send_mail, \
                patch.object(mail, 'create', autospec=True) as create_mail:
            self.assertFalse(job._process_one())
            self.assertEqual((job.state, job.attempt_count, job.last_error), ('retry', 1, 'INBOUND_FAILURE:JSONDecodeError'))
            self.assertEqual(job.case_id, case)
            self.assertFalse(job._process_one())
            self.assertEqual((job.state, job.attempt_count, job.last_error), ('dead', 2, 'INBOUND_FAILURE:JSONDecodeError'))
            case = job.case_id
            evidence = case.evidence_ids.filtered(lambda item: item.category == 'inbound_failure')
            self.assertEqual(len(evidence), 2)
            self.assertEqual({item.source_reference for item in evidence}, {'%s:%s' % (job.source_key, index) for index in range(2)})
            self.assertEqual({item.status for item in evidence}, {'invalid'})
            self.assertEqual({json.loads(item.payload)['status'] for item in evidence}, {'stored'})
            self.assertEqual({json.loads(item.payload)['manifest_identity'] for item in evidence}, {
                hashlib.sha256(attachment.raw).hexdigest() for attachment in attachments
            })
            with self.assertRaises(UserError):
                evidence[0].write({'status': 'valid'})
            self.assertEqual(self.env['logistics.idp.exception'].search_count([
                ('case_id', '=', case.id), ('exception_type', '=', 'integration_failure'), ('state', '=', 'open')]), 1)
            self.assertEqual(self.env['mail.activity'].search_count([
                ('res_model', '=', case._name), ('res_id', '=', case.id),
                ('activity_type_id', '=', self.env.ref('mail.mail_activity_data_todo').id)]), 1)
            self.assertFalse(job._process_one())
        send_mail.assert_not_called()
        create_mail.assert_not_called()
        self.assertEqual(self.env['mail.mail'].sudo().search_count([]), mail_count)
        self.assertNotIn(secret, job.last_error + ''.join(evidence.mapped('payload')) + case.next_action)
        self.assertEqual(len(case.evidence_ids.filtered(lambda item: item.category == 'inbound_failure')), 2)

    def test_inbound_job_success_records_case_and_completion(self):
        job = self.env['logistics.idp.inbound.job'].create({
            'profile_code': 'local-test', 'source_system': 'synthetic-erp', 'source_key': 'VALID-1',
            'source_version': 'v1', 'payload': json.dumps({
                'name': 'SYNTHETIC-JOB-1', 'provenance': 'fixture:development', 'effective_date': '2026-01-01',
            }),
        })
        self.assertTrue(job._process_one())
        self.assertEqual((job.state, job.attempt_count), ('done', 1))
        self.assertTrue(job.case_id and job.completed_at)
        self.assertFalse(job.last_error)

    def test_inbound_cron_isolates_failure_and_processes_healthy_peer(self):
        jobs = self.env['logistics.idp.inbound.job'].create([{
            'profile_code': 'local-test', 'source_system': 'synthetic-erp',
            'source_key': unique_fixture('INVALID-BATCH'), 'source_version': 'v1', 'payload': '{}',
            'next_attempt_at': '2000-01-01 00:00:00',
        }, {
            'profile_code': 'local-test', 'source_system': 'synthetic-erp',
            'source_key': unique_fixture('VALID-BATCH'), 'source_version': 'v1',
            'next_attempt_at': '2000-01-01 00:00:01',
            'payload': json.dumps({
                'name': 'SYNTHETIC-BATCH-PEER', 'provenance': 'fixture:development',
                'effective_date': '2026-01-01',
            }),
        }])
        jobs._cron_process(limit=2)
        self.assertEqual(jobs.mapped('state'), ['retry', 'done'])
        self.assertFalse(jobs[0].case_id)
        self.assertTrue(jobs[1].case_id)

    def test_exc001_canonical_types_are_company_scoped_and_idempotent(self):
        exceptions = self.env['logistics.idp.exception']
        expected = {
            'missing_document', 'unclassified_document', 'extraction_failure', 'low_confidence',
            'ambiguous_case', 'ambiguous_line_match', 'supplier_mismatch', 'quantity_mismatch',
            'price_mismatch', 'currency_mismatch', 'duplicate_invoice', 'material_code_invalid_missing',
            'hs_code_mismatch', 'customs_regime_mismatch', 'missing_master_data_dsnvl',
            'restricted_party_hit',
            'output_template_failure', 'integration_failure', 'sla_overdue',
        }
        self.assertEqual(set(exceptions._CONDITION_TYPES), expected)
        self.assertEqual({key for key, _label in exceptions._fields['exception_type'].selection}, expected)
        created = {condition: exceptions.open_or_reuse(self.case, condition) for condition in expected}
        sibling = self.env['logistics.idp.case'].create({
            'name': 'EXC001 sibling', 'company_id': self.case.company_id.id,
            'source_system': 'synthetic-erp', 'source_key': unique_fixture('exc001-sibling'),
            'source_version': 'v1', 'provenance': 'fixture:development', 'effective_date': '2026-01-01',
        })
        company = self.env['res.company'].create({'name': 'EXC001 other company'})
        foreign = self.env['logistics.idp.case'].create({
            'name': 'EXC001 foreign', 'company_id': company.id,
            'source_system': 'synthetic-erp', 'source_key': unique_fixture('exc001-foreign'),
            'source_version': 'v1', 'provenance': 'fixture:development', 'effective_date': '2026-01-01',
        })
        sibling_created = {condition: exceptions.open_or_reuse(sibling, condition) for condition in expected}
        foreign_created = {condition: exceptions.open_or_reuse(foreign, condition) for condition in expected}
        self.assertEqual({exception.exception_type for exception in created.values()}, expected)
        self.assertEqual({exception.case_id for exception in created.values()}, {self.case})
        self.assertEqual({exception.company_id for exception in created.values()}, {self.case.company_id})
        self.assertEqual({exception.case_id for exception in sibling_created.values()}, {sibling})
        self.assertEqual({exception.company_id for exception in sibling_created.values()}, {self.case.company_id})
        self.assertEqual({exception.case_id for exception in foreign_created.values()}, {foreign})
        self.assertEqual({exception.company_id for exception in foreign_created.values()}, {company})
        self.assertEqual(
            {condition: exceptions.open_or_reuse(self.case, condition).id for condition in expected},
            {condition: exception.id for condition, exception in created.items()},
        )
        self.assertTrue(all(created[condition] != sibling_created[condition] != foreign_created[condition]
                            for condition in expected))
        self.assertFalse(expected & {'policy', 'supplier', 'unmatched', 'inbound_failure'})

    def test_mes_snapshot_has_no_outbound_operation(self):
        self.assertFalse({'write_back', 'execute', 'plan'} & set(self.env['logistics.idp.mes.reference']._fields))

    def test_exception_sla_skips_weekend_and_company_holiday(self):
        calendar = self.env.company.resource_calendar_id
        self.env['resource.calendar.leaves'].create({
            'name': 'Synthetic Monday holiday', 'calendar_id': calendar.id,
            'date_from': '2026-01-05 00:00:00', 'date_to': '2026-01-05 23:59:59',
        })
        with patch('odoo.fields.Datetime.now', return_value='2026-01-02 16:00:00'):
            exception = self.env['logistics.idp.exception'].create({
                'case_id': self.case.id, 'exception_type': 'policy', 'severity': 'critical',
            })
        expected = calendar.plan_hours(2, fields.Datetime.to_datetime('2026-01-02 16:00:00'), compute_leaves=True)
        self.assertEqual(exception.due_at, fields.Datetime.to_datetime(expected))
        self.assertEqual(exception.due_at.date().isoformat(), '2026-01-06')

    def test_dashboard_metric_contract_empty_state_and_no_iap(self):
        empty_company = self.env['res.company'].create({'name': 'Dashboard Empty Company'})
        with patch('odoo.addons.insilos_logistics_idp.services.document_processor.IAPDocumentProcessor.process') as process:
            data = self.env['logistics.idp.case'].with_company(empty_company).get_dashboard_data({
                'company_ids': [empty_company.id], 'date_from': '2099-01-01', 'date_to': '2099-01-31',
            })
        process.assert_not_called()
        self.assertIsNone(data['metrics']['stp_rate']['value'])
        self.assertIsNone(data['metrics']['stp_rate']['delta'])
        self.assertEqual(data['metrics']['received_today']['value'], 0)
        self.assertEqual(data['metrics']['processed_today']['value'], 0)
        for key in ('waiting_supplier', 'waiting_broker_customs'):
            self.assertTrue(data['metrics'][key]['available'])
            self.assertEqual(data['metrics'][key]['value'], 0)
        self.assertTrue(data['metrics']['duplicate_detected']['available'])
        self.assertEqual(data['metrics']['duplicate_detected']['value'], 0)
        for key in ('overdue_import_declaration_cases', 'low_confidence_extraction_queue',
                    'average_processing_duration', 'ai_credits_per_document_case', 'source_erp_api_mes'):
            self.assertFalse(data['metrics'][key]['available'])
            self.assertEqual(data['metrics'][key]['contract']['status'], 'unavailable')
        self.assertEqual(data['strategic_panels']['status'], 'not_configured')
        onboarding = data['onboarding']
        self.assertEqual((onboarding['close_model'], onboarding['close_method']),
                         ('onboarding.onboarding', 'action_close_panel_logistics_idp'))
        self.assertEqual(len(onboarding['steps']), 6)
        self.assertEqual([step['action'] for step in onboarding['steps']], [
            'action_open_logistics_intake', 'action_open_logistics_collecting',
            'action_open_logistics_processing', 'action_open_logistics_review',
            'action_open_logistics_exceptions', 'action_open_logistics_completion',
        ])
        self.assertEqual({row['state'] for row in data['case_states']}, set())
        self.assertLessEqual(len(data['live_queue']), 50)
        for metric in data['metrics'].values():
            self.assertEqual(set(metric['contract']), {
                'formula', 'source_model', 'domain', 'date_field', 'unit', 'drilldown_available', 'status',
            })

    def test_dashboard_duplicate_metric_drilldown_and_unavailable_semantics(self):
        before_count = self.env['logistics.idp.evidence'].search_count([
            ('case_id.company_id', '=', self.env.company.id),
            ('category', '=', 'exact_document_duplicate')
        ])
        evidence = self.env['logistics.idp.evidence']._controlled_create({
            'case_id': self.case.id, 'category': 'exact_document_duplicate',
            'source_reference': 'sha256:' + unique_fixture('dashboard-duplicate'), 'status': 'valid', 'payload': {},
        }, 'dashboard_duplicate_test')
        foreign_company = self.env['res.company'].create({'name': unique_fixture('dashboard-duplicate-foreign')})
        foreign_case = self.env['logistics.idp.case'].with_company(foreign_company).create({
            'name': 'DASH-DUPLICATE-FOREIGN', 'company_id': foreign_company.id, 'source_system': 'fixture',
            'source_key': unique_fixture('dashboard-duplicate-foreign'), 'source_version': 'v1',
            'provenance': 'synthetic', 'effective_date': '2026-01-01',
        })
        self.env['logistics.idp.evidence'].with_company(foreign_company)._controlled_create({
            'case_id': foreign_case.id, 'category': 'exact_document_duplicate',
            'source_reference': 'sha256:' + unique_fixture('dashboard-duplicate-foreign'), 'status': 'valid', 'payload': {},
        }, 'dashboard_duplicate_test')
        filters = {'company_ids': [self.env.company.id]}
        data = self.env['logistics.idp.case'].get_dashboard_data(filters)
        metric = data['metrics']['duplicate_detected']
        self.assertEqual((metric['value'], metric['contract']['source_model'], metric['contract']['status']),
                         (before_count + 1, 'logistics.idp.evidence', 'available'))
        action = self.env['logistics.idp.case'].dashboard_drilldown('duplicate_detected', filters)
        records = self.env[action['res_model']].search(action['domain'])
        self.assertIn(evidence, records)
        self.assertNotIn(foreign_case.id, records.mapped('case_id').ids)
        for key in ('overdue_import_declaration_cases', 'low_confidence_extraction_queue',
                    'average_processing_duration', 'ai_credits_per_document_case'):
            metric = data['metrics'][key]
            self.assertEqual((metric['value'], metric['available'], metric['contract']['status']),
                             (None, False, 'unavailable'))
            with self.assertRaises(ValidationError):
                self.env['logistics.idp.case'].dashboard_drilldown(key, filters)

    def test_srs_section_7_3_case_workspace_cockpit_navigation_contract(self):
        case = self.case
        actions = {
            'action_open_documents': ('logistics.idp.document', [('case_id', '=', case.id)]),
            'action_open_runs': ('logistics.idp.extraction.run', [('case_id', '=', case.id)]),
            'action_open_checks': ('logistics.idp.check.result', [('case_id', '=', case.id)]),
            'action_open_outputs': ('logistics.idp.output', [('case_id', '=', case.id)]),
            'action_open_overrides': ('logistics.idp.override', [('case_id', '=', case.id)]),
            'action_open_exceptions': ('logistics.idp.exception', [('case_id', '=', case.id)]),
            'action_open_shipping_plan_outputs': ('logistics.idp.output', [
                ('case_id', '=', case.id), ('output_type', '=', 'shipping_plan')]),
            'action_open_gate_pass_outputs': ('logistics.idp.output', [
                ('case_id', '=', case.id), ('output_type', '=', 'gate_pass')]),
        }
        for method, (model, domain) in actions.items():
            action = getattr(case, method)()
            self.assertEqual((action['res_model'], action['domain'], action['context']),
                             (model, domain, {'create': False}))

        arch = self.env.ref('insilos_logistics_idp.view_logistics_case_form').arch
        for page in ('Documents', 'Runs', 'Checks', 'Outputs', 'Overrides', 'Audit and Compliance'):
            self.assertIn('<page string="%s">' % page, arch)
        for method in actions:
            self.assertIn('name="%s"' % method, arch)
        for absent in ('Purchase Orders', 'Email Source', 'Broker', 'Import'):
            self.assertNotIn('<page string="%s">' % absent, arch)

    def test_srs_section_7_1_control_tower_available_kpis_and_required_drilldowns(self):
        document = self.env['logistics.idp.document'].create({
            'case_id': self.case.id, 'document_type': 'invoice', 'source_channel': 'email',
            'content_hash': unique_fixture('control-tower-document'), 'mimetype': 'application/pdf', 'status': 'valid',
        })
        exception = self.env['logistics.idp.exception'].create({
            'case_id': self.case.id, 'exception_type': 'price', 'severity': 'high', 'state': 'open',
        })
        output = self.env['logistics.idp.output'].generate(self.case, 'e11', {})
        foreign_company = self.env['res.company'].create({'name': unique_fixture('control-tower-foreign-company')})
        foreign_case = self.env['logistics.idp.case'].with_company(foreign_company).create({
            'name': 'CONTROL-TOWER-FOREIGN', 'company_id': foreign_company.id, 'source_system': 'fixture',
            'source_key': unique_fixture('control-tower-foreign-case'), 'source_version': 'v1',
            'provenance': 'synthetic', 'effective_date': '2026-01-01', 'supplier_reference': self.case.supplier_reference,
        })
        foreign_document = self.env['logistics.idp.document'].with_company(foreign_company).create({
            'case_id': foreign_case.id, 'document_type': 'invoice', 'source_channel': 'email',
            'content_hash': unique_fixture('control-tower-foreign-document'), 'mimetype': 'application/pdf', 'status': 'valid',
        })
        foreign_exception = self.env['logistics.idp.exception'].with_company(foreign_company).create({
            'case_id': foreign_case.id, 'exception_type': 'price', 'severity': 'high', 'state': 'open',
        })
        foreign_output = self.env['logistics.idp.output'].with_company(foreign_company).generate(foreign_case, 'e11', {})
        filters = {'company_ids': [self.env.company.id]}
        dashboard = self.env['logistics.idp.case'].get_dashboard_data(filters)
        for key in ('open_cases', 'received_today', 'processed_today', 'waiting_supplier',
                    'waiting_broker_customs', 'duplicate_detected', 'stp_rate'):
            self.assertTrue(dashboard['metrics'][key]['available'], key)
        self.assertIn('high', {row['severity'] for row in dashboard['exception_severity']})
        self.assertEqual((dashboard['metrics']['overdue_import_declaration_cases']['value'],
                          dashboard['metrics']['overdue_import_declaration_cases']['available']), (None, False))
        for key in ('low_confidence_extraction_queue', 'average_processing_duration',
                    'ai_credits_per_document_case', 'source_erp_api_mes'):
            self.assertEqual((dashboard['metrics'][key]['value'], dashboard['metrics'][key]['available'],
                              dashboard['metrics'][key]['contract']['status']), (None, False, 'unavailable'))
            with self.assertRaises(ValidationError):
                self.env['logistics.idp.case'].dashboard_drilldown(key, filters)
        for metric, value, record, foreign_record in (
            ('supplier', self.case.supplier_reference, exception, foreign_exception),
            ('document_type', 'invoice', document, foreign_document),
            ('document_throughput', dashboard['drilldown_domains']['document_history'] + [
                ('create_date', '>=', document.create_date),
                ('create_date', '<', document.create_date + timedelta(days=1))], document, foreign_document),
            ('case_state', self.case.state, self.case, foreign_case),
            ('exception', 'price', exception, foreign_exception),
            ('customs_regime', 'e11', output, foreign_output),
            ('source', 'email', document, foreign_document),
        ):
            action = self.env['logistics.idp.case'].dashboard_drilldown(metric, filters, value)
            records = self.env[action['res_model']].search(action['domain'])
            self.assertIn(record, records, metric)
            self.assertNotIn(foreign_record, records, metric)

    def test_dashboard_ocr_aggregates_history_and_read_only_provider_boundary(self):
        document = self.env['logistics.idp.document'].create({
            'case_id': self.case.id, 'document_type': 'invoice',
            'content_hash': unique_fixture('dashboard-ocr'), 'mimetype': 'application/pdf',
            'page_count': 3, 'status': 'valid',
        })
        runs = self.env['logistics.idp.extraction.run']
        now = fields.Datetime.now()
        first = runs._controlled_create({
            'case_id': self.case.id, 'document_id': document.id, 'provider': 'provider-a',
            'model_version': 'model-a', 'schema_version': 'schema-v1', 'prompt_version': 'prompt-v1',
            'started_at': now, 'completed_at': now,
            'duration_seconds': 2, 'confidence': .9, 'status': 'valid', 'payload': {},
        }, 'dashboard_ocr_test')
        second = runs._controlled_create({
            'case_id': self.case.id, 'document_id': document.id, 'provider': 'provider-b',
            'model_version': 'model-b', 'schema_version': 'schema-v1', 'prompt_version': 'prompt-v1',
            'started_at': now, 'completed_at': now,
            'duration_seconds': 5, 'confidence': .4, 'status': 'review', 'payload': {},
        }, 'dashboard_ocr_test')
        filters = {'company_ids': [self.env.company.id]}
        before = {run.id: (run.company_id.id, run.document_type, run.page_count, run.write_date)
                  for run in (first | second)}
        with patch('odoo.addons.insilos_logistics_idp.services.document_processor.IAPDocumentProcessor.process') as process, \
             patch.object(type(self.env['openrouter.router']), 'complete', autospec=True) as complete, \
             patch.object(type(self.env['iap.charge']), 'hold', autospec=True) as hold, \
             patch.object(type(self.env['iap.charge']), 'consume', autospec=True) as consume, \
             patch.object(type(self.env['iap.charge']), 'release', autospec=True) as release, \
             patch.object(type(self.env['iap.charge']), 'refund', autospec=True) as refund:
            data = self.env['logistics.idp.case'].get_dashboard_data(filters)
            action = self.env['logistics.idp.case'].dashboard_drilldown('ocr_history', filters)
        self.env.flush_all()
        self.assertEqual(before, {run.id: (run.company_id.id, run.document_type, run.page_count, run.write_date)
                                  for run in (first | second)})
        process.assert_not_called()
        complete.assert_not_called()
        hold.assert_not_called()
        consume.assert_not_called()
        release.assert_not_called()
        refund.assert_not_called()
        by_provider = {row['provider']: row for row in data['ocr']}
        self.assertEqual((by_provider['provider-a']['__count'], by_provider['provider-a']['confidence:avg'],
                          by_provider['provider-a']['duration_seconds:avg'], by_provider['provider-a']['page_count:sum']),
                         (1, .9, 2, 3))
        self.assertEqual((by_provider['provider-b']['__count'], by_provider['provider-b']['confidence:avg'],
                          by_provider['provider-b']['duration_seconds:avg'], by_provider['provider-b']['page_count:sum']),
                         (1, .4, 5, 3))
        self.assertTrue({'model-a', 'model-b'} <= {row['model_version'] for row in data['ocr_by_model']})
        self.assertTrue({'valid', 'review'} <= {row['status'] for row in data['ocr_by_status']})
        self.assertEqual((action['res_model'], action['domain']),
                         ('logistics.idp.extraction.run', data['drilldown_domains']['ocr_history']))
        self.assertTrue(first | second <= self.env[action['res_model']].search(action['domain']))

    def test_dashboard_exception_dimensions_ranking_drilldowns_and_company_scope(self):
        date = '2042-01-01'
        opened_at = date + ' 00:00:00'
        company = self.env['res.company'].create({'name': unique_fixture('Dashboard Exception')})
        supplier = unique_fixture('dashboard-exception-supplier')
        case = self.env['logistics.idp.case'].with_company(company).create({
            'name': 'DASH-EXCEPTION', 'company_id': company.id, 'source_system': 'fixture',
            'source_key': unique_fixture('dashboard-exception'), 'source_version': 'v1',
            'provenance': 'synthetic', 'effective_date': date, 'supplier_reference': supplier,
        })
        exceptions = self.env['logistics.idp.exception'].with_company(company).create([
            {'case_id': case.id, 'exception_type': 'price', 'severity': 'high', 'state': 'waiting', 'opened_at': opened_at},
            {'case_id': case.id, 'exception_type': 'policy', 'severity': 'critical', 'state': 'open', 'opened_at': opened_at},
        ])
        other = self.env['logistics.idp.case'].with_company(company).create({
            'name': 'DASH-EXCEPTION-OTHER', 'company_id': company.id, 'source_system': 'fixture',
            'source_key': unique_fixture('dashboard-exception-other'), 'source_version': 'v1',
            'provenance': 'synthetic', 'effective_date': date,
            'supplier_reference': unique_fixture('dashboard-exception-other-supplier'),
        })
        self.env['logistics.idp.exception'].with_company(company).create({
            'case_id': other.id, 'exception_type': 'price', 'severity': 'low', 'state': 'resolved', 'opened_at': opened_at,
        })
        foreign_company = self.env['res.company'].create({'name': 'Dashboard Exception Foreign'})
        foreign = self.env['logistics.idp.case'].with_company(foreign_company).create({
            'name': 'DASH-EXCEPTION-FOREIGN', 'company_id': foreign_company.id, 'source_system': 'fixture',
            'source_key': unique_fixture('dashboard-exception-foreign'), 'source_version': 'v1',
            'provenance': 'synthetic', 'effective_date': date,
            'supplier_reference': supplier,
        })
        self.env['logistics.idp.exception'].with_company(foreign_company).create({
            'case_id': foreign.id, 'exception_type': 'policy', 'severity': 'critical', 'state': 'open', 'opened_at': opened_at,
        })
        env = self.env(context=dict(self.env.context, allowed_company_ids=[company.id]))
        filters = {'company_ids': [company.id], 'date_from': date, 'date_to': date}
        data = env['logistics.idp.case'].get_dashboard_data(filters)
        self.assertEqual([(row['severity'], row['__count']) for row in data['exception_severity']],
                         [('critical', 1), ('high', 1), ('low', 1)])
        self.assertEqual([(row['state'], row['__count']) for row in data['exception_states']],
                         [('open', 1), ('waiting', 1), ('resolved', 1)])
        self.assertEqual([(row['supplier_reference'], row['__count']) for row in data['top_suppliers'][:2]],
                         [(supplier, 2), (other.supplier_reference, 1)])
        self.assertEqual(next(row['__count'] for row in data['top_suppliers']
                              if row['supplier_reference'] == supplier), 2)
        for metric, value in (('exception', 'policy'), ('exception_severity', 'critical'),
                              ('exception_state', 'open'), ('supplier', supplier)):
            action = env['logistics.idp.case'].dashboard_drilldown(metric, filters, value)
            records = env[action['res_model']].sudo().search(action['domain'])
            self.assertTrue(exceptions & records)
            self.assertNotIn(foreign.id, records.mapped('case_id').ids)

    def test_dashboard_ocr_model_and_status_drilldowns(self):
        document = self.env['logistics.idp.document'].create({
            'case_id': self.case.id, 'document_type': 'invoice',
            'content_hash': unique_fixture('dashboard-ocr'), 'mimetype': 'application/pdf', 'status': 'valid',
        })
        runs = self.env['logistics.idp.extraction.run']
        created = {}
        for model_version, status in (('vision-v1', 'valid'), ('vision-v2', 'review')):
            created[model_version] = runs._controlled_create({
                'case_id': self.case.id, 'document_id': document.id, 'provider': 'test',
                'model_version': model_version, 'schema_version': 'v1', 'prompt_version': 'v1',
                'started_at': '2026-01-01 10:00:00', 'status': status, 'payload': {},
            }, 'dashboard_ocr_test')
        filters = {'company_ids': [self.env.company.id]}
        for metric, value, run in (('ocr_model', 'vision-v1', created['vision-v1']),
                                   ('ocr_status', 'review', created['vision-v2'])):
            action = self.env['logistics.idp.case'].dashboard_drilldown(metric, filters, value)
            self.assertEqual(action['res_model'], 'logistics.idp.extraction.run')
            self.assertIn(run, self.env[action['res_model']].search(action['domain']))

    def test_dashboard_filters_formulas_drilldown_and_company_validation(self):
        completed = self.env['logistics.idp.case'].create({
            'name': 'DASH-STP', 'source_system': 'fixture', 'source_key': 'DASH-STP',
            'source_version': 'v1', 'provenance': 'synthetic', 'effective_date': '2026-02-10',
            'supplier_reference': 'DASH-SUP',
        })
        self.cr.execute("UPDATE logistics_idp_case SET state = 'completed', verdict = 'pass' WHERE id = %s", [completed.id])
        completed.invalidate_recordset(['state', 'verdict'])
        blocked = self.env['logistics.idp.case'].create({
            'name': 'DASH-BLOCK', 'source_system': 'fixture', 'source_key': 'DASH-BLOCK',
            'source_version': 'v1', 'provenance': 'synthetic', 'effective_date': '2026-02-11',
            'supplier_reference': 'DASH-SUP', 'state': 'blocked', 'verdict': 'block',
        })
        document = self.env['logistics.idp.document'].create({
            'case_id': completed.id, 'document_type': 'invoice', 'content_hash': unique_fixture('dashboard-document'),
            'mimetype': 'application/pdf', 'page_count': 2, 'status': 'valid',
        })
        self.cr.execute('UPDATE logistics_idp_document SET create_date = %s WHERE id = %s',
                        ['2026-02-10 12:00:00', document.id])
        document.invalidate_recordset(['create_date'])
        data = self.env['logistics.idp.case'].get_dashboard_data({
            'company_ids': [self.env.company.id], 'date_from': '2026-02-01', 'date_to': '2026-02-28',
            'supplier': 'DASH-SUP', 'document_type': 'invoice',
        })
        self.assertEqual(data['metrics']['documents']['value'], 1)
        self.assertEqual(data['metrics']['stp_rate']['value'], 1.0)
        self.assertEqual(data['metrics']['stp_rate']['contract']['formula'],
                         'completed PASS cases with no review/block check, review/error extraction run, or override history in period / all completed cases in period; null when denominator=0')
        self.assertIn('company resource calendar plan_hours', data['metrics']['overdue_cases']['contract']['formula'])
        self.assertEqual(data['metrics']['blocked_cases']['value'], 0)
        action = self.env['logistics.idp.case'].dashboard_drilldown('documents', {
            'company_ids': [self.env.company.id], 'date_from': '2026-02-01', 'date_to': '2026-02-28',
            'supplier': 'DASH-SUP', 'document_type': 'invoice',
        })
        self.assertEqual(action['res_model'], 'logistics.idp.document')
        self.assertEqual(action['views'], [(False, 'list'), (False, 'form')])
        self.assertEqual(self.env[action['res_model']].search_count(action['domain']), 1)
        foreign_company = self.env['res.company'].create({'name': 'Dashboard Forbidden Company'})
        restricted = new_test_user(self.env, login='dashboard-company-user', groups='insilos_logistics_idp.group_logistics_operator')
        with self.assertRaises(ValidationError):
            self.env['logistics.idp.case'].with_user(restricted).get_dashboard_data({'company_ids': [foreign_company.id]})
        self.assertTrue(blocked)
        self.assertTrue(document)

    def test_dashboard_document_metric_applies_type_company_and_inclusive_datetime_dates(self):
        boundary = self.env['logistics.idp.case'].create({
            'name': 'DASH-BOUNDARY', 'source_system': 'dashboard-boundary',
            'source_key': unique_fixture('dashboard-boundary'), 'source_version': 'v1',
            'provenance': 'synthetic', 'effective_date': '2026-04-30',
        })
        documents = self.env['logistics.idp.document'].create([
            {'case_id': boundary.id, 'document_type': 'invoice', 'content_hash': unique_fixture('dashboard-start'),
             'mimetype': 'application/pdf', 'status': 'valid'},
            {'case_id': boundary.id, 'document_type': 'invoice', 'content_hash': unique_fixture('dashboard-end'),
             'mimetype': 'application/pdf', 'status': 'valid'},
            {'case_id': boundary.id, 'document_type': 'packing_list', 'content_hash': unique_fixture('dashboard-type'),
             'mimetype': 'application/pdf', 'status': 'valid'},
        ])
        for document, created_at in zip(documents, ('2026-04-01 00:00:00', '2026-04-30 23:59:59', '2026-04-30 23:59:59')):
            self.cr.execute('UPDATE logistics_idp_document SET create_date = %s WHERE id = %s', [created_at, document.id])
        data = self.env['logistics.idp.case'].get_dashboard_data({
            'company_ids': [self.env.company.id], 'date_from': '2026-04-01', 'date_to': '2026-04-30',
            'document_type': 'invoice',
        })
        self.assertEqual(data['metrics']['documents']['value'], 2)
        self.assertEqual(sum(row['__count'] for row in data['document_throughput']), 2)
        action = self.env['logistics.idp.case'].dashboard_drilldown('documents', {
            'company_ids': [self.env.company.id], 'date_from': '2026-04-01', 'date_to': '2026-04-30',
            'document_type': 'invoice',
        })
        self.assertEqual(self.env[action['res_model']].search_count(action['domain']), 2)

    def test_dashboard_document_distributions_and_drilldowns_are_company_isolated(self):
        document = self.env['logistics.idp.document'].create({
            'case_id': self.case.id, 'document_type': 'invoice', 'source_channel': 'email',
            'content_hash': unique_fixture('dashboard-document-distribution'), 'mimetype': 'application/pdf', 'status': 'valid',
        })
        output = self.env['logistics.idp.output'].generate(self.case, 'e11', {})
        foreign_company = self.env['res.company'].create({'name': unique_fixture('Dashboard Distribution Foreign')})
        foreign_case = self.env['logistics.idp.case'].with_company(foreign_company).create({
            'name': 'DASH-DISTRIBUTION-FOREIGN', 'company_id': foreign_company.id, 'source_system': 'fixture',
            'source_key': unique_fixture('dashboard-distribution-foreign'), 'source_version': 'v1',
            'provenance': 'synthetic', 'effective_date': '2026-01-01',
        })
        foreign_document = self.env['logistics.idp.document'].with_company(foreign_company).create({
            'case_id': foreign_case.id, 'document_type': 'invoice', 'source_channel': 'email',
            'content_hash': unique_fixture('dashboard-document-distribution-foreign'), 'mimetype': 'application/pdf', 'status': 'valid',
        })
        foreign_output = self.env['logistics.idp.output'].with_company(foreign_company).generate(foreign_case, 'e11', {})
        filters = {'company_ids': [self.env.company.id]}
        data = self.env['logistics.idp.case'].get_dashboard_data(filters)
        self.assertIn('invoice', {row['document_type'] for row in data['document_types']})
        for metric, value, record, foreign_record in (('source', 'email', document, foreign_document),
                                                       ('customs_regime', 'e11', output, foreign_output),
                                                       ('document_type', 'invoice', document, foreign_document)):
            action = self.env['logistics.idp.case'].dashboard_drilldown(metric, filters, value)
            records = self.env[action['res_model']].search(action['domain'])
            self.assertIn(record, records)
            self.assertNotIn(foreign_record, records)
        with self.assertRaises(ValidationError):
            self.env['logistics.idp.case'].dashboard_drilldown('document_type', filters, 'not_a_document_type')

    def test_dashboard_native_onboarding_queue_and_case_state_drilldown(self):
        supplier_reference = unique_fixture('dashboard-onboarding-supplier')
        self.case.write({
            'state': 'review', 'severity': 'high', 'next_action': 'Verify invoice',
            'supplier_reference': supplier_reference,
        })
        data = self.env['logistics.idp.case'].get_dashboard_data({
            'company_ids': [self.env.company.id], 'supplier': supplier_reference,
        })
        queue = next(row for row in data['live_queue'] if row['id'] == self.case.id)
        self.assertEqual({key: queue[key] for key in (
            'state', 'severity', 'next_action', 'document_count', 'check_count', 'output_count')}, {
                'state': 'review', 'severity': 'high', 'next_action': 'Verify invoice',
                'document_count': 0, 'check_count': 0, 'output_count': 0,
            })
        self.assertIn('review', {row['state'] for row in data['case_states']})
        action = self.env['logistics.idp.case'].dashboard_drilldown(
            'case_state', {'company_ids': [self.env.company.id]}, 'review')
        self.assertEqual(action['res_model'], 'logistics.idp.case')
        self.assertIn(self.case, self.env[action['res_model']].search(action['domain']))

    def test_native_onboarding_steps_open_existing_action_without_progress_and_close_persists(self):
        onboarding = self.env.ref('insilos_logistics_idp.onboarding_logistics_idp')
        onboarding._search_or_create_progress()
        steps = onboarding.step_ids.sorted('sequence')
        before = steps.mapped('current_step_state')
        expected_domains = [
            [('state', '=', 'intake')], [('state', '=', 'collecting')],
            [('state', '=', 'processing')], [('state', '=', 'review')],
            [('state', 'in', ('blocked', 'waiting_external'))],
            [('state', 'in', ('ready', 'completed'))],
        ]
        for step, domain in zip(steps, expected_domains):
            action = getattr(self.env['onboarding.onboarding.step'], step.panel_step_open_action_name)()
            self.assertEqual(action['res_model'], 'logistics.idp.case')
            self.assertEqual(action['domain'], domain)
        self.assertEqual(steps.mapped('current_step_state'), before)
        self.env['onboarding.onboarding'].action_close_panel_logistics_idp()
        onboarding.invalidate_recordset(['current_progress_id', 'is_onboarding_closed'])
        self.assertTrue(onboarding.current_progress_id.is_onboarding_closed)
        self.assertTrue(onboarding.is_onboarding_closed)

    def test_dashboard_stp_history_exclusions(self):
        today = fields.Date.today()
        today_string = fields.Date.to_string(today)
        today_datetime = '%s 10:00:00' % today_string
        completed_datetime = '%s 10:00:01' % today_string

        def completed_case(suffix):
            case = self.env['logistics.idp.case'].create({
                'name': 'DASH-STP-%s' % suffix, 'source_system': 'stp-test', 'source_key': suffix,
                'source_version': 'v1', 'provenance': 'synthetic', 'effective_date': today,
                'supplier_reference': 'STP-%s' % suffix,
            })
            self.cr.execute("UPDATE logistics_idp_case SET state = 'completed', verdict = 'pass' WHERE id = %s", [case.id])
            case.invalidate_recordset(['state', 'verdict'])
            return case

        def stp_rate(case):
            return self.env['logistics.idp.case'].get_dashboard_data({
                'company_ids': [self.env.company.id], 'date_from': today_string, 'date_to': today_string,
                'supplier': case.supplier_reference,
            })['metrics']['stp_rate']['value']

        for status, expected in (('valid', 1.0), ('review', 0.0), ('error', 0.0)):
            case = completed_case('RUN-%s' % status)
            document = self.env['logistics.idp.document'].create({
                'case_id': case.id, 'document_type': 'invoice', 'content_hash': unique_fixture('stp-%s' % status),
                'mimetype': 'application/pdf', 'status': status,
            })
            self.env['logistics.idp.extraction.run']._controlled_create({
                'case_id': case.id, 'document_id': document.id, 'provider': 'test',
                'model_version': 'v1', 'schema_version': 'v1', 'prompt_version': 'v1',
                'started_at': today_datetime, 'completed_at': completed_datetime,
                'status': status, 'payload': {},
            }, 'stp_test')
            self.assertEqual(stp_rate(case), expected)

        for verdict in ('review', 'block'):
            case = completed_case('CHECK-%s' % verdict)
            self.env['logistics.idp.check.result'].create({
                'case_id': case.id, 'code': 'stp-%s' % verdict, 'verdict': verdict,
                'rationale': 'synthetic', 'payload': {},
            })
            self.assertEqual(stp_rate(case), 0.0)

        case = completed_case('OVERRIDE')
        check = self.env['logistics.idp.check.result']._controlled_create({
            'case_id': case.id, 'code': 'stp-override', 'verdict': 'pass',
            'rationale': 'synthetic', 'payload': {},
        }, 'stp_test')
        exception = self.env['logistics.idp.exception'].create({
            'case_id': case.id, 'exception_type': 'missing_document'})
        case.request_override(exception, 'evidence_gap', 'synthetic', checks=check)
        self.assertEqual(stp_rate(case), 0.0)

    def test_fr104_supplier_email_filename_is_audit_only_until_content_classification(self):
        self.case.owner_id = self.operator
        content = b'%PDF-1.4\n/Type /Page\n%%EOF'
        job = self.env['logistics.idp.inbound.job'].intake_supplier_email({
            'message_id': '<fr104-%s@example.test>' % unique_fixture('supplier-email'),
            'subject': 'Invoice PO-FR104', 'po_reference': 'PO-FR104',
            'company_id': self.env.company.id,
        }, [{'filename': 'invoice-fr104.pdf', 'mimetype': 'application/pdf', 'content': content}])
        result = {'document_type': 'invoice', 'payload': {}, 'confidence': 1,
                  'classification_confidence': 1, 'validation_warnings': [], 'source_spans': []}
        with patch('odoo.addons.insilos_logistics_idp.models.logistics_idp.IAPDocumentProcessor.process',
                   return_value=result) as provider:
            self.assertTrue(job._process_one())
            document = job.case_id.document_ids
            self.assertEqual((document.document_type, document.classification_confidence, document.status),
                             ('unknown', 0, 'review'))
            evidence = job.case_id.evidence_ids.filtered(lambda item: item.category == 'email_attachment')
            payload = json.loads(evidence.payload)
            self.assertEqual((payload['filename'], job.case_id.email_subject),
                             ('invoice-fr104.pdf', 'Invoice PO-FR104'))
            document.with_user(self.operator).process(explicit=True)
        provider.assert_called_once()
        self.assertEqual(document.document_type, 'invoice')

    def test_supplier_email_intake_is_private_idempotent_and_keeps_pdf_hint_nonterminal(self):
        pdf = b'%PDF-1.4\n/Type /Page\n%%EOF'
        message = {'message_id': '<queue-test@example.test>', 'subject': 'Packing list PO-Q',
                   'po_reference': 'PO-Q', 'company_id': self.env.company.id}
        attachments = [
            {'filename': 'packing list.pdf', 'mimetype': 'application/pdf', 'content': pdf},
            {'filename': 'packing list.pdf', 'mimetype': 'application/pdf', 'content': pdf},
            {'filename': 'unsupported.bin', 'mimetype': 'application/octet-stream', 'content': b'unsupported'},
        ]
        with patch('odoo.addons.insilos_logistics_idp.services.document_processor.IAPDocumentProcessor.process') as process:
            job = self.env['logistics.idp.inbound.job'].intake_supplier_email(message, attachments)
            self.assertEqual(self.env['logistics.idp.inbound.job'].intake_supplier_email(message, attachments), job)
            self.assertNotIn(pdf.hex(), job.payload)
            self.assertEqual(len(job.attachment_ids), 2)
            self.assertTrue(job._process_one())
        process.assert_not_called()
        self.assertEqual(len(job.case_id.document_ids), 1)
        document = job.case_id.document_ids
        self.assertEqual((document.status, document.document_type, document.classification_confidence),
                         ('review', 'unknown', 0))
        evidence = job.case_id.evidence_ids.filtered(lambda item: item.category == 'email_attachment')
        self.assertEqual(len(evidence), 3)
        self.assertEqual(len(evidence.mapped('source_reference')), len(set(evidence.mapped('source_reference'))))
        self.assertTrue(job._process_one())
        self.assertEqual(self.env['logistics.idp.exception'].search_count([
            ('case_id', '=', job.case_id.id), ('state', 'in', ('open', 'waiting'))]), 1)
        self.assertEqual(self.env['mail.activity'].search_count([
            ('res_model', '=', job.case_id._name), ('res_id', '=', job.case_id.id)]), 1)
        other = self.env['res.company'].create({'name': 'Supplier Email Other Company'})
        self.env.user.company_ids |= other
        other_job = self.env['logistics.idp.inbound.job'].intake_supplier_email(
            {**message, 'company_id': other.id}, [])
        self.assertNotEqual(other_job, job)

    def test_supplier_email_unsupported_attachment_has_safe_terminal_manifest_evidence(self):
        filename = 'secret-name.bin'
        sender = 'secret-sender@example.test'
        job = self.env['logistics.idp.inbound.job'].intake_supplier_email({
            'message_id': '<unsupported-%s@example.test>' % unique_fixture('inbound'),
            'source_sender': sender, 'company_id': self.env.company.id,
        }, [
            {'filename': filename, 'mimetype': 'application/octet-stream', 'content': b'unsupported'},
            {'filename': filename, 'content': b''},
            {'filename': filename, 'content': b'x' * (MAX_BYTES + 1)},
            {'filename': filename, 'content': None},
        ])
        with patch('odoo.addons.insilos_logistics_idp.models.logistics_idp.batch_structure', side_effect=ValidationError('secret detail')):
            for _attempt in range(job.max_attempts):
                self.assertFalse(job._process_one())
        evidence = job.case_id.evidence_ids.filtered(lambda item: item.category == 'inbound_failure')
        self.assertEqual((job.state, len(evidence)), ('dead', 4))
        payloads = [json.loads(item.payload) for item in evidence]
        self.assertEqual({item.source_reference for item in evidence}, {'%s:%s' % (job.source_key, index) for index in range(4)})
        self.assertEqual({payload['status'] for payload in payloads}, {'unsupported', 'empty', 'oversized', 'missing'})
        self.assertEqual({payload['manifest_identity'] for payload in payloads}, {
            hashlib.sha256(b'unsupported').hexdigest(), hashlib.sha256(b'').hexdigest(),
            hashlib.sha256(b'x' * (MAX_BYTES + 1)).hexdigest(),
            hashlib.sha256(json.dumps({'index': 3, 'status': 'missing'}, separators=(',', ':'), sort_keys=True).encode()).hexdigest(),
        })
        for item in evidence:
            self.assertEqual(item.status, 'invalid')
            self.assertNotIn(filename, item.payload)
            self.assertNotIn(sender, item.payload)
            self.assertNotIn('secret detail', item.payload)

    def test_generic_inbound_attachment_payload_is_rejected(self):
        with self.assertRaisesRegex(ValidationError, 'cannot carry attachment metadata'):
            self.env['logistics.idp.inbound.job'].create({
                'profile_code': 'local-test', 'source_system': 'synthetic-erp',
                'source_key': unique_fixture('generic-attachment'), 'source_version': 'v1',
                'payload': json.dumps({'name': 'Synthetic', 'attachments': [{'index': 0}]}),
            })
        job = self.env['logistics.idp.inbound.job'].create({
            'profile_code': 'local-test', 'source_system': 'synthetic-erp',
            'source_key': unique_fixture('generic-empty'), 'source_version': 'v1',
            'payload': json.dumps({'name': 'Synthetic', 'provenance': 'fixture:development', 'effective_date': '2026-01-01'}),
        })
        self.assertTrue(job._process_one())

    def test_generic_inbound_attachment_m2m_is_rejected_on_create_and_write(self):
        attachment = self.env['ir.attachment'].create({
            'name': 'generic-bypass.pdf', 'raw': b'%PDF-1.4', 'mimetype': 'application/pdf',
        })
        values = {
            'profile_code': 'local-test', 'source_system': 'synthetic-erp',
            'source_key': unique_fixture('generic-m2m-create'), 'source_version': 'v1',
            'payload': json.dumps({'name': 'Synthetic', 'provenance': 'fixture:development', 'effective_date': '2026-01-01'}),
        }
        with self.assertRaisesRegex(ValidationError, 'controlled supplier-email intake'):
            self.env['logistics.idp.inbound.job'].create({**values, 'attachment_ids': [(6, 0, attachment.ids)]})
        job = self.env['logistics.idp.inbound.job'].create({**values, 'source_key': unique_fixture('generic-m2m-write')})
        with self.assertRaisesRegex(ValidationError, 'controlled supplier-email intake'):
            job.write({'attachment_ids': [(6, 0, attachment.ids)]})
        with self.assertRaisesRegex(ValidationError, 'controlled supplier-email intake'):
            job.with_context(_logistics_inbound_write=True).write({'attachment_ids': [(6, 0, attachment.ids)]})

    def test_dossier_exact_duplicate_disposition_is_durable_and_company_safe(self):
        content = unique_fixture('DUP-001-003').encode()
        canonical = self.env['logistics.idp.document'].intake_content(
            self.case, content, 'application/pdf', {'filename': 'canonical.pdf'}, process=False)
        other_case = self.env['logistics.idp.case'].create({
            'name': 'DUP-CASE', 'source_system': 'fixture', 'source_key': unique_fixture('dup-case'),
            'source_version': 'v1', 'provenance': 'synthetic', 'effective_date': '2026-01-01',
        })
        duplicate = self.env['logistics.idp.document'].intake_content(
            other_case, content, 'application/pdf', {'filename': 'copy.pdf'}, process=False)
        self.assertEqual(duplicate, canonical)
        self.assertFalse(other_case.document_ids)
        dispositions = other_case.evidence_ids.filtered(lambda item: item.category == 'exact_document_duplicate')
        self.assertEqual(len(dispositions), 1)
        self.assertEqual(json.loads(dispositions.payload), {
            'canonical_case_id': self.case.id, 'canonical_document_id': canonical.id,
            'content_hash': canonical.content_hash, 'disposition': 'excluded_from_aggregation',
        })
        self.env['logistics.idp.document'].intake_content(
            other_case, content, 'application/pdf', {'filename': 'copy-again.pdf'}, process=False)
        self.assertEqual(len(other_case.evidence_ids.filtered(
            lambda item: item.category == 'exact_document_duplicate')), 1)

    def test_document_duplicate_marker_is_durable_idempotent_and_does_not_mutate_case(self):
        canonical = self.env['logistics.idp.document'].intake_content(
            self.case, unique_fixture('document-duplicate-canonical').encode(), 'application/pdf', process=False)
        duplicate_case = self.env['logistics.idp.case'].create({
            'name': 'DOCUMENT-DUPLICATE', 'source_system': 'fixture', 'source_key': unique_fixture('document-duplicate'),
            'source_version': 'v1', 'provenance': 'synthetic', 'effective_date': '2026-01-01',
        })
        duplicate = self.env['logistics.idp.document'].intake_content(
            duplicate_case, unique_fixture('document-duplicate-source').encode(), 'application/pdf', process=False)
        before = (duplicate_case.state, duplicate_case.canonical_case_id, duplicate.case_id, duplicate.document_type,
                  duplicate.current_run_id, duplicate.extraction_run_ids)
        self.assertEqual(duplicate.with_user(self.operator).action_mark_duplicate(canonical), duplicate)
        self.assertEqual(duplicate.duplicate_of_id, canonical)
        self.assertEqual(duplicate.duplicate_marked_by_id, self.operator)
        evidence = duplicate_case.evidence_ids.filtered(lambda item: item.category == 'document_duplicate_mark')
        self.assertEqual(len(evidence), 1)
        payload = json.loads(evidence.payload)
        self.assertEqual((payload['source_document_id'], payload['canonical_document_id'], payload['actor_id']),
                         (duplicate.id, canonical.id, self.operator.id))
        self.assertIn('marked_at', payload)
        self.assertEqual(duplicate.with_user(self.operator).action_mark_duplicate(canonical), duplicate)
        self.assertEqual(len(duplicate_case.evidence_ids.filtered(
            lambda item: item.category == 'document_duplicate_mark')), 1)
        self.assertEqual((duplicate_case.state, duplicate_case.canonical_case_id, duplicate.case_id, duplicate.document_type,
                          duplicate.current_run_id, duplicate.extraction_run_ids), before)

    def test_dossier_waiting_lifecycle_history_and_dashboard_metrics(self):
        since = fields.Datetime.now()
        self.case.action_set_waiting('supplier', 'Corrected invoice required', since=since)
        self.assertEqual((self.case.state, self.case.waiting_party, self.case.waiting_reason,
                          self.case.waiting_since),
                         ('waiting_external', 'supplier', 'Corrected invoice required', since))
        self.case._derive_lifecycle()
        self.assertEqual(self.case.state, 'waiting_external')
        data = self.env['logistics.idp.case'].get_dashboard_data({
            'company_ids': [self.env.company.id], 'supplier': self.case.supplier_reference})
        self.assertTrue(data['metrics']['waiting_supplier']['available'])
        self.assertEqual(data['metrics']['waiting_supplier']['value'], 1)
        self.case.action_resume()
        self.assertFalse(self.case.waiting_party)
        self.assertNotEqual(self.case.state, 'waiting_external')
        history = self.case.evidence_ids.filtered(lambda item: item.category == 'waiting_transition')
        self.assertEqual(len(history), 2)
        with self.assertRaises(UserError):
            history.unlink()

    def test_fr102_upload_case_inbox_dedup_and_boundaries(self):
        self.case.owner_id = self.operator
        document_model = self.env['logistics.idp.document'].with_user(self.operator)
        upload = self.case.with_user(self.operator).action_open_upload_wizard()
        self.assertEqual(upload['context']['default_case_id'], self.case.id)
        content = b'%PDF-1.4\n/Type /Page\n%%EOF\nfr102'
        values = {'case_id': self.case.id, 'filename': 'fr102.pdf',
                  'content': base64.b64encode(content), 'mimetype': 'application/pdf'}
        with patch('odoo.addons.insilos_logistics_idp.models.logistics_idp.IAPDocumentProcessor.process') as provider:
            wizard = self.env['logistics.idp.document.upload.wizard'].with_user(self.operator).create(values)
            self.assertEqual(wizard.action_upload()['type'], 'is.actions.act_window_close')
            document = self.case.document_ids
            self.assertEqual((len(document), document.status, document.document_type), (1, 'processing', 'unknown'))
            self.assertEqual(document_model.upload_intake(self.case, content, 'application/pdf', 'copy.pdf'), document)
            inbox = self.env['logistics.idp.document.upload.wizard'].with_user(self.operator).create({
                'filename': 'inbox.pdf', 'content': base64.b64encode(content + b'-inbox'), 'mimetype': 'application/pdf'})
            self.assertEqual(inbox.action_upload()['type'], 'is.actions.act_window_close')
            inbox_case = self.env['logistics.idp.case'].search([('source_system', '=', 'manual_upload')], order='id desc', limit=1)
            self.assertEqual((inbox_case.state, inbox_case.document_ids.status), ('collecting', 'processing'))
            with self.assertRaisesRegex(ValidationError, 'valid base64'):
                self.env['logistics.idp.document.upload.wizard'].with_user(self.operator).create({
                    'case_id': self.case.id, 'filename': 'bad.pdf', 'content': b'%', 'mimetype': 'application/pdf'}).action_upload()
            with self.assertRaisesRegex(ValidationError, 'maximum supported size'):
                document_model.upload_intake(self.case, b'x' * (MAX_BYTES + 1), 'application/pdf', 'large.pdf')
            self.case.with_context(_logistics_case_lifecycle=_INTERNAL_CASE_LIFECYCLE_TOKEN).write({'state': 'completed', 'completion_reason': 'fixture'})
            with self.assertRaisesRegex(ValidationError, 'terminal cases'):
                document_model.upload_intake(self.case, b'terminal', 'application/pdf', 'terminal.pdf')
        provider.assert_not_called()

    def test_ops008_authorized_inbox_workflow_preserves_source_and_audit_evidence(self):
        message_reference = '<ops-008-%s@example.test>' % unique_fixture('message')
        message = {
            'message_id': message_reference, 'source_sender': 'supplier@example.test',
            'subject': 'OPS-008 invoice', 'company_id': self.env.company.id,
        }
        with patch('odoo.addons.insilos_logistics_idp.services.document_processor.IAPDocumentProcessor.process') as provider:
            job = self.env['logistics.idp.inbound.job'].intake_supplier_email(message, [{
                'filename': 'ops-008.pdf', 'mimetype': 'application/pdf',
                'content': ('%%PDF-1.4\n/Type /Page\n%%EOF\n%s' % unique_fixture('ops008-content')).encode(),
            }])
            self.assertTrue(job._process_one())
        provider.assert_not_called()
        document = job.case_id.document_ids
        self.assertEqual((document.source_sender, document.source_message_reference),
                         ('supplier@example.test', message_reference))
        self.assertTrue(document.attachment_id.exists() and document.document_id.exists())
        document.sudo().write({'document_type': 'invoice'})
        with patch.object(type(document), 'process', autospec=True, return_value=document) as process:
            self.assertEqual(document.with_user(self.operator).action_reclassify(), document)
        process.assert_called_once_with(document.with_user(self.operator), {
            'document_type': 'invoice', 'document_type_trusted': True}, explicit=True)
        with patch.object(type(document), 'process', autospec=True, return_value=document) as process:
            self.assertEqual(document.with_user(self.operator).action_reextract(), document)
        process.assert_called_once_with(document.with_user(self.operator), explicit=True)
        target = self.env['logistics.idp.case'].create({
            'name': 'OPS-008 TARGET', 'source_system': 'fixture', 'source_key': unique_fixture('ops008-target'),
            'source_version': 'v1', 'provenance': 'synthetic', 'effective_date': '2026-01-01',
            'owner_id': self.operator.id,
        })
        movable = self.env['logistics.idp.document'].intake_content(
            job.case_id, unique_fixture('ops008-move').encode(), 'application/pdf', process=False)
        source = movable.case_id
        source.owner_id = self.operator
        movable.with_user(self.operator).action_assign_case(target)
        movable.with_user(self.operator).action_assign_case(source)
        self.assertEqual(movable.case_id, source)
        self.assertEqual(len(source.evidence_ids.filtered(
            lambda item: item.category == 'document_assignment')), 2)
        self.assertEqual(len(target.evidence_ids.filtered(
            lambda item: item.category == 'document_assignment')), 2)
        duplicate = self.env['logistics.idp.case'].create({
            'name': 'OPS-008 DUP', 'source_system': 'fixture', 'source_key': unique_fixture('ops008-duplicate'),
            'source_version': 'v1', 'provenance': 'synthetic', 'effective_date': '2026-01-01',
        })
        manager = new_test_user(
            self.env, login='ops008-manager-' + unique_fixture('manager'),
            context={'no_reset_password': True}, groups='insilos_logistics_idp.group_logistics_manager')
        self.assertTrue(duplicate.with_user(manager).mark_semantic_duplicate(source))
        self.assertEqual((duplicate.canonical_case_id, duplicate.state), (source, 'closed_duplicate'))
        duplicate_evidence = (duplicate | source).evidence_ids.filtered(
            lambda item: item.category == 'semantic_duplicate')
        self.assertEqual(len(duplicate_evidence), 2)
        self.assertTrue(all(item.audit_actor_id == manager for item in duplicate_evidence))
        email_evidence = source.evidence_ids.filtered(lambda item: item.category == 'email_attachment')
        self.assertEqual(len(email_evidence), 1)
        self.assertEqual(email_evidence.source_reference.split(':')[0], job.source_key)

    def test_inbox_wizard_combines_unprocessed_assignment_and_explicit_reclassification(self):
        target = self.env['logistics.idp.case'].create({
            'name': 'WIZARD TARGET', 'source_system': 'fixture', 'source_key': unique_fixture('wizard-target'),
            'source_version': 'v1', 'provenance': 'synthetic', 'effective_date': '2026-01-01',
            'owner_id': self.operator.id,
        })
        document = self.env['logistics.idp.document'].intake_content(
            self.case, b'{"document_type":"invoice","confidence":1,"payload":{}}',
            'application/json', process=False)
        action = document.with_user(self.operator).action_open_inbox_wizard()
        self.assertEqual((action['res_model'], action['target']),
                         ('logistics.idp.document.inbox.wizard', 'new'))
        wizard = self.env[action['res_model']].with_user(self.operator).create({
            'document_id': document.id, 'case_id': target.id, 'document_type': 'invoice'})
        self.assertEqual(wizard.action_apply()['type'], 'is.actions.act_window_close')
        self.assertEqual((document.case_id, document.document_type), (target, 'invoice'))
        self.assertTrue(document.current_run_id)
        self.assertTrue(target.evidence_ids.filtered(lambda item: item.category == 'document_assignment'))

    def test_ws1_negative_lifecycle_assignment_evidence_and_company_identity(self):
        target = self.env['logistics.idp.case'].create({
            'name': 'WS1-TARGET', 'source_system': 'fixture', 'source_key': unique_fixture('ws1-target'),
            'source_version': 'v1', 'provenance': 'synthetic', 'effective_date': '2026-01-01',
        })
        document = self.env['logistics.idp.document'].intake_content(
            self.case, b'{"document_type":"invoice","confidence":1,"payload":{}}',
            'application/json', process=False)
        document.with_user(self.operator).process()
        with self.assertRaises(ValidationError):
            document.with_user(self.operator).action_assign_case(target)
        self.case.with_context(_logistics_case_lifecycle=_INTERNAL_CASE_LIFECYCLE_TOKEN).write({
            'state': 'completed', 'completion_reason': 'fixture'})
        with self.assertRaises(UserError):
            self.case.write({'waiting_reason': 'mutated'})
        duplicate = self.env['logistics.idp.case'].create({
            'name': 'WS1-DUP', 'source_system': 'fixture', 'source_key': unique_fixture('ws1-dup'),
            'source_version': 'v1', 'provenance': 'synthetic', 'effective_date': '2026-01-01',
        })
        manager = new_test_user(
            self.env, login='ws1-duplicate-manager-' + unique_fixture('duplicate'),
            context={'no_reset_password': True}, groups='insilos_logistics_idp.group_logistics_manager')
        duplicate.with_user(manager).mark_semantic_duplicate(target)
        with self.assertRaises(ValidationError):
            self.case.with_user(manager).mark_semantic_duplicate(duplicate)
        with self.assertRaises(ValidationError):
            self.env['logistics.idp.evidence']._controlled_create({
                'case_id': target.id, 'category': 'unknown_ws1', 'source_reference': 'unknown',
                'payload': {},
            }, 'ws1_test')
        other = self.env['res.company'].create({'name': 'WS1 Other Company'})
        self.env.user.company_ids |= other
        identity = unique_fixture('ws1-company-inbound')
        values = {'profile_code': 'local-test', 'source_system': 'ws1', 'source_key': identity,
                  'source_version': '1', 'payload': '{}'}
        first = self.env['logistics.idp.inbound.job'].create(values)
        second = self.env['logistics.idp.inbound.job'].with_company(other).create({**values, 'company_id': other.id})
        self.assertNotEqual(first.company_id, second.company_id)

    def test_ws1_blocked_waiting_and_datetime_end_date(self):
        self.case.write({'document_status': 'block'})
        self.case.with_user(self.operator).action_set_waiting('supplier', 'Need correction')
        self.assertEqual((self.case.state, self.case.verdict), ('blocked', 'block'))
        self.cr.execute('UPDATE logistics_idp_case SET effective_date = %s WHERE id = %s',
                        ['2026-03-31', self.case.id])
        self.case.invalidate_recordset(['effective_date'])
        document = self.env['logistics.idp.document'].create({
            'case_id': self.case.id, 'content_hash': unique_fixture('ws1-end-date'),
            'mimetype': 'application/pdf', 'status': 'valid',
        })
        self.cr.execute('UPDATE logistics_idp_document SET create_date = %s WHERE id = %s',
                        ['2026-03-31 23:59:59', document.id])
        data = self.env['logistics.idp.case'].get_dashboard_data({
            'company_ids': [self.env.company.id], 'date_from': '2026-03-31', 'date_to': '2026-03-31'})
        self.assertEqual(data['metrics']['documents']['value'], 1)

    def test_malformed_pdf_is_quarantined_locally_until_content_changes(self):
        self.case.owner_id = self.operator
        malformed = b'%PDF-1.4\nnot-a-page\n%%EOF'
        document = self.env['logistics.idp.document'].intake_content(
            self.case, malformed, 'application/pdf', {'filename': 'private.pdf'}, process=False)
        with patch('odoo.addons.insilos_logistics_idp.models.logistics_idp.IAPDocumentProcessor.process') as provider:
            document.with_user(self.operator).process()
            document.with_user(self.operator).reprocess()
        provider.assert_not_called()
        self.assertEqual(document.status, 'error')
        runs = self.case.extraction_run_ids.filtered(lambda run: run.document_id == document)
        self.assertEqual(len(runs), 2)
        self.assertTrue(all(run.error == 'invalid_pdf_structure' for run in runs))
        self.assertNotIn('private.pdf', ''.join(runs.mapped('payload')))
        valid = b'%PDF-1.4\n/Type /Page\n%%EOF'
        changed = self.env['logistics.idp.document'].intake_content(
            self.case, valid, 'application/pdf', {'filename': 'replacement.pdf'}, process=False)
        with patch('odoo.addons.insilos_logistics_idp.models.logistics_idp.IAPDocumentProcessor.process',
                   side_effect=ValueError('provider boundary reached')) as provider:
            changed.with_user(self.operator).process()
        provider.assert_called_once()

    def test_malformed_trust_boundary_matrix_never_calls_provider(self):
        xlsx = io.BytesIO()
        with zipfile.ZipFile(xlsx, 'w') as archive:
            archive.writestr('[Content_Types].xml', '<Types/>')
            archive.writestr('xl/worksheets/sheet1.xml', '<worksheet/>')
            archive.writestr('xl/embeddings/object.bin', b'active')
        samples = [
            ('encrypted.pdf', b'%PDF-1.4\n/Encrypt\n/Type /Page\n%%EOF', 'application/pdf', 'encrypted_document'),
            ('polyglot.pdf', b'%PDF-1.4\n/Type /Page\n%%EOF<script>', 'application/pdf', 'unsafe_active_content'),
            ('embedded.pdf', b'%PDF-1.4\n/EmbeddedFile\n/Type /Page\n%%EOF', 'application/pdf', 'unsafe_active_content'),
            ('embedded.xlsx', xlsx.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'unsafe_embedded_content'),
            ('embedded.eml', b'From: a@example.test\nSubject: x\nContent-Disposition: attachment; filename=x.pdf\n\nbody', 'message/rfc822', 'unsafe_embedded_content'),
            ('unsupported.bin', b'not supported', 'application/octet-stream', 'unsupported_mime'),
        ]
        with patch('odoo.addons.insilos_logistics_idp.models.logistics_idp.IAPDocumentProcessor.process') as provider:
            for name, content, mimetype, category in samples:
                document = self.env['logistics.idp.document'].intake_content(
                    self.case, content, mimetype, {'filename': name}, process=False)
                document.with_user(self.operator).process()
                self.assertEqual((document.status, document.current_run_id.error), ('error', category))
                self.assertEqual(json.loads(document.current_run_id.payload)['human_action'],
                                 'Replace or re-upload the source document.')
            oversized = self.env['logistics.idp.document'].intake_content(
                self.case, b'x' * (20 * 1024 * 1024 + 1), 'application/pdf',
                {'filename': 'oversized.pdf'}, process=False)
            oversized.with_user(self.operator).process()
            self.assertEqual(oversized.current_run_id.error, 'invalid_document_size')
        provider.assert_not_called()

    def test_only_iap_ocr_geometry_persists_in_immutable_run_payload(self):
        self.case.owner_id = self.operator
        geometry = {'words': {'0': [{'content': 'INV-1', 'coords': [.5, .25, .2, .1, 0]}]},
                    'numbers': {}, 'dates': {}}
        content = b'%PDF-1.4\n/Type /Page\n%%EOF'
        document = self.env['logistics.idp.document'].intake_content(
            self.case, content, 'application/pdf', {'filename': 'trusted.pdf'}, process=False)
        result = {'document_type': 'invoice', 'payload': {}, 'confidence': 1,
                  'classification_confidence': 1, 'validation_warnings': [], 'source_spans': [],
                  'ocr_results': geometry}
        with patch('odoo.addons.insilos_logistics_idp.models.logistics_idp.IAPDocumentProcessor.process',
                   return_value=result):
            document.with_user(self.operator).process()
        payload = json.loads(document.current_run_id.payload)
        self.assertEqual((document.extracted_words, document.extracted_numbers or {}, document.extracted_dates or {}),
                         tuple(payload['ocr_results'][field] for field in ('words', 'numbers', 'dates')))
        self.assertEqual(payload['ocr_results'], geometry)
        self.assertEqual(document.get_boxes()['word']['0'][0], {
            'id': 1, 'text': 'INV-1', 'page': '0', 'minX': .4, 'midX': .5, 'maxX': .6,
            'minY': .2, 'midY': .25, 'maxY': .3, 'width': .2, 'height': .1, 'angle': 0,
        })
        with self.assertRaises(UserError):
            document.current_run_id.write({'payload': {}})

    def test_local_provider_response_cannot_persist_ocr_geometry(self):
        geometry = {'words': {'0': [{'content': 'LOCAL', 'coords': [.5, .25, .2, .1, 0]}]},
                    'numbers': {}, 'dates': {}}
        document = self.env['logistics.idp.document'].intake_content(
            self.case, json.dumps({'document_type': 'invoice', 'confidence': 1, 'payload': {},
                                   'ocr_results': geometry}).encode(), 'application/json', process=False)
        document.with_user(self.operator).process()
        self.assertFalse(document.extracted_words)
        self.assertNotIn('ocr_results', json.loads(document.current_run_id.payload))

    def test_document_type_hint_does_not_select_local_processor(self):
        self.case.owner_id = self.operator
        content = b'%PDF-1.4\n/Type /Page\n%%EOF'
        attachment = self.env['ir.attachment'].create({
            'name': 'hint-iap.pdf', 'raw': content, 'mimetype': 'application/pdf',
            'res_model': self.case._name, 'res_id': self.case.id,
        })
        document = self.env['logistics.idp.document'].create({
            'case_id': self.case.id, 'attachment_id': attachment.id, 'document_type': 'packing_list',
            'content_hash': hashlib.sha256(content).hexdigest(), 'mimetype': 'application/pdf', 'status': 'processing',
        })
        with patch('odoo.addons.insilos_logistics_idp.models.logistics_idp.IAPDocumentProcessor.process',
                   side_effect=ValueError('stop after processor selection')) as iap, patch(
                   'odoo.addons.insilos_logistics_idp.models.logistics_idp.SyntheticDocumentProcessor.process') as local:
            document.with_user(self.operator).process({'document_type': 'packing_list'})
        iap.assert_called_once()
        local.assert_not_called()

    def test_remediation_activity_and_reviewer_unavailable_are_company_safe_idempotent(self):
        company = self.env['res.company'].create({'name': 'Review Workflow Isolated'})
        other_company = self.env['res.company'].create({'name': 'Review Workflow Foreign'})
        service = new_test_user(
            self.env, login='logistics-review-service', company_id=company.id,
            company_ids=[(6, 0, [company.id])],
            groups='insilos_logistics_idp.group_logistics_reviewer',
        )
        new_test_user(
            self.env, login='logistics-foreign-human-reviewer', company_id=other_company.id,
            company_ids=[(6, 0, [other_company.id])],
            groups='insilos_logistics_idp.group_logistics_reviewer',
        )
        case = self.env['logistics.idp.case'].with_company(company).create({
            'name': 'REVIEW-WORKFLOW', 'company_id': company.id, 'owner_id': service.id,
            'source_system': 'synthetic-review', 'source_key': self._testMethodName,
            'source_version': 'v1', 'provenance': 'fixture:development', 'effective_date': '2026-01-01',
        })
        document = self.env['logistics.idp.document'].with_company(company).create({
            'case_id': case.id, 'document_type': 'invoice',
            'content_hash': hashlib.sha256(self._testMethodName.encode()).hexdigest(),
            'mimetype': 'application/pdf', 'status': 'error',
        })
        mail_count = self.env['mail.mail'].sudo().search_count([])
        document._sync_review_activity()
        document._sync_review_activity()
        activities = self.env['mail.activity'].sudo().search([
            ('res_model', '=', case._name), ('res_id', '=', case.id),
            ('summary', '=', 'IDP document %s: remediation' % document.id),
        ])
        exceptions = self.env['logistics.idp.exception'].sudo().search([
            ('case_id', '=', case.id), ('message_ids.body', 'ilike', '[reviewer_unavailable]'),
        ])
        self.assertEqual((len(activities), len(exceptions)), (1, 1))
        self.assertEqual((activities.user_id, case.owner_id), (service, service))
        self.assertEqual(exceptions.company_id, company)
        self.assertEqual(self.env['mail.mail'].sudo().search_count([]), mail_count)

    def test_198_detailed_and_full_mapped_golden_48_orm_pipeline(self):
        company = self.env['res.company'].create({'name': 'Logistics UAT Isolated'})
        operator = new_test_user(
            self.env, login='logistics-uat-operator', context={'no_reset_password': True}, company_id=company.id,
            company_ids=[(6, 0, [company.id])],
            groups='insilos_logistics_idp.group_logistics_operator',
        )
        policies = self.env['logistics.idp.policy.source'].search(
            [('code', '=', 'TRADE_COMPLIANCE_VN_REFERENCE')], order='effective_from desc, id desc')
        policy = next((item for item in policies if isinstance(json.loads(item.payload).get('horizontal', {}).get(
            'extraction', {}).get('critical_fields'), dict)), self.env['logistics.idp.policy.source'])
        self.assertTrue(policy)
        self.env['logistics.idp.policy.source']._controlled_create({
            'code': policy.code, 'version': policy.version, 'company_id': company.id,
            'jurisdiction': policy.jurisdiction, 'regime': policy.regime,
            'source_tier': policy.source_tier, 'citation': policy.citation, 'provenance': policy.provenance,
            'effective_from': policy.effective_from, 'effective_to': policy.effective_to,
            'expected_refresh_hours': policy.expected_refresh_hours, 'stale_after_hours': policy.stale_after_hours,
            'last_successful_sync_at': policy.last_successful_sync_at, 'stale_disposition': policy.stale_disposition,
            'state': 'active', 'payload': policy.payload,
        }, 'test_fixture')
        result = run_orm_corpus(self.env['logistics.idp.case'].with_company(company).env, operator, write_result=False)
        self.assertEqual(result['detailed_case_counts'], {'passed': 198, 'failed': 0, 'not_run': 0, 'skipped': 0})
        self.assertEqual(result['golden_counts'], {'passed': 46, 'partial': 2, 'failed': 0, 'not_run': 0, 'skipped': 0})
        self.assertTrue(all(item['mapped_detailed_ids'] for item in result['golden_scenarios']))

    def test_fr202_batch_chunks_larger_than_100_and_reports_created(self):
        references = self.env['logistics.idp.reference.snapshot']
        for reference_type in ('master_data', 'dsnavl'):
            references.import_upsert({
                'reference_type': reference_type, 'source_system': 'synthetic-erp',
                'source_key': '%s-%s' % (reference_type, self._testMethodName), 'source_version': 'v1',
                'provenance': 'fixture:development', 'effective_date': '2026-01-01',
                'payload': {'material_codes': ['MAT-1']},
            })
        records = [{
            'po_reference': 'PO-CHUNK-%03d' % index, 'snapshot_date': '2026-01-01',
            'source_system': 'synthetic-erp', 'source_key': 'PO-CHUNK-%03d' % index,
            'source_version': 'v1', 'provenance': 'fixture:development',
            'lines': [{'line_key': '10', 'material_code': 'MAT-1', 'ordered_quantity': 1, 'remaining_quantity': 1}],
        } for index in range(101)]
        result = self.env['logistics.idp.po.snapshot'].import_batch(records)
        self.assertEqual((len(result['snapshots']), result['report']['created']), (101, 101))

    def test_fr203_latest_snapshot_closes_absent_material_without_mutating_prior(self):
        for reference_type in ('master_data', 'dsnavl'):
            self.env['logistics.idp.reference.snapshot'].import_upsert({
                'reference_type': reference_type, 'source_system': 'synthetic-erp',
                'source_key': '%s-%s' % (reference_type, self._testMethodName), 'source_version': 'v1',
                'provenance': 'fixture:development', 'effective_date': '2026-01-01',
                'payload': {'material_codes': ['MAT-1', 'MAT-2']},
            })
        model = self.env['logistics.idp.po.snapshot']
        base = {'po_reference': 'PO-CLOSE', 'source_system': 'synthetic-erp', 'provenance': 'fixture:development'}
        first = model.import_upsert({**base, 'snapshot_date': '2026-01-01', 'source_key': 'PO-CLOSE', 'source_version': 'v1', 'lines': [
            {'line_key': '10', 'material_code': 'MAT-1', 'ordered_quantity': 2, 'remaining_quantity': 2},
            {'line_key': '20', 'material_code': 'MAT-2', 'ordered_quantity': 3, 'remaining_quantity': 3},
        ]})
        result = model.import_batch([{**base, 'snapshot_date': '2026-01-02', 'source_key': 'PO-CLOSE', 'source_version': 'v2', 'lines': [
            {'line_key': '10', 'material_code': 'MAT-1', 'ordered_quantity': 2, 'remaining_quantity': 1},
        ]}])
        self.assertEqual((first.line_ids.filtered(lambda line: line.line_key == '20').remaining_quantity, result['snapshots'].line_ids.filtered(lambda line: line.line_key == '20').remaining_quantity, result['report']['closed']), (3, 0, 1))

    def test_fr204_snapshot_retry_is_ignored_without_duplicates(self):
        values = {
            'po_reference': 'PO-RETRY', 'snapshot_date': '2026-01-01', 'source_system': 'synthetic-erp',
            'source_key': 'PO-RETRY', 'source_version': 'v1', 'provenance': 'fixture:development',
            'lines': [{'line_key': '10', 'no_material_code': True, 'ordered_quantity': 2, 'remaining_quantity': 1}],
        }
        model = self.env['logistics.idp.po.snapshot']
        self.assertEqual(model.import_batch([values, values])['report'], {'created': 1, 'updated': 0, 'stale': 0, 'closed': 0, 'ignored': 1, 'failed': 0, 'failures': []})
        self.assertEqual(model.search_count([('source_key', '=', 'PO-RETRY')]), 1)

    def test_fr204_backfill_is_stale_not_latest_snapshot(self):
        model = self.env['logistics.idp.po.snapshot']
        base = {'po_reference': 'PO-BACKFILL', 'source_system': 'synthetic-erp', 'provenance': 'fixture:development'}
        latest = model.import_upsert({**base, 'snapshot_date': '2026-01-02', 'source_key': 'PO-BACKFILL', 'source_version': 'v2', 'lines': [
            {'line_key': '10', 'no_material_code': True, 'ordered_quantity': 2, 'remaining_quantity': 2},
            {'line_key': '20', 'no_material_code': True, 'ordered_quantity': 3, 'remaining_quantity': 3},
        ]})
        result = model.import_batch([{**base, 'snapshot_date': '2026-01-01', 'source_key': 'PO-BACKFILL-OLD', 'source_version': 'v1', 'lines': [
            {'line_key': '10', 'no_material_code': True, 'ordered_quantity': 2, 'remaining_quantity': 1},
        ]}])
        self.assertEqual((result['report']['stale'], result['report']['closed'], len(result['snapshots'].line_ids)), (1, 0, 1))
        self.assertEqual(latest.line_ids.filtered(lambda line: line.line_key == '20').remaining_quantity, 3)

    def test_fr204_retry_with_dropped_line_is_ignored_without_closure(self):
        model = self.env['logistics.idp.po.snapshot']
        base = {'po_reference': 'PO-RETRY-DROP', 'snapshot_date': '2026-01-01', 'source_system': 'synthetic-erp', 'source_key': 'PO-RETRY-DROP', 'source_version': 'v1', 'provenance': 'fixture:development'}
        model.import_upsert({**base, 'lines': [
            {'line_key': '10', 'no_material_code': True, 'ordered_quantity': 2, 'remaining_quantity': 2},
            {'line_key': '20', 'no_material_code': True, 'ordered_quantity': 3, 'remaining_quantity': 3},
        ]})
        result = model.import_batch([{**base, 'lines': [
            {'line_key': '10', 'no_material_code': True, 'ordered_quantity': 2, 'remaining_quantity': 2},
        ]}])
        self.assertEqual((result['report']['ignored'], result['report']['closed']), (1, 0))

    def test_fr205_direct_import_rejects_missing_material(self):
        values = {
            'po_reference': 'PO-DIRECT-VALIDATE', 'snapshot_date': '2026-01-01', 'source_system': 'synthetic-erp',
            'source_key': 'PO-DIRECT-VALIDATE', 'source_version': 'v1', 'provenance': 'fixture:development',
            'lines': [{'line_key': '10', 'material_code': 'MISSING', 'ordered_quantity': 1, 'remaining_quantity': 1}],
        }
        with self.assertRaisesRegex(ValidationError, 'PO snapshot material validation failed'):
            self.env['logistics.idp.po.snapshot'].import_upsert(values)
        self.assertFalse(self.env['logistics.idp.po.snapshot'].search([('source_key', '=', 'PO-DIRECT-VALIDATE')]))

    def test_fr205_future_material_reference_does_not_validate_historical_po(self):
        material_code = 'MAT-FUTURE-' + self._testMethodName
        for reference_type in ('master_data', 'dsnavl'):
            self.env['logistics.idp.reference.snapshot'].import_upsert({
                'reference_type': reference_type, 'source_system': 'synthetic-erp',
                'source_key': '%s-%s' % (reference_type, self._testMethodName), 'source_version': 'v1',
                'provenance': 'fixture:development', 'effective_date': '2026-01-02',
                'payload': {'material_codes': [material_code]},
            })
        result = self.env['logistics.idp.po.snapshot'].import_batch([{
            'po_reference': 'PO-FUTURE-REF', 'snapshot_date': '2026-01-01', 'source_system': 'synthetic-erp',
            'source_key': 'PO-FUTURE-REF', 'source_version': 'v1', 'provenance': 'fixture:development',
            'lines': [{'line_key': '10', 'material_code': material_code, 'ordered_quantity': 1, 'remaining_quantity': 1}],
        }])
        self.assertEqual((len(result['snapshots']), result['report']['failures'][0]['missing']), (0, ['master_data', 'dsnavl']))

    def test_fr205_material_reference_failures_and_no_material_line(self):
        references = self.env['logistics.idp.reference.snapshot']
        for reference_type in ('master_data', 'dsnavl'):
            references.import_upsert({
                'reference_type': reference_type, 'source_system': 'synthetic-erp',
                'source_key': '%s-%s' % (reference_type, self._testMethodName), 'source_version': 'v1',
                'provenance': 'fixture:development', 'effective_date': '2026-01-01',
                'payload': {'material_codes': ['MAT-1']},
            })
        result = self.env['logistics.idp.po.snapshot'].import_batch([{
            'po_reference': 'PO-VALIDATE', 'snapshot_date': '2026-01-01', 'source_system': 'synthetic-erp',
            'source_key': 'PO-VALIDATE', 'source_version': 'v1', 'provenance': 'fixture:development', 'lines': [
                {'line_key': '10', 'material_code': 'MISSING', 'ordered_quantity': 1, 'remaining_quantity': 1},
                {'line_key': '20', 'no_material_code': True, 'ordered_quantity': 1, 'remaining_quantity': 1},
            ],
        }])
        self.assertEqual((len(result['snapshots']), result['report']['failed'], result['report']['failures'][0]['missing']), (0, 1, ['master_data', 'dsnavl']))
        allowed = self.env['logistics.idp.po.snapshot'].import_batch([{
            'po_reference': 'PO-NO-MATERIAL', 'snapshot_date': '2026-01-01', 'source_system': 'synthetic-erp',
            'source_key': 'PO-NO-MATERIAL', 'source_version': 'v1', 'provenance': 'fixture:development',
            'lines': [{'line_key': '10', 'no_material_code': True, 'ordered_quantity': 1, 'remaining_quantity': 1}],
        }])
        self.assertEqual((len(allowed['snapshots']), allowed['report']['failed']), (1, 0))

    def test_versioned_imports_are_idempotent(self):
        supplier_values = {
            'supplier_reference': 'SUP-IDEMP', 'profile_code': 'DEFAULT', 'version': '1',
            'source_system': 'synthetic-erp', 'source_key': 'SUP-IDEMP', 'source_version': 'v1',
            'provenance': 'fixture:development', 'effective_from': '2026-01-01', 'payload': {'approved': True},
        }
        supplier = self.env['logistics.idp.supplier.profile'].import_upsert(supplier_values)
        self.assertEqual(self.env['logistics.idp.supplier.profile'].import_upsert(supplier_values), supplier)
        po_values = {
            'po_reference': 'PO-IDEMP', 'snapshot_date': '2026-01-01', 'source_system': 'synthetic-erp',
            'source_key': 'PO-IDEMP', 'source_version': 'v1', 'provenance': 'fixture:development',
            'lines': [{'line_key': '10', 'no_material_code': True, 'ordered_quantity': 2, 'remaining_quantity': 1}],
        }
        snapshot = self.env['logistics.idp.po.snapshot'].import_upsert(po_values)
        self.assertEqual(self.env['logistics.idp.po.snapshot'].import_upsert(po_values), snapshot)
        self.assertEqual(len(snapshot.line_ids), 1)

    def test_fr507_variance_blocks_before_artifact(self):
        from decimal import Decimal
        self._sap_sources([{'line_key': '1', 'material_code': 'MAT-1'}], [{'line_key': '1', 'quantity': '1', 'unit_price': '2'}])
        with patch.object(type(self.case), '_sap_erp_amount', side_effect=[Decimal('2'), Decimal('3')]), self.assertRaisesRegex(ValidationError, 'total variance'):
            self.case.generate_sap_erp_output()
        self.assertFalse(self.case.output_ids)

    def test_fr601_template_positions_sheet_and_style_preserved(self):
        from openpyxl import load_workbook
        self._sap_sources([{'line_key': '1', 'material_code': 'MAT-1'}], [{'line_key': '1', 'quantity': '1', 'unit_price': '2'}])
        output = self.case.generate_sap_erp_output()
        sheet = load_workbook(io.BytesIO(output.attachment_id.raw))['SAP']
        self.assertEqual((sheet['A1'].value, sheet.cell(2, 1).value, sheet.cell(2, 3).value), ('template', 'SAP-MAT-1', 'MAT-1'))

    def test_fr602_field_level_configured_default_sources(self):
        for kind, values in (('dsnavl', {'vietnamese_name': 'Tên Việt', 'regime': 'E13'}),
                             ('master_data', {'uom': 'EA', 'mrp_controller': 'MRP', 'material_group': 'GROUP'})):
            self.env['logistics.idp.reference.snapshot'].import_upsert({'reference_type': kind, 'source_system': 'fixture', 'source_key': kind + self._testMethodName, 'source_version': '1', 'provenance': 'fixture', 'effective_date': '2026-01-01', 'payload': {'lines': [{'material_code': 'MAT-1', **values}]}})
        config = json.loads(self.sap_profile.payload)
        stage = config['sap_erp_output']['stages'][str(self.case.task_id.stage_id.id)]
        stage['source_precedence'] = {'*': ['invoice_evidence', 'po_snapshot', 'dsnavl', 'master_data'], 'uom': ['master_data'], 'vietnamese_name': ['dsnavl'], 'regime': ['dsnavl'], 'mrp_controller': ['master_data'], 'material_group': ['master_data']}
        stage['columns'] += [{'name': name, 'source': source, 'position': position} for position, (name, source) in enumerate((('quantity', 'line.quantity'), ('invoice_number', 'line.invoice_number'), ('invoice_date', 'line.invoice_date'), ('uom', 'line.uom'), ('po_terms', 'line.payment_terms'), ('dsnavl_name', 'dsnavl.vietnamese_name'), ('regime', 'dsnavl.regime'), ('mrp', 'master_data.mrp_controller'), ('material_group', 'master_data.material_group')), 4)]
        self._write_sap_config(config)
        self._sap_sources([{'line_key': '1', 'material_code': 'MAT-1', 'quantity': '9', 'uom': 'PO-UOM', 'payment_terms': 'NET-30'}], [{'line_key': '1', 'material_code': 'MAT-1', 'quantity': '1', 'unit_price': '2', 'invoice_number': 'INV-602', 'invoice_date': '2026-01-01', 'uom': 'INV-UOM', 'payment_terms': 'INV-TERMS'}])
        row = json.loads(self.case.generate_sap_erp_output().payload)['rows'][0]
        self.assertEqual(row['values'], {'material': 'MAT-1', 'custom': 'SAP-MAT-1', 'amount': '2', 'quantity': '1', 'invoice_number': 'INV-602', 'invoice_date': '2026-01-01', 'uom': 'EA', 'po_terms': 'NET-30', 'dsnavl_name': 'Tên Việt', 'regime': 'E13', 'mrp': 'MRP', 'material_group': 'GROUP'})

    def test_sap_erp_source_precedence_is_deterministic(self):
        self._sap_sources([{'line_key': '1', 'material_code': 'PO-MAT'}], [{'line_key': '1', 'material_code': 'INV-MAT', 'quantity': '1', 'unit_price': '2'}])
        config = json.loads(self.sap_profile.payload)
        stage = config['sap_erp_output']['stages'][str(self.case.task_id.stage_id.id)]
        stage['columns'][0]['source'] = 'line.material_code'
        self._write_sap_config(config)
        self.assertEqual(json.loads(self.case.generate_sap_erp_output().payload)['rows'][0]['values']['material'], 'PO-MAT')
        stage['source_precedence'] = ['invoice_evidence', 'po_snapshot']
        self._write_sap_config(config)
        self.assertEqual(json.loads(self.case.generate_sap_erp_output().payload)['rows'][0]['values']['material'], 'INV-MAT')

    def test_sap_erp_missing_material_has_no_reference_provenance(self):
        snapshot = self.env['logistics.idp.reference.snapshot'].import_upsert({'reference_type': 'dsnavl', 'source_system': 'fixture', 'source_key': self._testMethodName, 'source_version': '1', 'provenance': 'fixture', 'effective_date': '2026-01-01', 'payload': {'lines': [{'material_code': 'OTHER', 'vietnamese_name': 'Other'}]}})
        config = json.loads(self.sap_profile.payload)
        config['sap_erp_output']['stages'][str(self.case.task_id.stage_id.id)]['columns'].append({'name': 'dsnavl_name', 'source': 'dsnavl.vietnamese_name', 'position': 4, 'required': False})
        self._write_sap_config(config)
        self._sap_sources([{'line_key': '1', 'material_code': 'MAT-1'}], [{'line_key': '1', 'quantity': '1', 'unit_price': '2'}])
        payload = json.loads(self.case.generate_sap_erp_output().payload)
        source = payload['rows'][0]['sources']['dsnavl_name']
        self.assertEqual((source['record_type'], source['record_id'], source['record_hash']), (False, False, False))
        self.assertNotIn({'id': snapshot.id, 'hash': snapshot.payload_hash}, payload['source_records'])
        config['sap_erp_output']['stages'][str(self.case.task_id.stage_id.id)]['columns'][-1]['required'] = True
        self._write_sap_config(config)
        with self.assertRaisesRegex(ValidationError, 'dsnavl_name .*no matching material'):
            self.case.generate_sap_erp_output()

    def test_fr603_missing_source_blocks_without_output(self):
        self._sap_sources([{'line_key': '1', 'material_code': 'MAT-1'}], [{'line_key': '2', 'quantity': '1', 'unit_price': '2'}])
        with self.assertRaisesRegex(ValidationError, 'no PO match'):
            self.case.generate_sap_erp_output()
        self.assertFalse(self.case.output_ids)

    def test_fr604_blank_invoice_and_po_material_use_configured_po_code_template(self):
        config = json.loads(self.sap_profile.payload)
        stage = config['sap_erp_output']['stages'][str(self.case.task_id.stage_id.id)]
        stage['custom_code_template'] = 'C-{po_number}-{po_item}'
        stage['columns'][0]['required'] = False
        self._write_sap_config(config)
        self._sap_sources([{'line_key': '10', 'material_code': '', 'custom_code': ''}], [{'line_key': '10', 'material_code': '', 'quantity': '1', 'unit_price': '2', 'custom_code': ''}])
        self.assertEqual(json.loads(self.case.generate_sap_erp_output().payload)['rows'][0]['values']['custom'], 'C-PO-1-10')

    def test_fr605a_configured_rounding_is_persisted_with_versioned_evidence(self):
        config = json.loads(self.sap_profile.payload)
        stage = config['sap_erp_output']['stages'][str(self.case.task_id.stage_id.id)]
        stage['normalization'] = {'version': 'supplier-currency-field-v1', 'rounding': [
            {'id': 'fallback', 'version': 'v1', 'field': 'amount', 'precision': '1', 'mode': 'down'},
            {'id': 'supplier-usd-template', 'version': 'v2', 'supplier_profile_code': 'DEFAULT', 'currency': 'USD', 'field': 'amount', 'template_version': 'sap-v1', 'precision': '0.01', 'mode': 'half_up'},
        ]}
        self._write_sap_config(config)
        self._sap_sources([{'line_key': '1', 'material_code': 'MAT-1'}], [{'line_key': '1', 'quantity': '1', 'unit_price': '1.005', 'currency': 'USD'}])
        payload = json.loads(self.case.generate_sap_erp_output().payload)
        self.assertEqual(payload['rows'][0]['normalization'], {'rounding': {'precision': '0.01', 'mode': 'half_up'}, 'rule': {'rule_id': 'supplier-usd-template', 'rule_version': 'v2', 'selectors': {'supplier_profile_code': 'DEFAULT', 'currency': 'USD', 'field': 'amount', 'template_version': 'sap-v1'}}, 'result': '1.01'})
        self._sap_sources([{'line_key': '2', 'material_code': 'MAT-2'}], [{'line_key': '2', 'quantity': '1', 'unit_price': '1.9', 'currency': 'EUR'}])
        payload = json.loads(self.case.generate_sap_erp_output().payload)
        self.assertEqual(payload['rows'][0]['normalization'], {'rounding': {'precision': '1', 'mode': 'down'}, 'rule': {'rule_id': 'fallback', 'rule_version': 'v1', 'selectors': {'field': 'amount'}}, 'result': '1'})

    def test_fr606_stable_po_order_and_idempotency(self):
        self.test_sap_erp_output_uses_authoritative_sources_and_persists_provenance()

    def test_operational_exposure_is_separate_candidate_evidence_and_idempotent(self):
        self.case.write({
            'document_status': 'pass', 'reconciliation_status': 'pass',
            'compliance_status': 'block', 'verdict': 'review',
        })
        before = {field: self.case[field] for field in (
            'document_status', 'reconciliation_status', 'compliance_status', 'verdict')}
        active_before = self.env['logistics.idp.policy.source'].search([
            ('company_id', 'in', (False, self.case.company_id.id)), ('state', '=', 'active')]).ids
        evidence = self.case.record_operational_exposure(
            'critical', 'tier4_candidate', 'tier4:signal:case:%s' % self.case.id,
            'fixture:tier4-candidate', {'candidate_id': 'tier4-1'})
        replay = self.case.record_operational_exposure(
            'critical', 'tier4_candidate', 'tier4:signal:case:%s' % self.case.id,
            'fixture:tier4-candidate', {'candidate_id': 'tier4-1'})
        self.case.invalidate_recordset()
        self.assertEqual((evidence, replay), (evidence, evidence))
        self.assertEqual((evidence.category, evidence.status), ('operational_exposure_signal', 'review'))
        self.assertEqual(json.loads(evidence.payload)['source_class'], 'tier4_candidate')
        self.assertEqual(self.case.operational_exposure, 'critical')
        self.assertEqual({field: self.case[field] for field in before}, before)
        self.assertEqual(self.env['logistics.idp.policy.source'].search([
            ('company_id', 'in', (False, self.case.company_id.id)), ('state', '=', 'active')]).ids, active_before)

    def test_fr607_configured_custom_code_format(self):
        self.test_sap_erp_blank_material_custom_code_and_price_normalization()

    def test_fr608_source_mapping_version_and_provenance(self):
        self.test_sap_erp_output_uses_authoritative_sources_and_persists_provenance()

    def test_fr609_policy_blocks_missing_ambiguous_and_unbound_before_evidence(self):
        config = json.loads(self.sap_profile.payload)
        config.pop('authoritative_reference_policy')
        self._write_sap_config(config)
        self._sap_sources([{'line_key': '1', 'material_code': 'MAT-1'}], [{'line_key': '1', 'quantity': '1', 'unit_price': '2'}])
        with self.assertRaisesRegex(ValidationError, 'policy is missing'):
            self.case.generate_sap_erp_output()
        self.assertFalse(self.case.output_ids)
        config['authoritative_reference_policy'] = {'policy_version': 'v1', 'rules': [
            {'id': 'one', 'stage': str(self.case.task_id.stage_id.id), 'field': 'sap_erp_output', 'source_type': 'output', 'output_types': ['sap_erp_output']},
            {'id': 'two', 'stage': str(self.case.task_id.stage_id.id), 'field': 'sap_erp_output', 'source_type': 'output', 'output_types': ['sap_erp_output']},
        ]}
        self._write_sap_config(config)
        with self.assertRaisesRegex(ValidationError, 'unambiguous'):
            self.case.generate_sap_erp_output()
        config['authoritative_reference_policy']['rules'] = [{
            'id': 'invalid', 'stage': str(self.case.task_id.stage_id.id), 'field': 'sap_erp_output',
            'source_type': 'caller_payload', 'output_types': ['sap_erp_output'],
        }]
        self._write_sap_config(config)
        with self.assertRaisesRegex(ValidationError, 'valid unambiguous'):
            self.case.generate_sap_erp_output()
        self.assertFalse(self.case.output_ids)

    def test_fr609_legacy_sap_requires_explicit_rule(self):
        config = json.loads(self.sap_profile.payload)
        rule = config['authoritative_reference_policy']['rules'][0]
        rule['output_types'] = ['e13']
        self._write_sap_config(config)
        self._sap_sources([{'line_key': '1', 'material_code': 'MAT-1'}], [{'line_key': '1', 'quantity': '1', 'unit_price': '2'}])
        with self.assertRaisesRegex(ValidationError, 'explicit SAP'):
            self.case.generate_sap_erp_output()
        rule['output_types'] = ['sap_erp_output']
        self._write_sap_config(config)
        self.assertEqual(self.case.generate_sap_erp_output().output_type, 'sap_erp_output')

    def test_fr609_changed_profile_hash_rejects_stale_customs_outputs(self):
        line = {'material_code': 'MAT-1', 'quantity': '1', 'unit_price': '2'}
        document = self.env['logistics.idp.document'].with_user(self.operator)._intake_synthetic_content(
            self.case, json.dumps({'document_type': 'import_declaration', 'confidence': 1, 'payload': {
                'declaration_number': 'TK-609', 'lines': [{**line, 'regime': 'E13'}, {**line, 'regime': 'E15'}]}}, sort_keys=True).encode(),
            'application/json', {'document_type': 'import_declaration', 'confidence': 1})
        config = json.loads(self.sap_profile.payload)
        for rule in config['authoritative_reference_policy']['rules']:
            rule['stage'] = str(self.case.task_id.stage_id.id)
        self._write_sap_config(config)
        self.case.generate_customs_output('E13', [line])
        self.case.generate_customs_output('E15', [line])
        config['authoritative_reference_policy']['policy_version'] = 'fr609-v2'
        self._write_sap_config(config)
        with self.assertRaisesRegex(ValidationError, 'unbound or stale'):
            self.case.reconcile_import_declaration(document)
        self.assertFalse(self.case.evidence_ids.filtered(lambda item: item.category == 'import_declaration'))

    def test_fr609_legacy_sap_reconciliation_requires_explicit_contract(self):
        config = json.loads(self.sap_profile.payload)
        rules = config['authoritative_reference_policy']['rules']
        declaration_rule = next(rule for rule in rules if rule['field'] == 'import_declaration')
        declaration_rule['output_types'] = ['sap_erp_output']
        for rule in rules:
            rule['stage'] = str(self.case.task_id.stage_id.id)
        stage = config['sap_erp_output']['stages'][str(self.case.task_id.stage_id.id)]
        stage['columns'] += [
            {'name': 'quantity', 'source': 'line.quantity', 'position': 4},
            {'name': 'uom', 'source': 'line.uom', 'position': 5},
        ]
        self._write_sap_config(config)
        self.assertEqual(self.case._authoritative_reference('import_declaration')[1]['output_types'], ['sap_erp_output'])
        document = self.env['logistics.idp.document'].with_user(self.operator)._intake_synthetic_content(
            self.case, json.dumps({'document_type': 'import_declaration', 'confidence': 1, 'payload': {
                'declaration_number': 'TK-LEGACY', 'lines': [{'material_code': 'MAT-1', 'quantity': '1', 'uom': 'EA', 'line_total': '2'}]}}, sort_keys=True).encode(),
            'application/json', {'document_type': 'import_declaration', 'confidence': 1})
        for rule in rules:
            rule['stage'] = str(self.case.task_id.stage_id.id)
        config['sap_erp_output']['stages'] = {str(self.case.task_id.stage_id.id): stage}
        self._write_sap_config(config)
        with self.assertRaisesRegex(ValidationError, 'missing generated SAP_ERP_OUTPUT'):
            self.case.reconcile_import_declaration(document)
        self._sap_sources([{'line_key': '1', 'material_code': 'MAT-1'}], [
            {'line_key': '1', 'material_code': 'MAT-1', 'quantity': '1', 'unit_price': '2', 'uom': 'EA'},
        ])
        sap = self.case.generate_sap_erp_output()
        self.assertEqual(sap.output_type, 'sap_erp_output')
        self.assertEqual(self.case.reconcile_import_declaration(document).status, 'valid')
        declaration_rule['output_types'] = ['e13']
        self._write_sap_config(config)
        with self.assertRaisesRegex(ValidationError, 'E13/E15'):
            self.case.reconcile_import_declaration(document)

    def test_fr609_policy_rule_selection_and_binding_tamper_fail_closed(self):
        config = json.loads(self.sap_profile.payload)
        rules = config['authoritative_reference_policy']['rules']
        rules[1]['id'] = 'wildcard-fallback'
        rules[1]['output_types'] = ['e13']
        self._write_sap_config(config)
        self.assertEqual(self.case._authoritative_reference('sap_erp_output')[1]['id'], 'sap-downstream')
        self.assertEqual(self.case._authoritative_reference('unknown')[1]['id'], 'wildcard-fallback')
        rules[0]['output_types'] = ['sap_erp_output', 'sap_erp_output']
        self._write_sap_config(config)
        with self.assertRaisesRegex(ValidationError, 'valid unambiguous'):
            self.case._authoritative_reference('sap_erp_output')
        rules[0]['output_types'] = ['sap_erp_output']
        self._write_sap_config(config)
        e13 = self.case.generate_customs_output('E13', [dict(self.po['lines'][0], quantity='1')])
        with self.assertRaisesRegex(Exception, 'Generated outputs are immutable'), self.env.cr.savepoint():
            self.env.cr.execute("UPDATE logistics_idp_output SET payload = %s WHERE id = %s", [
                json.dumps({**json.loads(e13.payload), 'authoritative_policy_binding': {'field': 'e15'}}), e13.id])

    def test_fr609_profile_effectivity_fails_closed(self):
        self.case.supplier_reference = 'SUP-NO-EFFECTIVE'
        with self.assertRaisesRegex(ValidationError, 'one effective'):
            self.case._authoritative_reference('sap_erp_output')
        self.case.supplier_reference = 'SUP-1'
        self.env['logistics.idp.supplier.profile'].import_upsert({
            'supplier_reference': 'SUP-1', 'profile_code': 'OVERLAP', 'version': '2', 'source_system': 'synthetic-erp',
            'source_key': unique_fixture('fr609-overlap'), 'source_version': 'v1', 'provenance': 'fixture:development',
            'effective_from': '2026-01-01', 'payload': json.loads(self.sap_profile.payload),
        })
        with self.assertRaisesRegex(ValidationError, 'one effective'):
            self.case._authoritative_reference('sap_erp_output')

    def test_fr902_full_import_declaration_canonical_evidence_and_mismatches(self):
        line = {'material_code': 'MAT-1', 'custom_code': 'PO-1-10', 'hs_code': '711311', 'quantity': '2', 'uom': 'EA',
                'line_total': '20', 'supplier': 'SUP-1', 'importer': 'IMP-1', 'currency': 'USD'}
        self._retag_authoritative_policy_rules()
        e13 = self.case.generate_customs_output('E13', [{**line, 'regime': 'E13'}], 'vn-2026')
        e15 = self.case.generate_customs_output('E15', [{**line, 'regime': 'E15'}], 'vn-2026')
        self.assertEqual((json.loads(e13.payload)['authoritative_policy_binding']['rule_id'], json.loads(e15.payload)['authoritative_policy_binding']['rule_id']), ('e13-downstream', 'e15-downstream'))
        def declaration(payload):
            return self.env['logistics.idp.document'].with_user(self.operator)._intake_synthetic_content(
                self.case, json.dumps({'document_type': 'import_declaration', 'confidence': 1, 'payload': payload}, sort_keys=True).encode(), 'application/json',
                {'document_type': 'import_declaration', 'confidence': 1, 'payload': payload})

        base = {'declaration_number': 'TK-902', 'supplier': 'SUP-1', 'importer': 'IMP-1', 'currency': 'USD',
                'total_invoice_value': '40', 'lines': [{**line, 'regime': 'E13'}, {**line, 'regime': 'E15'}]}
        def reconcile(item):
            config = json.loads(self.sap_profile.payload)
            for rule in config['authoritative_reference_policy']['rules']:
                rule['stage'] = str(self.case.task_id.stage_id.id)
            self._write_sap_config(config)
            self.case.generate_customs_output('E13', [{**line, 'regime': 'E13'}], 'vn-2026-reconciled')
            self.case.generate_customs_output('E15', [{**line, 'regime': 'E15'}], 'vn-2026-reconciled')
            return self.case.reconcile_import_declaration(item)

        evidence = reconcile(declaration(base))
        payload = json.loads(evidence.payload)
        self.assertEqual((evidence.status, payload['verdict'], payload['reference']), ('valid', 'pass', 'E13/E15'))
        self.assertEqual(payload['declaration_document']['id'], evidence.case_id.document_ids.filtered(lambda item: item.document_type == 'import_declaration')[-1].id)
        self.assertTrue(payload['declaration_document']['run_hash'] and payload['output_records'])
        for field in ('supplier', 'importer', 'currency', 'total_invoice_value', 'regime', 'hs_code', 'quantity', 'uom', 'line_total'):
            self.assertTrue(any(item['field'] == field and item['result'] == 'pass' for item in payload['results']))
        for field, value in (('supplier', 'SUP-X'), ('importer', 'IMP-X'), ('currency', 'EUR'), ('total_invoice_value', '41')):
            failed = reconcile(declaration({**base, field: value}))
            self.assertEqual((failed.status, json.loads(failed.payload)['verdict']), ('invalid', 'block'))
        for field, value in (('material_code', 'MAT-X'), ('hs_code', '999999'), ('quantity', '3'), ('uom', 'KG'), ('line_total', '21')):
            changed = {**base, 'lines': [{**base['lines'][0], field: value}, base['lines'][1]]}
            failed = reconcile(declaration(changed))
            self.assertEqual((failed.status, json.loads(failed.payload)['verdict']), ('invalid', 'block'))
        self.assertEqual(evidence.source_reference, 'TK-902')
        with self.assertRaisesRegex(ValidationError, 'canonical import-declaration'):
            self.case.reconcile_import_declaration(base)
